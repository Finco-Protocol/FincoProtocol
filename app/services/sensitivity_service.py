"""V4-3: Sensitivity analysis service.

Pure computation layer — no persistence side effects, no model changes.
Applies configurable shocks to ProjectInputs and re-runs the canonical engine.
"""
from __future__ import annotations

import csv
import io
import hashlib
import math
from dataclasses import dataclass, replace, fields
from typing import Any, Optional

from app.services.rate_shock import RateShockRefusal, apply_rate_shock_bps, rate_shock_evidence

# Shock type registry (key -> human label, unit)
SHOCK_REGISTRY: dict[str, tuple[str, str]] = {
    "capex": ("CAPEX", "%"),
    "opex": ("OPEX", "%"),
    "ppa_price": ("PPA Price", "%"),
    "merchant_price": ("Merchant Price", "%"),
    "yield": ("Yield (P50 Hours)", "%"),
    "availability": ("Availability", "%"),
    "interest_rate": ("Interest Rate", "bps"),
    "tax_rate": ("Tax Rate", "%"),
    # Data Center native drivers (Opus NEW-M-1) — vertical-gated; never
    # accepted for Solar/Wind/EV.
    "dc_service_price": ("DC Service Price (EUR/kW/month)", "%"),
    "dc_occupancy": ("DC Occupancy", "%"),
    "dc_pue": ("DC PUE", "%"),
    "dc_electricity_price": ("DC Electricity Price", "%"),
    # V4-7 BESS shocks
    "bess_arbitrage_spread": ("BESS Arbitrage Spread", "%"),
    "bess_cycles": ("BESS Cycles/Year", "%"),
    "bess_rte": ("BESS Round-Trip Efficiency", "%"),
    "bess_ancillary_price": ("BESS Ancillary Price", "%"),
    "bess_capacity_price": ("BESS Capacity Price", "%"),
}

# Relative percentage changes for non-rate shocks; rates use basis-point deltas.
DEFAULT_SHOCK_LEVELS = [-15.0, -10.0, -5.0, 5.0, 10.0, 15.0]
DEFAULT_RATE_SHOCK_LEVELS = [-200.0, -100.0, -50.0, 50.0, 100.0, 200.0]


class DCDriverContextRequiredError(ValueError):
    """Raised when a dc_* sensitivity shock is requested without the
    project's resolved canonical DataCenterDrivers context.

    Correction C (PR #155): dc_* shocks must start from the ACTUAL
    project/scenario source drivers (via ``drivers_from_snapshot`` on the
    resolved runtime snapshot).  Falling back to the generic reference
    driver set here would silently replace real user assumptions with
    reference defaults, so the seam fails closed instead.
    """

    def __init__(self, shock_type: str):
        self.shock_type = shock_type
        super().__init__(
            f"Sensitivity driver {shock_type!r} requires the resolved "
            "Data Center source-driver context (dc_drivers) for this "
            "project/scenario; none was provided. Generic reference "
            "defaults are NOT substituted silently."
        )


class UnsupportedSensitivityDriverError(ValueError):
    """Raised when a sensitivity driver is not supported for a vertical.

    Fail-closed per Opus finding NEW-M-1: an unsupported driver is never
    silently ignored, converted to another driver, or returned as a valid
    no-op result.  ``vertical`` and ``shock_type`` identify the exact
    rejected combination.
    """

    def __init__(self, vertical: str, shock_type: str):
        self.vertical = vertical
        self.shock_type = shock_type
        super().__init__(
            f"Sensitivity driver {shock_type!r} is not supported for "
            f"vertical {vertical!r}."
        )


# Vertical-specific driver capability (Opus NEW-M-1).
#
# Canonicity of each mapping: every driver key corresponds to exactly one
# existing canonical input path in _apply_shock.
#   - Data Center drivers map to the Data Center input contract introduced
#     with the Generic Data Center Reference (service price in
#     revenue.service_price_eur_kw_month lineage; occupancy as the
#     utilisation authority; PUE; purchased-electricity cost).  Renewable
#     revenue semantics (PPA price, merchant price, yield/P50 hours) are
#     economically meaningless for an IT-load capacity business and are
#     FORBIDDEN for Data Center.
#   - Solar/Wind/EV Charging keep their existing driver sets unchanged.
#   - EV Charging offers only the drivers that reach its authoritative economics
#     (FIN-V2-D02; see the EV entry below).
_VERTICAL_SENSITIVITY_DRIVERS: dict[str, frozenset[str]] = {
    "solar": frozenset({
        "capex", "opex", "ppa_price", "merchant_price", "yield",
        "availability", "interest_rate", "tax_rate",
    }),
    "wind": frozenset({
        "capex", "opex", "ppa_price", "merchant_price", "yield",
        "availability", "interest_rate", "tax_rate",
    }),
    # Data Center: occupancy / IT MW / PUE / electricity price / service
    # price / CAPEX / OPEX-style maintenance cost drivers.
    "data_center": frozenset({
        "capex", "opex", "dc_service_price", "dc_occupancy", "dc_it_mw",
        "dc_pue", "dc_electricity_price", "interest_rate", "tax_rate",
    }),
    # EV Charging (FIN-V2-D02): revenue is an energy-throughput schedule built by the EV runtime adapter
    # (``app.ev_charging_economics.apply_ev_charging_runtime_adapter``): ``first_merchant_operating_period_index = 0``
    # (no PPA volume) and an explicit per-calendar-year price schedule that overrides the merchant curve.  A ``ppa_price``
    # or ``merchant_price`` shock therefore changes a field the engine never reads (proved: revenue delta 0.0), and
    # ``availability`` is informational (EFLH are net of availability; the adapter pins it at 1.0, so +x% clamps to no
    # change and -x% moves revenue without the derived electricity cost).  None of the three is offered.  ``yield`` moves
    # the stabilised EFLH AND the derived electricity procurement cost together (see ``_scale_ev_energy_dependent_opex``).
    "ev_charging": frozenset({
        "capex", "opex", "yield", "interest_rate", "tax_rate",
    }),
}

# Aliases accepted as vertical keys (project_type strings seen at runtime).
_VERTICAL_ALIASES: dict[str, str] = {
    "solar": "solar", "solar_pv": "solar",
    "wind": "wind", "wind_onshore": "wind",
    "data center": "data_center", "data_center": "data_center",
    "datacenter": "data_center",
    "ev charging": "ev_charging", "ev_charging": "ev_charging",
    "ev": "ev_charging",
}


def resolve_vertical(vertical: str | None) -> str:
    """Normalize a vertical key; unknown values resolve to themselves so the
    capability lookup fails closed rather than matching a real vertical."""
    key = (vertical or "").strip().lower()
    return _VERTICAL_ALIASES.get(key, key)


def supported_sensitivity_drivers(vertical: str | None) -> frozenset[str]:
    """Return the supported sensitivity driver set for a vertical."""
    return _VERTICAL_SENSITIVITY_DRIVERS.get(resolve_vertical(vertical), frozenset())


def assert_driver_supported(vertical: str | None, shock_type: str) -> None:
    """Fail closed when a driver is unsupported for the vertical."""
    if shock_type not in supported_sensitivity_drivers(vertical):
        raise UnsupportedSensitivityDriverError(
            resolve_vertical(vertical), shock_type)

# KPI definitions: (attr_on_WaterfallResult, display_label, format_hint)
KPI_DEFS: list[tuple[str, str, str]] = [
    ("total_revenue_keur", "Revenue (kEUR)", "keur"),
    ("total_ebitda_keur", "EBITDA (kEUR)", "keur"),
    ("_cfads_keur", "CFADS (kEUR)", "keur"),
    ("total_tax_keur", "Corp. Tax (kEUR)", "keur"),
    ("total_senior_ds_keur", "Senior DS (kEUR)", "keur"),
    ("total_distribution_keur", "Equity Distribution (kEUR)", "keur"),
    ("project_irr", "Project IRR", "pct"),
    # H-3 truth: equity_irr is the pure share-capital return (EQUITY_ONLY);
    # shareholder-loan flows are NOT included in this number.
    ("equity_irr", "Equity IRR (share capital only)", "pct"),
    ("equity_npv", "Equity NPV (kEUR)", "keur"),
    ("actual_avg_dscr", "Avg DSCR", "x"),
    ("min_dscr", "Min DSCR", "x"),
    ("min_llcr", "Min LLCR", "x"),
]


class SensitivitySourceError(ValueError):
    """Typed read-only refusal; details never reveal another tenant's records."""

    def __init__(self, code: str, status: int = 409):
        self.code, self.status = code, status
        super().__init__(code)


@dataclass(frozen=True)
class SensitivitySource:
    inputs: Any
    label: str
    vertical: str
    snapshot: Any = None
    context: Any = None

    def validate_current(self):
        if self.context is not None:
            from app.v2.post_run_context import PostRunContextChanged
            try:
                self.context.validate_current()
            except (PostRunContextChanged, PermissionError, ValueError) as exc:
                raise SensitivitySourceError("SENSITIVITY_SOURCE_CHANGED") from exc

    @property
    def provenance(self):
        c = self.context
        return {"basis": "WORKING_COPY_DRAFT" if c else "EXPLICIT_FACTORY_ALIAS",
                "project_id": c.project_id if c else "",
                "scenario_id": (c.workspace.active_scenario_id or "") if c else "",
                "input_hash": hashlib.sha256(repr(self.inputs).encode()).hexdigest(),
                "working_copy_hash": c.identity.composite_hash if c else "", "label": self.label}


# Compatibility is a finite alias contract, never a catch-all for saved codes.
_FACTORY_SENSITIVITY_ALIASES = {
    "generic_wind_reference": ("create_generic_wind_reference", "wind"),
    "generic_solar_reference": ("create_generic_solar_reference", "solar"),
    "generic_data_center_reference": ("create_generic_data_center_reference", "data_center"),
    "generic_data_center": ("create_generic_data_center_reference", "data_center"),
    "generic_solar": ("create_default_solar_project", "solar"),
    "solar": ("create_default_solar_project", "solar"),
    "test-solar-1": ("create_default_solar_project", "solar"),
    "generic_wind": ("create_default_wind_project", "wind"),
    "wind": ("create_default_wind_project", "wind"),
    "test-wind-1": ("create_default_wind_project", "wind"),
}


def resolve_sensitivity_source(user: Any, project: str, scenario_id: str = "") -> SensitivitySource:
    """Current Working Copy + active scenario, or a named immutable factory.

    No provisioning, selection, backfill, Save or Run-history write. A named
    scenario must still be the project's selected scenario; stale links refuse.
    Materialisation reuses Run's existing cost, revenue, financing and profile
    authorities. Unsupported scenario configurations refuse before execution.
    """
    from app.persistence.projects_repository import resolve_accessible_project
    from app.ui.protected_reference_service import is_protected_reference
    from app.v2.post_run_context import PostRunRequestContext, PostRunContextChanged
    from app.workbook.workbook_identity import WorkbookIdentityError
    from app.workbook.service import WorkbookService
    from app.services.sensitivity_matrix import Basis, pre_financing_inputs, materialise_financing

    key = (project or "").strip()
    if not getattr(user, "user_id", None):
        raise SensitivitySourceError("AUTH_REQUIRED", 401)
    record, owner = resolve_accessible_project(user.user_id, key)
    alias = _FACTORY_SENSITIVITY_ALIASES.get(key.lower())
    if record is None and alias is not None and not scenario_id:
        from app import project_factories
        factory, vertical = alias
        return SensitivitySource(getattr(project_factories, factory)(),
                                 "Explicit factory alias: " + key.lower(), vertical)
    if record is None:
        raise SensitivitySourceError("SENSITIVITY_PROJECT_UNAVAILABLE", 404)
    if (record.archived or is_protected_reference(record)
            or record.project_role not in {"working_copy", "user_project"}
            or owner != user.user_id):
        raise SensitivitySourceError("SENSITIVITY_PROJECT_UNAVAILABLE", 404)
    try:
        context = PostRunRequestContext.capture(owner_id=owner, project_id=record.project_id)
        context.require_scope(owner, record.project_id)
        pr, ws = context.project_record, context.workspace
        if (pr.archived or is_protected_reference(pr)
                or pr.project_role not in {"working_copy", "user_project"}):
            raise SensitivitySourceError("SENSITIVITY_PROJECT_UNAVAILABLE", 404)
        if scenario_id and scenario_id != ws.active_scenario_id:
            raise SensitivitySourceError("SENSITIVITY_SCENARIO_UNAVAILABLE", 404)
        vertical = resolve_vertical(pr.project_type)
        snapshot = dict(context.draft_inputs.snapshot_origin)
        if resolve_vertical(snapshot.get("project_type")) != vertical:
            raise SensitivitySourceError("SENSITIVITY_INPUT_SCOPE_INVALID")
        pis = context.draft_inputs
        if pis.coercion_errors:
            raise SensitivitySourceError("SENSITIVITY_INPUT_UNAVAILABLE")
        overrides, scope, name = None, "base", "Base"
        if ws.active_scenario_id:
            from app.persistence.scenarios_repository import get_scenario
            from app.workbook.scenario_revenue_authority import bind_scenario_tariff
            sc = get_scenario(ws.active_scenario_id, owner)
            if (sc is None or sc.archived or sc.project_id != pr.project_id
                    or sc.user_id != owner):
                raise SensitivitySourceError("SENSITIVITY_SCENARIO_UNAVAILABLE", 404)
            overrides, name = dict(sc.overrides or {}), sc.scenario_name
            admitted = {"tariff_eur_mwh", "rev_multistream_config_json",
                        "_capex_sub_line_overrides", "_opex_sub_line_overrides"}
            if set(overrides) - admitted or (sc.replay_metadata or {}).get("action") == "q4_finding_whatif_v1":
                raise SensitivitySourceError("SENSITIVITY_SCENARIO_INPUT_UNAVAILABLE")
            # A stale deleted line cannot be silently ignored by the fold.
            from app.persistence.capex_sub_lines import get_active_sub_lines_for_project as capex_lines
            from app.persistence.opex_sub_lines import get_active_sub_lines_for_project as opex_lines
            for kind, read_lines in (("capex", capex_lines), ("opex", opex_lines)):
                raw = overrides.get("_" + kind + "_sub_line_overrides", {})
                if not isinstance(raw, dict) or set(raw) - {r.sub_line_id for r in read_lines(pr.project_id)}:
                    raise SensitivitySourceError("SENSITIVITY_SCENARIO_INPUT_UNAVAILABLE")
            if not sc.is_base_case:
                scope = sc.scenario_id
                snapshot = bind_scenario_tariff(snapshot, overrides)
                pis = WorkbookService.build_input_set(snapshot)
        basis = Basis(pis, pr.project_id, overrides, scope, name, ws.active_scenario_id)
        pi, effective_pis, _ = pre_financing_inputs(basis, [])
        pi = materialise_financing(basis, pi, effective_pis)
        source = SensitivitySource(pi, "Working Copy draft — " + name + " (not a committed Run)",
                                   vertical, snapshot, context)
        source.validate_current()
        return source
    except SensitivitySourceError:
        raise
    except (ValueError, TypeError, KeyError, PermissionError, WorkbookIdentityError, PostRunContextChanged) as exc:
        raise SensitivitySourceError("SENSITIVITY_INPUT_UNAVAILABLE") from exc


def _resolve_sensitivity_project(user: Any, project: str, scenario_id: str = "") -> tuple[Any, str]:
    source = resolve_sensitivity_source(user, project, scenario_id)
    return source.inputs, source.label


def _extract_kpis(result: Any, *, base_cfads: Any = None) -> dict[str, Optional[float]]:
    """Map one execution's KPIs; CFADS is the sum of its canonical base vector.

    ``base_cfads`` is the full-period tax_and_cfads.cfads_keur vector used by
    the clean Senior base-DSCR numerator, including signed construction flows.
    Never substitute the bank-sizing/P90 vector, post-reserve cash, or absent
    legacy row fields. Missing or incomplete evidence remains unavailable.
    """
    kpis: dict[str, Optional[float]] = {}
    for attr, label, _ in KPI_DEFS:
        if attr == "_cfads_keur":
            try:
                values = tuple(base_cfads) if base_cfads is not None else ()
                kpis[attr] = (sum(values) if values and all(
                    v is not None and math.isfinite(v) for v in values
                ) else None)
            except (TypeError, ValueError, OverflowError):
                kpis[attr] = None
        else:
            v = getattr(result, "actual_min_dscr" if attr == "min_dscr" else attr, None)
            kpis[attr] = float(v) if v is not None else None
    return kpis


# One refusal sentence shared by every price-shock surface (single-driver, two-driver matrix, Goal Seek, service): explicit revenue contracts carry
# their own prices, so a shock on the scalar tariff / legacy price fields would be inactive. Modify the actual contracted revenue assumptions instead.
CONTRACTED_PRICE_SHOCK_REFUSAL = (
    "This project prices revenue from explicit revenue contracts, so a shock on the scalar tariff would change nothing and is refused rather than "
    "shown as a flat result. To test price risk, change the contracted prices in the Revenue contracts (or a scenario's contracts) and run the model.")


class InactiveShockDriver(ValueError):
    """The requested driver does not reach the project's authoritative economics (FIN-V2-D01/D02)."""

    code = "SHOCK_DRIVER_INACTIVE"

    def __init__(self, detail: str):
        super().__init__(f"{self.code}: {detail}")


class ShockNotApplied(ValueError):
    """A non-zero shock left the typed inputs unchanged; it is refused, never reported as a flat result."""

    code = "SHOCK_NOT_APPLIED"

    def __init__(self, shock_type: str, level_pct: float):
        super().__init__(f"{self.code}: {shock_type} {level_pct:+g} did not change the model inputs")


def _ppa_price_is_active(revenue: Any) -> bool:
    """A PPA tariff shock reaches revenue only when the PPA carries volume for at least one operating period."""
    return (
        float(getattr(revenue, "ppa_production_share", 0.0) or 0.0) > 0.0
        and float(getattr(revenue, "ppa_term_years", 0.0) or 0.0) > 0.0
        and getattr(revenue, "first_merchant_operating_period_index", None) != 0
    )


class EVRuntimeAuthorityUnavailable(ValueError):
    """The project handed to an EV sensitivity is not an EV-adapted project (it would be shocked on the wrong economics)."""

    code = "EV_RUNTIME_AUTHORITY_UNAVAILABLE"

    def __init__(self, detail: str):
        super().__init__(f"{self.code}: {detail}")


def assert_ev_runtime_authority(proj: Any) -> None:
    """An EV vertical must run on EV-adapted inputs: all-merchant calendar price schedule + derived electricity line.

    Guards the factory fallback of the legacy resolver (an unresolved project code falls back to the Wind reference),
    which would otherwise be sensitised under an EV label with Wind economics.
    """
    revenue = proj.revenue
    if (getattr(revenue, "first_merchant_operating_period_index", None) != 0
            or not getattr(revenue, "market_prices_by_calendar_year_eur_mwh", ())
            or not any(str(getattr(o, "name", "")) == "Electricity Procurement" for o in proj.opex)):
        raise EVRuntimeAuthorityUnavailable(
            "the resolved inputs are not the EV Charging runtime authority; no sensitivity was run")


def _scale_ev_energy_dependent_opex(proj: Any, factor: float) -> Any:
    """EV yield shock: purchased energy scales with delivered energy (``delivered / efficiency * price``).

    The EV runtime adapter derives the ``Electricity Procurement`` OPEX line from the same full-load hours that drive
    revenue, so scaling hours alone would cut revenue and leave the electricity bill untouched (an overstated downside).
    Every hour in the ramp scales by ``factor`` -> y1 amount and every step amount scale by ``factor``.  Identity with
    re-running the adapter on factor-scaled EFLH drivers is pinned in ``tests/test_finv2_d01_d02_sensitivity_surfaces.py``.
    """
    items = tuple(proj.opex)
    if not any(str(getattr(o, "name", "")) == "Electricity Procurement" for o in items):
        raise InactiveShockDriver("EV yield authority unavailable: no Electricity Procurement line to keep consistent")
    scaled = tuple(
        replace(
            o,
            y1_amount_keur=o.y1_amount_keur * factor,
            step_changes=tuple((year, amount * factor) for year, amount in o.step_changes),
        ) if str(getattr(o, "name", "")) == "Electricity Procurement" else o
        for o in items
    )
    return replace(proj, opex=scaled)


def _apply_shock(proj: Any, shock_type: str, level_pct: float,
                 dc_drivers: Any | None = None, vertical: str | None = None) -> Any:
    """Return a new ProjectInputs with the given shock applied.

    level_pct is a signed percentage: +10.0 = +10%, -5.0 = -5%.
    For interest_rate shock the level is a basis-point change applied to the effective Senior all-in rate vector
    (``app.services.rate_shock``): 5.5 % with level 100.0 -> 6.5 %.  A shock that cannot be applied faithfully raises
    a typed refusal instead of returning an unchanged model.
    """
    factor = 1.0 + level_pct / 100.0

    if shock_type == "capex":
        from finco_core.inputs._models import CapexItem
        capex = proj.capex
        updates: dict[str, Any] = {}
        for f in fields(capex):
            v = getattr(capex, f.name)
            if isinstance(v, CapexItem):
                updates[f.name] = replace(v, amount_keur=v.amount_keur * factor)
            elif isinstance(v, (int, float)):
                updates[f.name] = v * factor
        return replace(proj, capex=replace(capex, **updates))

    elif shock_type == "opex":
        new_opex = tuple(
            replace(o, y1_amount_keur=o.y1_amount_keur * factor)
            for o in proj.opex
        )
        return replace(proj, opex=new_opex)

    elif shock_type == "ppa_price":
        if getattr(proj.revenue, "multistream_config_json", ""):
            raise ValueError("REVENUE_V2_SENSITIVITY_UNSUPPORTED: " + CONTRACTED_PRICE_SHOCK_REFUSAL)
        if level_pct != 0.0 and not _ppa_price_is_active(proj.revenue):
            raise InactiveShockDriver(
                "ppa_price: this project's PPA carries no volume (zero share/term or merchant from the first period), "
                "so a tariff shock would change nothing")
        return replace(
            proj,
            revenue=replace(proj.revenue, ppa_base_tariff=proj.revenue.ppa_base_tariff * factor),
        )

    elif shock_type == "merchant_price":
        if getattr(proj.revenue, "multistream_config_json", ""):
            raise ValueError("REVENUE_V2_SENSITIVITY_UNSUPPORTED: " + CONTRACTED_PRICE_SHOCK_REFUSAL)
        if level_pct != 0.0 and getattr(proj.revenue, "market_prices_by_calendar_year_eur_mwh", ()):
            raise InactiveShockDriver(
                "merchant_price: an explicit per-calendar-year price schedule overrides the merchant curve, so scaling "
                "the curve would change nothing")
        orig = proj.revenue.market_prices_curve
        if orig:
            new_curve = tuple(v * factor for v in orig)
            return replace(
                proj,
                revenue=replace(proj.revenue, market_prices_curve=new_curve),
            )
        if level_pct != 0.0:
            raise InactiveShockDriver("merchant_price: the project has no merchant price curve to shock")
        return proj

    elif shock_type == "yield":
        shocked = replace(
            proj,
            technical=replace(
                proj.technical,
                operating_hours_p50=proj.technical.operating_hours_p50 * factor,
            ),
        )
        if resolve_vertical(vertical) == "ev_charging":
            # The EV adapter pins the P90-10y lender hours to the stabilised EFLH, so they move together.
            shocked = replace(shocked, technical=replace(
                shocked.technical,
                operating_hours_p90_10y=shocked.technical.operating_hours_p90_10y * factor))
            shocked = _scale_ev_energy_dependent_opex(shocked, factor)
        return shocked

    elif shock_type == "availability":
        new_avail = min(1.0, max(0.0, proj.technical.plant_availability * factor))
        return replace(
            proj,
            technical=replace(proj.technical, plant_availability=new_avail),
        )

    elif shock_type == "interest_rate":
        # FIN-V2-D01: level_pct is basis points on the EFFECTIVE Senior all-in rate vector.  ``financing.base_rate`` is
        # auxiliary state under an explicit all-in schedule and was never read by the engine (proved inert).
        return apply_rate_shock_bps(proj, level_pct)

    elif shock_type == "tax_rate":
        new_rate = max(0.0, min(1.0, proj.tax.corporate_rate * factor))
        return replace(
            proj,
            tax=replace(proj.tax, corporate_rate=new_rate),
        )

    elif shock_type in ("dc_service_price", "dc_occupancy", "dc_pue",
                        "dc_electricity_price", "dc_it_mw"):
        # Correction C: DC source drivers recompute through the CANONICAL
        # Data Center runtime adapter (app.data_center_authority), starting
        # from the project's ACTUAL resolved drivers (passed in explicitly
        # by the route/workspace authority).  The generic reference driver
        # set is NEVER substituted for a user-edited project.
        if dc_drivers is None:
            raise DCDriverContextRequiredError(shock_type)
        from app.data_center_authority import apply_data_center_runtime_adapter
        capacity = proj.technical.capacity_mw
        drivers = dc_drivers
        if shock_type == "dc_service_price":
            drivers = replace(drivers,
                              service_price_eur_kw_month=drivers.service_price_eur_kw_month * factor)
        elif shock_type == "dc_occupancy":
            drivers = replace(drivers,
                              occupancy_y1=min(1.0, drivers.occupancy_y1 * factor),
                              occupancy_y2=min(1.0, drivers.occupancy_y2 * factor),
                              stabilized_occupancy=min(1.0, drivers.stabilized_occupancy * factor))
        elif shock_type == "dc_pue":
            drivers = replace(drivers, pue=drivers.pue * factor)
        elif shock_type == "dc_electricity_price":
            drivers = replace(drivers,
                              electricity_price_eur_mwh=drivers.electricity_price_eur_mwh * factor)
        elif shock_type == "dc_it_mw":
            capacity = capacity * factor

        shocked = apply_data_center_runtime_adapter(
            replace(proj, technical=replace(proj.technical, capacity_mw=capacity)),
            drivers,
        )
        return shocked

    elif shock_type in ("bess_arbitrage_spread", "bess_cycles", "bess_rte",
                        "bess_ancillary_price", "bess_capacity_price"):
        bess = getattr(proj.technical, "bess", None)
        if bess is None:
            return proj
        if shock_type == "bess_arbitrage_spread":
            new_bess = replace(bess, arbitrage_spread_eur_mwh=bess.arbitrage_spread_eur_mwh * factor)
        elif shock_type == "bess_cycles":
            new_bess = replace(bess, cycles_per_year=bess.cycles_per_year * factor)
        elif shock_type == "bess_rte":
            new_rte = min(1.0, max(0.0, bess.round_trip_efficiency * factor))
            new_bess = replace(bess, round_trip_efficiency=new_rte)
        elif shock_type == "bess_ancillary_price":
            new_bess = replace(bess, ancillary_revenue_eur_mw_year=bess.ancillary_revenue_eur_mw_year * factor)
        elif shock_type == "bess_capacity_price":
            new_bess = replace(bess, capacity_revenue_eur_mw_year=bess.capacity_revenue_eur_mw_year * factor)
        else:
            new_bess = bess
        return replace(proj, technical=replace(proj.technical, bess=new_bess))

    return proj


def _run_once(proj: Any) -> dict[str, Optional[float]]:
    # PR-8 correction pass: every sensitivity shock executes through the ONE
    # shared production authority seam — clean G2C for clean-ready projects
    # (exactly one clean calculation per shocked snapshot), explicitly
    # classified legacy for blocked projects. Same engine regardless of
    # which route asked. No route-dependent authority.
    from app.services.production_waterfall_seam import execute_production_waterfall

    execution = execute_production_waterfall(proj)
    clean_run = getattr(execution, "clean_run", None)
    model = clean_run.g2c_result.financing_result.project_model_result if clean_run is not None else None
    tax_cfads = getattr(model, "tax_and_cfads", None)
    return _extract_kpis(execution.result, base_cfads=getattr(tax_cfads, "cfads_keur", None))


def run_sensitivity(
    proj: Any,
    shock_types: list[str],
    shock_levels: list[float] | None = None,
    vertical: str | None = None,
    dc_drivers: Any | None = None,
) -> dict[str, Any]:
    """Run full sensitivity matrix.

    ``vertical`` gates the driver set (Opus NEW-M-1): when provided, every
    requested shock type must be supported for that vertical or
    UnsupportedSensitivityDriverError is raised BEFORE any model execution.
    Existing callers that omit ``vertical`` keep the legacy un-gated
    behaviour.
    """
    if vertical is not None:
        for shock_type in shock_types:
            assert_driver_supported(vertical, shock_type)
            if (
                shock_type.startswith("dc_")
                and dc_drivers is None
            ):
                raise DCDriverContextRequiredError(shock_type)
        if resolve_vertical(vertical) == "ev_charging":
            assert_ev_runtime_authority(proj)
    """Run full sensitivity matrix.

    Returns:
      base_kpis: dict[str, float]   — KPIs at base
      rows: list[SensitivityRow]    — one per (shock_type, level) combination
      kpi_defs: list                — ordered KPI metadata
      shock_registry: dict          — human labels for shock types
    """
    base_kpis = _run_once(proj)

    rows = []
    for stype in shock_types:
        levels = shock_levels if shock_levels is not None else (
            DEFAULT_RATE_SHOCK_LEVELS if stype == "interest_rate" else DEFAULT_SHOCK_LEVELS)
        for level in levels:
            try:
                shocked_proj = _apply_shock(proj, stype, level, dc_drivers, vertical)
                if level != 0 and shocked_proj == proj:
                    # FIN-V2-D01: an unapplied shock is refused, never reported as a valid flat KPI row.
                    raise ShockNotApplied(stype, level)
                evidence = (rate_shock_evidence(proj, shocked_proj, level)
                            if stype == "interest_rate" and level != 0 else None)
                shocked_kpis = _run_once(shocked_proj)
                deltas = {
                    k: (shocked_kpis[k] - base_kpis[k])
                    if shocked_kpis[k] is not None and base_kpis[k] is not None
                    else None
                    for k in base_kpis
                }
            except Exception as exc:
                shocked_kpis = {k: None for k in base_kpis}
                deltas = {k: None for k in base_kpis}
                rows.append({
                    "shock_type": stype,
                    "shock_label": SHOCK_REGISTRY.get(stype, (stype, "%"))[0],
                    "level_pct": level,
                    "kpis": shocked_kpis,
                    "deltas": deltas,
                    "error": str(exc),
                    "error_code": getattr(exc, "code", None),
                    "evidence": None,
                })
                continue
            rows.append({
                "shock_type": stype,
                "shock_label": SHOCK_REGISTRY.get(stype, (stype, "%"))[0],
                "level_pct": level,
                "kpis": shocked_kpis,
                "deltas": deltas,
                "error": None,
                "error_code": None,
                "evidence": evidence,
            })

    return {
        "base_kpis": base_kpis,
        "rows": rows,
        "kpi_defs": KPI_DEFS,
        "shock_registry": SHOCK_REGISTRY,
    }


def build_tornado_data(
    sensitivity_result: dict[str, Any],
    kpi_key: str = "equity_irr",
    n_top: int = 10,
) -> list[dict[str, Any]]:
    """Compute tornado chart data: range of impact per shock type.

    For each shock type, computes max_delta and min_delta across all levels,
    then ranks by absolute range (|max - min|), largest first.

    Returns a list of dicts sorted descending by impact range.
    """
    base_kpis = sensitivity_result["base_kpis"]
    base_val = base_kpis.get(kpi_key)
    rows = sensitivity_result["rows"]

    impact_by_shock: dict[str, dict[str, Any]] = {}
    for row in rows:
        stype = row["shock_type"]
        delta = row["deltas"].get(kpi_key)
        if delta is None:
            continue
        if stype not in impact_by_shock:
            impact_by_shock[stype] = {
                "label": row["shock_label"],
                "min_delta": delta,
                "max_delta": delta,
                "min_level": row["level_pct"],
                "max_level": row["level_pct"],
            }
        else:
            if delta < impact_by_shock[stype]["min_delta"]:
                impact_by_shock[stype]["min_delta"] = delta
                impact_by_shock[stype]["min_level"] = row["level_pct"]
            if delta > impact_by_shock[stype]["max_delta"]:
                impact_by_shock[stype]["max_delta"] = delta
                impact_by_shock[stype]["max_level"] = row["level_pct"]

    tornado = []
    for stype, d in impact_by_shock.items():
        impact_range = d["max_delta"] - d["min_delta"]
        tornado.append({
            "shock_type": stype,
            "label": d["label"],
            "min_delta": d["min_delta"],
            "max_delta": d["max_delta"],
            "min_level": d["min_level"],
            "max_level": d["max_level"],
            "impact_range": impact_range,
            "abs_range": abs(impact_range),
        })

    tornado.sort(key=lambda x: x["abs_range"], reverse=True)
    return tornado[:n_top]


def export_sensitivity_csv(sensitivity_result: dict[str, Any]) -> str:
    """Export sensitivity table to CSV string."""
    kpi_defs = sensitivity_result["kpi_defs"]
    base_kpis = sensitivity_result["base_kpis"]

    buf = io.StringIO()
    writer = csv.writer(buf)
    source = sensitivity_result.get("source", {})
    source_keys = list(source)
    source_values = [source[k] for k in source_keys]

    kpi_attrs = [k for k, _, _ in kpi_defs]
    kpi_labels = [lbl for _, lbl, _ in kpi_defs]

    # Header
    writer.writerow(["Shock Type", "Level", "Unit", "Error"] + kpi_labels + [f"Δ {lbl}" for lbl in kpi_labels] + source_keys)

    # Base row
    writer.writerow(
        ["Base", "0", "", ""]
        + [_fmt_csv(base_kpis.get(k)) for k in kpi_attrs]
        + [""] * len(kpi_attrs) + source_values
    )

    for row in sensitivity_result["rows"]:
        writer.writerow(
            [row["shock_label"], row["level_pct"],
             SHOCK_REGISTRY.get(row["shock_type"], ("", "%"))[1], row["error"] or ""]
            + [_fmt_csv(row["kpis"].get(k)) for k in kpi_attrs]
            + [_fmt_csv(row["deltas"].get(k)) for k in kpi_attrs] + source_values
        )

    return buf.getvalue()


def export_sensitivity_xlsx(sensitivity_result: dict[str, Any]) -> bytes:
    """Export sensitivity table to Excel bytes."""
    import openpyxl
    from openpyxl.styles import Font, PatternFill, Alignment
    from openpyxl.utils import get_column_letter

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Sensitivity"

    kpi_defs = sensitivity_result["kpi_defs"]
    base_kpis = sensitivity_result["base_kpis"]
    kpi_attrs = [k for k, _, _ in kpi_defs]
    kpi_labels = [lbl for _, lbl, _ in kpi_defs]

    hdr_fill = PatternFill("solid", fgColor="1F3864")
    hdr_font = Font(color="FFFFFF", bold=True)
    base_fill = PatternFill("solid", fgColor="D9E1F2")
    pos_fill = PatternFill("solid", fgColor="C6EFCE")
    neg_fill = PatternFill("solid", fgColor="FFC7CE")

    # Header row
    headers = ["Shock Type", "Level", "Unit", "Error"] + kpi_labels + [f"Δ {lbl}" for lbl in kpi_labels]
    for col, h in enumerate(headers, 1):
        cell = ws.cell(row=1, column=col, value=h)
        cell.fill = hdr_fill
        cell.font = hdr_font
        cell.alignment = Alignment(horizontal="center")

    # Base row
    base_row = 2
    ws.cell(row=base_row, column=1, value="Base")
    ws.cell(row=base_row, column=2, value=0)
    for col, k in enumerate(kpi_attrs, 5):
        v = base_kpis.get(k)
        cell = ws.cell(row=base_row, column=col, value=round(v, 6) if v is not None else None)
        cell.fill = base_fill

    # Sensitivity rows
    for r_idx, row in enumerate(sensitivity_result["rows"], base_row + 1):
        ws.cell(row=r_idx, column=1, value=row["shock_label"])
        ws.cell(row=r_idx, column=2, value=row["level_pct"])
        ws.cell(row=r_idx, column=3, value=SHOCK_REGISTRY.get(row["shock_type"], ("", "%"))[1])
        ws.cell(row=r_idx, column=4, value=row["error"] or "")
        # KPI values
        for col, k in enumerate(kpi_attrs, 5):
            v = row["kpis"].get(k)
            ws.cell(row=r_idx, column=col, value=round(v, 6) if v is not None else None)
        # Delta values
        delta_start = 5 + len(kpi_attrs)
        for col, k in enumerate(kpi_attrs, delta_start):
            d = row["deltas"].get(k)
            cell = ws.cell(row=r_idx, column=col, value=round(d, 6) if d is not None else None)
            if d is not None:
                cell.fill = pos_fill if d > 0 else (neg_fill if d < 0 else PatternFill())

    # Auto-width
    for col in range(1, len(headers) + 1):
        ws.column_dimensions[get_column_letter(col)].width = 16

    if sensitivity_result.get("source"):
        provenance = wb.create_sheet("Source")
        for key, value in sensitivity_result["source"].items():
            provenance.append([key, value])

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def _fmt_csv(v: Any) -> str:
    if v is None:
        return ""
    try:
        return f"{float(v):.6f}"
    except (TypeError, ValueError):
        return str(v)
