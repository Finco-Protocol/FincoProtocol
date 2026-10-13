"""Format validated immutable F4 Run evidence; no financial execution."""
from io import BytesIO

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill


def export_investor_evidence(evidence, *, project_name, historical=False):
    wb = Workbook()
    lineage = wb.active
    lineage.title = "Run Authority"
    # User-supplied text is always literal, never an Excel formula.
    def text(cell, value):
        cell.value = str(value)
        cell.data_type = "s"
    for key, value in (("Project", project_name), ("Authority", "IMMUTABLE_RUN_HISTORY" if historical else "CANONICAL_LAST_RUN"),
                       ("Currency / units", "EUR / kEUR"), ("Policy", "RESIDUAL_ONLY"),
                       ("ROC", "Unsupported - legal policy not approved"), ("NAV", "Unsupported - realized cash multiples only"),
                       *evidence["binding"].items(), ("Agreement ID", evidence["agreement"]["agreement_id"]),
                       ("Agreement version", evidence["agreement"]["version"]),
                       ("Agreement digest", evidence["agreement_digest"]), ("Evidence digest", evidence["evidence_digest"]),
                       ("Receipt source", evidence["source"])):
        row = lineage.max_row + 1 if lineage.cell(1, 1).value is not None else 1
        text(lineage.cell(row, 1), key)
        text(lineage.cell(row, 2), value if value is not None else "Base (legacy scenario ID absent)")
    accounts = wb.create_sheet("Investor Accounts")
    accounts.append(["Investor", "Ownership", "Commitment kEUR", "Paid-in kEUR", "Uncalled kEUR",
                     "Returned capital kEUR", "Residual cash kEUR", "Unreturned paid-in kEUR", "DPI", "Realized MOIC", "XIRR", "XIRR status"])
    flows = wb.create_sheet("Dated Investor Cash")
    flows.append(["Investor", "Cash date", "Paid-in kEUR", "Returned capital kEUR", "Residual cash kEUR", "Net cash kEUR", "Unreturned paid-in kEUR"])
    for i in evidence["investors"]:
        accounts.append([None, i["ownership_fraction"], i["commitment_keur"], i["total_paid_in_keur"],
            i["uncalled_capital_keur"], i["returned_capital_keur"], i["residual_distributions_keur"],
            i["unreturned_paid_in_keur"], i["dpi"], i["realized_moic"], i["xirr"], i["xirr_status"]])
        text(accounts.cell(accounts.max_row, 1), i["investor_id"])
        accounts.cell(accounts.max_row, 2).number_format = "0.00%"
        accounts.cell(accounts.max_row, 11).number_format = "0.00%"
        for event in i["cash_events"]:
            flows.append([None, event["date"], event["contribution_keur"], event["returned_capital_keur"],
                event["residual_distribution_keur"], event["net_cashflow_keur"], event["unreturned_paid_in_keur"]])
            text(flows.cell(flows.max_row, 1), i["investor_id"])
    receipts = wb.create_sheet("Canonical Equity Receipts")
    receipts.append(["Date", "Actual paid-in kEUR", "Gross dividend kEUR", "Dividend WHT kEUR", "Net equity receipt kEUR"])
    for row in evidence["canonical_receipts"]:
        receipts.append([row["date"], row["paid_in_keur"], row["gross_dividend_keur"], row["dividend_wht_keur"], row["net_equity_receipt_keur"]])
    for sheet in wb:
        sheet.freeze_panes = "A2"
        sheet.auto_filter.ref = sheet.dimensions
        for cell in sheet[1]:
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = PatternFill("solid", fgColor="365C50")
        for column in sheet.columns:
            sheet.column_dimensions[column[0].column_letter].width = 26
        for row in sheet.iter_rows(min_row=2):
            for cell in row:
                if type(cell.value) in (int, float):
                    if cell.number_format == "General":
                        cell.number_format = '#,##0.00;(#,##0.00);"-"'
                    cell.alignment = Alignment(horizontal="right")
    output = BytesIO()
    wb.save(output)
    return output.getvalue()
