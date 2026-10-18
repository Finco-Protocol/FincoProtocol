"""Institutional reporting V2 -- run-bound lender/investor review sheets (RC-7 / Workstream E).

READ-ONLY presentation over a frozen :class:`~app.export.run_evidence.RunEvidence`.
This module is not a workbook engine:

- it never imports or calls a financial engine and never recalculates a result;
- every number is written exactly as persisted (no rounding, no scaling);
- evidence that was not persisted is shown as ``n/a`` / ``NOT_AVAILABLE`` and
  listed on the Disclosures sheet -- it is never replaced by zero or rebuilt;
- the few "checks" shown are consistency comparisons between persisted values,
  labelled as such; they do not change any reported figure.

Two entry points:
``write_reporting_v2_sheets`` appends the new sheets to the existing institutional
workbook (Last Run); ``build_run_evidence_workbook`` builds a self-contained,
deterministic workbook for ANY committed Run (including a historical Run History
row) using that Run's own evidence only.
"""
from __future__ import annotations

import io
import re
import zipfile
from datetime import datetime, timezone
from typing import Any, Sequence

from openpyxl import Workbook
from openpyxl.comments import Comment
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from app.export.reporting_common import (
    DATE_FORMAT, FLOW, KEUR_FORMAT, KEUR_FULL_FORMAT, MULTIPLE_FORMAT, NA_CELL, PCT_FORMAT,
    RATIO, STOCK, classify_metric, display_total, parse_date, period_basis, period_label,
)
from app.export.run_evidence import (
    NOT_AVAILABLE, STATE_CURRENT, STATE_HISTORICAL, STATE_STALE, RunEvidence, number,
)

REPORT_VERSION = "INSTITUTIONAL_REPORTING_V2"

_THIN = Side(style="thin", color="D5D9E2")
_BORDER = Border(left=_THIN, right=_THIN, top=_THIN, bottom=_THIN)
_TITLE_FILL = PatternFill("solid", fgColor="0F274A")
_SECTION_FILL = PatternFill("solid", fgColor="DCE6F2")
_HEADER_FILL = PatternFill("solid", fgColor="2E75B6")
_META_FILL = PatternFill("solid", fgColor="EEF3F8")
_TOTAL_FILL = PatternFill("solid", fgColor="F2F4F8")
_STATE_FILLS = {
    STATE_CURRENT: PatternFill("solid", fgColor="E7F3E8"),
    STATE_STALE: PatternFill("solid", fgColor="FFF2CC"),
    STATE_HISTORICAL: PatternFill("solid", fgColor="FDEBD0"),
}
_WARN_FILL = PatternFill("solid", fgColor="FFF2CC")
_FAIL_FILL = PatternFill("solid", fgColor="FCE4E4")
_H = Font(bold=True, size=10, name="Calibri")
_B = Font(size=10, name="Calibri")
_W = Font(bold=True, color="FFFFFF", size=10, name="Calibri")
_NA_FONT = Font(size=10, name="Calibri", italic=True, color="7F7F7F")

SHEET_REPORT_GUIDE = "Report_Guide"
SHEET_PERIODS = "Periods_Dates"
SHEET_SOURCES_USES = "Sources_Uses"
SHEET_CONSTRUCTION = "Construction_Funding"
SHEET_DEBT = "Debt_DSCR"
SHEET_RESERVES = "Reserves"
SHEET_TAX = "Tax_Bridge"
SHEET_INTEGRITY = "Run_Integrity"
SHEET_DISCLOSURES = "Disclosures"
# Only in the self-contained evidence workbook (the institutional workbook already has these)
SHEET_PNL = "P&L"
SHEET_CASHFLOW = "Cash Flow"
SHEET_BALANCE = "Balance Sheet"
SHEET_RETURNS = "Returns"

COMMON_SHEETS = (SHEET_PERIODS, SHEET_SOURCES_USES, SHEET_CONSTRUCTION, SHEET_DEBT,
                 SHEET_RESERVES, SHEET_TAX, SHEET_INTEGRITY, SHEET_DISCLOSURES)

#: documented consistency tolerance for persisted 2-dp period values (each value is rounded to
#: 0.005; a sum of two plus a third can differ by up to 0.015)
PERSISTED_2DP_TOLERANCE_KEUR = 0.02


# ═══════════════════════════ low-level sheet helpers ═══════════════════════════

class _Cursor:
    def __init__(self, sheet, ev: RunEvidence, title: str, subtitle: str = ""):
        self.sheet, self.ev, self.row = sheet, ev, 1
        sheet.sheet_view.showGridLines = False
        c = sheet.cell(row=1, column=1, value=title)
        c.fill, c.font = _TITLE_FILL, Font(bold=True, color="FFFFFF", size=14, name="Calibri")
        sheet.cell(row=2, column=1, value=_state_banner(ev)).font = _H
        for col in range(1, 9):
            sheet.cell(row=2, column=col).fill = _STATE_FILLS.get(ev.state, _META_FILL)
        sheet.cell(row=3, column=1, value=(
            f"Run {ev.snapshot_id} | ran {ev.ran_at or NOT_AVAILABLE} | engine "
            f"{ev.engine_version or NOT_AVAILABLE} | evidence digest {ev.evidence_digest[:16]}")).font = _B
        sheet.cell(row=4, column=1, value=subtitle or "Values-only persisted Run evidence; no recalculation.").font = _B
        self.row = 6

    def section(self, text: str) -> None:
        c = self.sheet.cell(row=self.row, column=1, value=text)
        c.font, c.fill = _H, _SECTION_FILL
        self.row += 1

    def blank(self, n: int = 1) -> None:
        self.row += n

    def note(self, text: str, *, warn: bool = False) -> None:
        c = self.sheet.cell(row=self.row, column=1, value=text)
        c.font = _B
        c.alignment = Alignment(wrap_text=True, vertical="top")
        self.sheet.merge_cells(start_row=self.row, start_column=1, end_row=self.row, end_column=8)
        self.sheet.row_dimensions[self.row].height = max(15, 15 * (1 + len(text) // 140))
        if warn:
            c.fill = _WARN_FILL
        self.row += 1

    def kv(self, rows: Sequence[tuple], headers=("Field", "Value", "Source / basis", "Notes")) -> None:
        for i, h in enumerate(headers, start=1):
            c = self.sheet.cell(row=self.row, column=i, value=h)
            c.font, c.fill = _W, _HEADER_FILL
        self.row += 1
        for row in rows:
            label, value, source, notes = (list(row) + ["", "", ""])[:4]
            fmt = row[4] if len(row) > 4 else None
            self.sheet.cell(row=self.row, column=1, value=label).font = _H
            vc = self.sheet.cell(row=self.row, column=2, value=_cell(value))
            vc.font = _NA_FONT if vc.value in (NA_CELL, NOT_AVAILABLE) else _B
            vc.alignment = Alignment(wrap_text=True, vertical="top", horizontal="right"
                                     if isinstance(value, (int, float)) else "left")
            if fmt and isinstance(value, (int, float)) and not isinstance(value, bool):
                vc.number_format = fmt
            self.sheet.cell(row=self.row, column=3, value=source).font = _B
            nc = self.sheet.cell(row=self.row, column=4, value=notes)
            nc.font, nc.alignment = _B, Alignment(wrap_text=True, vertical="top")
            for col in range(1, 5):
                self.sheet.cell(row=self.row, column=col).border = _BORDER
            self.row += 1
        self.row += 1

    def table(self, headers: Sequence[str], rows: Sequence[Sequence[Any]],
              formats: dict[int, str] | None = None, fills: dict[int, PatternFill] | None = None) -> None:
        for i, h in enumerate(headers, start=1):
            c = self.sheet.cell(row=self.row, column=i, value=h)
            c.font, c.fill, c.alignment = _W, _HEADER_FILL, Alignment(wrap_text=True, vertical="center")
        self.row += 1
        for r_i, row in enumerate(rows):
            for i, v in enumerate(row, start=1):
                c = self.sheet.cell(row=self.row, column=i, value=_cell(v))
                c.font = _NA_FONT if c.value in (NA_CELL, NOT_AVAILABLE) else _B
                c.border = _BORDER
                c.alignment = Alignment(wrap_text=True, vertical="top")
                if formats and i in formats and isinstance(v, (int, float)) and not isinstance(v, bool):
                    c.number_format = formats[i]
            if fills and r_i in fills:
                for i in range(1, len(headers) + 1):
                    self.sheet.cell(row=self.row, column=i).fill = fills[r_i]
            self.row += 1
        self.row += 1

    def grid(self, title: str, labels: Sequence[str], phases: Sequence[str] | None,
             metrics: Sequence[tuple[str, str, Sequence[Any], str]], *, note: str = "",
             default_kind: str | None = None) -> int:
        """Horizontal period table: metric rows x period columns.

        ``metrics`` = (label, unit, values-by-period, number-format).  Columns:
        A Metric | B Unit | C Basis | D Total (flows) | E Coverage | F.. one column per period.
        Returns the row of the header (for freeze panes).
        """
        self.section(title)
        if note:
            self.note(note)
        first_period_col = 6
        sh = self.sheet
        if phases:
            sh.cell(row=self.row, column=1, value="Phase").font = _H
            for j, ph in enumerate(phases):
                c = sh.cell(row=self.row, column=first_period_col + j, value=ph)
                c.font, c.alignment = _H, Alignment(horizontal="center")
            self.row += 1
        for i, h in enumerate(("Metric", "Unit", "Basis", "Total (flows only)", "Periods persisted")):
            c = sh.cell(row=self.row, column=1 + i, value=h)
            c.font, c.fill, c.alignment = _W, _HEADER_FILL, Alignment(wrap_text=True, vertical="center")
        for j, lab in enumerate(labels):
            c = sh.cell(row=self.row, column=first_period_col + j, value=lab)
            c.font, c.fill, c.alignment = _W, _HEADER_FILL, Alignment(wrap_text=True, horizontal="center")
            sh.column_dimensions[get_column_letter(first_period_col + j)].width = 15
        header_row = self.row
        self.row += 1
        for label, unit, values, fmt in metrics:
            kind = classify_metric(label, default_kind)
            total, cover = display_total(values, kind)
            persisted = sum(1 for v in values if number(v) is not None)
            sh.cell(row=self.row, column=1, value=label).font = _H
            sh.cell(row=self.row, column=2, value=unit).font = _B
            sh.cell(row=self.row, column=3, value=kind).font = _B
            tc = sh.cell(row=self.row, column=4, value=total)
            tc.fill = _TOTAL_FILL
            if isinstance(total, (int, float)):
                tc.number_format = fmt
                tc.font = _H
            else:
                tc.font, tc.alignment = _NA_FONT, Alignment(horizontal="right")
            cc = sh.cell(row=self.row, column=5, value=f"{persisted}/{len(values)}")
            cc.font = _B
            if cover:
                cc.fill = _WARN_FILL
                tc.comment = Comment(cover, "FINCO")
            for j, v in enumerate(values):
                n = number(v)
                c = sh.cell(row=self.row, column=first_period_col + j, value=n if n is not None else NA_CELL)
                if n is not None:
                    c.number_format, c.font = fmt, _B
                else:
                    c.font, c.alignment = _NA_FONT, Alignment(horizontal="center")
            for col in range(1, first_period_col + len(labels)):
                sh.cell(row=self.row, column=col).border = _BORDER
            self.row += 1
        self.row += 1
        return header_row


def _cell(v: Any) -> Any:
    if v is None:
        return NA_CELL
    return v


def _state_banner(ev: RunEvidence) -> str:
    src = "Last Run" if ev.source == "LAST_RUN" else "Run History"
    return f"RUN STATE: {ev.state}  |  Source: {src}  |  {ev.state_reason}"


def _setup(sheet, widths: dict[str, float], freeze: str | None = None) -> None:
    for col, w in widths.items():
        sheet.column_dimensions[col].width = w
    if freeze:
        sheet.freeze_panes = freeze
    sheet.page_setup.orientation = "landscape"
    sheet.page_setup.fitToWidth = 1
    sheet.page_setup.fitToHeight = 0
    sheet.sheet_properties.pageSetUpPr.fitToPage = True
    sheet.print_title_rows = "1:4"


# ═══════════════════════════ evidence helpers ═══════════════════════════

def _col(periods: Sequence[dict], key: str) -> list:
    return [p.get(key) for p in periods]


def _axis(periods: Sequence[dict], index_key: str, date_keys: Sequence[str]) -> tuple[list[str], list]:
    labels, dates = [], []
    for p in periods:
        d = next((parse_date(p.get(k)) for k in date_keys if parse_date(p.get(k))), None)
        labels.append(period_label(p.get(index_key), d))
        dates.append(d)
    return labels, dates


def _phases(periods: Sequence[dict]) -> list[str] | None:
    if not any("is_operation" in p for p in periods):
        return None
    return ["Operation" if p.get("is_operation") else ("Construction" if "is_operation" in p else NA_CELL)
            for p in periods]


def _metadata_sources(ev: RunEvidence):
    yield "financial_statements", ev.financial_statements
    yield "debt_schedule", ev.debt_schedule
    yield "tax_schedule", ev.tax_schedule
    yield "runtime_summary", ev.runtime_summary
    yield "identity", ev.identity


def _basis(ev: RunEvidence):
    debt = ev.schedule_periods("debt_schedule")
    dates = [parse_date(p.get("date")) for p in debt] or [parse_date(p.get("date")) for p in ev.statement_periods("pnl")]
    operating = [bool(p.get("is_operation")) for p in debt] if debt and any("is_operation" in p for p in debt) else None
    return period_basis(dates, _metadata_sources(ev), operating)


def _present(periods: Sequence[dict], key: str) -> bool:
    return any(number(p.get(key)) is not None for p in periods)


def _metrics_present(periods, specs):
    """Only rows with >=1 persisted value are drawn; the rest are listed as absent."""
    keep, absent = [], []
    for label, unit, key, fmt in specs:
        (keep if _present(periods, key) else absent).append((label, unit, key, fmt))
    return keep, [a[0] for a in absent]


# ═══════════════════════════ sheets ═══════════════════════════

def _write_report_guide(wb: Workbook, ev: RunEvidence, sheets: Sequence[str], *, index: int | None = None,
                        inputs_note: str = "") -> None:
    sh = wb.create_sheet(SHEET_REPORT_GUIDE, index) if index is not None else wb.create_sheet(SHEET_REPORT_GUIDE)
    cur = _Cursor(sh, ev, "Institutional Report Guide -- run-bound lender / investor review pack",
                  "How to read this workbook: scope, units, period basis, run context and non-claims.")
    cur.section("1. Run identity and context")
    cur.kv([
        ("Project", ev.project_name or NOT_AVAILABLE, "project record", ""),
        ("Project code / ID", f"{ev.project_code} / {ev.project_id}", "project record", ""),
        ("Scenario", ev.scenario_name or ev.scenario_id or "Base / not recorded", "Run identity",
         "Scenario the Run was produced under" + ("" if ev.scenario_name else " (name not persisted for this Run)")),
        ("Run state", ev.state, "Run freshness authority",
         ev.state_reason),
        ("Evidence source", "Last Run" if ev.source == "LAST_RUN" else "Run History (append-only)", "persistence", ""),
        ("Run History ID", ev.history_id or "not recorded (legacy Run)", "run_history", "Immutable UUID of the Run record"),
        ("Runtime snapshot ID", ev.snapshot_id, "committed Run", "Immutable Run identifier"),
        ("Run timestamp", ev.ran_at or NOT_AVAILABLE, "committed Run", "When the engine produced this Run"),
        ("Engine version", ev.engine_version or NOT_AVAILABLE, "committed Run identity",
         "As recorded at commit; not the version running now"),
        ("Input composite hash", ev.composite_hash or NOT_AVAILABLE, "committed Run identity",
         "Identity of the inputs this Run used (source input identity)"),
        ("Run origin", ev.runtime_origin or NOT_AVAILABLE, "committed Run", ""),
        ("Evidence digest (SHA-256)", ev.evidence_digest, "this report",
         "Equal digests = identical persisted evidence = reproducible report"),
    ])
    cur.section("2. Run state definitions")
    cur.kv([
        (STATE_CURRENT, "Last Run; Working Copy still equals the inputs it used", "", ""),
        (STATE_STALE, "Last Run, but the Working Copy has changed since -- numbers describe the PRIOR inputs", "", ""),
        (STATE_HISTORICAL, "Older immutable Run superseded by a later Run; frozen evidence, never edited", "", ""),
    ], headers=("State", "Meaning", "", ""))
    cur.section("3. Conventions")
    cur.kv([
        ("Currency / units", "kEUR unless a row states otherwise (ratios are x or %)", "persisted Run", ""),
        ("Period labels", "P<index> | <persisted period-end date>", "persisted dates",
         "Never a quarter/half-year name unless the Run persisted that metadata"),
        ("Stock vs flow", "Balances (Stock) are point-in-time and never totalled; Flows have a Total", "classification",
         "Ratios (DSCR etc.) are never totalled"),
        ("Missing vs zero", "n/a = not persisted for that period. 0.0 = persisted zero.", "persisted Run",
         "A Total covers only periods that were persisted (see 'Periods persisted')"),
        ("Rounding", "Period grids are persisted at 2 decimals by the Run; full-precision facts "
                     "(Sources & Uses, construction funding, senior debt integrity evidence) are shown unrounded",
         "persisted Run", "This workbook performs no rounding of its own"),
        ("Calculation", "None. No engine is executed to produce this workbook.", "", ""),
    ], headers=("Topic", "Convention", "Basis", "Notes"))
    b = _basis(ev)
    cur.section("4. Period basis")
    cur.kv([
        ("Period frequency", b.label, b.declared_source or "observed from persisted dates", b.disclosure),
        ("Observed spacing (months)", ", ".join(map(str, b.observed_months)) or NOT_AVAILABLE, "persisted period ends", ""),
        *([("Metadata warning", b.warning, "", "")] if b.warning else []),
    ])
    cur.section("5. Contents")
    cur.table(("Sheet", "Purpose"), [(s, _SHEET_PURPOSE.get(s, "")) for s in sheets])
    cur.section("6. Non-claims")
    for line in (
        "Credit covenant compliance is NOT declared: no contractual covenant, lock-up or default thresholds are "
        "persisted with this Run. DSCR values are model outputs; targets are sizing targets, not covenants.",
        "Not a lender approval, credit decision, audit certification or independent model review.",
        "Not an editable calculation model: values only, taken from the committed Run.",
        "A Working Copy edit made after this Run never changes this report; a changed Working Copy is flagged STALE.",
    ):
        cur.note(line, warn=True)
    if inputs_note:
        cur.note(inputs_note)
    _setup(sh, {"A": 34, "B": 52, "C": 34, "D": 60})


_SHEET_PURPOSE = {
    SHEET_REPORT_GUIDE: "Run identity, context state, conventions, non-claims",
    SHEET_PERIODS: "Period axis, observed/declared frequency, key dates",
    SHEET_SOURCES_USES: "Engine-audited Sources & Uses (full precision) and totals",
    SHEET_CONSTRUCTION: "Construction funding by period incl. Senior draws (where persisted)",
    SHEET_DEBT: "Senior debt balance, service and DSCR by period with consistency checks",
    SHEET_RESERVES: "DSRA / MRA balances and contributions",
    SHEET_TAX: "Tax bridge by period from the persisted tax schedule",
    SHEET_INTEGRITY: "Existing Run Integrity checks over the Run's full-precision evidence",
    SHEET_DISCLOSURES: "Methodology and unavailable-data register",
    SHEET_PNL: "Income statement by period (persisted statements)",
    SHEET_CASHFLOW: "Project-finance cash waterfall by period (persisted statements)",
    SHEET_BALANCE: "Balance sheet by period (persisted statements)",
    SHEET_RETURNS: "Project / share-capital / sponsor returns and coverage KPIs",
}


def _write_periods(wb: Workbook, ev: RunEvidence, inputs=None) -> None:
    sh = wb.create_sheet(SHEET_PERIODS)
    cur = _Cursor(sh, ev, "Periods and Dates", "Actual persisted period dates; frequency is declared only if persisted.")
    b = _basis(ev)
    cur.section("Period basis")
    cur.kv([
        ("Frequency", b.label, b.declared_source or "observed", b.disclosure),
        ("Observed spacing (months)", ", ".join(map(str, b.observed_months)) or NOT_AVAILABLE, "persisted dates", ""),
        *([("Warning", b.warning, "", "")] if b.warning else []),
    ])
    axes = [
        ("debt_schedule", ev.schedule_periods("debt_schedule"), "period", ("date",)),
        ("income statement", ev.statement_periods("pnl"), "period", ("date",)),
        ("cash waterfall", ev.statement_periods("pf_cash_waterfall"), "period_index", ("date",)),
        ("balance sheet", ev.statement_periods("balance_sheet"), "period_index", ("date",)),
        ("tax_schedule", ev.schedule_periods("tax_schedule"), "period", ("date",)),
        ("construction funding", ev.construction_periods, "period_index", ("period_end", "cashflow_date")),
    ]
    rows = []
    for name, periods, idx_key, date_keys in axes:
        _, dates = _axis(periods, idx_key, date_keys)
        real = [d for d in dates if d]
        sp = sorted({(b2.year - a.year) * 12 + b2.month - a.month for a, b2 in zip(real, real[1:])})
        rows.append((name, len(periods) or NOT_AVAILABLE, real[0].isoformat() if real else NA_CELL,
                     real[-1].isoformat() if real else NA_CELL, ", ".join(map(str, sp)) or NA_CELL))
    cur.section("Persisted period axes")
    cur.table(("Series", "Periods", "First period end", "Last period end", "Spacing (months)"), rows)
    debt = ev.schedule_periods("debt_schedule")
    op = [parse_date(p.get("date")) for p in debt if p.get("is_operation")]
    con = ev.construction_periods
    key_dates = [
        ("First funding period start (construction)", (parse_date(con[0].get("period_start")) if con else None),
         "financing_evidence.construction_funding", "Persisted period axis; not an input date"),
        ("Last funding period end (construction)", (parse_date(con[-1].get("period_end")) if con else None),
         "financing_evidence.construction_funding", ""),
        ("First operating period end", min((d for d in op if d), default=None),
         "debt_schedule (is_operation)", "First period flagged operating by the Run"),
        ("Last operating period end", max((d for d in op if d), default=None), "debt_schedule (is_operation)", ""),
    ]
    info = getattr(inputs, "info", None) if inputs is not None else None
    if info is not None:
        key_dates = [
            ("Financial Close", getattr(info, "financial_close", None), "Last Run input snapshot", "Input date"),
            ("Commercial Operation Date (COD)", getattr(info, "cod_date", None), "Last Run input snapshot", "Input date"),
            ("Construction months", getattr(info, "construction_months", None), "Last Run input snapshot", ""),
            ("Input period frequency", getattr(getattr(info, "period_frequency", None), "name", None),
             "Last Run input snapshot", "Input setting; shown for reference, not used to relabel periods"),
        ] + key_dates
    else:
        key_dates = [
            ("Financial Close", None, "not persisted in Run History", "Run History stores results, not the input snapshot"),
            ("Commercial Operation Date (COD)", None, "not persisted in Run History", ""),
        ] + key_dates
    cur.section("Key dates")
    cur.kv([(a, d.isoformat() if hasattr(d, "isoformat") else d, c, n) for a, d, c, n in key_dates])
    _setup(sh, {"A": 44, "B": 22, "C": 36, "D": 60, "E": 22})


_SU_USES = (
    ("Base project CAPEX", "base_project_capex_keur"), ("Capitalised IDC", "capitalized_idc_keur"),
    ("Commitment fee", "commitment_fee_keur"), ("Structuring fee", "structuring_fee_keur"),
    ("Other financing costs", "other_financing_costs_keur"), ("Initial DSRA funding", "initial_dsra_funding_keur"),
    ("Other uses", "other_uses_keur"), ("Development cost reimbursement", "development_cost_reimbursement_keur"),
    ("Developer fee", "developer_fee_keur"),
)
_SU_SOURCES = (
    ("Senior debt", "senior_debt_keur"), ("Junior / other main funding", "junior_or_other_keur"),
    ("Share capital and other equity", "share_capital_and_other_equity_keur"),
    ("Shareholder loan (cash)", "shareholder_loan_cash_keur"),
)
_SU_TOL = 1e-6


def _su_source(ev: RunEvidence) -> tuple[dict, str]:
    fe = ev.financing_evidence.get("sources_uses")
    if isinstance(fe, dict) and fe:
        return fe, "runtime_summary.financing_evidence.sources_uses (full precision)"
    ie = (ev.integrity_evidence.get("sources_uses") or {}).get("summary")
    if isinstance(ie, dict) and ie:
        return ie, "integrity_evidence.sources_uses.summary (full precision)"
    return {}, ""


def _write_sources_uses(wb: Workbook, ev: RunEvidence) -> None:
    sh = wb.create_sheet(SHEET_SOURCES_USES)
    cur = _Cursor(sh, ev, "Sources & Uses", "Engine-audited, full precision, kEUR. Totals are the persisted totals.")
    su, src = _su_source(ev)
    if not su:
        cur.note("NOT_AVAILABLE: this Run did not persist an engine Sources & Uses result "
                 "(legacy Run). It is not reconstructed from inputs.", warn=True)
        kp = ev.runtime_summary
        legacy = [
            ("Total project uses", kp.get("total_project_uses_keur")),
            ("Senior debt", kp.get("senior_debt_keur")),
            ("Shareholder loan (cash)", kp.get("derived_shl_cash_keur")),
            ("Sponsor equity", kp.get("sponsor_equity_sources_keur")),
        ]
        cur.section("Run summary components (partial; where persisted)")
        cur.kv([(a, number(v), "runtime_summary", "", KEUR_FULL_FORMAT) for a, v in legacy])
        _setup(sh, {"A": 40, "B": 22, "C": 40, "D": 50})
        return
    def line(label, key):
        return (label, number(su.get(key)), src, "", KEUR_FULL_FORMAT)
    cur.section("Uses")
    cur.kv([line(a, k) for a, k in _SU_USES] + [line("TOTAL USES (persisted)", "total_uses_keur")])
    cur.section("Sources")
    cur.kv([line(a, k) for a, k in _SU_SOURCES] + [line("TOTAL SOURCES (persisted)", "total_sources_keur")])
    uses = [number(su.get(k)) for _, k in _SU_USES]
    sources = [number(su.get(k)) for _, k in _SU_SOURCES]
    checks = []
    for name, parts, total_key in (("Sum of listed Uses vs total", uses, "total_uses_keur"),
                                   ("Sum of listed Sources vs total", sources, "total_sources_keur")):
        total = number(su.get(total_key))
        if total is None or any(p is None for p in parts):
            checks.append((name, "NOT_CHECKED", None, "one or more components not persisted"))
        else:
            dev = abs(sum(parts) - total)
            checks.append((name, "PASS" if dev <= _SU_TOL else "FAIL", dev, f"tolerance {_SU_TOL:g} kEUR"))
    diff, tu, ts = number(su.get("difference_keur")), number(su.get("total_uses_keur")), number(su.get("total_sources_keur"))
    checks.append(("Persisted Sources - Uses difference", ("PASS" if abs(diff) <= _SU_TOL else "FAIL")
                   if diff is not None else "NOT_CHECKED", diff, "engine-audited residual; no balancing plug is applied"))
    if tu is not None and ts is not None:
        dev = abs(ts - tu)
        checks.append(("Total Sources vs Total Uses", "PASS" if dev <= _SU_TOL else "FAIL", dev, f"tolerance {_SU_TOL:g} kEUR"))
    cur.section("Reconciliation (comparison of persisted values; nothing is adjusted)")
    fills = {i: _FAIL_FILL for i, c in enumerate(checks) if c[1] == "FAIL"}
    cur.table(("Check", "Result", "Abs. deviation (kEUR)", "Basis"), checks, {3: KEUR_FULL_FORMAT}, fills)
    gearing = [("Senior debt / gearing basis", number(ev.runtime_summary.get("actual_gearing_pct")), PCT_FORMAT),
               ("Requested gearing cap", number(ev.runtime_summary.get("gearing_cap_pct")), PCT_FORMAT)]
    cur.section("Gearing (persisted)")
    cur.kv([(a, v, "runtime_summary", "", f) for a, v, f in gearing])
    bank = ev.financing_evidence.get("bankability")
    if isinstance(bank, dict) and bank:
        cur.section("Senior sizing evidence (persisted)")
        cur.kv([(k, bank.get(k), "financing_evidence.bankability", "", KEUR_FULL_FORMAT) for k in bank])
    _setup(sh, {"A": 42, "B": 24, "C": 52, "D": 54})


_CF_SPECS = (
    ("Project cash uses", "kEUR", "project_cash_uses_keur", KEUR_FORMAT),
    ("Senior draw", "kEUR", "senior_draw_keur", KEUR_FORMAT),
    ("Junior / other main funding draw", "kEUR", "junior_or_other_main_funding_draw_keur", KEUR_FORMAT),
    ("Share capital draw", "kEUR", "share_capital_draw_keur", KEUR_FORMAT),
    ("Share premium draw", "kEUR", "share_premium_draw_keur", KEUR_FORMAT),
    ("Other committed equity draw", "kEUR", "other_committed_equity_draw_keur", KEUR_FORMAT),
    ("Additional equity draw", "kEUR", "additional_equity_draw_keur", KEUR_FORMAT),
    ("Shareholder loan cash draw", "kEUR", "shl_cash_draw_keur", KEUR_FORMAT),
    ("Total sponsor cash draw", "kEUR", "total_sponsor_cash_draw_keur", KEUR_FORMAT),
    ("Total sources", "kEUR", "total_sources_keur", KEUR_FORMAT),
    ("Sources - uses difference", "kEUR", "sources_uses_difference_keur", KEUR_FORMAT),
    ("Cumulative project cash uses", "kEUR", "cumulative_project_cash_uses_keur", KEUR_FORMAT),
    ("Cumulative senior draw", "kEUR", "cumulative_senior_draw_keur", KEUR_FORMAT),
    ("Cumulative total sources", "kEUR", "cumulative_total_sources_keur", KEUR_FORMAT),
    ("Cumulative sources - uses difference", "kEUR", "cumulative_sources_uses_difference_keur", KEUR_FORMAT),
)


def _write_construction(wb: Workbook, ev: RunEvidence) -> None:
    sh = wb.create_sheet(SHEET_CONSTRUCTION)
    cur = _Cursor(sh, ev, "Construction Funding and Senior Draws",
                  "Per-period construction funding as persisted with the Run (kEUR).")
    periods = ev.construction_periods
    if not periods:
        cur.note("NOT_AVAILABLE: this Run did not persist construction funding periods "
                 "(financing_evidence.construction_funding). Senior draws are not reconstructed.", warn=True)
        _setup(sh, {"A": 44})
        return
    labels, _ = _axis(periods, "period_index", ("period_end", "cashflow_date"))
    keep, absent = _metrics_present(periods, _CF_SPECS)
    header = cur.grid("Construction funding by period", labels, None,
                      [(a, u, _col(periods, k), f) for a, u, k, f in keep],
                      note="Cumulative rows are balances (not totalled). Flow totals cover persisted periods only.")
    cf = ev.financing_evidence.get("construction_funding") or {}
    cur.section("Funding audit (persisted)")
    cur.kv([
        ("Funding policy", cf.get("policy"), "financing_evidence.construction_funding", ""),
        ("Max period Sources-Uses difference", number(cf.get("maximum_period_difference_keur")), "persisted", "", KEUR_FULL_FORMAT),
        ("Max cumulative difference", number(cf.get("maximum_cumulative_difference_keur")), "persisted", "", KEUR_FULL_FORMAT),
        ("Total audit uses", number(cf.get("total_audit_uses_keur")), "persisted", "Construction + non-construction uses", KEUR_FULL_FORMAT),
        ("Total audit sources", number(cf.get("total_audit_sources_keur")), "persisted", "", KEUR_FULL_FORMAT),
        ("Total audit residual", number(cf.get("total_audit_residual_keur")), "persisted", "", KEUR_FULL_FORMAT),
    ])
    if absent:
        cur.note("Not persisted for this Run: " + "; ".join(absent), warn=True)
    _setup(sh, {"A": 40, "B": 8, "C": 8, "D": 20, "E": 14}, freeze=f"F{header + 1}")


_DEBT_SPECS = (
    ("Senior closing balance", "kEUR", "senior_balance_keur", KEUR_FORMAT),
    ("Senior principal", "kEUR", "senior_principal_keur", KEUR_FORMAT),
    ("Senior interest", "kEUR", "senior_interest_keur", KEUR_FORMAT),
    ("Senior debt service", "kEUR", "senior_ds_keur", KEUR_FORMAT),
    ("DSCR", "x", "dscr", MULTIPLE_FORMAT),
    ("DSRA balance", "kEUR", "dsra_balance_keur", KEUR_FORMAT),
    ("DSRA contribution", "kEUR", "dsra_contribution_keur", KEUR_FORMAT),
)
_FULL_DEBT_SPECS = (
    ("Senior opening balance", "kEUR", "opening", KEUR_FULL_FORMAT),
    ("Senior interest", "kEUR", "interest", KEUR_FULL_FORMAT),
    ("Senior principal", "kEUR", "principal", KEUR_FULL_FORMAT),
    ("Senior debt service", "kEUR", "debt_service", KEUR_FULL_FORMAT),
    ("Senior closing balance", "kEUR", "closing", KEUR_FULL_FORMAT),
    ("Annual rate", "%", "annual_rate", PCT_FORMAT),
    ("Bank CFADS", "kEUR", "bank_cfads", KEUR_FULL_FORMAT),
    ("Base CFADS", "kEUR", "base_cfads", KEUR_FULL_FORMAT),
    ("Reported DSCR", "x", "reported_dscr", MULTIPLE_FORMAT),
    ("Target DSCR (sizing)", "x", "target_dscr", MULTIPLE_FORMAT),
)


def _debt_checks(periods: Sequence[dict]) -> list[tuple]:
    out = []
    devs = []
    for p in periods:
        ds, pr, it = number(p.get("senior_ds_keur")), number(p.get("senior_principal_keur")), number(p.get("senior_interest_keur"))
        if None not in (ds, pr, it):
            devs.append(abs(ds - (pr + it)))
    out.append(("Debt service = principal + interest", len(devs), max(devs) if devs else None))
    rf = []
    for a, b in zip(periods, periods[1:]):
        c0, c1, pr = number(a.get("senior_balance_keur")), number(b.get("senior_balance_keur")), number(b.get("senior_principal_keur"))
        if None not in (c0, c1, pr):
            rf.append(abs(c0 - pr - c1))
    out.append(("Balance roll-forward: prior closing - principal = closing", len(rf), max(rf) if rf else None))
    return out


def _write_debt(wb: Workbook, ev: RunEvidence) -> None:
    sh = wb.create_sheet(SHEET_DEBT)
    cur = _Cursor(sh, ev, "Senior Debt Schedule and DSCR",
                  "Persisted per-period senior debt (2-dp display precision) and full-precision Run Integrity evidence.")
    periods = ev.schedule_periods("debt_schedule")
    header = None
    if not periods:
        cur.note("NOT_AVAILABLE: no debt schedule was persisted with this Run.", warn=True)
    else:
        labels, _ = _axis(periods, "period", ("date",))
        keep, absent = _metrics_present(periods, _DEBT_SPECS)
        header = cur.grid("Persisted debt schedule (2-dp)", labels, _phases(periods),
                          [(a, u, _col(periods, k), f) for a, u, k, f in keep],
                          note="Balances are closing balances at period end (not totalled); principal, interest and "
                               "service are per-period flows.")
        rows = []
        for name, n, dev in _debt_checks(periods):
            status = "NOT_CHECKED" if dev is None else ("PASS" if dev <= PERSISTED_2DP_TOLERANCE_KEUR else "FAIL")
            rows.append((name, status, n, dev, f"tolerance {PERSISTED_2DP_TOLERANCE_KEUR} kEUR (values persisted at 2 dp)"))
        cur.section("Consistency of persisted values (comparison only; nothing is recalculated)")
        cur.table(("Check", "Result", "Periods compared", "Max abs. deviation (kEUR)", "Basis"), rows,
                  {4: KEUR_FULL_FORMAT}, {i: _FAIL_FILL for i, r in enumerate(rows) if r[1] == "FAIL"})
        if absent:
            cur.note("Not persisted for this Run: " + "; ".join(absent), warn=True)
    s = (ev.debt_schedule.get("summary") or {}) if isinstance(ev.debt_schedule.get("summary"), dict) else {}
    km = ev.runtime_summary
    def pick(key, alt=None):
        v = number(s.get(key))
        return v if v is not None else number(km.get(alt or key))
    cur.section("Coverage summary (persisted Run outputs)")
    cur.kv([
        ("Minimum DSCR", pick("actual_min_dscr", "min_dscr"), "debt_schedule.summary / kpis", "Model output over operating periods", MULTIPLE_FORMAT),
        ("Average DSCR", pick("actual_avg_dscr", "avg_dscr"), "debt_schedule.summary / kpis", "", MULTIPLE_FORMAT),
        ("Target DSCR (sizing)", pick("target_dscr"), "debt_schedule.summary / kpis",
         "Sizing target used by the engine -- not a contractual covenant", MULTIPLE_FORMAT),
        ("Minimum LLCR", pick("min_llcr"), "debt_schedule.summary / kpis", "", MULTIPLE_FORMAT),
        ("Total senior debt service", pick("total_senior_ds_keur"), "debt_schedule.summary / kpis", "", KEUR_FORMAT),
        ("Periods in lock-up", s.get("periods_in_lockup", km.get("periods_in_lockup")), "persisted", "Engine output"),
        ("Covenant compliance", "NOT DECLARED", "no contractual covenant terms persisted",
         "Report states model outputs only; compliance requires contractual thresholds not held with the Run"),
    ])
    full = (ev.integrity_evidence.get("senior_debt") or {}).get("periods")
    full = [p for p in (full or []) if isinstance(p, dict)]
    if full:
        date_by_idx = {p.get("period"): parse_date(p.get("date")) for p in periods}
        labels = [period_label(p.get("period_index"), date_by_idx.get(p.get("period_index"))) for p in full]
        keep, _ = _metrics_present(full, _FULL_DEBT_SPECS)
        g = cur.grid("Full-precision senior debt (Run Integrity evidence)", labels, None,
                     [(a, u, _col(full, k), f) for a, u, k, f in keep],
                     note="Unrounded values recorded at commit. Opening balance and CFADS are available here only.")
        header = header or g
    else:
        cur.note("Full-precision senior debt evidence is not persisted for this Run (legacy Run).", warn=True)
    _setup(sh, {"A": 40, "B": 8, "C": 8, "D": 20, "E": 14}, freeze=f"F{header + 1}" if header else None)


def _write_reserves(wb: Workbook, ev: RunEvidence) -> None:
    sh = wb.create_sheet(SHEET_RESERVES)
    cur = _Cursor(sh, ev, "Reserves", "DSRA / MRA as persisted (kEUR).")
    debt, dist = ev.schedule_periods("debt_schedule"), ev.schedule_periods("distribution_schedule")
    bs = ev.statement_periods("balance_sheet")
    su, _ = _su_source(ev)
    cur.section("Initial reserve funding")
    cur.kv([("Initial DSRA funding (Uses)", number(su.get("initial_dsra_funding_keur")),
             "engine-audited Sources & Uses", "Funded at Financial Close/COD per the Run", KEUR_FULL_FORMAT)])
    header = None
    for title, periods, idx, specs in (
        ("DSRA (debt schedule)", debt, "period", (("DSRA balance", "kEUR", "dsra_balance_keur", KEUR_FORMAT),
                                                  ("DSRA contribution", "kEUR", "dsra_contribution_keur", KEUR_FORMAT))),
        ("DSRA / MRA (distribution schedule)", dist, "period", (
            ("DSRA balance", "kEUR", "dsra_balance_keur", KEUR_FORMAT),
            ("DSRA contribution", "kEUR", "dsra_contribution_keur", KEUR_FORMAT),
            ("MRA balance", "kEUR", "mra_balance_keur", KEUR_FORMAT),
            ("MRA contribution", "kEUR", "mra_contribution_keur", KEUR_FORMAT))),
        ("DSRA (balance sheet)", bs, "period_index", (("DSRA balance", "kEUR", "dsra_balance_keur", KEUR_FORMAT),)),
    ):
        keep, _ = _metrics_present(periods, specs)
        if not keep:
            cur.section(title)
            cur.note("NOT_AVAILABLE: no reserve values persisted in this series.")
            continue
        labels, _ = _axis(periods, idx, ("date",))
        g = cur.grid(title, labels, None, [(a, u, _col(periods, k), f) for a, u, k, f in keep])
        header = header or g
    _setup(sh, {"A": 40, "B": 8, "C": 8, "D": 20, "E": 14}, freeze=f"F{header + 1}" if header else None)


_TAX_SPECS = (
    ("Taxable profit", "kEUR", "taxable_profit_keur", KEUR_FORMAT),
    ("Tax (period)", "kEUR", "tax_keur", KEUR_FORMAT),
    ("Corporate tax cash", "kEUR", "corporate_tax_cash_keur", KEUR_FORMAT),
    ("Tax depreciation", "kEUR", "tax_depreciation_audit_keur", KEUR_FORMAT),
    ("Taxable income before losses", "kEUR", "taxable_income_before_losses_audit_keur", KEUR_FORMAT),
    ("Opening tax losses", "kEUR", "tax_loss_opening_audit_keur", KEUR_FORMAT),
    ("Tax losses used", "kEUR", "tax_loss_used_audit_keur", KEUR_FORMAT),
    ("Closing tax losses", "kEUR", "tax_loss_closing_audit_keur", KEUR_FORMAT),
    ("Taxable profit after losses", "kEUR", "taxable_profit_after_losses_audit_keur", KEUR_FORMAT),
    ("CIT accrual", "kEUR", "cit_accrual_audit_keur", KEUR_FORMAT),
    ("Cash tax current period", "kEUR", "cash_tax_current_period_audit_keur", KEUR_FORMAT),
)


def _write_tax(wb: Workbook, ev: RunEvidence) -> None:
    sh = wb.create_sheet(SHEET_TAX)
    cur = _Cursor(sh, ev, "Tax Bridge", "Per-period tax bridge from the persisted tax schedule (kEUR).")
    periods = ev.schedule_periods("tax_schedule")
    if not periods:
        cur.note("NOT_AVAILABLE: no tax schedule was persisted with this Run. Tax totals are not estimated.", warn=True)
        _setup(sh, {"A": 40})
        return
    labels, _ = _axis(periods, "period", ("date",))
    keep, absent = _metrics_present(periods, _TAX_SPECS)
    header = cur.grid("Tax bridge by period", labels, _phases(periods),
                      [(a, u, _col(periods, k), f) for a, u, k, f in keep],
                      note="Opening/closing tax-loss rows are balances (not totalled).")
    tot = number((ev.tax_schedule.get("summary") or {}).get("total_tax_keur")) \
        if isinstance(ev.tax_schedule.get("summary"), dict) else None
    cur.kv([("Total tax (persisted summary)", tot if tot is not None else number(ev.runtime_summary.get("total_tax_keur")),
             "tax_schedule.summary / kpis", "Run output", KEUR_FORMAT)])
    if absent:
        cur.note("Not persisted for this Run: " + "; ".join(absent), warn=True)
    _setup(sh, {"A": 40, "B": 8, "C": 8, "D": 20, "E": 14}, freeze=f"F{header + 1}")


def _write_integrity(wb: Workbook, ev: RunEvidence) -> None:
    from app.run_integrity import run_integrity_checks

    sh = wb.create_sheet(SHEET_INTEGRITY)
    cur = _Cursor(sh, ev, "Run Integrity (existing checks)",
                  "Result of the existing Run Integrity authority over this Run's persisted full-precision evidence.")
    report = run_integrity_checks(ev.integrity_evidence or None)
    cur.kv([("Overall", report.overall.value, "app.run_integrity", report.scope),
            ("Counts", ", ".join(f"{k}={v}" for k, v in report.counts().items()), "", ""),
            ("Evidence digest", report.evidence_digest or NOT_AVAILABLE, "persisted", "")])
    rows = [(c.title, c.status.value, c.reason_code or "", c.evaluated, c.max_abs_deviation, c.tolerance, c.detail)
            for c in report.checks]
    fills = {i: _FAIL_FILL for i, r in enumerate(rows) if r[1] == "FAIL"}
    fills.update({i: _WARN_FILL for i, r in enumerate(rows) if r[1] == "UNAVAILABLE"})
    cur.table(("Check", "Status", "Reason", "Evaluated", "Max abs. deviation", "Tolerance", "Detail"), rows,
              {5: "0.000000000", 6: "0.000000000"}, fills)
    cur.note("Integrity checks test internal consistency of the Run. They are not a covenant test, "
             "credit decision or external model review.", warn=True)
    _setup(sh, {"A": 48, "B": 16, "C": 34, "D": 12, "E": 18, "F": 14, "G": 70})


def _expected_items(ev: RunEvidence) -> list[tuple[str, bool, str]]:
    pnl, cf, bs = (ev.statement_periods(k) for k in ("pnl", "pf_cash_waterfall", "balance_sheet"))
    kp = ev.runtime_summary
    return [
        ("Income statement periods", bool(pnl), "Section omitted when absent"),
        ("Cash flow statement periods", bool(cf), "Section omitted when absent"),
        ("Balance sheet periods", bool(bs), "Section omitted when absent"),
        ("Balance sheet totals (total assets / liabilities + equity)",
         any(number(p.get("total_assets_keur")) is not None for p in bs),
         "Runtime does not publish aggregate totals; none are constructed. The runtime balance check is shown instead"),
        ("Tax bridge periods (financial statements)", bool(ev.statement_periods("tax_bridge")),
         "Tax is reported from the persisted tax schedule instead"),
        ("Tax schedule periods", bool(ev.schedule_periods("tax_schedule")), "Tax_Bridge shows NOT_AVAILABLE when absent"),
        ("Debt schedule periods", bool(ev.schedule_periods("debt_schedule")), "Debt_DSCR shows NOT_AVAILABLE when absent"),
        ("Full-precision senior debt evidence", bool((ev.integrity_evidence.get("senior_debt") or {}).get("periods")),
         "Opening balances / CFADS shown only when persisted"),
        ("Engine Sources & Uses", bool(_su_source(ev)[0]), "Legacy Runs show partial KPI components only"),
        ("Construction funding periods (Senior draws)", bool(ev.construction_periods), "Section shows NOT_AVAILABLE when absent"),
        ("Run Integrity evidence", bool(ev.integrity_evidence), "Checks report UNAVAILABLE; none are assumed to pass"),
        ("Engine version", bool(ev.engine_version), "Shown as NOT_AVAILABLE for legacy Runs"),
        ("Explicit period frequency metadata", _basis(ev).declared_months is not None,
         "Not required: the report uses actual persisted dates; no frequency is assumed or relabelled"),
        ("Project IRR", number(kp.get("project_irr")) is not None, "n/a when not published (never 0)"),
        ("Share-capital IRR", number(kp.get("equity_irr")) is not None, "n/a when not published (never 0)"),
        ("Total Sponsor XIRR", number(kp.get("total_sponsor_xirr")) is not None, "n/a when not published (never 0)"),
        ("Contractual covenant thresholds", False, "Not persisted for any Run: covenant compliance is never declared"),
    ]


def _write_disclosures(wb: Workbook, ev: RunEvidence) -> None:
    sh = wb.create_sheet(SHEET_DISCLOSURES)
    cur = _Cursor(sh, ev, "Methodology and Unavailable-Data Disclosures")
    cur.section("Methodology")
    for line in (
        "Source of truth: the immutable committed Run (Last Run or one append-only Run History row). Nothing else is read.",
        "No engine is executed, no figure is recalculated, rebuilt, interpolated or plugged during export.",
        "Period grids are the Run's persisted tables (2-dp). Full-precision facts are shown unrounded where persisted.",
        "Totals appear only for Flow rows and cover persisted periods; balances and ratios are never summed.",
        "Reconciliation rows compare persisted values with a stated tolerance; they never alter a reported number.",
        "The Working Copy is not read for values. Later Working Copy edits do not change this report.",
        "Assumption provenance (inputs) is shown only where the Run's own input snapshot is available (Last Run); "
        "Run History persists results and identity, not the input snapshot.",
        "Run Integrity results are those of the existing authority (app.run_integrity); no new credit analysis is performed here.",
    ):
        cur.note("- " + line)
    cur.blank()
    cur.section("Evidence register")
    items = _expected_items(ev)
    cur.table(("Evidence item", "Status", "Treatment when absent"),
              [(a, "PERSISTED" if ok else "NOT PERSISTED", c) for a, ok, c in items],
              fills={i: _WARN_FILL for i, (_, ok, _c) in enumerate(items) if not ok})
    _setup(sh, {"A": 56, "B": 20, "C": 90})


# ── statements / returns (self-contained evidence workbook only) ──

_PNL_SPECS = (
    ("Revenue", "kEUR", "revenues_keur"), ("OPEX", "kEUR", "operating_expenses_keur"),
    ("Depreciation", "kEUR", "depreciation_keur"), ("EBIT", "kEUR", "ebit_keur"),
    ("Senior interest", "kEUR", "senior_interest_expense_keur"), ("Junior interest", "kEUR", "junior_interest_expense_keur"),
    ("SHL interest", "kEUR", "shl_interest_expense_keur"), ("EBT", "kEUR", "earnings_before_tax_keur"),
    ("CIT accrual", "kEUR", "cit_accrual_keur"), ("Net income", "kEUR", "net_income_keur"),
    ("Net dividends", "kEUR", "net_dividends_keur"), ("Retained earnings (closing)", "kEUR", "retained_earnings_keur"),
)
_CASH_SPECS = (
    ("Revenue cash", "kEUR", "revenue_cash_keur"), ("OPEX cash", "kEUR", "opex_cash_keur"),
    ("EBITDA cash", "kEUR", "ebitda_cash_keur"), ("Cash tax", "kEUR", "cash_tax_keur"),
    ("FCF banks", "kEUR", "fcf_banks_keur"), ("Senior debt service", "kEUR", "senior_total_ds_keur"),
    ("Junior cash interest", "kEUR", "junior_cash_interest_keur"), ("Junior principal", "kEUR", "junior_principal_keur"),
    ("Junior debt service", "kEUR", "junior_debt_service_keur"), ("Post-junior cash", "kEUR", "post_junior_cash_keur"),
    ("DSRA funding", "kEUR", "dsra_funding_keur"), ("DSRA release", "kEUR", "dsra_release_keur"),
    ("FCF junior", "kEUR", "fcf_junior_keur"), ("FCF for distribution", "kEUR", "fcf_for_distribution_keur"),
    ("Net dividends", "kEUR", "net_dividends_keur"),
)
_BS_SPECS = (
    ("Net fixed assets", "kEUR", "net_fixed_assets_keur"), ("DSRA balance", "kEUR", "dsra_balance_keur"),
    ("Cash", "kEUR", "cash_keur"), ("Total assets", "kEUR", "total_assets_keur"),
    ("Share capital", "kEUR", "share_capital_keur"), ("Retained earnings", "kEUR", "retained_earnings_keur"),
    ("SHL balance", "kEUR", "shl_balance_keur"), ("Senior debt balance", "kEUR", "senior_balance_keur"),
    ("Junior debt balance", "kEUR", "junior_balance_keur"), ("Total liabilities + equity", "kEUR", "total_liabilities_equity_keur"),
    ("Balance check (runtime residual)", "kEUR", "balance_check_keur"),
)


def _write_statement(wb, ev, sheet_name, title, key, specs, note) -> None:
    sh = wb.create_sheet(sheet_name)
    cur = _Cursor(sh, ev, title, "Persisted statements of the Run (kEUR); no totals or plugs are constructed.")
    periods = ev.statement_periods(key)
    if not periods:
        cur.note(f"NOT_AVAILABLE: the Run did not persist {title.lower()} periods.", warn=True)
        _setup(sh, {"A": 40})
        return
    idx = "period" if key == "pnl" else "period_index"
    labels, _ = _axis(periods, idx, ("date",))
    keep, absent = _metrics_present(periods, [(a, u, k, KEUR_FORMAT) for a, u, k in specs])
    header = cur.grid(title, labels, None, [(a, u, _col(periods, k), f) for a, u, k, f in keep], note=note,
                      default_kind=STOCK if key == "balance_sheet" else None)
    if absent:
        cur.note("Not persisted for this Run (shown as unavailable, not zero): " + "; ".join(absent), warn=True)
    _setup(sh, {"A": 40, "B": 8, "C": 8, "D": 20, "E": 14}, freeze=f"F{header + 1}")


def _write_returns(wb: Workbook, ev: RunEvidence) -> None:
    sh = wb.create_sheet(SHEET_RETURNS)
    cur = _Cursor(sh, ev, "Returns and Coverage KPIs", "Persisted Run outputs; typed status shown where the engine published one.")
    k = ev.runtime_summary
    cur.section("Returns")
    cur.kv([
        ("Project IRR", number(k.get("project_irr")), "kpis.project_irr", "Unlevered project XIRR", PCT_FORMAT),
        ("Share-capital IRR (equity only)", number(k.get("equity_irr")), "kpis.equity_irr",
         "Status: " + str(k.get("share_capital_irr_status") or NOT_AVAILABLE) + ". Excludes shareholder-loan flows", PCT_FORMAT),
        ("Total Sponsor XIRR", number(k.get("total_sponsor_xirr")), "kpis.total_sponsor_xirr",
         "Status: " + str(k.get("total_sponsor_xirr_status") or NOT_AVAILABLE) + ". Share capital + shareholder loan", PCT_FORMAT),
        ("Project NPV", number(k.get("project_npv_keur")), "kpis.project_npv_keur", "kEUR", KEUR_FORMAT),
        ("Equity NPV", number(k.get("equity_npv_keur")), "kpis.equity_npv_keur", "kEUR", KEUR_FORMAT),
    ])
    cur.section("Totals and coverage")
    cur.kv([(a, number(k.get(key)), f"kpis.{key}", unit, fmt) for a, key, unit, fmt in (
        ("Total revenue", "total_revenue_keur", "kEUR", KEUR_FORMAT), ("Total EBITDA", "total_ebitda_keur", "kEUR", KEUR_FORMAT),
        ("Total OPEX", "total_opex_keur", "kEUR", KEUR_FORMAT), ("Total senior debt service", "total_senior_ds_keur", "kEUR", KEUR_FORMAT),
        ("Total SHL service", "total_shl_service_keur", "kEUR", KEUR_FORMAT), ("Total tax", "total_tax_keur", "kEUR", KEUR_FORMAT),
        ("Total distributions", "total_distributions_keur", "kEUR", KEUR_FORMAT),
        ("Minimum DSCR", "min_dscr", "x", MULTIPLE_FORMAT), ("Average DSCR", "avg_dscr", "x", MULTIPLE_FORMAT),
        ("Minimum LLCR", "min_llcr", "x", MULTIPLE_FORMAT))])
    _setup(sh, {"A": 38, "B": 22, "C": 30, "D": 70})


# ═══════════════════════════ public API ═══════════════════════════

def write_reporting_v2_sheets(wb: Workbook, ev: RunEvidence, *, inputs=None) -> list[str]:
    """Append the V2 sheets to the existing institutional workbook; returns the new sheet names.

    The Report_Guide is placed after Export_Metadata / Workbook_Index so it is the first thing a
    reviewer reads; statements and returns already exist in that workbook and are not duplicated.
    """
    names = [SHEET_REPORT_GUIDE, *COMMON_SHEETS]
    _write_report_guide(wb, ev, names, index=2,
                        inputs_note="Existing sheets (Cover ... Validation Status) follow; they combine this Run's "
                                    "persisted results with the Last Run input snapshot.")
    _write_periods(wb, ev, inputs)
    _write_sources_uses(wb, ev)
    _write_construction(wb, ev)
    _write_debt(wb, ev)
    _write_reserves(wb, ev)
    _write_tax(wb, ev)
    _write_integrity(wb, ev)
    _write_disclosures(wb, ev)
    return names


def build_run_evidence_workbook(ev: RunEvidence) -> bytes:
    """Self-contained, deterministic workbook for ONE committed Run (used for historical Runs).

    Reads only ``ev``.  No clock, no Working Copy, no engine: equal evidence -> equal bytes.
    """
    wb = Workbook()
    wb.remove(wb.active)
    names = [SHEET_REPORT_GUIDE, SHEET_PERIODS, SHEET_SOURCES_USES, SHEET_CONSTRUCTION, SHEET_DEBT, SHEET_RESERVES,
             SHEET_PNL, SHEET_CASHFLOW, SHEET_BALANCE, SHEET_TAX, SHEET_RETURNS, SHEET_INTEGRITY, SHEET_DISCLOSURES]
    _write_report_guide(wb, ev, names, inputs_note=(
        "This workbook was produced from the Run's own persisted evidence only. Input-assumption sheets "
        "(Inputs, CAPEX, OPEX, Revenue assumptions) are intentionally not included: Run History does not "
        "persist the input snapshot, and the current Working Copy must not be substituted for it."
        if ev.source != "LAST_RUN" else ""))
    _write_periods(wb, ev)
    _write_sources_uses(wb, ev)
    _write_construction(wb, ev)
    _write_debt(wb, ev)
    _write_reserves(wb, ev)
    _write_statement(wb, ev, SHEET_PNL, "Income Statement", "pnl", _PNL_SPECS,
                     "Flows have a Total over persisted periods; retained earnings is a balance.")
    _write_statement(wb, ev, SHEET_CASHFLOW, "Cash Flow Statement", "pf_cash_waterfall", _CASH_SPECS,
                     "Project-finance cash waterfall as persisted by the Run.")
    _write_statement(wb, ev, SHEET_BALANCE, "Balance Sheet", "balance_sheet", _BS_SPECS,
                     "Balances are point-in-time and never totalled. The balance check is the runtime's own residual; "
                     "totals the runtime did not publish are not constructed.")
    _write_tax(wb, ev)
    _write_returns(wb, ev)
    _write_integrity(wb, ev)
    _write_disclosures(wb, ev)
    for sh in wb.worksheets:
        sh.sheet_properties.tabColor = {"Report_Guide": "1F4E79", "Disclosures": "7F7F7F"}.get(sh.title, "2E75B6")
    return finalize_workbook_bytes(wb, ev.ran_at)


def finalize_workbook_bytes(wb: Workbook, ran_at: str) -> bytes:
    """Serialize with fixed document properties and zip timestamps (byte-reproducible)."""
    stamp = _stamp(ran_at)
    wb.properties.creator = "FINCO Model"
    wb.properties.lastModifiedBy = "FINCO Model"
    wb.properties.created = stamp
    wb.properties.modified = stamp
    buf = io.BytesIO()
    wb.save(buf)
    out = io.BytesIO()
    with zipfile.ZipFile(io.BytesIO(buf.getvalue())) as src, \
            zipfile.ZipFile(out, "w", zipfile.ZIP_DEFLATED) as dst:
        for info in sorted(src.infolist(), key=lambda i: i.filename):
            zi = zipfile.ZipInfo(info.filename, date_time=(1980, 1, 1, 0, 0, 0))
            zi.compress_type, zi.external_attr = zipfile.ZIP_DEFLATED, info.external_attr
            data = src.read(info.filename)
            if info.filename == "docProps/core.xml":
                # openpyxl stamps 'modified' with the wall clock at save; pin it to the Run time
                data = re.sub(rb"(<dcterms:modified[^>]*>)[^<]*(</dcterms:modified>)",
                              rb"\g<1>" + stamp.strftime("%Y-%m-%dT%H:%M:%SZ").encode() + rb"\g<2>", data)
            dst.writestr(zi, data)
    return out.getvalue()


def _stamp(ran_at: str) -> datetime:
    try:
        d = datetime.fromisoformat(str(ran_at).replace("Z", "+00:00"))
        d = d.astimezone(timezone.utc) if d.tzinfo else d
        return d.replace(tzinfo=None, microsecond=0)
    except ValueError:
        return datetime(1980, 1, 1)
