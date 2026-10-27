"""Institutional reporting V3 -- lender & investor evidence package (RC reporting / Workstream A).

Extends :mod:`app.export.institutional_reporting_v2`; it does not replace it and it is not a second workbook
engine.  READ-ONLY presentation over a frozen :class:`~app.export.run_evidence.RunEvidence`:

- no financial engine is imported or run and no result is created, corrected, inferred or re-solved;
- every headline number is written exactly as persisted; a handful of clearly labelled presentation sums or
  ratios over persisted values are marked ``PRESENTATION_ARITHMETIC`` and never replace the persisted figure;
- evidence that was not persisted is shown as ``n/a`` / ``NOT_AVAILABLE`` (never zero) and listed on
  ``Evidence_Limits``;
- contractual assessments (covenant compliance, thresholds, default events) and any investment recommendation are
  explicitly NOT provided.

V3 adds the executive/lender front section (``Report_Contents``, ``Executive_Overview``, ``Lender_Overview``),
the metric-level ``Evidence_Limits`` register and print/navigation polish.  The V2 schedules, statements,
Run Integrity and disclosures are reused unchanged.
"""
from __future__ import annotations

from typing import Any, Sequence

from openpyxl import Workbook
from openpyxl.styles import Alignment

from app.export import institutional_reporting_v2 as V2
from app.export.reporting_common import KEUR_FORMAT, KEUR_FULL_FORMAT, MULTIPLE_FORMAT, NA_CELL, PCT_FORMAT, parse_date
from app.export.run_evidence import NOT_AVAILABLE, RunEvidence, number

REPORT_VERSION_V3 = "INSTITUTIONAL_REPORTING_V3"

SHEET_CONTENTS = "Report_Contents"
SHEET_EXECUTIVE = "Executive_Overview"
SHEET_LENDER = "Lender_Overview"
SHEET_LIMITS = "Evidence_Limits"
V3_FRONT_SHEETS = (SHEET_CONTENTS, SHEET_EXECUTIVE, SHEET_LENDER)

# Evidence classes (A4).  Exactly one per reported metric.
PERSISTED = "PERSISTED"                                   # persisted canonical value, written as stored
ARITHMETIC = "PRESENTATION_ARITHMETIC"                    # sum/ratio of persisted values, shown beside the stored figure
MISSING = "NOT_AVAILABLE"                                 # not persisted for this Run: never zero, never rebuilt
CHECK = "OBSERVATIONAL_CHECK"                             # comparison of persisted values; changes no figure
CONTRACTUAL = "CONTRACTUAL_ASSESSMENT_NOT_AVAILABLE"      # needs contractual terms not held with the Run
EVIDENCE_CLASSES = (PERSISTED, ARITHMETIC, MISSING, CHECK, CONTRACTUAL)

_SUMMARY_TOL_PER_PERIOD = 0.005   # each persisted period value is rounded to 2 dp
_SUMMARY_TOL_BASE = 0.02


class _Ledger:
    """Collects (metric, sheet, class, source) so the register can never drift from what the sheets show."""

    def __init__(self) -> None:
        self.entries: list[tuple[str, str, str, str]] = []

    def add(self, sheet: str, metric: str, cls: str, source: str) -> None:
        self.entries.append((metric, sheet, cls, source))


def _fmt(value: Any, fmt: str | None):
    return fmt if number(value) is not None else None


def _rows(ledger: _Ledger, sheet: str, items: Sequence[tuple]) -> list[tuple]:
    """items: (label, value, source, notes, fmt[, cls]).  Missing values are classed NOT_AVAILABLE."""
    out = []
    for it in items:
        label, value, source, notes, fmt = it[:5]
        cls = it[5] if len(it) > 5 else PERSISTED
        if cls in (PERSISTED, ARITHMETIC) and value is None:
            cls = MISSING
        ledger.add(sheet, label, cls, source)
        out.append((label, value if value is not None else None, f"{cls} | {source}", notes, fmt))
    return out


def _state_text(ev: RunEvidence) -> str:
    return {"CURRENT": "CURRENT -- Last Run and the Working Copy still agree",
            "STALE": "STALE -- Last Run; the Working Copy has changed since it ran, numbers describe the PRIOR inputs",
            "HISTORICAL": "HISTORICAL -- superseded immutable Run record"}.get(ev.state, ev.state)


def _integrity(ev: RunEvidence):
    from app.run_integrity import run_integrity_checks
    return run_integrity_checks(ev.integrity_evidence or None)


def _terminal(ev: RunEvidence) -> dict:
    sp = ev.sponsor_schedule.get("summary") if isinstance(ev.sponsor_schedule, dict) else None
    t = (sp or {}).get("terminal_financial_state") if isinstance(sp, dict) else None
    return t if isinstance(t, dict) else {}


def _sponsor_summary(ev: RunEvidence) -> dict:
    sp = ev.sponsor_schedule.get("summary") if isinstance(ev.sponsor_schedule, dict) else None
    return sp if isinstance(sp, dict) else {}


# ═══════════════════════════ sheets ═══════════════════════════

def _write_executive(wb: Workbook, ev: RunEvidence, ledger: _Ledger) -> None:
    sh = wb.create_sheet(SHEET_EXECUTIVE)
    cur = V2._Cursor(sh, ev, "Executive Overview -- project and investment",
                     "Headline outputs of the committed Run. kEUR unless stated; percentages are annual returns; no recalculation.")
    rs, sp = ev.runtime_summary, _sponsor_summary(ev)
    b = V2._basis(ev)
    report = _integrity(ev)
    S = SHEET_EXECUTIVE
    cur.section("1. Project and Run identity")
    cur.kv(_rows(ledger, S, [
        ("Project", ev.project_name or None, "project record", "", None),
        ("Project code", ev.project_code or None, "project record", "", None),
        ("Scenario", ev.scenario_name or ev.scenario_id or None, "Run identity",
         "Scenario of the Run" + ("" if ev.scenario_name else " (name not persisted)"), None),
        ("Reporting basis", "Committed Run -- " + ("Last Run" if ev.source == "LAST_RUN" else "Run History record"),
         "Run provenance", "Values are read from the immutable Run; the Working Copy is never read for numbers", None),
        ("Run state", _state_text(ev), "Run freshness authority", ev.state_reason, None),
        ("Run timestamp", ev.ran_at or None, "committed Run", "", None),
        ("Engine version", ev.engine_version or None, "committed Run identity", "As recorded at commit", None),
        ("Runtime snapshot ID", ev.snapshot_id or None, "committed Run", "Immutable Run identifier", None),
        ("Run History ID", ev.history_id or None, "run_history", "Absent for a legacy Run", None),
        ("Input composite hash", ev.composite_hash or None, "committed Run identity", "", None),
        ("Evidence digest (SHA-256)", ev.evidence_digest, "this report", "Equal digests = identical persisted evidence", None),
        ("Period basis", b.label, b.declared_source or "observed from persisted dates", b.disclosure, None),
        ("Run Integrity", report.overall.value, "app.run_integrity", report.scope, None),
        ("Data availability", f"{sum(1 for _, ok, _c in V2._expected_items(ev) if ok)} of {len(V2._expected_items(ev))} "
                              "evidence items persisted", "Evidence register (Disclosures)",
         "See Evidence_Limits for each metric", None),
    ]))
    cur.section("2. Investment returns")
    cur.kv(_rows(ledger, S, [
        ("Project IRR", number(rs.get("project_irr")), "runtime_summary.project_irr", "Unlevered, pre-financing", PCT_FORMAT),
        ("Total Sponsor XIRR", number(rs.get("total_sponsor_xirr")), "runtime_summary.total_sponsor_xirr",
         f"Status: {rs.get('total_sponsor_xirr_status') or NOT_AVAILABLE}", PCT_FORMAT),
        ("Share-capital (pure equity) IRR", number(rs.get("equity_irr")), "runtime_summary.equity_irr",
         f"Secondary metric; status: {rs.get('share_capital_irr_status') or NOT_AVAILABLE}; denominator excludes SHL", PCT_FORMAT),
        ("Pure-equity multiple (MOIC)", number(sp.get("pure_equity_moic")), "sponsor_schedule.summary.pure_equity_moic",
         "Multiple of invested capital as persisted by the sponsor-returns authority", MULTIPLE_FORMAT),
        ("Project NPV", number(rs.get("project_npv_keur")), "runtime_summary.project_npv_keur",
         "n/a when the Run published no NPV (never shown as 0)", KEUR_FORMAT),
        ("Equity NPV", number(rs.get("equity_npv_keur")), "runtime_summary.equity_npv_keur", "", KEUR_FORMAT),
    ]))
    cur.section("3. Capital and financing")
    su, su_src = V2._su_source(ev)
    cur.kv(_rows(ledger, S, [
        ("Total CAPEX", number(rs.get("total_capex_keur")), "runtime_summary.total_capex_keur", "Base project CAPEX", KEUR_FORMAT),
        ("Total project uses", number(rs.get("total_project_uses_keur")), "runtime_summary.total_project_uses_keur",
         "CAPEX plus capitalised financing costs and reserve funding", KEUR_FORMAT),
        ("Senior debt (committed)", number(rs.get("senior_debt_keur")), "runtime_summary.senior_debt_keur", "", KEUR_FORMAT),
        ("Shareholder loan (cash)", number(rs.get("derived_shl_cash_keur")), "runtime_summary.derived_shl_cash_keur", "", KEUR_FORMAT),
        ("Sponsor equity (sources)", number(rs.get("sponsor_equity_sources_keur")),
         "runtime_summary.sponsor_equity_sources_keur", "", KEUR_FORMAT),
        ("Achieved gearing", number(rs.get("actual_gearing_pct")), "runtime_summary.actual_gearing_pct",
         "Senior debt over the gearing basis", PCT_FORMAT),
        ("Total Sources (engine-audited)", number(su.get("total_sources_keur")), su_src or "engine Sources & Uses",
         "Full precision; reconciled on Sources_Uses", KEUR_FULL_FORMAT),
        ("Total Uses (engine-audited)", number(su.get("total_uses_keur")), su_src or "engine Sources & Uses", "", KEUR_FULL_FORMAT),
    ]))
    cur.section("4. Operating performance (totals over the persisted horizon)")
    revenue, ebitda = number(rs.get("total_revenue_keur")), number(rs.get("total_ebitda_keur"))
    margin = (ebitda / revenue) if revenue not in (None, 0.0) and ebitda is not None else None
    cur.kv(_rows(ledger, S, [
        ("Revenue", revenue, "runtime_summary.total_revenue_keur", "Flow total", KEUR_FORMAT),
        ("Operating expenses", number(rs.get("total_opex_keur")), "runtime_summary.total_opex_keur", "Flow total", KEUR_FORMAT),
        ("EBITDA", ebitda, "runtime_summary.total_ebitda_keur", "Flow total", KEUR_FORMAT),
        ("EBITDA margin (total EBITDA / total revenue)", margin, "runtime_summary totals",
         "Shown beside the persisted totals; computed here only for presentation", PCT_FORMAT, ARITHMETIC),
        ("Corporate tax", number(rs.get("total_tax_keur")), "runtime_summary.total_tax_keur", "Flow total", KEUR_FORMAT),
        ("Senior debt service", number(rs.get("total_senior_ds_keur")), "runtime_summary.total_senior_ds_keur", "Flow total", KEUR_FORMAT),
        ("Shareholder-loan service", number(rs.get("total_shl_service_keur")), "runtime_summary.total_shl_service_keur",
         "Flow total", KEUR_FORMAT),
        ("Distributions", number(rs.get("total_distributions_keur")), "runtime_summary.total_distributions_keur",
         "Flow total, full precision", KEUR_FORMAT),
    ]))
    ocf = [number(p.get("fcf_for_distribution_keur")) for p in ev.statement_periods("pf_cash_waterfall")]
    ocf_total = sum(v for v in ocf if v is not None) if any(v is not None for v in ocf) else None
    cur.section("5. Operating cash flow (persisted period grid, 2 dp)")
    cur.kv(_rows(ledger, S, [
        ("Cash flow available for distribution (sum of periods)", ocf_total,
         "financial_statements.pf_cash_waterfall.fcf_for_distribution_keur",
         f"Sum over {sum(1 for v in ocf if v is not None)} of {len(ocf)} persisted periods; each period value is rounded to 2 dp",
         KEUR_FORMAT, ARITHMETIC),
    ]))
    cur.section("6. Reading notes")
    for line in (
        "Units: kEUR for amounts; IRR/gearing as annual percentages; DSCR and MOIC as multiples (x).",
        "Totals are flow sums over the persisted horizon; balances are never summed (see Lender_Overview and the schedules).",
        "No investment recommendation, credit approval or covenant opinion is expressed anywhere in this report.",
    ):
        cur.note("- " + line)
    V2._setup(sh, {"A": 44, "B": 26, "C": 58, "D": 62})


def _debt_summary(periods: Sequence[dict]) -> dict:
    prin = [number(p.get("senior_principal_keur")) for p in periods]
    intr = [number(p.get("senior_interest_keur")) for p in periods]
    ds = [number(p.get("senior_ds_keur")) for p in periods]
    bal = [number(p.get("senior_balance_keur")) for p in periods]
    dsra = [number(p.get("dsra_balance_keur")) for p in periods]
    dsra_c = [number(p.get("dsra_contribution_keur")) for p in periods]

    def total(vs):
        got = [v for v in vs if v is not None]
        return (sum(got), len(got)) if got else (None, 0)

    repay = [parse_date(p.get("date")) for p, v in zip(periods, prin) if v is not None and v > 0]
    peak = max((v for v in bal if v is not None), default=None)
    last_bal = next((v for v in reversed(bal) if v is not None), None)
    last_dsra = next((v for v in reversed(dsra) if v is not None), None)
    return {"principal": total(prin), "interest": total(intr), "service": total(ds), "dsra_contribution": total(dsra_c),
            "peak_balance": peak, "closing_balance": last_bal,
            "peak_dsra": max((v for v in dsra if v is not None), default=None), "closing_dsra": last_dsra,
            "first_repayment": min((d for d in repay if d), default=None), "last_repayment": max((d for d in repay if d), default=None),
            "n": len(periods)}


def _write_senior_debt_evidence(cur, ledger: _Ledger, rs: dict) -> None:
    """FIN-V2-C1-FIX-1a: maturity vs payoff, binding constraint and DSCR bases, exactly as recorded with the Run.

    Runs committed before this evidence existed show NOT_RECORDED; nothing is rebuilt or inferred.
    """
    from app.run_integrity.senior_debt_evidence import STATUS_RECORDED, project_senior_debt_evidence
    S = SHEET_LENDER
    e = project_senior_debt_evidence(rs)
    cur.section("1b. Senior debt maturity, payoff and sizing-limit evidence (read-only; recorded with the Run)")
    if e.get("status") != STATUS_RECORDED:
        cur.note(f"{e.get('status')}: {e.get('reason')}", warn=True)
    rp, sz = e.get("repayment") or {}, e.get("sizing") or {}
    dscr = e.get("dscr") or {}
    ds, dr = dscr.get("sizing") or {}, dscr.get("reported") or {}
    src = "financing_evidence.senior_debt_evidence"
    rows = _rows(ledger, S, [
        ("Contractual repayment start (period index / date)",
         f"{rp['contractual_repayment_start_period_index']} / {rp['contractual_repayment_start_date']}" if rp else None,
         src + ".repayment", "Model period index and period start", None),
        ("Contractual maturity (period index / date)",
         f"{rp['contractual_maturity_period_index']} / {rp['contractual_maturity_date']}" if rp else None,
         src + ".repayment", "Last period in which principal can be repaid", None),
        ("Actual loan extinguishment (period index / date)",
         (f"{rp['actual_payoff_period_index']} / {rp['actual_payoff_date']}" if rp.get("actual_payoff_period_index") is not None
          else rp.get("payoff_status")) if rp else None,
         src + ".repayment", "Period in which the Senior balance reaches zero", None),
        ("Maturity minus extinguishment (periods / days)",
         f"{rp['periods_maturity_minus_payoff']} / {rp['days_maturity_minus_payoff']}" if rp.get("periods_maturity_minus_payoff") is not None else None,
         src + ".repayment", rp.get("payoff_status") or "", None),
        ("Binding sizing constraint (engine)", sz.get("binding_constraint"), src + ".sizing",
         f"Nature: {sz.get('limiting_nature') or NOT_AVAILABLE}; early payoff: {sz.get('payoff_explanation') or NOT_AVAILABLE}", None),
        ("DSCR debt capacity (funded-pass diagnostic)", number(sz.get("dscr_debt_capacity_keur")), src + ".sizing",
         "Capacity if only the DSCR budget limited the loan; the loan is the minimum of this and the gearing capacity", KEUR_FULL_FORMAT),
        ("DSCR debt capacity (standalone capacity pass)", number(sz.get("dscr_debt_capacity_standalone_pass_keur")), src + ".sizing",
         "Separate DSCR-only pass, as reported in section 1; may differ slightly from the funded-pass value", KEUR_FULL_FORMAT),
        ("Gearing debt capacity (sizing evidence)", number(sz.get("gearing_debt_capacity_keur")), src + ".sizing",
         "Capacity if only the gearing cap limited the loan", KEUR_FULL_FORMAT),
        ("Limiting row(s) for DSCR capacity (period index)",
         ", ".join(str(r["period_index"]) for r in sz.get("limiting_rows") or []) or (None if not sz else "none (gearing sets the amount)"),
         src + ".sizing.limiting_rows",
         "Rows where the DSCR budget covers interest only (no amortisation headroom)", None),
        ("Sizing DSCR target / basis", ds.get("target_dscr"), src + ".dscr.sizing",
         f"{ds.get('basis') or NOT_AVAILABLE}; requested yield scenario {ds.get('production_yield_scenario_requested') or NOT_AVAILABLE}", None),
        ("Minimum realised sizing-case DSCR", number(ds.get("min_realised_dscr")), src + ".dscr.sizing", ds.get("definition") or "", None),
        ("Reported DSCR basis", dr.get("basis"), src + ".dscr.reported", dr.get("production_basis") or "", None),
        ("Minimum reported DSCR", number(dr.get("min_dscr")), src + ".dscr.reported", dr.get("definition") or "", None),
        ("Sizing case differs from base case", dscr.get("bank_case_distinct_from_base_case"), src + ".dscr",
         "False = the requested sizing case produced the same CFADS as the base case for this Run", None),
    ])
    cur.kv(rows)
    cur.note("Loan life for any future LLCR is an Owner-defined convention: both contractual maturity and actual "
             "extinguishment are recorded above. No LLCR is computed here.")


def _write_lender(wb: Workbook, ev: RunEvidence, ledger: _Ledger) -> None:
    sh = wb.create_sheet(SHEET_LENDER)
    cur = V2._Cursor(sh, ev, "Lender Overview -- senior debt, coverage and reserves",
                     "Debt sizing, repayment, coverage and reserve evidence of the committed Run. kEUR unless stated; no recalculation.")
    rs = ev.runtime_summary
    S = SHEET_LENDER
    bank = ev.financing_evidence.get("bankability") if isinstance(ev.financing_evidence.get("bankability"), dict) else {}
    diag = ((ev.integrity_evidence.get("senior_debt") or {}).get("diagnostics")) or {}
    sd_int = ev.integrity_evidence.get("senior_debt") or {}
    cur.section("1. Senior facility and sizing evidence")
    cur.kv(_rows(ledger, S, [
        ("Final senior commitment", number(bank.get("final_senior_commitment_keur")), "financing_evidence.bankability",
         "Committed senior debt of the Run", KEUR_FULL_FORMAT),
        ("Binding sizing constraint", bank.get("binding_senior_constraint") or diag.get("binding_constraint"),
         "financing_evidence.bankability / integrity senior_debt.diagnostics", "Which sizing limit set the debt amount", None),
        ("DSCR debt capacity", number(bank.get("dscr_debt_capacity_keur")), "financing_evidence.bankability", "", KEUR_FULL_FORMAT),
        ("Gearing debt capacity", number(bank.get("gearing_debt_capacity_keur")), "financing_evidence.bankability", "", KEUR_FULL_FORMAT),
        ("Gearing ratio applied", number(bank.get("gearing_ratio")), "financing_evidence.bankability", "", PCT_FORMAT),
        ("Gearing basis", number(bank.get("gearing_basis_keur")), "financing_evidence.bankability", "", KEUR_FULL_FORMAT),
        ("Sizing solver converged", diag.get("converged"), "integrity senior_debt.diagnostics",
         f"Termination: {diag.get('termination_reason') or NOT_AVAILABLE}", None),
        ("Interest day-count", sd_int.get("day_count"), "integrity senior_debt", "", None),
        ("Terminal balloon permitted", sd_int.get("permit_terminal_balloon"), "integrity senior_debt",
         "Engine policy flag as persisted", None),
    ]))
    _write_senior_debt_evidence(cur, ledger, rs)
    periods = ev.schedule_periods("debt_schedule")
    d = _debt_summary(periods)
    sched_note = (f"Sum over {d['principal'][1]} of {d['n']} persisted periods; each period value is rounded to 2 dp"
                  if periods else "No debt schedule persisted")
    cur.section("2. Repayment profile (persisted debt schedule, 2 dp)")
    cur.kv(_rows(ledger, S, [
        ("Principal repaid (sum of periods)", d["principal"][0], "debt_schedule.senior_principal_keur", sched_note, KEUR_FORMAT, ARITHMETIC),
        ("Interest paid (sum of periods)", d["interest"][0], "debt_schedule.senior_interest_keur", sched_note, KEUR_FORMAT, ARITHMETIC),
        ("Senior debt service (sum of periods)", d["service"][0], "debt_schedule.senior_ds_keur", sched_note, KEUR_FORMAT, ARITHMETIC),
        ("Senior debt service (Run total as persisted)", number(rs.get("total_senior_ds_keur")),
         "runtime_summary.total_senior_ds_keur", "Run-level total (full precision)", KEUR_FORMAT),
        ("Peak senior balance", d["peak_balance"], "debt_schedule.senior_balance_keur",
         "Maximum of the persisted closing balances (a stock, not a sum)", KEUR_FORMAT, ARITHMETIC),
        ("Closing senior balance, last persisted period", d["closing_balance"], "debt_schedule.senior_balance_keur",
         "Stock at the last persisted period end", KEUR_FORMAT),
        ("First repayment period end", d["first_repayment"].isoformat() if d["first_repayment"] else None,
         "debt_schedule.date where principal > 0", "", None),
        ("Last repayment period end", d["last_repayment"].isoformat() if d["last_repayment"] else None,
         "debt_schedule.date where principal > 0", "", None),
    ]))
    if d["service"][0] is not None and number(rs.get("total_senior_ds_keur")) is not None:
        dev = abs(d["service"][0] - number(rs["total_senior_ds_keur"]))
        tol = _SUMMARY_TOL_BASE + _SUMMARY_TOL_PER_PERIOD * d["service"][1]
        ledger.add(S, "Debt service: sum of periods vs Run total", CHECK, "persisted values")
        cur.section("Observational consistency check (comparison only; nothing is adjusted)")
        cur.table(("Check", "Result", "Abs. deviation (kEUR)", "Tolerance (kEUR)", "Basis"),
                  [("Sum of period debt service vs persisted Run total", "PASS" if dev <= tol else "FAIL", dev, tol,
                    "Period values are persisted at 2 dp; tolerance = 0.02 + 0.005 per period")],
                  {3: KEUR_FULL_FORMAT, 4: KEUR_FULL_FORMAT}, {0: V2._FAIL_FILL} if dev > tol else None)
    else:
        ledger.add(S, "Debt service: sum of periods vs Run total", MISSING, "persisted values")
    cs = ev.construction_periods
    draws = [number(p.get("senior_draw_keur")) for p in cs]
    cur.section("3. Drawdowns (construction funding, persisted)")
    if cs:
        total_draw = sum(v for v in draws if v is not None) if any(v is not None for v in draws) else None
        cur.kv(_rows(ledger, S, [("Senior drawn during construction (sum of periods)", total_draw,
                                  "financing_evidence.construction_funding.periods.senior_draw_keur",
                                  f"{sum(1 for v in draws if v is not None)} of {len(draws)} construction periods",
                                  KEUR_FULL_FORMAT, ARITHMETIC)]))
        cur.table(("Construction period", "Period start", "Period end", "Senior draw (kEUR)", "Cumulative senior draw (kEUR)"),
                  [(p.get("period_index"), p.get("period_start"), p.get("period_end"), number(p.get("senior_draw_keur")),
                    number(p.get("cumulative_senior_draw_keur"))) for p in cs],
                  {4: KEUR_FULL_FORMAT, 5: KEUR_FULL_FORMAT})
        ledger.add(S, "Construction drawdown table", PERSISTED, "financing_evidence.construction_funding.periods")
    else:
        ledger.add(S, "Senior drawn during construction", MISSING, "financing_evidence.construction_funding")
        cur.note("NOT_AVAILABLE: no construction funding periods were persisted with this Run (legacy Run).", warn=True)
    cur.section("4. Coverage (model outputs)")
    summ = ev.debt_schedule.get("summary") if isinstance(ev.debt_schedule.get("summary"), dict) else {}

    def pick(key, alt=None):
        v = number(summ.get(key))
        return v if v is not None else number(rs.get(alt or key))
    cur.kv(_rows(ledger, S, [
        ("Minimum DSCR", pick("actual_min_dscr", "min_dscr"), "debt_schedule.summary / runtime_summary", "Model output over operating periods", MULTIPLE_FORMAT),
        ("Average DSCR", pick("actual_avg_dscr", "avg_dscr"), "debt_schedule.summary / runtime_summary", "", MULTIPLE_FORMAT),
        ("Target DSCR (sizing)", pick("target_dscr"), "debt_schedule.summary / runtime_summary",
         "Sizing target used by the engine -- not a contractual covenant", MULTIPLE_FORMAT),
        ("Minimum LLCR", pick("min_llcr"), "debt_schedule.summary / runtime_summary", "n/a when not published", MULTIPLE_FORMAT),
        ("Periods in lock-up (engine output)", summ.get("periods_in_lockup", rs.get("periods_in_lockup")),
         "debt_schedule.summary / runtime_summary", "Engine output, not a contractual test", None),
        ("Covenant compliance", "NOT DECLARED", "no contractual covenant terms persisted",
         "Compliance needs contractual thresholds that are not held with the Run", None, CONTRACTUAL),
    ]))
    su, su_src = V2._su_source(ev)
    cur.section("5. Debt service reserve (DSRA)")
    cur.kv(_rows(ledger, S, [
        ("Initial DSRA funding (a Project Use)", number(su.get("initial_dsra_funding_keur")), su_src or "engine Sources & Uses",
         "Full precision", KEUR_FULL_FORMAT),
        ("Peak DSRA balance", d["peak_dsra"], "debt_schedule.dsra_balance_keur", "Maximum of the persisted balances (a stock)", KEUR_FORMAT, ARITHMETIC),
        ("Closing DSRA balance, last persisted period", d["closing_dsra"], "debt_schedule.dsra_balance_keur",
         "Reserve available at the last persisted period end (a stock)", KEUR_FORMAT),
        ("DSRA contributions (sum of periods)", d["dsra_contribution"][0], "debt_schedule.dsra_contribution_keur",
         f"{d['dsra_contribution'][1]} persisted periods", KEUR_FORMAT, ARITHMETIC),
    ]))
    t = _terminal(ev)
    senior, dsra, acct = (t.get("senior") or {}), (t.get("senior_dsra") or {}), (t.get("distribution_account") or {})
    shl = t.get("shareholder_loan") or {}
    cur.section("6. Financial maturity and settlement (terminal state as persisted)")
    cur.kv(_rows(ledger, S, [
        ("Senior contractual maturity date", senior.get("contractual_maturity_date"),
         "sponsor_schedule.summary.terminal_financial_state.senior", "Maturity date used by the Run's terminal-state authority", None),
        ("Senior balance at contractual maturity", number(senior.get("balance_at_contractual_maturity_keur")),
         "sponsor_schedule.summary.terminal_financial_state.senior", "", KEUR_FULL_FORMAT),
        ("Senior terminal status", senior.get("status"), "sponsor_schedule.summary.terminal_financial_state.senior",
         "Model terminal classification, not a legal opinion", None),
        ("Senior terminal balance at model horizon", number(senior.get("terminal_model_horizon_balance_keur")),
         "sponsor_schedule.summary.terminal_financial_state.senior", "", KEUR_FULL_FORMAT),
        ("Senior DSRA terminal status", dsra.get("status"), "sponsor_schedule.summary.terminal_financial_state.senior_dsra", "", None),
        ("Distribution account terminal status", acct.get("status"),
         "sponsor_schedule.summary.terminal_financial_state.distribution_account", "", None),
        ("Shareholder-loan terminal status", shl.get("status"), "sponsor_schedule.summary.terminal_financial_state.shareholder_loan",
         f"Repayment mode: {shl.get('repayment_mode') or NOT_AVAILABLE}", None),
    ]))
    V2._setup(sh, {"A": 48, "B": 26, "C": 60, "D": 58, "E": 26})


def _register_purposes() -> None:
    for name, text in _V3_PURPOSE.items():
        V2._SHEET_PURPOSE.setdefault(name, text)


def _write_contents(wb: Workbook, ev: RunEvidence, sheets: Sequence[str]) -> None:
    sh = wb.create_sheet(SHEET_CONTENTS)
    cur = V2._Cursor(sh, ev, "Report Contents -- institutional lender & investor evidence package",
                     "Navigate with the links. Section order: overview, lender, schedules, statements, evidence.")
    groups = (
        ("A. Overview", (SHEET_EXECUTIVE, SHEET_LENDER)),
        ("B. Capital and funding", (V2.SHEET_SOURCES_USES, V2.SHEET_CONSTRUCTION, V2.SHEET_PERIODS)),
        ("C. Debt, reserves, tax", (V2.SHEET_DEBT, V2.SHEET_RESERVES, V2.SHEET_TAX)),
        ("D. Financial statements and returns", (V2.SHEET_PNL, V2.SHEET_CASHFLOW, V2.SHEET_BALANCE, V2.SHEET_RETURNS,
                                                  "Income Statement", "Cash Flow Statement", "Balance Sheet")),
        ("E. Evidence, integrity and limitations", (SHEET_LIMITS, V2.SHEET_INTEGRITY, V2.SHEET_DISCLOSURES, V2.SHEET_REPORT_GUIDE)),
    )
    present = set(sheets)
    for title, names in groups:
        cur.section(title)
        for name in names:
            if name not in present:
                continue
            c = sh.cell(row=cur.row, column=1, value=name)
            c.hyperlink = f"#'{name}'!A1"
            c.font = V2.Font(size=10, name="Calibri", color="0563C1", underline="single")
            sh.cell(row=cur.row, column=2, value=V2._SHEET_PURPOSE.get(name) or _V3_PURPOSE.get(name, "")).font = V2._B
            cur.row += 1
        cur.blank()
    cur.note("Every sheet carries the Run state banner, the immutable Run identifier and the evidence digest. "
             "Read Evidence_Limits before relying on any single figure.")
    V2._setup(sh, {"A": 36, "B": 100})


_V3_PURPOSE = {
    SHEET_CONTENTS: "This index.",
    SHEET_EXECUTIVE: "Project and Run identity, returns, capital, operating totals (investor view).",
    SHEET_LENDER: "Senior sizing, repayment, drawdowns, coverage, DSRA and maturity/settlement evidence (lender view).",
    SHEET_LIMITS: "For every reported metric: persisted / presentation arithmetic / missing / observational check / contractual not available.",
    "Income Statement": "Persisted income statement.", "Cash Flow Statement": "Persisted cash flow.",
    "Balance Sheet": "Persisted balance sheet.",
}


_register_purposes()


def _write_limits(wb: Workbook, ev: RunEvidence, ledger: _Ledger) -> None:
    sh = wb.create_sheet(SHEET_LIMITS)
    cur = V2._Cursor(sh, ev, "Evidence & Limitations",
                     "Class of every metric on the overview sheets and what this report does NOT claim.")
    cur.section("Evidence classes")
    cur.table(("Class", "Meaning"), [
        (PERSISTED, "Value stored by the committed Run and written unchanged."),
        (ARITHMETIC, "A sum, maximum or ratio of persisted values computed only to present them; the persisted figure is shown beside it."),
        (MISSING, "Not persisted for this Run. Shown as n/a, never as zero, and never rebuilt."),
        (CHECK, "A comparison between persisted values with a stated tolerance; it never alters a reported number."),
        (CONTRACTUAL, "Needs contractual terms (covenants, thresholds, default events) that are not held with the Run."),
    ])
    cur.section("Metric register")
    rows = sorted(ledger.entries, key=lambda e: (e[1] != SHEET_EXECUTIVE, e[1], e[0]))
    fills = {i: V2._WARN_FILL for i, r in enumerate(rows) if r[2] in (MISSING, CONTRACTUAL)}
    cur.table(("Metric", "Sheet", "Evidence class", "Source"), rows, fills=fills)
    cur.section("Not provided by this report")
    for line in (
        "Credit covenant compliance, lock-up or default findings: no contractual thresholds are persisted with a Run.",
        "Any investment recommendation, credit approval, audit opinion or independent model review.",
        "Recalculated, corrected or inferred figures: every value is read from the committed Run.",
        "Input assumptions for a Run History record: Run History persists results and identity, not the input snapshot; the "
        "current Working Copy is never substituted.",
        "Quarterly or semiannual labelling unless the Run persisted that period metadata (see the period basis).",
    ):
        cur.note("- " + line, warn=True)
    V2._setup(sh, {"A": 52, "B": 24, "C": 40, "D": 70})


# ═══════════════════════════ assembly ═══════════════════════════

def _polish(wb: Workbook, ev: RunEvidence) -> None:
    """Deterministic navigation / print settings for every sheet (no values touched)."""
    name = (ev.project_name or ev.project_code or "project")[:40]
    for sh in wb.worksheets:
        sh.page_setup.orientation = "landscape"
        sh.page_setup.fitToWidth = 1
        sh.page_setup.fitToHeight = 0
        if sh.sheet_properties.pageSetUpPr is not None:
            sh.sheet_properties.pageSetUpPr.fitToPage = True
        sh.oddHeader.left.text = f"{name}  |  Run {ev.snapshot_id}"
        sh.oddHeader.right.text = f"{ev.state}  |  &A"
        sh.oddFooter.left.text = f"Digest {ev.evidence_digest[:16]}"
        sh.oddFooter.right.text = "Page &P of &N"
        if sh.title != SHEET_CONTENTS and SHEET_CONTENTS in wb.sheetnames and sh.max_column >= 8:
            link = sh.cell(row=1, column=8)
            if link.value is None:
                link.value = "<< Contents"
                link.hyperlink = f"#'{SHEET_CONTENTS}'!A1"
                link.font = V2.Font(size=10, name="Calibri", color="0563C1", underline="single")
    colours = {SHEET_CONTENTS: "1F4E79", SHEET_EXECUTIVE: "0F274A", SHEET_LENDER: "0F274A", SHEET_LIMITS: "7F7F7F",
               V2.SHEET_DISCLOSURES: "7F7F7F"}
    for sh in wb.worksheets:
        if sh.title in colours:
            sh.sheet_properties.tabColor = colours[sh.title]


def last_run_sheet_names() -> list[str]:
    """Sheet list the Report_Guide announces for the Last Run workbook (V2 sheets plus the V3 front section and register)."""
    return [V2.SHEET_REPORT_GUIDE, *V3_FRONT_SHEETS, V2.SHEET_PERIODS, V2.SHEET_SOURCES_USES, V2.SHEET_CONSTRUCTION, V2.SHEET_DEBT,
            V2.SHEET_RESERVES, V2.SHEET_TAX, SHEET_LIMITS, V2.SHEET_INTEGRITY, V2.SHEET_DISCLOSURES]


def write_reporting_v3_sheets(wb: Workbook, ev: RunEvidence, *, v2_names: Sequence[str]) -> list[str]:
    """Add the V3 sheets to a workbook that already carries the V2 sheets; returns the V3 sheet names.

    The front section is placed directly after the V2 Report_Guide, the register before Run_Integrity.
    """
    ledger = _Ledger()
    _write_executive(wb, ev, ledger)
    _write_lender(wb, ev, ledger)
    from app.export.financial_core_reporting import write_financial_core
    write_financial_core(wb, ev, ledger, _rows, SHEET_EXECUTIVE, SHEET_LENDER)
    _write_limits(wb, ev, ledger)
    _write_contents(wb, ev, [*wb.sheetnames])
    guide = wb.sheetnames.index(V2.SHEET_REPORT_GUIDE) if V2.SHEET_REPORT_GUIDE in wb.sheetnames else -1
    front = [SHEET_CONTENTS, SHEET_EXECUTIVE, SHEET_LENDER]
    for offset, name in enumerate(front, start=1):
        wb.move_sheet(name, offset=(guide + offset) - wb.sheetnames.index(name))
    if V2.SHEET_INTEGRITY in wb.sheetnames:
        wb.move_sheet(SHEET_LIMITS, offset=wb.sheetnames.index(V2.SHEET_INTEGRITY) - wb.sheetnames.index(SHEET_LIMITS))
    _polish(wb, ev)
    return [*front, SHEET_LIMITS]


def build_run_evidence_workbook_v3(ev: RunEvidence) -> bytes:
    """Self-contained, deterministic V3 package for ONE committed Run (used for historical Runs).

    Reads only ``ev``: no clock, no Working Copy, no engine.  Equal evidence -> equal bytes.
    """
    wb = Workbook()
    wb.remove(wb.active)
    names = [V2.SHEET_REPORT_GUIDE, SHEET_CONTENTS, SHEET_EXECUTIVE, SHEET_LENDER, V2.SHEET_PERIODS, V2.SHEET_SOURCES_USES,
             V2.SHEET_CONSTRUCTION, V2.SHEET_DEBT, V2.SHEET_RESERVES, V2.SHEET_PNL, V2.SHEET_CASHFLOW, V2.SHEET_BALANCE,
             V2.SHEET_TAX, V2.SHEET_RETURNS, SHEET_LIMITS, V2.SHEET_INTEGRITY, V2.SHEET_DISCLOSURES]
    V2._write_report_guide(wb, ev, names, inputs_note=(
        "This workbook was produced from the Run's own persisted evidence only. Input-assumption sheets "
        "(Inputs, CAPEX, OPEX, Revenue assumptions) are intentionally not included: Run History does not "
        "persist the input snapshot, and the current Working Copy must not be substituted for it."
        if ev.source != "LAST_RUN" else ""))
    V2._write_periods(wb, ev)
    V2._write_sources_uses(wb, ev)
    V2._write_construction(wb, ev)
    V2._write_debt(wb, ev)
    V2._write_reserves(wb, ev)
    V2._write_statement(wb, ev, V2.SHEET_PNL, "Income Statement", "pnl", V2._PNL_SPECS,
                        "Flows have a Total over persisted periods; retained earnings is a balance.")
    V2._write_statement(wb, ev, V2.SHEET_CASHFLOW, "Cash Flow Statement", "pf_cash_waterfall", V2._CASH_SPECS,
                        "Project-finance cash waterfall as persisted by the Run.")
    V2._write_statement(wb, ev, V2.SHEET_BALANCE, "Balance Sheet", "balance_sheet", V2._BS_SPECS,
                        "Balances are point-in-time and never totalled. The balance check is the runtime's own residual; "
                        "totals the runtime did not publish are not constructed.")
    V2._write_tax(wb, ev)
    V2._write_returns(wb, ev)
    V2._write_integrity(wb, ev)
    V2._write_disclosures(wb, ev)
    write_reporting_v3_sheets(wb, ev, v2_names=names)
    return V2.finalize_workbook_bytes(wb, ev.ran_at)
