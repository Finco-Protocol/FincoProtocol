"""Bounded, NON-EXECUTING structural inspection of .xlsx / .xlsm workbooks.

Everything in a workbook is untrusted data.  This module never evaluates a formula, never decompresses or
interprets VBA, never opens an external link/DDE/ActiveX part, never writes a file and never calls an engine.
It reads: workbook metadata (sheets, hidden state, defined names, calc settings), cell literals, formula
*text* and the cached result the producing application stored next to it, external-link *descriptors*, and the
OLE directory (stream names only) of a VBA project.

Cached formula results are evidence of what some spreadsheet application once displayed.  They are exposed here
as ``cached`` and are never an authoritative FINCO input.
"""
from __future__ import annotations

import hashlib
import io
import os
import re
import struct
import time
import zipfile
from collections import Counter
from dataclasses import dataclass, field
from datetime import date, datetime
from typing import Any

from defusedxml import ElementTree as SafeET
from defusedxml.common import DefusedXmlException

from app.model_import.intake import (
    MAX_CELLS, MAX_COLS, MAX_ROWS_PER_SHEET, MAX_SHEETS, MAX_UPLOAD_BYTES, IntakeError, _check_time,
    _check_xlsx_zip, _col, _safe_filename,
)

MAX_XML_PART_BYTES = 2 * 1024 * 1024
MAX_DEFINED_NAMES = 500
MAX_FORMULA_TEXT = 240
MAX_RANGE_EXPANSION = 2000
MAX_VBA_BYTES = 8 * 1024 * 1024
MAX_VBA_SECTORS = 8192
MAX_VBA_DIR_ENTRIES = 512
MAX_VBA_MODULES = 200

_NS_MAIN = "{http://schemas.openxmlformats.org/spreadsheetml/2006/main}"
_NS_REL = "{http://schemas.openxmlformats.org/officeDocument/2006/relationships}"
_CONTROL = re.compile("[\x00-\x1f\x7f-\x9f\u200b-\u200f\u202a-\u202e\u2066-\u2069\ufeff]")   # controls + bidi/zero-width spoofing

# Functions that need a live application, the network, a COM host or dynamic addressing.  Presence is
# reported; nothing is evaluated.
_ACTIVE_FUNCTIONS = frozenset({
    "WEBSERVICE", "FILTERXML", "HYPERLINK", "CALL", "REGISTER", "REGISTER.ID", "EXEC", "RTD", "DDE",
    "SQL.REQUEST", "EUROCONVERT",
})
_DYNAMIC_FUNCTIONS = frozenset({"INDIRECT", "OFFSET", "CELL", "INFO", "NOW", "TODAY", "RAND", "RANDBETWEEN"})
_SELECTOR_FUNCTIONS = frozenset({"CHOOSE", "INDEX", "MATCH", "VLOOKUP", "HLOOKUP", "XLOOKUP", "LOOKUP", "IF", "IFS",
                                 "SWITCH", "OFFSET", "INDIRECT"})

_STRING_LITERAL = re.compile(r'"(?:[^"]|"")*"')
_FUNCTION_CALL = re.compile(r"(?<![A-Za-z0-9_.])([A-Za-z][A-Za-z0-9_.]*)\(")
_CELL_REF = re.compile(
    r"(?:(?:'((?:[^']|'')+)'|([A-Za-z0-9_.]+))!)?\$?([A-Z]{1,3})\$?([0-9]{1,7})(?::\$?([A-Z]{1,3})\$?([0-9]{1,7}))?(?![A-Za-z0-9_(])")
_EXTERNAL_REF = re.compile(r"\[\d+\]|\[[^\]]+\.(?:xls|xlsx|xlsm|xlsb)\]", re.IGNORECASE)
_DDE_REF = re.compile(r"'[^']+'\|'?[^!]*'?!|\bDDE\s*\(", re.IGNORECASE)


class InspectionError(IntakeError):
    """Typed, fail-closed inspection rejection (shares the Import error vocabulary)."""


def clean_text(value: Any, limit: int = 120) -> str:
    """Printable, bounded text safe to display (Jinja autoescapes on top of this)."""
    return _CONTROL.sub("", str(value if value is not None else "")).strip()[:limit]


def col_to_index(col: str) -> int:
    n = 0
    for ch in col:
        n = n * 26 + (ord(ch) - 64)
    return n


@dataclass
class Inspection:
    """``report`` is JSON-safe and displayable.  The rest is in-memory working data and is never persisted."""
    report: dict
    tables: list = field(default_factory=list)            # intake-identical (formulas masked), for the mapping engine
    formulas: dict = field(default_factory=dict)          # (sheet, coord) -> {"text", "cached", "refs", "functions"}
    literals: dict = field(default_factory=dict)          # (sheet, coord) -> literal value
    names: dict = field(default_factory=dict)             # lower-case defined name -> descriptor
    sheet_state: dict = field(default_factory=dict)       # sheet title -> visible|hidden|veryHidden
    digest: str = ""


# ── workbook.xml / content types / rels ────────────────────────────────────────────────────────────────────
def _read_part(archive: zipfile.ZipFile, name: str, *, required: bool = False) -> bytes | None:
    try:
        info = archive.getinfo(name)
    except KeyError:
        if required:
            raise InspectionError("IMPORT_XLSX_INVALID_STRUCTURE", name) from None
        return None
    if info.file_size > MAX_XML_PART_BYTES:
        raise InspectionError("IMPORT_XLSX_PART_TOO_LARGE", name)
    return archive.read(name)


def _xml(raw: bytes | None):
    if raw is None:
        return None
    try:
        return SafeET.fromstring(raw, forbid_dtd=True, forbid_entities=True, forbid_external=True)
    except (DefusedXmlException, SafeET.ParseError, ValueError) as exc:
        raise InspectionError("IMPORT_XLSX_XML_UNSAFE_OR_MALFORMED") from exc


def _workbook_meta(archive: zipfile.ZipFile) -> dict:
    root = _xml(_read_part(archive, "xl/workbook.xml", required=True))
    sheets = []
    for el in root.iter(f"{_NS_MAIN}sheet"):
        sheets.append({
            "name": clean_text(el.get("name"), 120), "state": el.get("state") or "visible",
            "rid": el.get(f"{_NS_REL}id"),
        })
        if len(sheets) > MAX_SHEETS:
            raise InspectionError("IMPORT_SHEET_LIMIT_EXCEEDED")
    names = []
    for el in root.iter(f"{_NS_MAIN}definedName"):
        if len(names) >= MAX_DEFINED_NAMES:
            names.append({"name": "(truncated)", "refers_to": "", "scope": "workbook", "hidden": False, "kind": "truncated"})
            break
        sid = el.get("localSheetId")
        refers = (el.text or "").strip()
        names.append({
            "name": clean_text(el.get("name"), 80), "refers_to": clean_text(refers, 200),
            "scope": (sheets[int(sid)]["name"] if sid and sid.isdigit() and int(sid) < len(sheets) else "workbook"),
            "hidden": el.get("hidden") in ("1", "true"),
        })
    calc = root.find(f"{_NS_MAIN}calcPr")
    ext_refs = len(list(root.iter(f"{_NS_MAIN}externalReference")))
    return {
        "sheets": sheets, "defined_names": names, "external_references": ext_refs,
        "calc": {
            "mode": (calc.get("calcMode") if calc is not None else None) or "auto",
            "iterative": bool(calc is not None and calc.get("iterate") in ("1", "true")),
            "full_calc_on_load": bool(calc is not None and calc.get("fullCalcOnLoad") in ("1", "true")),
        },
    }


def _classify_name(refers_to: str) -> dict:
    ref = refers_to.lstrip("=")
    if _EXTERNAL_REF.search(ref):
        return {"kind": "external", "target": None}
    m = _CELL_REF.fullmatch(ref.replace("$", ""))
    if m and m.group(1) is None and m.group(2) is None:
        return {"kind": "formula", "target": None}
    if m:
        sheet = (m.group(1) or m.group(2) or "").replace("''", "'")
        a = f"{m.group(3)}{m.group(4)}"
        b = f"{m.group(5)}{m.group(6)}" if m.group(5) else None
        return {"kind": "range" if b else "cell", "target": {"sheet": sheet, "start": a, "end": b}}
    if re.fullmatch(r"-?[0-9.]+(?:E[+-]?[0-9]+)?|\"[^\"]*\"", ref, re.IGNORECASE):
        return {"kind": "constant", "target": None}
    return {"kind": "formula", "target": None}


def _external_links(archive: zipfile.ZipFile, names: set[str]) -> list[dict]:
    """Descriptors only.  Targets are never opened, resolved or fetched."""
    links = []
    for name in sorted(names):
        if not (name.startswith("xl/externalLinks/_rels/") and name.endswith(".rels")):
            continue
        root = _xml(_read_part(archive, name))
        for rel in (root.iter() if root is not None else ()):
            target = rel.get("Target")
            if target is None:
                continue
            low = target.casefold()
            kind = ("network_url" if re.match(r"[a-z][a-z0-9+.-]*://", low) and not low.startswith("file:")
                    else "file_path" if low.startswith("file:") or re.match(r"^[a-z]:[\\/]", low) or low.startswith(("/", "\\", ".."))
                    else "relative_workbook")
            links.append({"kind": kind, "target_mode": rel.get("TargetMode") or "Internal",
                          "target_hint": clean_text(os.path.basename(target.replace("\\", "/")), 60)})
        if len(links) >= 50:
            break
    return links


# ── VBA container: directory stream names only (no decompression, no source) ──────────────────────────────
def _sector_chain(fat: list[int], start: int, limit: int) -> list[int]:
    chain, seen = [], set()
    cur = start
    while 0 <= cur < 0xFFFFFFFA:
        if cur in seen or len(chain) >= limit:
            raise ValueError("CFB_CHAIN_LOOP_OR_LIMIT")
        seen.add(cur)
        chain.append(cur)
        if cur >= len(fat):
            raise ValueError("CFB_CHAIN_OUT_OF_RANGE")
        cur = fat[cur]
    return chain


def parse_vba_container(blob: bytes) -> dict:
    """Return module/stream NAMES from a vbaProject.bin OLE container.  Never raises, never executes.

    Module source is compressed inside the module streams; it is not read.  This is metadata only.
    """
    result: dict = {"status": "PARSED", "modules": [], "stream_count": 0, "module_stream_bytes": 0, "reason": None}
    try:
        if len(blob) > MAX_VBA_BYTES:
            raise ValueError("VBA_CONTAINER_TOO_LARGE")
        if len(blob) < 512 or blob[:8] != b"\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1":
            raise ValueError("VBA_CONTAINER_SIGNATURE_INVALID")
        shift = struct.unpack_from("<H", blob, 30)[0]
        if shift not in (9, 12):
            raise ValueError("VBA_CONTAINER_SECTOR_SIZE_INVALID")
        size = 1 << shift
        n_fat, first_dir = struct.unpack_from("<II", blob, 44)
        if n_fat == 0 or n_fat > 109:
            raise ValueError("VBA_CONTAINER_FAT_UNSUPPORTED")
        difat = struct.unpack_from("<109I", blob, 76)[:n_fat]
        total_sectors = (len(blob) - size) // size
        if total_sectors > MAX_VBA_SECTORS:
            raise ValueError("VBA_CONTAINER_TOO_MANY_SECTORS")

        def sector(i: int) -> bytes:
            off = (i + 1) * size
            if i < 0 or off + size > len(blob):
                raise ValueError("VBA_CONTAINER_SECTOR_OUT_OF_RANGE")
            return blob[off:off + size]

        fat: list[int] = []
        for sid in difat:
            fat.extend(struct.unpack_from(f"<{size // 4}I", sector(sid)))
        directory = b"".join(sector(s) for s in _sector_chain(fat, first_dir, MAX_VBA_SECTORS))
        count = len(directory) // 128
        if count > MAX_VBA_DIR_ENTRIES:
            raise ValueError("VBA_CONTAINER_DIRECTORY_TOO_LARGE")
        entries = []
        for i in range(count):
            raw = directory[i * 128:(i + 1) * 128]
            name_len = struct.unpack_from("<H", raw, 64)[0]
            kind = raw[66]
            if kind == 0:
                entries.append(None)
                continue
            if name_len < 2 or name_len > 64:
                raise ValueError("VBA_CONTAINER_NAME_INVALID")
            name = raw[:name_len - 2].decode("utf-16-le", errors="replace")
            left, right, child = struct.unpack_from("<III", raw, 68)
            stream_size = struct.unpack_from("<Q", raw, 120)[0] & 0xFFFFFFFF
            entries.append({"name": name, "type": kind, "left": left, "right": right, "child": child, "size": stream_size})
        if not entries or entries[0] is None or entries[0]["type"] != 5:
            raise ValueError("VBA_CONTAINER_ROOT_MISSING")

        def children(index: int) -> list[int]:
            """Red-black sibling tree walk, iterative and loop-safe."""
            out, stack, seen = [], [entries[index]["child"]], set()
            while stack:
                cur = stack.pop()
                if cur >= count or cur in seen or entries[cur] is None:
                    continue
                seen.add(cur)
                out.append(cur)
                stack.extend((entries[cur]["left"], entries[cur]["right"]))
            return out

        vba_storage = next((c for c in children(0) if entries[c]["name"].upper() == "VBA" and entries[c]["type"] == 1), None)
        streams = [c for c in children(0) if entries[c]["type"] == 2]
        result["stream_count"] = len(streams) + (len([c for c in children(vba_storage) if entries[c]["type"] == 2]) if vba_storage is not None else 0)
        if vba_storage is None:
            result["status"] = "NO_VBA_STORAGE"
            return result
        skip = {"DIR", "_VBA_PROJECT"}
        modules = []
        for c in children(vba_storage):
            e = entries[c]
            if e["type"] != 2 or e["name"].upper() in skip or e["name"].startswith("__SRP_"):
                continue
            modules.append({"name": clean_text(e["name"], 64), "stream_bytes": e["size"]})
            if len(modules) >= MAX_VBA_MODULES:
                result["status"] = "PARSED_TRUNCATED"
                break
        modules.sort(key=lambda m: m["name"].casefold())
        result["modules"] = modules
        result["module_stream_bytes"] = sum(m["stream_bytes"] for m in modules)
    except (ValueError, struct.error, IndexError, OverflowError) as exc:
        return {"status": "MALFORMED_CONTAINER", "modules": [], "stream_count": 0, "module_stream_bytes": 0,
                "reason": clean_text(exc, 80) or "VBA_CONTAINER_MALFORMED"}
    return result


# ── formulas ───────────────────────────────────────────────────────────────────────────────────────────────
def analyse_formula(text: str, home_sheet: str, names: dict) -> dict:
    """Static token scan.  Returns references, function names and risk flags; evaluates nothing."""
    body = text[1:] if text.startswith("=") else text
    stripped = _STRING_LITERAL.sub('""', body)
    functions = sorted({m.group(1).upper() for m in _FUNCTION_CALL.finditer(stripped)})
    refs, named = [], []
    for m in _CELL_REF.finditer(stripped):
        sheet = (m.group(1) or m.group(2) or home_sheet).replace("''", "'")
        refs.append({"sheet": sheet, "start": f"{m.group(3)}{m.group(4)}",
                     "end": f"{m.group(5)}{m.group(6)}" if m.group(5) else None})
        if len(refs) >= 64:
            break
    for token in set(re.findall(r"(?<![A-Za-z0-9_.!'])([A-Za-z_][A-Za-z0-9_.]{1,79})(?![A-Za-z0-9_.(!'])", stripped)):
        if token.lower() in names:
            named.append(token)
    flags = []
    if _EXTERNAL_REF.search(stripped):
        flags.append("EXTERNAL_WORKBOOK_REFERENCE")
    if _DDE_REF.search(stripped):
        flags.append("DDE_REFERENCE")
    if set(functions) & _ACTIVE_FUNCTIONS:
        flags.append("ACTIVE_OR_NETWORK_FUNCTION")
    if set(functions) & _DYNAMIC_FUNCTIONS:
        flags.append("DYNAMIC_OR_VOLATILE_FUNCTION")
    return {"refs": refs, "names": sorted(named)[:16], "functions": functions, "flags": flags}


def expand_ref(ref: dict, limit: int = MAX_RANGE_EXPANSION) -> list[tuple[str, str]] | None:
    """Cells of a reference, or None when the range exceeds the bound (recorded as unexpanded)."""
    if ref["end"] is None:
        return [(ref["sheet"], ref["start"])]
    a = re.fullmatch(r"([A-Z]+)([0-9]+)", ref["start"])
    b = re.fullmatch(r"([A-Z]+)([0-9]+)", ref["end"])
    if not a or not b:
        return None
    c1, c2 = sorted((col_to_index(a.group(1)), col_to_index(b.group(1))))
    r1, r2 = sorted((int(a.group(2)), int(b.group(2))))
    if (c2 - c1 + 1) * (r2 - r1 + 1) > limit:
        return None
    return [(ref["sheet"], f"{_col(c)}{r}") for r in range(r1, r2 + 1) for c in range(c1, c2 + 1)]


# ── main entry ─────────────────────────────────────────────────────────────────────────────────────────────
def _cell_formula_text(cell: Any) -> str | None:
    value = cell.value
    if getattr(cell, "data_type", None) == "f" or (isinstance(value, str) and value.startswith("=")):
        text = getattr(value, "text", value)
        return str(text if text is not None else "")
    if hasattr(value, "text") and hasattr(value, "ref"):          # array formula object
        return str(getattr(value, "text", "") or "")
    return None


def inspect_workbook(data: bytes, filename: str) -> Inspection:
    start = time.monotonic()
    deadline_start = start            # same PARSER_TIMEOUT_SECONDS bound as the apply path
    name = _safe_filename(filename)
    ext = os.path.splitext(name)[1].lower()
    if ext not in (".xlsx", ".xlsm"):
        raise InspectionError("IMPORT_FORMAT_UNSUPPORTED")
    if not isinstance(data, bytes) or not data or len(data) > MAX_UPLOAD_BYTES:
        raise InspectionError("IMPORT_FILE_SIZE_INVALID")
    flags = _check_xlsx_zip(data, deadline_start, inspect=True)
    names = set(flags["member_names"])
    warnings: list[str] = []
    with zipfile.ZipFile(io.BytesIO(data)) as archive:
        meta = _workbook_meta(archive)
        content_types = (_read_part(archive, "[Content_Types].xml", required=True) or b"").decode("utf-8", "replace")
        macro_content_type = "macroEnabled" in content_types
        links = _external_links(archive, names)
        vba_blob = None
        if flags["vba_members"]:
            vba_name = next((m for m in flags["vba_members"] if m.casefold().endswith("vbaproject.bin")), None)
            if vba_name:
                info = archive.getinfo(vba_name)
                vba_blob = archive.read(vba_name) if info.file_size <= MAX_VBA_BYTES else b""
    vba = parse_vba_container(vba_blob) if vba_blob is not None else {"status": "ABSENT", "modules": [], "reason": None,
                                                                       "stream_count": 0, "module_stream_bytes": 0}
    if vba_blob is not None and not vba_blob:
        vba = {"status": "MALFORMED_CONTAINER", "modules": [], "reason": "VBA_CONTAINER_TOO_LARGE", "stream_count": 0, "module_stream_bytes": 0}
    macro_bearing = bool(flags["vba_members"]) or macro_content_type or ext == ".xlsm"
    if ext == ".xlsx" and (flags["vba_members"] or macro_content_type):
        warnings.append("EXTENSION_MACRO_MISMATCH")
    if ext == ".xlsm" and not flags["vba_members"]:
        warnings.append("XLSM_WITHOUT_VBA_PROJECT")
    if vba["status"] == "MALFORMED_CONTAINER":
        warnings.append("VBA_CONTAINER_MALFORMED")
    if names & {"xl/vbaProjectSignature.bin"}:
        warnings.append("VBA_PROJECT_SIGNED")

    from openpyxl import load_workbook
    sheet_titles = [s["name"] for s in meta["sheets"]]
    state = {s["name"]: s["state"] for s in meta["sheets"]}
    defined = {}
    for item in meta["defined_names"]:
        item.update(_classify_name(item["refers_to"]))
        defined[item["name"].lower()] = item

    try:
        wb_f = load_workbook(io.BytesIO(data), read_only=True, data_only=False, keep_links=False, keep_vba=False)
    except Exception as exc:
        raise InspectionError("IMPORT_XLSX_PARSE_FAILED") from exc
    tables, formulas, literals, sheet_reports = [], {}, {}, []
    total_cells = 0
    try:
        if len(wb_f.worksheets) > MAX_SHEETS:
            raise InspectionError("IMPORT_SHEET_LIMIT_EXCEEDED")
        for sheet in wb_f.worksheets:
            _check_time(deadline_start)
            if (sheet.max_row and sheet.max_row > MAX_ROWS_PER_SHEET) or (sheet.max_column and sheet.max_column > MAX_COLS):
                raise InspectionError("IMPORT_TABLE_LIMIT_EXCEEDED")
            rows, n_lit, n_formula, max_r, max_c = [], 0, 0, 0, 0
            for i, cells in enumerate(sheet.iter_rows(), 1):
                _check_time(deadline_start)
                if i > MAX_ROWS_PER_SHEET or len(cells) > MAX_COLS:
                    raise InspectionError("IMPORT_TABLE_LIMIT_EXCEEDED")
                total_cells += len(cells)
                if total_cells > MAX_CELLS:
                    raise InspectionError("IMPORT_TABLE_LIMIT_EXCEEDED")
                vals = []
                for j, cell in enumerate(cells, 1):
                    coord = getattr(cell, "coordinate", None) or f"{_col(j)}{i}"
                    val = cell.value
                    fmt = str(getattr(cell, "number_format", "") or "")
                    ftext = _cell_formula_text(cell)
                    if ftext is not None:
                        n_formula += 1
                        max_r, max_c = max(max_r, i), max(max_c, j)
                        info = analyse_formula(ftext, sheet.title, defined)
                        formulas[(sheet.title, coord)] = {"text": clean_text(ftext, MAX_FORMULA_TEXT), "cached": None,
                                                          "has_cached": False, "percent_format": "%" in fmt, **info}
                        vals.append(("=FORMULA_NOT_AUTHORIZED", None, coord))
                        continue
                    if val is not None and not (isinstance(val, str) and not val.strip()):
                        n_lit += 1
                        max_r, max_c = max(max_r, i), max(max_c, j)
                        literals[(sheet.title, coord)] = val
                    inferred = "fraction" if "%" in fmt and isinstance(val, (float, int)) and not isinstance(val, bool) else None
                    vals.append((val, inferred, coord))
                rows.append(vals)
            tables.append((sheet.title, rows))
            sheet_reports.append({
                "name": clean_text(sheet.title, 120), "state": state.get(sheet.title, "visible"),
                "hidden": state.get(sheet.title, "visible") != "visible",
                "used_rows": max_r, "used_columns": max_c, "literal_cells": n_lit, "formula_cells": n_formula,
            })
    except IntakeError:
        raise
    except Exception as exc:
        raise InspectionError("IMPORT_XLSX_ROWS_INVALID") from exc
    finally:
        wb_f.close()

    # Pass 2: the cached result stored next to each formula (evidence only, never authoritative).
    try:
        wb_v = load_workbook(io.BytesIO(data), read_only=True, data_only=True, keep_links=False, keep_vba=False)
    except Exception as exc:
        raise InspectionError("IMPORT_XLSX_PARSE_FAILED") from exc
    try:
        for sheet in wb_v.worksheets:
            _check_time(deadline_start)
            for i, cells in enumerate(sheet.iter_rows(), 1):
                for j, cell in enumerate(cells, 1):
                    coord = getattr(cell, "coordinate", None) or f"{_col(j)}{i}"
                    entry = formulas.get((sheet.title, coord))
                    if entry is not None and cell.value is not None:
                        v = cell.value
                        entry["cached"] = v.isoformat() if isinstance(v, (date, datetime)) else v if isinstance(v, (int, float, bool)) else clean_text(v, 80)
                        entry["has_cached"] = True
    except IntakeError:
        raise
    except Exception as exc:
        raise InspectionError("IMPORT_XLSX_ROWS_INVALID") from exc
    finally:
        wb_v.close()

    per_sheet_formula = Counter(k[0] for k in formulas)
    per_sheet_cached = Counter(k[0] for k, v in formulas.items() if v["has_cached"])
    for rep in sheet_reports:
        rep["formula_with_cached_value"] = per_sheet_cached.get(rep["name"], 0)
        rep["formula_without_cached_value"] = per_sheet_formula.get(rep["name"], 0) - per_sheet_cached.get(rep["name"], 0)

    function_counts = Counter(fn for f in formulas.values() for fn in f["functions"])
    formula_flags = Counter(flag for f in formulas.values() for flag in f["flags"])
    if formula_flags.get("EXTERNAL_WORKBOOK_REFERENCE") or links:
        warnings.append("EXTERNAL_LINKS_PRESENT")
    if formula_flags.get("DDE_REFERENCE"):
        warnings.append("DDE_REFERENCE_PRESENT")
    if formula_flags.get("ACTIVE_OR_NETWORK_FUNCTION"):
        warnings.append("ACTIVE_FUNCTION_PRESENT")
    if meta["calc"]["iterative"]:
        warnings.append("ITERATIVE_CALCULATION_ENABLED")
    if any(r["hidden"] for r in sheet_reports):
        warnings.append("HIDDEN_SHEETS_PRESENT")
    if any(s["formula_without_cached_value"] for s in sheet_reports):
        warnings.append("FORMULAS_WITHOUT_CACHED_VALUES")
    if flags["active_members"]:
        warnings.append("ACTIVEX_OR_EMBEDDED_OBJECTS_PRESENT")
    if meta["calc"]["mode"] == "manual":
        warnings.append("MANUAL_CALCULATION_MODE_CACHED_VALUES_MAY_BE_STALE")

    digest = hashlib.sha256(data).hexdigest()
    report = {
        "filename": name, "format": ext[1:], "size_bytes": len(data), "sha256": digest,
        "container": {
            "member_count": len(names), "macro_enabled_content_type": macro_content_type,
            "vba_project_present": bool(flags["vba_members"]), "external_link_parts": len(flags["external_members"]),
            "activex_or_embedded_parts": len(flags["active_members"]),
        },
        "macro_bearing": macro_bearing,
        "vba": {"status": vba["status"], "reason": vba.get("reason"), "module_count": len(vba["modules"]),
                "modules": vba["modules"], "module_stream_bytes": vba["module_stream_bytes"],
                "source_interpreted": False, "executed": False},
        "sheets": sheet_reports,
        "defined_names": meta["defined_names"],
        "external_links": links,
        "calc": meta["calc"],
        "formulas": {
            "total": len(formulas), "with_cached_value": sum(v["has_cached"] for v in formulas.values()),
            "functions": dict(function_counts.most_common(40)), "flags": dict(formula_flags),
        },
        "literal_cells": len(literals),
        "warnings": sorted(set(warnings)),
        "executed": {"vba": False, "formulas": False, "external_links": False, "dde": False},
    }
    return Inspection(report=report, tables=tables, formulas=formulas, literals=literals,
                      names=defined, sheet_state=state, digest=digest)
