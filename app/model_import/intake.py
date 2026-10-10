"""Workflow C deterministic financial input extraction. No file persistence or engine.

A proposed source label is EVIDENCE, never write authority.  Mapping to an
actual editable field is revalidated by WorkbookUpdateService in C0.
"""
from __future__ import annotations

import csv
import io
import os
import re
import time
import zipfile
from collections import Counter, defaultdict
from datetime import date, datetime
from decimal import Decimal, InvalidOperation
from pathlib import PurePath
from typing import Any

from app.workbook.registry import WORKBOOK
from app.workbook.specs import FieldType
from app.workbook.update_service import (
    BatchApplyError, WorkbookUpdateService, _assert_batch_field_applicable,
)
from app.v2.register_path_map import register_path_for_field

MAX_UPLOAD_BYTES = 2 * 1024 * 1024
MAX_EXPANDED_BYTES = 24 * 1024 * 1024
MAX_SHEETS = 12
MAX_ROWS_PER_SHEET = 1500
MAX_COLS = 24
MAX_CELLS = 12000
MAX_ZIP_MEMBERS = 120
PARSER_TIMEOUT_SECONDS = 5.0

class IntakeError(ValueError):
    def __init__(self, code: str, message: str = "") -> None:
        self.code = code
        super().__init__(message or code)


def _check_time(start: float) -> None:
    if time.monotonic() - start > PARSER_TIMEOUT_SECONDS:
        raise IntakeError("IMPORT_PARSE_TIMEOUT")


def _safe_filename(filename: str) -> str:
    # No source path or uploaded customer identifier is ever used in storage.
    base = (filename or "").replace("\\", "/").rsplit("/", 1)[-1]
    if not base or len(base) > 160 or any(ord(ch) < 32 for ch in base):
        raise IntakeError("IMPORT_FILENAME_INVALID")
    return base


def _csv_rows(data: bytes, *, encoding: str, start: float):
    if encoding not in ("utf-8", "utf-8-sig", "cp1250", "cp1252"):
        raise IntakeError("IMPORT_ENCODING_UNSUPPORTED")
    try:
        text = data.decode(encoding, errors="strict")
    except UnicodeDecodeError as exc:
        raise IntakeError("IMPORT_CSV_DECODING_FAILED") from exc
    try:
        delimiter = csv.Sniffer().sniff(text[:3072], delimiters=",;\t").delimiter
    except csv.Error:
        delimiter = ","
    rows = []
    total_cells = 0
    try:
        for i, row in enumerate(csv.reader(io.StringIO(text), delimiter=delimiter), start=1):
            _check_time(start)
            total_cells += len(row)
            if i > MAX_ROWS_PER_SHEET or len(row) > MAX_COLS or total_cells > MAX_CELLS:
                raise IntakeError("IMPORT_TABLE_LIMIT_EXCEEDED")
            rows.append([(item, None, f"{_col(j + 1)}{i}") for j, item in enumerate(row)])
    except csv.Error as exc:
        raise IntakeError("IMPORT_CSV_MALFORMED") from exc
    return [("CSV", rows)]


def _col(index: int) -> str:
    result = ""
    while index:
        index, rem = divmod(index - 1, 26)
        result = chr(65 + rem) + result
    return result


def _check_xlsx_zip(data: bytes, start: float) -> None:
    try:
        with zipfile.ZipFile(io.BytesIO(data)) as archive:
            members = archive.infolist()
            if len(members) > MAX_ZIP_MEMBERS:
                raise IntakeError("IMPORT_ZIP_MEMBER_LIMIT_EXCEEDED")
            if sum(x.file_size for x in members) > MAX_EXPANDED_BYTES:
                raise IntakeError("IMPORT_XLSX_EXPANSION_LIMIT_EXCEEDED")
            names = set()
            for entry in members:
                _check_time(start)
                name = entry.filename.replace("\\", "/")
                if (name.startswith("/") or name.startswith("../") or "/../" in name
                        or name in names or entry.flag_bits & 1):
                    raise IntakeError("IMPORT_XLSX_UNSAFE_ARCHIVE")
                names.add(name)
                if entry.file_size > MAX_EXPANDED_BYTES:
                    raise IntakeError("IMPORT_XLSX_EXPANSION_LIMIT_EXCEEDED")
                if entry.file_size and entry.compress_size == 0:
                    raise IntakeError("IMPORT_XLSX_UNSAFE_COMPRESSION")
                if entry.compress_size and entry.file_size / entry.compress_size > 150:
                    raise IntakeError("IMPORT_XLSX_UNSAFE_COMPRESSION")
                low = name.casefold()
                if ("vbaproject" in low or "/externallinks/" in low
                        or "/activex/" in low or "/embeddings/" in low
                        or "/connections" in low or "querytable" in low):
                    raise IntakeError("IMPORT_XLSX_ACTIVE_OR_EXTERNAL_CONTENT")
            if "xl/workbook.xml" not in names or "[Content_Types].xml" not in names:
                raise IntakeError("IMPORT_XLSX_INVALID_STRUCTURE")
    except (zipfile.BadZipFile, EOFError) as exc:
        raise IntakeError("IMPORT_XLSX_MALFORMED") from exc


def _xlsx_rows(data: bytes, *, start: float):
    _check_xlsx_zip(data, start)
    from openpyxl import load_workbook
    try:
        book = load_workbook(io.BytesIO(data), read_only=True,
                             data_only=False, keep_links=False)
    except Exception as exc:
        raise IntakeError("IMPORT_XLSX_PARSE_FAILED") from exc
    try:
        if len(book.worksheets) > MAX_SHEETS:
            raise IntakeError("IMPORT_SHEET_LIMIT_EXCEEDED")
        tables = []
        total_cells = 0
        for sheet in book.worksheets:
            _check_time(start)
            if sheet.max_row and sheet.max_row > MAX_ROWS_PER_SHEET:
                raise IntakeError("IMPORT_TABLE_LIMIT_EXCEEDED")
            if sheet.max_column and sheet.max_column > MAX_COLS:
                raise IntakeError("IMPORT_TABLE_LIMIT_EXCEEDED")
            rows = []
            for i, cells in enumerate(sheet.iter_rows(), 1):
                _check_time(start)
                if i > MAX_ROWS_PER_SHEET or len(cells) > MAX_COLS:
                    raise IntakeError("IMPORT_TABLE_LIMIT_EXCEEDED")
                total_cells += len(cells)
                if total_cells > MAX_CELLS:
                    raise IntakeError("IMPORT_TABLE_LIMIT_EXCEEDED")
                vals = []
                for j, cell in enumerate(cells, 1):
                    fmt = str(getattr(cell, "number_format", "") or "")
                    val = cell.value
                    if getattr(cell, "data_type", None) == "f" or (
                        isinstance(val, str) and val.startswith("=")
                    ):
                        val = "=FORMULA_NOT_AUTHORIZED"
                    inferred_unit = "fraction" if "%" in fmt and isinstance(val, (float, int)) else None
                    vals.append((val, inferred_unit, getattr(cell, "coordinate", f"{_col(j)}{i}")))
                rows.append(vals)
            tables.append((sheet.title, rows))
        return tables
    except IntakeError:
        raise
    except Exception as exc:
        raise IntakeError("IMPORT_XLSX_ROWS_INVALID") from exc
    finally:
        book.close()


def _label_key(value: object) -> str:
    return re.sub(r"\s+", " ", str(value or "").strip()).casefold()


def _registry_index():
    exact_ids = {}
    labels: dict[str, list[str]] = defaultdict(list)
    fields = WORKBOOK.all_fields()
    for spec in fields:
        exact_ids[spec.field_id] = spec
        labels[_label_key(spec.label)].append(spec.field_id)
    return exact_ids, labels


def _numeric(value: object) -> Decimal:
    if isinstance(value, bool) or value is None:
        raise IntakeError("IMPORT_VALUE_INVALID")
    if isinstance(value, (int, float, Decimal)):
        try:
            n = Decimal(str(value))
        except InvalidOperation as exc:
            raise IntakeError("IMPORT_VALUE_INVALID") from exc
    else:
        text = str(value).strip().replace("\u00a0", "")
        if text.startswith("="):
            raise IntakeError("IMPORT_FORMULA_NOT_AUTHORITATIVE")
        if text.endswith("%"):
            text = text[:-1].strip()
        # No locale guessing: comma-only decimals and grouped numerals
        # remain unresolved until the user supplies an unambiguous value.
        if "," in text or " " in text:
            raise IntakeError("IMPORT_NUMBER_FORMAT_AMBIGUOUS")
        try:
            n = Decimal(text)
        except InvalidOperation as exc:
            raise IntakeError("IMPORT_VALUE_INVALID") from exc
    if not n.is_finite():
        raise IntakeError("IMPORT_VALUE_NON_FINITE")
    return n


def normalize_value(field_id: str, value: object, source_unit: str | None) -> str:
    """Strict explicit unit conversion; never infer percentage scale."""
    spec = WORKBOOK.field(field_id)
    unit = (source_unit or "").strip().casefold().replace("€", "eur")
    target = (spec.unit or "").strip().casefold()
    if isinstance(value, str) and value.startswith("="):
        raise IntakeError("IMPORT_FORMULA_NOT_AUTHORITATIVE")
    if spec.field_type == FieldType.DATE:
        if unit and unit not in ("date", "iso", "iso8601"):
            raise IntakeError("IMPORT_UNIT_INCOMPATIBLE")
        if isinstance(value, datetime):
            return value.date().isoformat()
        if isinstance(value, date):
            return value.isoformat()
        try:
            return date.fromisoformat(str(value).strip()).isoformat()
        except ValueError as exc:
            raise IntakeError("IMPORT_DATE_INVALID") from exc
    if spec.field_type == FieldType.BOOL:
        if isinstance(value, bool):
            return "true" if value else "false"
        v = str(value).strip().lower()
        if v in ("true", "false"):
            return v
        raise IntakeError("IMPORT_BOOL_AMBIGUOUS")
    if spec.field_type in (FieldType.TEXT, FieldType.SELECT):
        if unit not in ("", "n/a"):
            raise IntakeError("IMPORT_UNIT_INCOMPATIBLE")
        return str(value).strip()

    amount = _numeric(value)
    has_suffix_pct = isinstance(value, str) and value.strip().endswith("%")
    if spec.field_type == FieldType.PCT:
        if has_suffix_pct and unit in ("fraction", "decimal", "ratio"):
            raise IntakeError("IMPORT_PERCENT_SCALE_CONFLICT")
        if unit in ("fraction", "decimal", "ratio"):
            amount *= 100
        elif unit in ("%", "pct", "percent", "percentage") or has_suffix_pct:
        else:
            raise IntakeError("IMPORT_PERCENT_SCALE_AMBIGUOUS")
    elif spec.field_type == FieldType.KEUR:
        if unit in ("eur", "euro", "euros", "eur/yr"):
            amount /= 1000
        elif unit in ("keur", "k eur", "keur/yr", "thousand eur"):
            pass
        elif not unit:
            raise IntakeError("IMPORT_CURRENCY_UNIT_REQUIRED")
        else:
            raise IntakeError("IMPORT_UNIT_INCOMPATIBLE")
    else:
        # Numeric units MUST match the registry field's economic quantity;
        # the old permissive union accepted e.g. 'years' for a rate.
        numeric_aliases = {
            FieldType.MW: {"mw"},
            FieldType.MWH: {"h", "h/yr", "mwh"},
            FieldType.YEARS: {"year", "years"},
            FieldType.MONTHS: {"month", "months"},
            FieldType.INT: {""},
            FieldType.FLOAT: {target} if target else {"", "ratio"},
        }
        accepted = numeric_aliases.get(spec.field_type, {target} if target else {""})
        if unit and unit not in accepted:
            raise IntakeError("IMPORT_UNIT_INCOMPATIBLE")
    # Strict integer enforcement is delegated to canonical Workbook validator.
    return format(amount, "f")


def _tabular_candidates(tables: list[tuple[str, list]], *, start: float):
    """Read value-bearing cells from a label/value/unit (or two-column) table."""
    extracted = []
    for sheet_name, rows in tables:
        if len(sheet_name) > 120:
            raise IntakeError("IMPORT_SHEET_NAME_TOO_LONG")
        header = None
        for r, cells in enumerate(rows, 1):
            _check_time(start)
            if not cells or all(v is None or str(v).strip() == "" for v, _, _ in cells):
                continue
            vals = [_label_key(x[0]) for x in cells]
            if "assumption" in vals and "value" in vals:
                header = (vals.index("assumption"), vals.index("value"),
                          vals.index("unit") if "unit" in vals else None)
                continue
            if "field" in vals and "value" in vals:
                header = (vals.index("field"), vals.index("value"),
                          vals.index("unit") if "unit" in vals else None)
                continue
            if "label" in vals and "value" in vals:
                header = (vals.index("label"), vals.index("value"),
                          vals.index("unit") if "unit" in vals else None)
                continue
            li, vi, ui = header if header is not None else (0, 1, 2 if len(cells) > 2 else None)
            if max(li, vi) >= len(cells):
                continue
            label, _, label_cell = cells[li]
            value, fmt_unit, value_cell = cells[vi]
            if label is None or value is None or not str(label).strip() or not str(value).strip():
                continue
            explicit_unit = None
            if ui is not None and ui < len(cells):
                explicit_unit = str(cells[ui][0]).strip() if cells[ui][0] is not None else None
            extracted.append({
                "sheet": sheet_name, "source_cell": value_cell,
                "source_label_cell": label_cell,
                "label": str(label).strip()[:300],
                "value": str(value)[:300] if not isinstance(value, (date, datetime)) else value.isoformat(),
                "unit": explicit_unit or fmt_unit,
                "formula": bool(isinstance(value, str) and value.startswith("=")),
                "raw": value if isinstance(value, (int, float, str, bool)) or value is None else value.isoformat(),
            })
            if len(extracted) > MAX_CELLS:
                raise IntakeError("IMPORT_TABLE_LIMIT_EXCEEDED")
    return extracted


def extract_proposals(data: bytes, filename: str, *,
                      project_type: str, template_source: str,
                      encoding: str = "utf-8-sig") -> dict:
    """Never calls model engine, any persistence writer or external provider."""
    start = time.monotonic()
    name = _safe_filename(filename)
    if not isinstance(data, bytes) or not data or len(data) > MAX_UPLOAD_BYTES:
        raise IntakeError("IMPORT_FILE_SIZE_INVALID")
    ext = os.path.splitext(name)[1].lower()
    if ext == ".csv":
        tables = _csv_rows(data, encoding=encoding, start=start)
    elif ext == ".xlsx":
        tables = _xlsx_rows(data, start=start)
    else:
        raise IntakeError("IMPORT_FORMAT_UNSUPPORTED")

    cells = _tabular_candidates(tables, start=start)
    specs, names = _registry_index()
    proposals = []
    for index, source in enumerate(cells):
        label = source["label"]
        method = "unmapped"
        field_id = None
        hits = names.get(_label_key(label), [])
        if label in specs:
            field_id = label
            method = "exact_field_id"
        elif len(hits) == 1:
            field_id = hits[0]
            method = "unique_exact_registry_label"
        elif len(hits) > 1:
            method = "ambiguous_exact_registry_label"
        normalized = None
        reason = None
        status = "unsupported" if field_id is None else "needs_review"
        if source["formula"]:
            status = "invalid"
            reason = "IMPORT_FORMULA_NOT_AUTHORITATIVE"
        elif field_id:
            try:
                _assert_batch_field_applicable(
                    field_id, project_type=project_type,
                    template_source=template_source,
                )
                normalized = normalize_value(field_id, source["raw"], source["unit"])
                val = WorkbookUpdateService.validate_field_update(field_id, normalized)
                if not val.is_valid:
                    status, reason = "invalid", str(val.error)
                else:
                    status = "ready" if method == "exact_field_id" else "needs_review"
            except (IntakeError, BatchApplyError, ValueError, KeyError) as exc:
                status = "invalid" if not isinstance(exc, BatchApplyError) else "unsupported"
                reason = getattr(exc, "code", str(exc))
        if field_id is None and reason is None:
            reason = ("IMPORT_MAPPING_AMBIGUOUS" if hits else "IMPORT_UNSUPPORTED_LABEL")
        proposals.append({
            "id": index, "sheet": source["sheet"],
            "source_cell": source["source_cell"],
            "label_cell": source["source_label_cell"],
            "label": label, "original_value": source["value"],
            "source_unit": source["unit"], "field_id": field_id,
            "register_path": register_path_for_field(field_id) if field_id else None,
            "normalized_value": normalized, "mapping_method": method,
            "confidence": "exact" if method == "exact_field_id" else (
                "proposed" if method == "unique_exact_registry_label" else "unresolved"),
            "status": status, "reason": reason,
        })
    counts = Counter(p["field_id"] for p in proposals if p["field_id"])
    for p in proposals:
        if p["field_id"] and counts[p["field_id"]] > 1:
            p["status"] = "needs_review"
            p["reason"] = "IMPORT_DUPLICATE_TARGET_REQUIRES_REVIEW"
    return {
        "filename": name, "format": ext[1:], "source_sheets": len(tables),
        "detected": len(proposals), "proposals": proposals,
        "totals": dict(Counter(p["status"] for p in proposals)),
    }
