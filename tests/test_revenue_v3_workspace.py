"""Reuse the real Revenue writer, authorization, scenario CAS and immutable export."""

from tests.model_mutation_csrf import post_valid_mutation

from tests.run_clone_csrf import post_valid_csrf
import copy
import json
from io import BytesIO

import pytest

from tests.test_model_decision_workspace_v2 import env, tokens
from tests.test_wf07_revenue_workspace import contract_tokens
from tests.test_revenue_v3_products import candidate
from tests.revenue_v3_cases import book


@pytest.mark.parametrize("kind", ["solar", "wind"])
def test_candidate_save_scenario_stale_run_history_and_product_export(env, candidate, monkeypatch, kind):
    from app.persistence.workspace_repository import get_workspace_state
    from app.persistence.scenarios_repository import list_scenarios, get_scenario
    from app.persistence.run_history_repository import get_run_history
    from app.workbook.service import WorkbookService
    from app.workbook.revenue_multistream import SNAPSHOT_KEY
    from app.services.export_service import resolve_canonical_last_run_from_workspace
    from app.services.v2_export_service import build_canonical_last_run_institutional_workbook_export
    record = env.create(kind)
    project = {"project": record.project_code}
    def ws():
        return get_workspace_state("decision-user", record.project_id)
    def page():
        response = env.client.get("/v2/workbook", params=project)
        assert response.status_code == 200
        return response.text
    def save(payload, **extra):
        return env.client.post("/v2/workbook/revenue/contracts", headers={"HX-Request": "true"},
            data={**project, **contract_tokens(page()), "contracts": json.dumps(payload), **extra})
    def run():
        response = post_valid_csrf(env.client, "/v2/workbook/run", headers={"HX-Request": "true"}, data={**project, **tokens(page())})
        assert response.status_code == 200, response.text[:2000]
        assert 'data-testid="overview-status-current"' in response.text
        assert not ws().dirty
        assert 'data-testid="revenue-products-audit"' in response.text
        return copy.deepcopy(ws().last_runtime_summary)
    raw, axis = book(WorkbookService.build_draft_input_set_from_workspace(ws()).to_projectinputs(), kind)
    html = page()
    assert 'data-product-add' in html and axis[0].start_date.isoformat() in html
    assert save(raw).status_code == 200
    assert ws().dirty
    assert json.loads(ws().draft_snapshot[SNAPSHOT_KEY]) == raw
    baseline = run()
    assert baseline["revenue_multistream_audit"][0]["products"]
    assert post_valid_mutation(env.client, "/v2/workbook/scenarios/create", data={**project, "scenario_name": "Service contract alternative"}).status_code == 200
    sc = next(s for s in list_scenarios("decision-user", record.project_id) if not s.is_base_case)
    changed = copy.deepcopy(raw)
    changed["products"][0]["price_eur"] = 25
    draft_before = copy.deepcopy(ws().draft_snapshot)
    assert save(changed).status_code == 200
    assert ws().draft_snapshot == draft_before
    assert json.loads(get_scenario(sc.scenario_id, "decision-user").overrides[SNAPSHOT_KEY]) == changed
    result = run()
    assert result["total_revenue_keur"] > baseline["total_revenue_keur"]
    history = get_run_history("decision-user", record.project_id)
    later = copy.deepcopy(changed)
    later["products"][0]["price_eur"] = 30
    assert save(later).status_code == 200
    assert ws().dirty
    assert get_run_history("decision-user", record.project_id) == history
    authority = resolve_canonical_last_run_from_workspace(record, "decision-user", ws())
    assert authority.working_changed_since_run
    assert json.loads(authority.project_inputs.revenue.multistream_config_json) == changed
    def forbidden(*args, **kwargs):
        raise AssertionError("GET / export must not calculate finances")
    with monkeypatch.context() as patch:
        patch.setattr("app.services.production_financial_authority.run_clean_production", forbidden)
        patch.setattr("app.runtime.model_execution.run_model_process", forbidden)
        assert 'data-testid="revenue-products-audit"' in page()
        exported = build_canonical_last_run_institutional_workbook_export("generic_" + kind + "_reference",
            safe_project=record.project_code, project_record=record, user_id="decision-user")
        assert exported.status_code == 200, exported.error_content
        from openpyxl import load_workbook
        workbook = load_workbook(BytesIO(exported.bytes_data), read_only=True, data_only=True)
        try:
            rows = list(workbook["Revenue_Products"].iter_rows(min_row=8, values_only=True))
            assert any(r[2] == "capacity_market" and r[6] == 25 for r in rows)
            assert not any(r[2] == "capacity_market" and r[6] == 30 for r in rows)
            assert {r[5] for r in rows} == {"MW", "CERTIFICATE_MWH"}
        finally:
            workbook.close()


def test_closed_release_route_and_cas_csrf_owner(env, monkeypatch):
    from app.persistence.workspace_repository import get_workspace_state
    from app.workbook.service import WorkbookService
    from domain.revenue import advanced_products
    record = env.create()
    project = {"project": record.project_code}
    html = env.client.get("/v2/workbook", params=project).text
    initial = get_workspace_state("decision-user", record.project_id)
    raw, _ = book(WorkbookService.build_draft_input_set_from_workspace(initial).to_projectinputs())
    data = {**project, **contract_tokens(html), "contracts": json.dumps(raw)}
    assert 'data-product-add' not in html
    assert 'data-testid="revenue-v3-release-closed"' in html
    response = env.client.post("/v2/workbook/revenue/contracts", data=data)
    assert response.status_code == 422 and "RELEASE_CLOSED" in response.text
    assert get_workspace_state("decision-user", record.project_id) == initial
    monkeypatch.setattr(advanced_products, "ACTIVATION_ENABLED", True)
    for token in ("", "forged"):
        assert env.client.post("/v2/workbook/revenue/contracts", data={**data, "csrf_token": token}).status_code == 403
        assert get_workspace_state("decision-user", record.project_id) == initial
    assert env.client.post("/v2/workbook/revenue/contracts", data=data).status_code == 200
    saved = get_workspace_state("decision-user", record.project_id)
    assert env.client.post("/v2/workbook/revenue/contracts", data=data).status_code == 409
    assert get_workspace_state("decision-user", record.project_id) == saved
    from app.auth import COOKIE_NAME, create_session_token
    env.client.cookies.set(COOKIE_NAME, create_session_token(user_id="different-service-owner", username="different"))
    assert env.client.post("/v2/workbook/revenue/contracts", data=data).status_code == 404
    assert get_workspace_state("decision-user", record.project_id) == saved
