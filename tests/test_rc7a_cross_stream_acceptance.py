"""Combined RC-7A evidence: monthly funding remains bound across read-only consumers."""
from __future__ import annotations

import io
from unittest.mock import patch

import openpyxl
import pytest

from tests.test_model_decision_workspace_v2 import env  # noqa: F401
from tests.test_model_credit_decision_v3_real_runs import cenv, credit
from tests.test_model_rc7_capex_profile_run import (
    OWNER, _base_leaves, _hash, _latest, _row, _run, _save, _ws,
)


def test_profiled_run_credit_and_history_export_share_immutable_evidence(cenv):
    from app.persistence.run_history_repository import get_run_history
    from app.services.export_service import resolve_canonical_last_run_from_workspace
    from app.services.v2_export_service import (
        build_canonical_last_run_institutional_workbook_export,
        build_historical_run_workbook_export,
    )

    project = cenv.create("solar")
    assert _run(cenv, project).status_code == 200
    baseline = _latest(project)
    baseline_data = baseline.to_dict()
    baseline_export = build_historical_run_workbook_export(
        project_record=project, user_id=OWNER, history_id=baseline.history_id,
    )
    assert baseline_export.status_code == 200

    _save(project, [_row(leaf.item_key, "FRONT_LOADED") for leaf in _base_leaves(project)])
    assert _run(cenv, project).status_code == 200
    current = _latest(project)
    current_data = current.to_dict()
    construction = current.runtime_summary["financing_evidence"]["construction_funding"]
    assert len(construction["periods"]) == 14
    before_hash = _hash(project)

    def no_engine(*args, **kwargs):
        raise AssertionError("read-only evidence consumers must not recalculate")

    with patch("app.api.project_runner.run_project", side_effect=no_engine), \
         patch("app.services.production_financial_authority.run_clean_production", side_effect=no_engine), \
         patch("app.services.production_waterfall_seam.execute_production_waterfall", side_effect=no_engine):
        _, data = credit(cenv, project, baseline.history_id, current.history_id)
        assert {row["history_id"] for row in data["credit_v4"]["run_risks"]} == {
            baseline.history_id, current.history_id,
        }
        assert "Contractual covenant/default" in {row["metric"] for row in data["credit_v4"]["unavailable"]}

        # The old Last Run exporter cannot safely rebuild monthly construction inputs.
        with pytest.raises(ValueError, match="CAPEX_PROFILE_UNBOUND_SURFACE"):
            resolve_canonical_last_run_from_workspace(project, OWNER, _ws(project))
        refused = build_canonical_last_run_institutional_workbook_export(
            project.template_source, safe_project=project.project_code,
            project_record=project, user_id=OWNER,
        )
        assert refused.status_code == 400

        exported = build_historical_run_workbook_export(
            project_record=project, user_id=OWNER, history_id=current.history_id,
        )
        assert exported.status_code == 200
        book = openpyxl.load_workbook(io.BytesIO(exported.bytes_data))
        rows = list(book["Construction_Funding"].values)
        for label, key in (
            ("Total audit uses", "total_audit_uses_keur"),
            ("Total audit sources", "total_audit_sources_keur"),
            ("Total audit residual", "total_audit_residual_keur"),
        ):
            row = next(row for row in rows if row[0] == label)
            assert row[1] == pytest.approx(construction[key], abs=1e-8)
        old_export = build_historical_run_workbook_export(
            project_record=project, user_id=OWNER, history_id=baseline.history_id,
        )
        assert baseline_export.metadata["export_run_state"] == "CURRENT"
        assert old_export.metadata["export_run_state"] == "HISTORICAL"
        assert old_export.metadata["export_evidence_digest"] == baseline_export.metadata["export_evidence_digest"]
        before_book = openpyxl.load_workbook(io.BytesIO(baseline_export.bytes_data))
        after_book = openpyxl.load_workbook(io.BytesIO(old_export.bytes_data))
        for sheet in ("Sources_Uses", "Construction_Funding", "Debt_DSCR", "Tax_Bridge"):
            def numeric_rows(book):
                return [row for row in book[sheet].values
                        if any(isinstance(value, (int, float)) and not isinstance(value, bool) for value in row)]
            assert numeric_rows(before_book) == numeric_rows(after_book)
        repeated = build_historical_run_workbook_export(
            project_record=project, user_id=OWNER, history_id=baseline.history_id,
        )
        assert repeated.bytes_data == old_export.bytes_data

    assert _hash(project) == before_hash
    entries = {entry.history_id: entry.to_dict() for entry in get_run_history(OWNER, project.project_id)}
    assert entries == {baseline.history_id: baseline_data, current.history_id: current_data}
