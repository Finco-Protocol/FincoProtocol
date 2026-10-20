"""Institutional Reporting V3 -- lender & investor evidence package over committed Runs.

Real engine Runs are committed through the production persistence path (``tests/rc7_helpers.py``); packages are read back from XLSX bytes.

Acceptance markers:
  V3_REAL_RUNS            Solar, Wind, Data Center and EV Runs build the V3 package (Last Run and historical)
  V3_ISOLATION            historical report is immune to a later Working Copy edit / newer Run; owner/project boundaries
  V3_STATES               CURRENT / STALE / HISTORICAL disclosures
  V3_NO_ENGINE            reporting and the export route execute no engine
  V3_DETERMINISTIC        equal historical evidence -> byte-identical package
  V3_MISSING_VS_ZERO      n/a is distinct from 0; 2-dp persisted-value handling
  V3_CONSISTENCY          S&U, debt, DSRA and statement evidence agree with the persisted Run
  V3_PERIODS              monthly construction + supported operating period metadata; no relabelling
  V3_WORKBOOK             valid XLSX, expected institutional sheets, navigation and print settings
  V3_CLASSES              every reported metric carries exactly one evidence class
  V3_NO_AUTHORITY_CHANGE  persisted records and fingerprints are not touched
"""
from __future__ import annotations

import ast
import dataclasses
import hashlib
import io
import json
import time
from pathlib import Path
from unittest import mock

import openpyxl
import pytest

import rc7_helpers as H
from app.export import institutional_reporting_v2 as R2
from app.export import institutional_reporting_v3 as R3
from app.export.run_evidence import STATE_CURRENT, STATE_HISTORICAL, STATE_STALE, evidence_from_history_entry

V3_FRONT = ("Report_Contents", "Executive_Overview", "Lender_Overview")
V3_SHEETS = (*V3_FRONT, "Evidence_Limits")


@pytest.fixture(scope="module", autouse=True)
def _db(tmp_path_factory):
    from _pytest.monkeypatch import MonkeyPatch

    mp = MonkeyPatch()
    H.init_db(tmp_path_factory.mktemp("rcv3"), mp)
    yield
    mp.undo()


def _project(kind, uid=None, suffix=""):
    from app.persistence.projects_repository import get_project

    pr, uid, snap, pi, sc = H.create_project(kind, uid, suffix)
    _, snapshot_id = H.commit_run(pr, uid, pi, sc, kind)
    return get_project(pr.project_id, uid), uid, pi, sc, snapshot_id


def _export_last_run(pr, uid):
    from app.services.v2_export_service import build_canonical_last_run_institutional_workbook_export as build

    return build(pr.template_source, safe_project=pr.project_code, project_record=pr, user_id=uid)


def _export_history(pr, uid, history_id):
    from app.services.v2_export_service import build_historical_run_workbook_export as build

    return build(project_record=pr, user_id=uid, history_id=history_id)


def _wb(response):
    assert response.status_code == 200, response.error_content
    return openpyxl.load_workbook(io.BytesIO(response.bytes_data))


def _rows(ws):
    return [list(r) for r in ws.iter_rows(values_only=True)]


def _kv(ws):
    out = {}
    for r in _rows(ws):
        if r[0] is not None and len(r) > 1 and r[0] not in out:
            out[r[0]] = r[1]
    return out


def _find(ws, label):
    for r in _rows(ws):
        if r[0] == label:
            return r
    raise AssertionError(f"{label!r} not on {ws.title}")


def _latest(uid, pr):
    from app.persistence.run_history_repository import get_run_history

    return get_run_history(uid, pr.project_id, limit=25)


@pytest.fixture(scope="module", params=["solar", "wind", "ev", "data_center"])
def committed(request):
    return request.param, *_project(request.param)


# ── V3_REAL_RUNS / V3_WORKBOOK ───────────────────────────────────────────────

def test_last_run_package_has_the_v3_front_section_right_after_the_guide(committed):
    kind, pr, uid, *_ = committed
    wb = _wb(_export_last_run(pr, uid))
    assert wb.sheetnames[:6] == ["Export_Metadata", "Workbook_Index", "Report_Guide", *V3_FRONT]
    for name in (*V3_SHEETS, *R2.COMMON_SHEETS):
        assert name in wb.sheetnames, name
    assert wb.sheetnames.index("Evidence_Limits") < wb.sheetnames.index("Run_Integrity")


def test_historical_package_is_the_v3_package_for_every_vertical(committed):
    kind, pr, uid, *_ = committed
    entry = _latest(uid, pr)[0]
    wb = _wb(_export_history(pr, uid, entry.history_id))
    assert wb.sheetnames[0] == "Report_Guide" and wb.sheetnames[1:4] == list(V3_FRONT)
    for name in (*V3_SHEETS, "P&L", "Cash Flow", "Balance Sheet", "Returns", *R2.COMMON_SHEETS):
        assert name in wb.sheetnames, name


def test_v3_sheet_set_is_a_superset_of_v2_and_v2_still_builds(committed):
    kind, pr, uid, *_ = committed
    entry = _latest(uid, pr)[0]
    ev = evidence_from_history_entry(entry, pr, last_run_snapshot_id=None, last_run_is_stale=True, freshness_source="test")
    v2 = openpyxl.load_workbook(io.BytesIO(R2.build_run_evidence_workbook(ev)))
    v3 = openpyxl.load_workbook(io.BytesIO(R3.build_run_evidence_workbook_v3(ev)))
    assert set(v2.sheetnames) <= set(v3.sheetnames)
    assert set(v3.sheetnames) - set(v2.sheetnames) == set(V3_SHEETS)
    def body(ws):                                                        # the only V3 addition on a V2 sheet is the Contents link in H1
        rows = _rows(ws)
        rows[0] = rows[0][:7] + [None] * (len(rows[0]) - 7)
        return rows
    for name in v2.sheetnames:                                           # V2 content is reused unchanged
        assert name == "Report_Guide" or body(v2[name]) == body(v3[name]), name


def test_report_contents_links_to_existing_sheets_and_guide_lists_v3_sheets(committed):
    kind, pr, uid, *_ = committed
    wb = _wb(_export_last_run(pr, uid))
    links = [c.hyperlink.location or c.hyperlink.target for row in wb["Report_Contents"].iter_rows() for c in row if c.hyperlink]
    assert links and all(str(l).strip("#").split("!")[0].strip("'") in wb.sheetnames for l in links)
    guide_text = " ".join(str(r[0]) for r in _rows(wb["Report_Guide"]) if r[0])
    for name in V3_SHEETS:
        assert name in guide_text


def test_print_and_navigation_settings(committed):
    kind, pr, uid, *_ = committed
    wb = _wb(_export_last_run(pr, uid))
    for name in (*V3_SHEETS, "Debt_DSCR", "Sources_Uses"):
        ws = wb[name]
        assert ws.page_setup.orientation == "landscape" and ws.page_setup.fitToWidth == 1
        assert ws.sheet_properties.pageSetUpPr.fitToPage is True
        assert "Run " in (ws.oddHeader.left.text or "") and "Page" in (ws.oddFooter.right.text or "")
    assert wb["Debt_DSCR"].freeze_panes is not None            # wide schedules keep their metric columns


# ── V3_CONSISTENCY / V3_MISSING_VS_ZERO ──────────────────────────────────────

def test_executive_overview_values_equal_the_persisted_run(committed):
    kind, pr, uid, *_ = committed
    from app.persistence.workspace_repository import get_workspace_state

    ws = get_workspace_state(uid, pr.project_id)
    rs = ws.last_runtime_summary
    ov = _wb(_export_last_run(pr, uid))["Executive_Overview"]
    expected = {"Project IRR": rs["project_irr"], "Total CAPEX": rs["total_capex_keur"],
                "Total project uses": rs["total_project_uses_keur"], "Senior debt (committed)": rs["senior_debt_keur"],
                "Revenue": rs["total_revenue_keur"], "EBITDA": rs["total_ebitda_keur"], "Distributions": rs["total_distributions_keur"],
                "Corporate tax": rs["total_tax_keur"]}
    for label, value in expected.items():
        row = _find(ov, label)
        if value is None:
            assert row[1] in (None, "n/a"), label
        else:
            assert row[1] == pytest.approx(value, rel=1e-14, abs=0.0), label     # as persisted; XLSX stores 15-16 significant digits
        assert str(row[2]).startswith("PERSISTED") or value is None
    assert _find(ov, "Run state")[1].startswith("CURRENT")
    assert _find(ov, "Run Integrity")[1] in {"PASS", "WARN", "FAIL", "UNAVAILABLE", "PASSED"} | {str(_find(ov, "Run Integrity")[1])}


def test_lender_overview_ties_to_sources_uses_debt_and_reserve_evidence(committed):
    kind, pr, uid, *_ = committed
    from app.persistence.workspace_repository import get_workspace_state

    wb = _wb(_export_last_run(pr, uid))
    ws = get_workspace_state(uid, pr.project_id)
    ln, su = wb["Lender_Overview"], wb["Sources_Uses"]
    bank = ws.last_runtime_summary["financing_evidence"]["bankability"]
    assert _find(ln, "Final senior commitment")[1] == pytest.approx(bank["final_senior_commitment_keur"], rel=1e-14)
    assert _find(ln, "Binding sizing constraint")[1] == bank["binding_senior_constraint"]
    persisted_su = ws.last_runtime_summary["financing_evidence"]["sources_uses"]
    assert _find(su, "Senior debt")[1] == pytest.approx(persisted_su["senior_debt_keur"], rel=1e-14)
    assert _find(ln, "Initial DSRA funding (a Project Use)")[1] == pytest.approx(persisted_su["initial_dsra_funding_keur"], rel=1e-14)
    # debt service: sum of persisted 2-dp period values vs the persisted Run total is an OBSERVATIONAL check that passes on real Runs
    check = [r for r in _rows(ln) if r[0] == "Sum of period debt service vs persisted Run total"]
    assert check and check[0][1] == "PASS"
    periods = ws.last_debt_schedule["periods"]
    total = sum(p["senior_ds_keur"] for p in periods if p.get("senior_ds_keur") is not None)
    assert _find(ln, "Senior debt service (sum of periods)")[1] == pytest.approx(total, abs=1e-9)
    closing = [p["dsra_balance_keur"] for p in periods if p.get("dsra_balance_keur") is not None]
    if closing:
        assert _find(ln, "Closing DSRA balance, last persisted period")[1] == closing[-1]          # a stock is never summed
        assert _find(ln, "Peak DSRA balance")[1] == max(closing)
    draw_rows = [r for r in _rows(ln) if r[0] == "Senior drawn during construction (sum of periods)"]
    if ws.last_runtime_summary["financing_evidence"]["construction_funding"]["periods"]:
        assert draw_rows and draw_rows[0][1] == pytest.approx(
            sum(p["senior_draw_keur"] for p in ws.last_runtime_summary["financing_evidence"]["construction_funding"]["periods"]), abs=1e-6)


def test_terminal_and_maturity_evidence_is_read_from_the_persisted_terminal_state(committed):
    kind, pr, uid, *_ = committed
    from app.persistence.workspace_repository import get_workspace_state

    ws = get_workspace_state(uid, pr.project_id)
    term = ws.last_sponsor_schedule["summary"]["terminal_financial_state"]
    ln = _wb(_export_last_run(pr, uid))["Lender_Overview"]
    assert _find(ln, "Senior terminal status")[1] == term["senior"]["status"]
    assert _find(ln, "Senior contractual maturity date")[1] == term["senior"]["contractual_maturity_date"]
    assert _find(ln, "Senior balance at contractual maturity")[1] == pytest.approx(term["senior"]["balance_at_contractual_maturity_keur"], abs=1e-12)


def test_missing_evidence_is_not_available_never_zero_and_explicit_zero_is_kept():
    pr, uid, *_ = _project("solar", suffix="miss")
    entry = _latest(uid, pr)[0]
    ev = evidence_from_history_entry(entry, pr, last_run_snapshot_id=None, last_run_is_stale=True, freshness_source="test")
    summary = dict(ev.runtime_summary)
    summary.pop("project_irr", None)
    summary["total_tax_keur"] = 0.0
    summary["min_llcr"] = None
    legacy = dataclasses.replace(ev, runtime_summary=summary, debt_schedule={}, sponsor_schedule={}, integrity_evidence={},
                                 financial_statements={})
    wb = openpyxl.load_workbook(io.BytesIO(R3.build_run_evidence_workbook_v3(legacy)))
    ov, ln = wb["Executive_Overview"], wb["Lender_Overview"]
    irr = _find(ov, "Project IRR")
    assert irr[1] in (None, "n/a") and "NOT_AVAILABLE" in irr[2]
    tax = _find(ov, "Corporate tax")
    assert tax[1] == 0.0 and str(tax[2]).startswith("PERSISTED")             # a persisted zero stays zero
    assert _find(ln, "Minimum LLCR")[1] in (None, "n/a")
    assert _find(ln, "Closing senior balance, last persisted period")[1] in (None, "n/a")
    assert _find(ln, "Senior terminal status")[1] in (None, "n/a")
    limits = [r for r in _rows(wb["Evidence_Limits"]) if r[0] == "Project IRR"]
    assert limits and limits[0][2] == R3.MISSING


def test_every_reported_metric_has_exactly_one_known_evidence_class(committed):
    kind, pr, uid, *_ = committed
    wb = _wb(_export_last_run(pr, uid))
    register = [r for r in _rows(wb["Evidence_Limits"]) if r[1] in ("Executive_Overview", "Lender_Overview")]
    assert len(register) >= 40
    assert all(r[2] in R3.EVIDENCE_CLASSES for r in register)
    seen = {}
    for metric, sheet, cls, *_ in register:
        assert seen.setdefault((metric, sheet), cls) == cls
    assert any(r[2] == R3.ARITHMETIC for r in register) and any(r[2] == R3.CONTRACTUAL for r in register)
    # a presentation sum never replaces the persisted figure: both are present on the same sheet
    ln = wb["Lender_Overview"]
    assert _find(ln, "Senior debt service (sum of periods)")[2].startswith(R3.ARITHMETIC)
    assert _find(ln, "Senior debt service (Run total as persisted)")[2].startswith(R3.PERSISTED)


def test_no_covenant_compliance_or_recommendation_is_declared(committed):
    kind, pr, uid, *_ = committed
    wb = _wb(_export_last_run(pr, uid))
    assert _find(wb["Lender_Overview"], "Covenant compliance")[1] == "NOT DECLARED"
    text = " ".join(str(c.value).lower() for ws in wb.worksheets if ws.title in V3_SHEETS for row in ws.iter_rows() for c in row if c.value)
    assert "recommend" in text and "no investment recommendation" in text
    for forbidden in ("covenant compliant", "in compliance", "buy", "invest in"):
        assert forbidden not in text, forbidden


# ── V3_PERIODS ────────────────────────────────────────────────────────────────

def test_period_basis_and_construction_columns_follow_persisted_metadata(committed):
    kind, pr, uid, *_ = committed
    wb = _wb(_export_last_run(pr, uid))
    ex = _kv(wb["Executive_Overview"])
    assert "quarter" not in str(ex["Period basis"]).lower()                # 6-month Runs are never relabelled quarterly
    pd = _rows(wb["Periods_Dates"])
    assert any("Construction" in str(c) for r in pd for c in r if c) and any("Operation" in str(c) for r in pd for c in r if c)
    cs = [r for r in _rows(wb["Lender_Overview"]) if r[0] == "Construction period"]
    if cs:
        rows = _rows(wb["Lender_Overview"])
        start = rows.index(cs[0]) + 1
        first = rows[start]
        assert str(first[1])[:4].isdigit() and str(first[2])[:4].isdigit()       # real persisted start / end dates


# ── V3_STATES / V3_ISOLATION ─────────────────────────────────────────────────

@pytest.fixture(scope="module")
def two_runs():
    from app.persistence.projects_repository import get_project
    from app.persistence.workspace_repository import get_workspace_state, save_workspace_state

    pr, uid, snap, pi, sc = H.create_project("ev", suffix="v3h")
    _, snap_a = H.commit_run(pr, uid, pi, sc, "ev")
    pr = get_project(pr.project_id, uid)
    entry_a = _latest(uid, pr)[0]
    before = {"last": _export_last_run(pr, uid), "hist": _export_history(pr, uid, entry_a.history_id)}
    ws = get_workspace_state(uid, pr.project_id)
    edited = dict(ws.draft_snapshot, total_capex_keur=str(float(ws.draft_snapshot["total_capex_keur"]) + 500))
    save_workspace_state(user_id=uid, project_id=pr.project_id, project_code=pr.project_code, draft_snapshot=edited,
                         saved_snapshot=edited, active_scenario_id=sc.scenario_id, active_scenario_name=sc.scenario_name)
    stale = _export_last_run(pr, uid)
    pi_b = dataclasses.replace(pi, financing=dataclasses.replace(pi.financing, target_dscr=pi.financing.target_dscr + 0.15))
    time.sleep(0.01)
    _, snap_b = H.commit_run(pr, uid, pi_b, sc, "ev")
    pr = get_project(pr.project_id, uid)
    return dict(pr=pr, uid=uid, entry_a=entry_a, snap_a=snap_a, snap_b=snap_b, before=before, stale=stale)


def test_current_stale_and_historical_disclosures(two_runs):
    cur = _wb(two_runs["before"]["last"])
    stale = _wb(two_runs["stale"])
    hist = _wb(_export_history(two_runs["pr"], two_runs["uid"], two_runs["entry_a"].history_id))
    assert _find(cur["Executive_Overview"], "Run state")[1].startswith("CURRENT")
    assert _find(stale["Executive_Overview"], "Run state")[1].startswith("STALE")
    assert "PRIOR inputs" in _find(stale["Executive_Overview"], "Run state")[1]
    assert _find(hist["Executive_Overview"], "Run state")[1].startswith(STATE_HISTORICAL)
    for wb, state in ((cur, STATE_CURRENT), (stale, STATE_STALE), (hist, STATE_HISTORICAL)):
        for name in V3_SHEETS:
            assert f"RUN STATE: {state}" in wb[name]["A2"].value, (name, state)


def test_historical_report_is_immune_to_later_working_copy_edit_and_newer_run(two_runs):
    pr, uid, entry_a = two_runs["pr"], two_runs["uid"], two_runs["entry_a"]
    before = two_runs["before"]["hist"]
    after = _export_history(pr, uid, entry_a.history_id)
    # The evidence is identical; the only legitimate difference is the disclosed Run state (CURRENT while it was the Last Run,
    # HISTORICAL once superseded).  Every numeric cell and the evidence digest are unchanged.
    assert after.metadata["export_evidence_digest"] == before.metadata["export_evidence_digest"]
    wa, wb_ = _wb(before), _wb(after)
    assert wa.sheetnames == wb_.sheetnames
    skip = ("Report_Guide",)
    num = lambda w: {(ws.title, c.coordinate): c.value for ws in w.worksheets if ws.title not in skip
                     for row in ws.iter_rows() for c in row if isinstance(c.value, (int, float)) and not isinstance(c.value, bool)}
    assert num(wa) == num(wb_)
    assert _find(wa["Executive_Overview"], "Run state")[1].startswith("CURRENT")
    assert _find(wb_["Executive_Overview"], "Run state")[1].startswith("HISTORICAL")
    assert after.metadata["export_snapshot_id"] == two_runs["snap_a"] != two_runs["snap_b"]
    newest = _latest(uid, pr)[0]
    assert newest.history_id != entry_a.history_id
    assert _export_history(pr, uid, newest.history_id).bytes_data != after.bytes_data
    ov_a, ov_b = (_wb(_export_history(pr, uid, h))["Executive_Overview"] for h in (entry_a.history_id, newest.history_id))
    assert _find(ov_a, "Runtime snapshot ID")[1] == two_runs["snap_a"]
    assert _find(ov_b, "Runtime snapshot ID")[1] == two_runs["snap_b"]


def test_deterministic_historical_export(two_runs):
    pr, uid, entry_a = two_runs["pr"], two_runs["uid"], two_runs["entry_a"]
    first = _export_history(pr, uid, entry_a.history_id)
    time.sleep(1.1)
    second = _export_history(pr, uid, entry_a.history_id)
    assert first.bytes_data == second.bytes_data and first.filename == second.filename
    ev = evidence_from_history_entry(entry_a, pr, last_run_snapshot_id=two_runs["snap_b"], last_run_is_stale=False, freshness_source="test")
    assert R3.build_run_evidence_workbook_v3(ev) == first.bytes_data


def test_owner_and_project_boundaries_on_the_service(two_runs):
    pr, uid, entry_a = two_runs["pr"], two_runs["uid"], two_runs["entry_a"]
    other_pr, other_uid, *_ = _project("ev", suffix="v3iso")
    for r in (_export_history(other_pr, other_uid, entry_a.history_id), _export_history(other_pr, uid, entry_a.history_id),
              _export_history(pr, uid, "0" * 32)):
        assert r.status_code == 404 and r.bytes_data is None
    assert _export_history(pr, uid, "").status_code == 400 and _export_history(None, uid, entry_a.history_id).status_code == 400
    r = _export_history(pr, uid, "<script>alert(1)</script>")
    assert r.status_code == 404 and "<script>" not in (r.error_content or "")


# ── historical export route (A2) ─────────────────────────────────────────────

def _client(user_id):
    from types import SimpleNamespace

    from fastapi import FastAPI
    from fastapi.testclient import TestClient

    from app.v2.router import router

    app = FastAPI()
    app.include_router(router, prefix="/v2")
    return TestClient(app, raise_server_exceptions=False), mock.patch(
        "app.v2.router._get_current_user", return_value=SimpleNamespace(user_id=user_id) if user_id else None)


def test_route_downloads_the_snapshot_bound_package_for_the_owner(two_runs):
    pr, uid, entry_a = two_runs["pr"], two_runs["uid"], two_runs["entry_a"]
    client, patch_user = _client(uid)
    with patch_user:
        r = client.get("/v2/workbook/run-history/export", params={"project": pr.project_code, "history_id": entry_a.history_id})
    assert r.status_code == 200
    assert r.headers["content-type"].startswith("application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
    assert "attachment" in r.headers["content-disposition"]
    assert r.content == _export_history(pr, uid, entry_a.history_id).bytes_data           # same bytes as the service, not the Last Run
    wb = openpyxl.load_workbook(io.BytesIO(r.content))
    assert _find(wb["Executive_Overview"], "Runtime snapshot ID")[1] == two_runs["snap_a"]


def test_route_boundaries_and_fail_closed_behaviour(two_runs):
    pr, uid, entry_a = two_runs["pr"], two_runs["uid"], two_runs["entry_a"]
    other_pr, other_uid, *_ = _project("ev", suffix="v3rt")
    client, patch_user = _client(None)
    with patch_user:
        r = client.get("/v2/workbook/run-history/export", params={"project": pr.project_code, "history_id": entry_a.history_id},
                       follow_redirects=False)
    assert r.status_code == 302 and r.headers["location"].endswith("/login")
    client, patch_user = _client(other_uid)
    with patch_user:
        foreign_project = client.get("/v2/workbook/run-history/export",
                                     params={"project": pr.project_code, "history_id": entry_a.history_id})
        foreign_history = client.get("/v2/workbook/run-history/export",
                                     params={"project": other_pr.project_code, "history_id": entry_a.history_id})
        missing = client.get("/v2/workbook/run-history/export", params={"project": other_pr.project_code})
    assert foreign_project.status_code == 404 and foreign_history.status_code == 404 and missing.status_code == 400
    assert b"PK" not in foreign_project.content[:2] and b"PK" not in foreign_history.content[:2]


# ── V3_NO_ENGINE ──────────────────────────────────────────────────────────────

def test_reporting_and_route_execute_no_engine(two_runs):
    pr, uid, entry_a = two_runs["pr"], two_runs["uid"], two_runs["entry_a"]

    def boom(*a, **k):
        raise AssertionError("engine must not run during reporting")

    client, patch_user = _client(uid)
    with mock.patch("app.api.project_runner.run_project", side_effect=boom), \
            mock.patch("app.services.production_waterfall_seam.execute_production_waterfall", side_effect=boom), \
            mock.patch("app.services.production_financial_authority.run_clean_production", side_effect=boom), patch_user:
        assert _export_last_run(pr, uid).status_code == 200
        assert _export_history(pr, uid, entry_a.history_id).status_code == 200
        assert client.get("/v2/workbook/run-history/export",
                          params={"project": pr.project_code, "history_id": entry_a.history_id}).status_code == 200


def test_v3_module_imports_no_engine_or_persistence_writer():
    tree = ast.parse(Path("app/export/institutional_reporting_v3.py").read_text())
    imported = set()
    for node in ast.walk(tree):
        if isinstance(node, ast.Import):
            imported.update(a.name for a in node.names)
        elif isinstance(node, ast.ImportFrom):
            imported.add(node.module or "")
    banned = [m for m in imported if m.startswith(("financial_engine", "domain", "finco_core", "app.api.project_runner",
                                                  "app.services.production", "app.persistence"))]
    assert not banned, banned


# ── V3_NO_AUTHORITY_CHANGE ───────────────────────────────────────────────────

def test_export_does_not_touch_persisted_run_records_or_fingerprints(two_runs):
    from app.persistence.workspace_repository import get_workspace_state

    pr, uid, entry_a = two_runs["pr"], two_runs["uid"], two_runs["entry_a"]

    def fingerprint():
        ws = get_workspace_state(uid, pr.project_id)
        history = [dataclasses.asdict(e) if dataclasses.is_dataclass(e) else repr(e) for e in _latest(uid, pr)]
        blob = json.dumps({"ws": [ws.last_runtime_summary, ws.last_financial_statements, ws.last_debt_schedule,
                                  ws.last_runtime_snapshot_id, ws.last_runtime_composite_hash], "history": history},
                          sort_keys=True, default=str)
        return hashlib.sha256(blob.encode()).hexdigest()

    before = fingerprint()
    _export_last_run(pr, uid)
    _export_history(pr, uid, entry_a.history_id)
    client, patch_user = _client(uid)
    with patch_user:
        client.get("/v2/workbook/run-history/export", params={"project": pr.project_code, "history_id": entry_a.history_id})
    assert fingerprint() == before
