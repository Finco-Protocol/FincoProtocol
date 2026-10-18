"""RC-7 / Workstream E -- Institutional Reporting V2: run-bound lender / investor workbook.

Real Runs are produced by the real engine and committed through the production persistence path
(``tests/rc7_helpers.py``); exports are then read back from the XLSX bytes.

Acceptance markers:
  RC7_REAL_RUN_EXPORTS            Solar, Wind, EV and Data Center Runs export and reconcile
  RC7_STATEMENT_TOTALS            flow totals equal persisted values; S&U and senior principal reconcile
  RC7_STOCK_VS_FLOW               balances / ratios are never summed
  RC7_MISSING_VS_ZERO             n/a (missing) is distinct from an explicit 0
  RC7_DEBT_BALANCES_AND_SERVICE   roll-forward and DS = principal + interest comparisons
  RC7_HISTORICAL_IMMUTABLE        a later Working Copy edit / later Run never alters an older report
  RC7_REPRODUCIBLE                equivalent source snapshots -> byte-identical workbooks
  RC7_ISOLATION                   owner / project isolation of Run History exports
  RC7_NO_ENGINE                   export performs zero engine execution
  RC7_WORKBOOK_INTEGRITY          archive and workbook open cleanly
  RC7_WIDE_AND_LONG                wide schedules and long text stay usable
  RC7_PERIOD_BASIS                6-month data is never relabelled quarterly; future metadata is read
  RC7_NO_COVENANT_CLAIM           covenant compliance is never declared
"""
from __future__ import annotations

import dataclasses
import io
import math
import shutil
import subprocess
import time
import zipfile
from datetime import date
from unittest import mock

import openpyxl
import pytest

import rc7_helpers as H
from app.export import institutional_reporting_v2 as R
from app.export.reporting_common import (FLOW, RATIO, STOCK, classify_metric, display_total, period_basis)
from app.export.run_evidence import (STATE_CURRENT, STATE_HISTORICAL, STATE_STALE, RunEvidence,
                                     evidence_from_history_entry)

NEW_SHEETS = ("Report_Guide", "Periods_Dates", "Sources_Uses", "Construction_Funding", "Debt_DSCR",
              "Reserves", "Tax_Bridge", "Run_Integrity", "Disclosures")


@pytest.fixture(scope="module", autouse=True)
def _db(tmp_path_factory):
    from _pytest.monkeypatch import MonkeyPatch

    mp = MonkeyPatch()
    H.init_db(tmp_path_factory.mktemp("rc7"), mp)
    yield
    mp.undo()


# ── helpers ──────────────────────────────────────────────────────────────────

def _project(kind, uid=None, suffix=""):
    from app.persistence.projects_repository import get_project

    pr, uid, snap, pi, sc = H.create_project(kind, uid, suffix)
    result, snapshot_id = H.commit_run(pr, uid, pi, sc, kind)
    return get_project(pr.project_id, uid), uid, pi, sc, snapshot_id


def _export_last_run(pr, uid):
    from app.services.v2_export_service import build_canonical_last_run_institutional_workbook_export as build

    return build(pr.template_source, safe_project=pr.project_code, project_record=pr, user_id=uid)


def _export_history(pr, uid, history_id):
    from app.services.v2_export_service import build_historical_run_workbook_export as build

    return build(project_record=pr, user_id=uid, history_id=history_id)


def _wb(response):
    assert response.status_code == 200, response.error_content
    return openpyxl.load_workbook(io.BytesIO(response.bytes_data))


def _rows(ws):
    return [list(r) for r in ws.iter_rows(values_only=True)]


def _find(ws, label, col=1):
    for row in ws.iter_rows():
        if row[col - 1].value == label:
            return [c.value for c in row]
    raise AssertionError(f"{label!r} not found on {ws.title}")


def _kv(ws):
    return {r[0]: r[1] for r in _rows(ws) if r[0] is not None and len(r) > 1}


def _latest(uid, pr):
    from app.persistence.run_history_repository import get_run_history

    return get_run_history(uid, pr.project_id, limit=25)


# ── RC7_REAL_RUN_EXPORTS ─────────────────────────────────────────────────────

@pytest.fixture(scope="module", params=["solar", "wind", "ev", "data_center"])
def committed(request):
    return request.param, *_project(request.param)


def test_real_run_export_contains_run_bound_sheets(committed):
    kind, pr, uid, _pi, _sc, snapshot_id = committed
    from app.persistence.workspace_repository import get_workspace_state

    wb = _wb(_export_last_run(pr, uid))
    ws = get_workspace_state(uid, pr.project_id)
    for name in NEW_SHEETS:
        assert name in wb.sheetnames, name
    # the guide is the first thing read after the metadata/index tabs; existing sheets are retained
    assert wb.sheetnames[:3] == ["Export_Metadata", "Workbook_Index", "Report_Guide"]
    for existing in ("P&L", "Cash Flow", "Balance Sheet", "Tax", "Returns", "Run Identity", "Reconciliation"):
        assert existing in wb.sheetnames
    guide = _kv(wb["Report_Guide"])
    assert guide["Run state"] == STATE_CURRENT
    assert guide["Runtime snapshot ID"] == snapshot_id == ws.last_runtime_snapshot_id
    assert guide["Input composite hash"] == ws.last_runtime_composite_hash
    assert guide["Engine version"] == ws.last_runtime_identity["engine_version"]
    # a real immutable Run UUID is carried when the history ledger holds the Run
    assert guide["Run History ID"] == _latest(uid, pr)[0].history_id
    assert "RUN STATE: CURRENT" in wb["Debt_DSCR"]["A2"].value


def test_sources_uses_reconcile_and_match_persisted_evidence(committed):
    kind, pr, uid, *_ = committed
    from app.persistence.workspace_repository import get_workspace_state

    wb = _wb(_export_last_run(pr, uid))
    su = get_workspace_state(uid, pr.project_id).last_runtime_summary["financing_evidence"]["sources_uses"]
    ws = wb["Sources_Uses"]
    # unrounded, exactly as persisted
    # openpyxl serializes doubles with 16 significant digits (<= 1 ulp); no value is rounded by the export
    assert _find(ws, "TOTAL USES (persisted)")[1] == pytest.approx(su["total_uses_keur"], rel=1e-14, abs=0)
    assert _find(ws, "TOTAL SOURCES (persisted)")[1] == pytest.approx(su["total_sources_keur"], rel=1e-14, abs=0)
    assert _find(ws, "Senior debt")[1] == pytest.approx(su["senior_debt_keur"], rel=1e-14, abs=0)
    results = {r[0]: r[1] for r in _rows(ws) if r[0] in (
        "Sum of listed Uses vs total", "Sum of listed Sources vs total",
        "Persisted Sources - Uses difference", "Total Sources vs Total Uses")}
    assert set(results.values()) == {"PASS"}, results


# ── RC7_STATEMENT_TOTALS ─────────────────────────────────────────────────────

def test_flow_totals_equal_persisted_values_and_senior_principal_reconciles(committed):
    kind, pr, uid, *_ = committed
    from app.persistence.workspace_repository import get_workspace_state

    wb = _wb(_export_last_run(pr, uid))
    ws = get_workspace_state(uid, pr.project_id)
    debt = ws.last_debt_schedule["periods"]
    sheet = wb["Debt_DSCR"]
    interest = _find(sheet, "Senior interest")
    persisted = [p["senior_interest_keur"] for p in debt if p["senior_interest_keur"] is not None]
    assert math.isclose(interest[3], sum(persisted), rel_tol=0, abs_tol=1e-9)   # Total column
    assert interest[4] == f"{len(persisted)}/{len(debt)}"                       # coverage is explicit
    # full-precision block: scheduled principal repays exactly the senior commitment (S&U)
    senior_sources = ws.last_runtime_summary["financing_evidence"]["sources_uses"]["senior_debt_keur"]
    full_rows = [r for r in _rows(sheet) if r[0] == "Senior principal" and isinstance(r[3], float)]
    full_principal = full_rows[-1][3]                                          # the unrounded block
    assert math.isclose(full_principal, senior_sources, rel_tol=0, abs_tol=1e-6)


def test_existing_statement_sheets_total_only_flows(committed):
    kind, pr, uid, *_ = committed
    wb = _wb(_export_last_run(pr, uid))
    bs = wb["Balance Sheet"]
    header = next(r for r in _rows(bs) if r[0] == "Metric")
    total_col = header.index("Total (flows only)")
    assert header[1][:4].isdigit() and header[1][4] == "-"        # canonical ISO period date, unchanged for consumers
    assert "Run state" not in str(bs["B3"].value)                 # existing sheets carry no volatile run state
    for label in ("Senior Debt Balance", "SHL Balance", "Net Fixed Assets", "Retained Earnings"):
        assert _find(bs, label)[total_col] == "n/a (balance)", label


# ── RC7_STOCK_VS_FLOW ────────────────────────────────────────────────────────

@pytest.mark.parametrize("label,kind", [
    ("Senior closing balance", STOCK), ("Cumulative senior draw", STOCK), ("DSRA balance", STOCK),
    ("Closing tax losses", STOCK), ("Opening tax losses", STOCK), ("Retained earnings (closing)", STOCK),
    ("Cash", STOCK), ("DSCR", RATIO), ("Reported DSCR", RATIO), ("Annual rate", RATIO),
    ("Senior principal", FLOW), ("Senior interest", FLOW), ("Cash tax", FLOW), ("Corporate tax cash", FLOW),
    ("Fiscal reintegration", FLOW), ("Revenue cash", FLOW), ("Tax losses used", FLOW),
    ("Senior draw", FLOW), ("Share capital draw", FLOW),
])
def test_stock_flow_classification(label, kind):
    assert classify_metric(label) == kind


def test_stock_and_ratio_rows_are_never_summed():
    assert display_total([1.0, 2.0], STOCK)[0] == "n/a (balance)"
    assert display_total([1.25, 1.3], RATIO)[0] == "n/a (ratio)"
    assert display_total([1.0, 2.0, 3.5], FLOW)[0] == 6.5


# ── RC7_MISSING_VS_ZERO ──────────────────────────────────────────────────────

def _synthetic_evidence(**sections) -> RunEvidence:
    base = dict(source="RUN_HISTORY", state=STATE_HISTORICAL, state_reason="synthetic", project_id="p", project_code="pc",
                project_name="Synthetic", owner_id="o", history_id="h", snapshot_id="s1",
                ran_at="2026-01-02T03:04:05+00:00", runtime_origin="v2_run", engine_version="e1",
                workbook_version=None, composite_hash="c1", scenario_id=None, scenario_name=None,
                last_runtime_scenario_id=None, identity={}, runtime_summary={}, financial_statements={},
                debt_schedule={}, tax_schedule={}, distribution_schedule={}, sponsor_schedule={},
                integrity_evidence={}, evidence_digest="d" * 64)
    base.update(sections)
    return RunEvidence(**base)


def test_missing_is_not_zero_and_explicit_zero_is_kept():
    debt = {"periods": [
        {"period": 1, "date": "2030-06-30", "is_operation": True, "senior_balance_keur": 100.0,
         "senior_principal_keur": 0.0, "senior_interest_keur": 5.0, "senior_ds_keur": 5.0, "dscr": None},
        {"period": 2, "date": "2030-12-31", "is_operation": True, "senior_balance_keur": 100.0,
         "senior_principal_keur": None, "senior_interest_keur": 5.0, "senior_ds_keur": 5.0, "dscr": None},
    ]}
    wb = openpyxl.load_workbook(io.BytesIO(R.build_run_evidence_workbook(_synthetic_evidence(debt_schedule=debt))))
    ws = wb["Debt_DSCR"]
    principal = _find(ws, "Senior principal")
    assert principal[5] == 0 and isinstance(principal[5], (int, float))      # persisted zero stays a number
    assert principal[6] == "n/a"                                             # missing is not 0
    assert principal[3] == 0.0 and principal[4] == "1/2"                     # total over persisted periods only
    assert "DSCR" not in [r[0] for r in _rows(ws) if r[0] == "DSCR"]         # row with nothing persisted is not drawn
    assert any("DSCR" in str(c) for r in _rows(ws) for c in r if c and "Not persisted" in str(c))


def test_legacy_run_without_evidence_reports_not_available_and_builds():
    ev = _synthetic_evidence()
    wb = openpyxl.load_workbook(io.BytesIO(R.build_run_evidence_workbook(ev)))
    assert any("NOT_AVAILABLE" in str(r[0]) for r in _rows(wb["Sources_Uses"]) if r[0])
    assert any("NOT_AVAILABLE" in str(r[0]) for r in _rows(wb["Construction_Funding"]) if r[0])
    assert any("NOT_AVAILABLE" in str(r[0]) for r in _rows(wb["Tax_Bridge"]) if r[0])
    statuses = {r[1] for r in _rows(wb["Run_Integrity"]) if r[1] in ("PASS", "FAIL", "UNAVAILABLE")}
    assert statuses == {"UNAVAILABLE"}                       # nothing is assumed to pass
    returns = _kv(wb["Returns"])
    assert returns["Project IRR"] == "n/a" and returns["Total Sponsor XIRR"] == "n/a"


def test_tax_sheet_does_not_fabricate_zero_totals(committed):
    kind, pr, uid, *_ = committed
    wb = _wb(_export_last_run(pr, uid))
    tax = _kv(wb["Tax"])
    # the persisted statements carry no tax-bridge periods: totals are unavailable, not 0
    assert tax["Runtime total cash tax"] == "NOT_AVAILABLE"
    assert tax["Runtime total CIT accrual"] == "NOT_AVAILABLE"
    assert wb["Tax_Bridge"].max_row > 8                      # the persisted tax schedule is reported instead


# ── RC7_DEBT_BALANCES_AND_SERVICE ────────────────────────────────────────────

def test_debt_consistency_checks_pass_on_real_runs(committed):
    kind, pr, uid, *_ = committed
    ws = _wb(_export_last_run(pr, uid))["Debt_DSCR"]
    checks = {r[0]: r[1] for r in _rows(ws) if r[0] in (
        "Debt service = principal + interest", "Balance roll-forward: prior closing - principal = closing")}
    assert set(checks.values()) == {"PASS"}, checks
    assert _find(ws, "Covenant compliance")[1] == "NOT DECLARED"
    integrity = _wb(_export_last_run(pr, uid))["Run_Integrity"]
    assert _kv(integrity)["Overall"] in ("PASS", "INCOMPLETE")   # existing authority's verdict, never FAIL here


def test_debt_consistency_detects_a_broken_persisted_schedule():
    debt = {"periods": [
        {"period": 1, "date": "2030-06-30", "senior_balance_keur": 90.0, "senior_principal_keur": 10.0,
         "senior_interest_keur": 5.0, "senior_ds_keur": 99.0},
        {"period": 2, "date": "2030-12-31", "senior_balance_keur": 70.0, "senior_principal_keur": 10.0,
         "senior_interest_keur": 4.0, "senior_ds_keur": 14.0},
    ]}
    ws = openpyxl.load_workbook(io.BytesIO(R.build_run_evidence_workbook(_synthetic_evidence(debt_schedule=debt))))["Debt_DSCR"]
    assert _find(ws, "Debt service = principal + interest")[1] == "FAIL"
    assert _find(ws, "Balance roll-forward: prior closing - principal = closing")[1] == "FAIL"


# ── RC7_HISTORICAL_IMMUTABLE / RC7_REPRODUCIBLE ──────────────────────────────

@pytest.fixture(scope="module")
def two_runs():
    """Run A, a Working Copy edit, then a numerically different Run B on the same project."""
    from app.persistence.projects_repository import get_project
    from app.persistence.workspace_repository import get_workspace_state, save_workspace_state

    pr, uid, snap, pi, sc = H.create_project("ev", suffix="hist")
    _, snap_a = H.commit_run(pr, uid, pi, sc, "ev")
    pr = get_project(pr.project_id, uid)
    entry_a = _latest(uid, pr)[0]
    before = {
        "last_run_export": _export_last_run(pr, uid),
        "history_export": _export_history(pr, uid, entry_a.history_id),
    }
    # Working Copy edit AFTER Run A (draft only, no rerun)
    ws = get_workspace_state(uid, pr.project_id)
    edited = dict(ws.draft_snapshot, total_capex_keur=str(float(ws.draft_snapshot["total_capex_keur"]) + 500))
    save_workspace_state(user_id=uid, project_id=pr.project_id, project_code=pr.project_code, draft_snapshot=edited,
                         saved_snapshot=edited, active_scenario_id=sc.scenario_id, active_scenario_name=sc.scenario_name)
    stale = {"last_run_export": _export_last_run(pr, uid)}
    # Run B on different financing inputs, committed under the edited Working Copy identity
    pi_b = dataclasses.replace(pi, financing=dataclasses.replace(pi.financing, target_dscr=pi.financing.target_dscr + 0.15))
    time.sleep(0.01)
    _, snap_b = H.commit_run(pr, uid, pi_b, sc, "ev")
    pr = get_project(pr.project_id, uid)
    return dict(pr=pr, uid=uid, entry_a=entry_a, snap_a=snap_a, snap_b=snap_b, before=before, stale=stale)


def _numeric_cells(wb, skip=("Report_Guide",)):
    out = {}
    for ws in wb.worksheets:
        if ws.title in skip:
            continue
        for row in ws.iter_rows():
            for c in row:
                if isinstance(c.value, (int, float)) and not isinstance(c.value, bool):
                    out[(ws.title, c.coordinate)] = c.value
    return out


def test_working_copy_edit_marks_last_run_stale_without_changing_numbers(two_runs):
    before = _wb(two_runs["before"]["last_run_export"])
    stale = _wb(two_runs["stale"]["last_run_export"])
    assert _kv(before["Report_Guide"])["Run state"] == STATE_CURRENT
    assert _kv(stale["Report_Guide"])["Run state"] == STATE_STALE
    assert "RUN STATE: STALE" in stale["Sources_Uses"]["A2"].value
    assert two_runs["stale"]["last_run_export"].metadata["export_run_state"] == STATE_STALE
    for sheet in ("Sources_Uses", "Debt_DSCR", "Construction_Funding", "Tax_Bridge", "Reserves"):
        a = {k: v for k, v in _numeric_cells(before).items() if k[0] == sheet}
        b = {k: v for k, v in _numeric_cells(stale).items() if k[0] == sheet}
        assert a == b, sheet                                     # the Working Copy edit changed nothing
    assert (two_runs["before"]["last_run_export"].metadata["export_evidence_digest"]
            == two_runs["stale"]["last_run_export"].metadata["export_evidence_digest"])


def test_historical_run_export_is_immutable_after_edit_and_later_run(two_runs):
    pr, uid, entry_a = two_runs["pr"], two_runs["uid"], two_runs["entry_a"]
    now = _export_history(pr, uid, entry_a.history_id)
    wb_now, wb_before = _wb(now), _wb(two_runs["before"]["history_export"])
    # superseded by Run B -> HISTORICAL (it was the Last Run, CURRENT, when first exported)
    assert _kv(wb_before["Report_Guide"])["Run state"] == STATE_CURRENT
    assert _kv(wb_now["Report_Guide"])["Run state"] == STATE_HISTORICAL
    assert now.metadata["export_run_state"] == STATE_HISTORICAL
    assert "RUN STATE: HISTORICAL" in wb_now["Debt_DSCR"]["A2"].value
    # identity preserved: snapshot, time, engine version, composite hash, history id
    guide = _kv(wb_now["Report_Guide"])
    assert guide["Runtime snapshot ID"] == two_runs["snap_a"] != two_runs["snap_b"]
    assert guide["Run History ID"] == entry_a.history_id
    assert guide["Engine version"] == entry_a.engine_version
    assert guide["Input composite hash"] == entry_a.composite_hash
    assert guide["Run timestamp"] == entry_a.ran_at
    # every number is unchanged and the evidence digest is unchanged
    assert _numeric_cells(wb_now) == _numeric_cells(wb_before)
    assert now.metadata["export_evidence_digest"] == two_runs["before"]["history_export"].metadata["export_evidence_digest"]
    # Run B is numerically different, so the Last Run report differs from the historical one
    b = _wb(_export_last_run(pr, uid))
    assert _kv(b["Report_Guide"])["Runtime snapshot ID"] == two_runs["snap_b"]
    assert _find(b["Debt_DSCR"], "Target DSCR (sizing)")[1] != _find(wb_now["Debt_DSCR"], "Target DSCR (sizing)")[1]


def test_historical_export_matches_stored_history_row_exactly(two_runs):
    pr, uid, entry_a = two_runs["pr"], two_runs["uid"], two_runs["entry_a"]
    wb = _wb(_export_history(pr, uid, entry_a.history_id))
    su = entry_a.runtime_summary["financing_evidence"]["sources_uses"]
    assert _find(wb["Sources_Uses"], "TOTAL USES (persisted)")[1] == pytest.approx(su["total_uses_keur"], rel=1e-14, abs=0)
    periods = entry_a.debt_schedule["periods"]
    ds = _find(wb["Debt_DSCR"], "Senior debt service")
    assert ds[3] == pytest.approx(sum(p["senior_ds_keur"] for p in periods if p["senior_ds_keur"] is not None), abs=1e-9)
    # input-assumption sheets are NOT produced: Run History holds no input snapshot to bind them to
    for absent in ("Inputs", "CAPEX", "OPEX", "Revenue"):
        assert absent not in wb.sheetnames
    assert any("not persisted in Run History" in str(r[2]) for r in _rows(wb["Periods_Dates"]) if len(r) > 2 and r[2])
    for name in ("P&L", "Cash Flow", "Balance Sheet", "Returns", "Tax_Bridge"):
        assert name in wb.sheetnames


def test_equivalent_source_snapshots_produce_byte_identical_workbooks(two_runs):
    pr, uid, entry_a = two_runs["pr"], two_runs["uid"], two_runs["entry_a"]
    first = _export_history(pr, uid, entry_a.history_id)
    time.sleep(1.1)                                              # crosses a clock-second boundary
    second = _export_history(pr, uid, entry_a.history_id)
    assert first.bytes_data == second.bytes_data
    assert first.filename == second.filename
    # building from an independently decoded copy of the same row is also identical
    ev = evidence_from_history_entry(entry_a, pr, last_run_snapshot_id=two_runs["snap_b"], last_run_is_stale=False,
                                     freshness_source="test")
    assert R.build_run_evidence_workbook(ev) == first.bytes_data
    # a different Run is a different report
    entry_b = _latest(uid, pr)[0]
    assert entry_b.history_id != entry_a.history_id
    assert _export_history(pr, uid, entry_b.history_id).bytes_data != first.bytes_data


def test_last_run_history_row_state_follows_existing_semantics(two_runs):
    pr, uid = two_runs["pr"], two_runs["uid"]
    entry_b = _latest(uid, pr)[0]
    wb = _wb(_export_history(pr, uid, entry_b.history_id))
    # Run B is the Last Run, committed under the edited Working Copy -> CURRENT
    assert _kv(wb["Report_Guide"])["Run state"] == STATE_CURRENT


# ── RC7_ISOLATION ────────────────────────────────────────────────────────────

def test_owner_and_project_isolation(two_runs):
    from app.persistence.projects_repository import get_project

    pr, uid, entry_a = two_runs["pr"], two_runs["uid"], two_runs["entry_a"]
    other_pr, other_uid, *_ = _project("ev", suffix="iso")
    # another owner cannot read this history row, even through their own project
    r = _export_history(other_pr, other_uid, entry_a.history_id)
    assert r.status_code == 404 and r.bytes_data is None
    # the owner's id with a different owner's project record
    r = _export_history(other_pr, uid, entry_a.history_id)
    assert r.status_code == 404 and r.bytes_data is None
    # unknown id, missing id, missing project / owner
    assert _export_history(pr, uid, "0" * 32).status_code == 404
    assert _export_history(pr, uid, "").status_code == 400
    assert _export_history(None, uid, entry_a.history_id).status_code == 400
    assert _export_history(pr, None, entry_a.history_id).status_code == 400
    # hostile id is never reflected
    r = _export_history(pr, uid, "<script>alert(1)</script>")
    assert r.status_code == 404 and "<script>" not in (r.error_content or "")
    # the other owner's Last Run export carries only their own Run
    mine = _kv(_wb(_export_last_run(other_pr, other_uid))["Report_Guide"])
    assert mine["Runtime snapshot ID"] != two_runs["snap_a"]
    assert get_project(other_pr.project_id, uid) is None


# ── RC7_NO_ENGINE ────────────────────────────────────────────────────────────

def test_exports_execute_no_engine(two_runs):
    pr, uid, entry_a = two_runs["pr"], two_runs["uid"], two_runs["entry_a"]

    def boom(*a, **k):
        raise AssertionError("engine must not run during export")

    with mock.patch("app.api.project_runner.run_project", side_effect=boom), \
            mock.patch("app.services.production_waterfall_seam.execute_production_waterfall", side_effect=boom), \
            mock.patch("app.services.production_financial_authority.run_clean_production", side_effect=boom):
        assert _export_last_run(pr, uid).status_code == 200
        assert _export_history(pr, uid, entry_a.history_id).status_code == 200


def test_reporting_modules_import_no_engine():
    import ast
    from pathlib import Path

    for name in ("institutional_reporting_v2", "run_evidence", "reporting_common"):
        tree = ast.parse(Path(f"app/export/{name}.py").read_text())
        imported = set()
        for node in ast.walk(tree):
            if isinstance(node, ast.Import):
                imported.update(a.name for a in node.names)
            elif isinstance(node, ast.ImportFrom):
                imported.add(node.module or "")
        banned = [m for m in imported if m.startswith(("financial_engine", "domain", "finco_core", "app.api.project_runner",
                                                      "app.services.production"))]
        assert not banned, (name, banned)


# ── RC7_WORKBOOK_INTEGRITY ───────────────────────────────────────────────────

def test_workbooks_open_without_corruption(committed, two_runs):
    kind, pr, uid, *_ = committed
    exports = [_export_last_run(pr, uid), _export_history(two_runs["pr"], two_runs["uid"], two_runs["entry_a"].history_id)]
    for export in exports:
        with zipfile.ZipFile(io.BytesIO(export.bytes_data)) as z:
            assert z.testzip() is None
            assert "[Content_Types].xml" in z.namelist() and "xl/workbook.xml" in z.namelist()
        wb = openpyxl.load_workbook(io.BytesIO(export.bytes_data))
        assert all(len(s.title) <= 31 for s in wb.worksheets)
        for ws in wb.worksheets:
            for row in ws.iter_rows():
                for c in row:
                    assert not (isinstance(c.value, str) and c.value.startswith("#")), (ws.title, c.coordinate)
                    assert not (isinstance(c.value, str) and c.value.startswith("=")), (ws.title, c.coordinate)  # values only
                    assert not (isinstance(c.value, float) and not math.isfinite(c.value))


@pytest.mark.skipif(shutil.which("soffice") is None, reason="LibreOffice not installed")
def test_workbook_opens_in_libreoffice(two_runs, tmp_path):
    export = _export_history(two_runs["pr"], two_runs["uid"], two_runs["entry_a"].history_id)
    src = tmp_path / "in" / "run.xlsx"
    src.parent.mkdir()
    src.write_bytes(export.bytes_data)
    out = tmp_path / "out"
    proc = subprocess.run(["soffice", "--headless", "--convert-to", "pdf", "--outdir", str(out), str(src)],
                          capture_output=True, timeout=240, env={"HOME": str(tmp_path), "PATH": "/usr/bin:/bin"})
    assert proc.returncode == 0 and (out / "run.pdf").stat().st_size > 1000, proc.stderr


# ── RC7_WIDE_AND_LONG ────────────────────────────────────────────────────────

def test_wide_schedules_are_usable(committed):
    kind, pr, uid, *_ = committed
    from openpyxl.utils import get_column_letter

    wb = _wb(_export_last_run(pr, uid))
    for name in ("Debt_DSCR", "Tax_Bridge", "Reserves"):
        ws = wb[name]
        assert ws.freeze_panes and ws.freeze_panes.startswith("F"), (name, ws.freeze_panes)   # metric columns stay visible
        assert ws.max_column >= 20 or kind in ("ev",) or name == "Reserves"
        for col in range(6, ws.max_column + 1):
            assert (ws.column_dimensions[get_column_letter(col)].width or 0) >= 12, (name, col)
    for name in ("P&L", "Cash Flow", "Balance Sheet"):        # existing wide sheets: period columns beyond E are sized
        ws = wb[name]
        for col in range(6, ws.max_column + 1):
            assert (ws.column_dimensions[get_column_letter(col)].width or 0) >= 12, (name, col)
    assert wb["Debt_DSCR"].page_setup.orientation == "landscape"


def test_long_text_wraps_and_rows_are_tall_enough(two_runs):
    wb = _wb(_export_history(two_runs["pr"], two_runs["uid"], two_runs["entry_a"].history_id))
    ws = wb["Report_Guide"]
    long_rows = [c for row in ws.iter_rows() for c in row if isinstance(c.value, str) and len(c.value) > 140]
    assert long_rows
    for c in long_rows:
        assert c.alignment.wrap_text
    merged_notes = [c for row in ws.iter_rows() for c in row if isinstance(c.value, str) and len(c.value) > 200]
    for c in merged_notes:
        assert (ws.row_dimensions[c.row].height or 15) > 15


# ── RC7_PERIOD_BASIS ─────────────────────────────────────────────────────────

def _dates(*pairs):
    return [date(y, m, d) for y, m, d in pairs]


def test_half_year_data_is_never_labelled_quarterly():
    dates = _dates((2031, 6, 30), (2031, 12, 31), (2032, 6, 30), (2032, 12, 31))
    b = period_basis(dates)
    assert b.declared_months is None and b.observed_months == (6,)
    assert "6-month" in b.label and "uarter" not in b.label + b.disclosure
    # construction stubs do not make operation look irregular
    stubs = _dates((2030, 7, 1), (2031, 1, 1), (2031, 3, 1)) + dates
    b2 = period_basis(stubs, operating=[False, False, False, True, True, True, True])
    assert "6-month spacing in operation" in b2.label and "uarter" not in b2.label


def test_future_frequency_metadata_is_read_not_assumed():
    half = _dates((2031, 6, 30), (2031, 12, 31), (2032, 6, 30))
    assert period_basis(half, [("debt_schedule", {"period_months": 6})]).label.startswith("Half-year")
    assert period_basis(half, [("x", {"frequency": "Semiannual"})]).declared_months == 6
    quarterly = _dates((2031, 3, 31), (2031, 6, 30), (2031, 9, 30))
    q = period_basis(quarterly, [("debt_schedule", {"period_months": 3})])
    assert q.label == "Quarterly" and not q.warning
    # metadata that contradicts the persisted dates is flagged, dates stay as persisted
    clash = period_basis(half, [("debt_schedule", {"period_months": 3})])
    assert "FREQUENCY_METADATA_MISMATCH" in clash.warning
    # without metadata nothing is assumed
    assert period_basis([None, None]).label == "Not determinable"


def test_period_labels_use_actual_dates_and_no_quarter_names(committed):
    kind, pr, uid, *_ = committed
    wb = _wb(_export_last_run(pr, uid))
    header = next(r for r in _rows(wb["Debt_DSCR"]) if r[0] == "Metric")
    assert header[5].startswith("P") and "| 20" in header[5] and header[5].count("-") == 2   # 'P<n> | yyyy-mm-dd'
    # real Runs persist no frequency metadata: the basis is described from the dates, never as quarterly
    assert "uarter" not in str(_kv(wb["Report_Guide"])["Period frequency"])
    assert "uarter" not in str(_kv(wb["Periods_Dates"])["Frequency"])
    assert _kv(wb["Periods_Dates"])["Financial Close"] not in (None, "n/a")     # Last Run: input snapshot date


# ── RC7_NO_COVENANT_CLAIM ────────────────────────────────────────────────────

def test_covenant_compliance_is_never_declared(committed, two_runs):
    kind, pr, uid, *_ = committed
    exports = [_export_last_run(pr, uid), _export_history(two_runs["pr"], two_runs["uid"], two_runs["entry_a"].history_id)]
    for export in exports:
        wb = _wb(export)
        text = "\n".join(str(v) for ws in wb.worksheets for r in ws.iter_rows(values_only=True) for v in r if v)
        lowered = text.lower()
        for forbidden in ("covenant compliant", "in compliance", "covenant satisfied", "no breach", "covenant met",
                          "passes covenant"):
            assert forbidden not in lowered
        assert "NOT DECLARED" in text or "NOT declared" in text
        assert "not a lender approval" in lowered
