"""Revenue V2 advanced exact-period prices: canonical run, no new kernel or golden.

All fixture prices and project identities are synthetic. This suite supplements
rather than changes WF07/F0/F3.8 tests or financial tolerances.
"""
from __future__ import annotations
import copy
import json
from decimal import Decimal

import pytest

from app import project_factories as factories
from app.workbook.revenue_multistream import apply_config, canonical_json
from app.workbook.revenue_commercial_stress import (
    CommercialStressUnavailable, changed_book, compare_commercial_cases,
)
from domain.revenue.multistream_runtime import parse_config
from financial_engine.adapters.project_inputs import from_project_inputs
from financial_engine.orchestrator import _build_period_engine
from finco_core.revenue.generation import revenue_decomposition_schedule
from tests.test_wf07_revenue_multistream import contracts
from tests.test_model_decision_workspace_v2 import env


def book(kind="solar", *, with_cfd=True, pi=None):
    if pi is None:
        pi = getattr(factories, "create_generic_" + kind + "_reference")()
    axis = _build_period_engine(from_project_inputs(pi))
    operating = tuple(p for p in axis.periods() if p.is_operation)
    assert len(operating) > 2
    starts = [p.start_date.isoformat() for p in operating]
    common = dict(start_date=starts[0], end_date=operating[-1].end_date.isoformat(),
                  settlement="SAME_PERIOD", source_ref="synthetic-independent-contract")
    rows = [
        dict(common, id="alpha", type="ppa", volume_share=.45,
             price_eur_mwh=50, indexation_rate=0,
             price_curve=[{"period_start": day, "price_eur_mwh": 55 if i < 2 else 47}
                          for i, day in enumerate(starts)]),
        dict(common, id="beta", type="ppa", volume_share=.35,
             price_eur_mwh=70, indexation_rate=0),
        dict(common, id="market", type="merchant", volume_share=None,
             price_eur_mwh=80, indexation_rate=0, capture_rate=.8,
             price_curve=[{"period_start": day, "price_eur_mwh": 40 if i < 2 else 95}
                          for i, day in enumerate(starts)]),
    ]
    if with_cfd:
        rows.append(dict(common, id="settlement", type="cfd", volume_share=.20,
                         price_eur_mwh=65, indexation_rate=0,
                         reference_stream_id="market"))
    return pi, axis, operating, {"version": 2, "technology": kind,
                                 "indexation_basis": "OPERATING_YEAR", "streams": rows}


@pytest.mark.parametrize("kind", ["solar", "wind"])
def test_exact_operating_curve_points_feed_canonical_price_and_two_way_cfd(kind):
    pi, axis, operating, payload = book(kind)
    # This executes the actual canonical production/revenue adapter with
    # calculated generation, allocation, capture and CfD — not a cloned model.
    project = apply_config(pi, json.dumps(payload), project_type=kind)
    out = revenue_decomposition_schedule(project, _build_period_engine(from_project_inputs(project)))
    found_positive = found_negative = False
    for i, period in enumerate(operating):
        period_result = out[period.index]
        streams = {s["id"]: s for s in period_result["multistream_evidence"]["streams"]}
        assert set(streams) == {"alpha", "beta", "market", "settlement"}
        ppa = Decimal("55") if i < 2 else Decimal("47")
        market = Decimal("40") if i < 2 else Decimal("95")
        assert streams["alpha"]["price_eur_mwh"] == pytest.approx(float(ppa), abs=1e-12)
        assert streams["market"]["price_eur_mwh"] == pytest.approx(float(market), abs=1e-12)
        production = Decimal(str(period_result["generation_mwh"]))
        expected_alpha = production * Decimal(".45") * ppa / 1000
        expected_beta = production * Decimal(".35") * Decimal("70") / 1000
        expected_market = production * Decimal(".20") * market * Decimal(".8") / 1000
        assert streams["alpha"]["revenue_keur"] == pytest.approx(float(expected_alpha), abs=1e-8)
        assert streams["beta"]["revenue_keur"] == pytest.approx(float(expected_beta), abs=1e-8)
        assert streams["market"]["revenue_keur"] == pytest.approx(float(expected_market), abs=1e-8)
        expected_cfd = production * Decimal(".20") * (Decimal("65") - market) / 1000
        assert streams["settlement"]["settlement_keur"] == pytest.approx(float(expected_cfd), abs=1e-8)
        assert sum(s["quantity_mwh"] for s in streams.values() if s["type"] != "cfd") == pytest.approx(
            period_result["generation_mwh"], abs=1e-8)
        assert streams["settlement"]["quantity_mwh"] == pytest.approx(
            .20 * period_result["generation_mwh"], abs=1e-8)
        found_positive |= expected_cfd > 0
        found_negative |= expected_cfd < 0
    assert found_positive and found_negative


@pytest.mark.parametrize("kind", ["solar", "wind"])
def test_v1_and_v2_no_curve_have_identical_complete_canonical_revenue(kind):
    pi, _, _, _ = book(kind)
    old = contracts(pi, kind, price=85, merchant_price=70, strike=90, index=.02)
    v2 = copy.deepcopy(old)
    v2["version"] = 2
    first = apply_config(pi, json.dumps(old), project_type=kind)
    second = apply_config(pi, json.dumps(v2), project_type=kind)
    def revenue(proj):
        return revenue_decomposition_schedule(proj, _build_period_engine(from_project_inputs(proj)))
    assert revenue(first) == revenue(second)


@pytest.mark.parametrize("fault", [
    "missing_point", "duplicate_point", "reversed_order", "outside_tenor",
    "negative_price", "nonfinite_price", "non_iso_date", "indexed_curve",
    "cfd_curve", "v1_curve", "unknown_key", "duplicate_json_key",
])
def test_bad_curves_rejected_without_partial_runtime(fault):
    pi, axis, operating, payload = book("solar")
    curve = payload["streams"][0]["price_curve"]
    if fault == "missing_point":
        curve.pop()
    elif fault == "duplicate_point":
        curve[1]["period_start"] = curve[0]["period_start"]
    elif fault == "reversed_order":
        curve[0], curve[1] = curve[1], curve[0]
    elif fault == "outside_tenor":
        curve[0]["period_start"] = "2020-01-01"
    elif fault == "negative_price":
        curve[0]["price_eur_mwh"] = -3
    elif fault == "nonfinite_price":
        curve[0]["price_eur_mwh"] = float("nan")
    elif fault == "non_iso_date":
        curve[0]["period_start"] = "1-2-2030"
    elif fault == "indexed_curve":
        payload["streams"][0]["indexation_rate"] = .01
    elif fault == "cfd_curve":
        payload["streams"][-1]["price_curve"] = curve.copy()
    elif fault == "v1_curve":
        payload["version"] = 1
    elif fault == "unknown_key":
        curve[0]["interpolate"] = "linear"
    if fault == "duplicate_json_key":
        raw = json.dumps(payload).replace('"period_start": ', '"period_start": "2001-01-01", "period_start": ', 1)
    else:
        raw = json.dumps(payload)
    with pytest.raises(ValueError):
        apply_config(pi, raw, project_type="solar")


def test_version2_curve_canonical_serialization_and_price_policy():
    pi, axis, _, payload = book()
    raw = canonical_json(json.dumps(payload))
    parsed = json.loads(raw)
    assert parsed["version"] == 2
    assert all(row["indexation_rate"] == 0 for row in parsed["streams"])
    assert parse_config(raw).canonical_json == raw
    config = parse_config(raw)
    config.validate_axis(axis.periods())


def test_stress_changes_only_declared_price_authority_and_never_mutates_base():
    pi, axis, operating, payload = book()
    old = json.dumps(payload)
    original = copy.deepcopy(payload)
    lower = json.loads(changed_book(old, "merchant_down", 10))
    higher = json.loads(changed_book(old, "merchant_up", 10))
    assert payload == original
    assert lower["streams"][0] == payload["streams"][0]
    assert lower["streams"][2]["price_curve"][0]["price_eur_mwh"] == pytest.approx(36)
    assert higher["streams"][2]["price_curve"][0]["price_eur_mwh"] == pytest.approx(44)
    assert lower["streams"][2]["capture_rate"] == payload["streams"][2]["capture_rate"]
    assert lower["streams"][3] == payload["streams"][3]
    assert json.loads(changed_book(old, "ppa_down", 10))["streams"][0]["price_curve"][0][
        "price_eur_mwh"] == pytest.approx(49.5)
    assert apply_config(pi, json.dumps(lower), project_type="solar")


def test_stress_early_expiry_requires_residual_and_explicit_axis_boundary():
    pi, axis, operating, payload = book()
    day = operating[2].start_date.isoformat()
    # No selected contract ID is exposed in the V1 expiry UI: this is
    # intentionally a single-PPA fixture. Multi-PPA books MUST fail closed.
    payload["streams"] = [row for row in payload["streams"] if row["id"] != "beta"]
    raw = json.dumps(payload)
    shocked = json.loads(changed_book(raw, "ppa_expiry", 0, expiry_date=day))
    assert shocked["streams"][0]["end_date"] == day
    assert len(shocked["streams"][0]["price_curve"]) == 2
    assert apply_config(pi, json.dumps(shocked), project_type="solar")
    # The remaining explicit Merchant still conserves 100% physical MWh,
    # but cannot serve as a residual allocation for earlier expiry.
    next(row for row in payload["streams"] if row["id"] == "market")["volume_share"] = .55
    with pytest.raises(CommercialStressUnavailable, match="RESIDUAL_REQUIRED"):
        changed_book(json.dumps(payload), "ppa_expiry", 0, expiry_date=day)


@pytest.mark.parametrize("problem", ["invalid_stress", "no_match", "outside_limit", "unsupported_vertical"])
def test_stress_boundaries_fail_closed(problem):
    pi, _, _, payload = book()
    if problem == "invalid_stress":
        with pytest.raises(CommercialStressUnavailable):
            changed_book(json.dumps(payload), "bank_specific_discount", 10)
    elif problem == "no_match":
        # Remove the CfD referencing this merchant too: the target here is
        # NO eligible physical Merchant, NOT an invalid orphaned CfD.
        payload["streams"] = [r for r in payload["streams"] if r["type"] == "ppa"]
        next(r for r in payload["streams"] if r["id"] == "beta")["volume_share"] = .55
        with pytest.raises(CommercialStressUnavailable, match="NO_ELIGIBLE_CONTRACT"):
            changed_book(json.dumps(payload), "merchant_down", 10)
    elif problem == "outside_limit":
        with pytest.raises(CommercialStressUnavailable):
            changed_book(json.dumps(payload), "ppa_down", -101)
    else:
        with pytest.raises(ValueError, match="TECHNOLOGY"):
            apply_config(pi, json.dumps(payload), project_type="data_center")


@pytest.mark.parametrize("kind", ["solar", "wind"])
def test_canonical_financing_read_only_stress_and_persisted_input_unchanged(kind):
    pi, axis, _, payload = book(kind)
    before = apply_config(pi, json.dumps(payload), project_type=kind)
    original = before.revenue.multistream_config_json
    result = compare_commercial_cases(before, shock="merchant_down", amount_pct=10)
    assert result["status"] == "COMPARISON"
    assert result["persistence"] == "NONE"
    assert result["bank_case_status"] == "UNAVAILABLE"  # no inferred lender case
    assert result["metrics"]
    assert before.revenue.multistream_config_json == original
    data = {row["key"]: row for row in result["metrics"]}
    assert data["total_revenue_keur"]["stress"] < data["total_revenue_keur"]["base"]
    assert data["total_ebitda_keur"]["stress"] < data["total_ebitda_keur"]["base"]
    assert data["total_senior_ds_keur"]["base"] is not None
    assert data["project_irr"]["base"] is not None


def test_v2_curve_last_run_and_export_are_immutable_after_unsaved_change(env, monkeypatch):
    """A new curve Save makes CURRENT stale but cannot rewrite Last Run/export."""
    import copy
    from io import BytesIO
    from openpyxl import load_workbook
    from app.persistence.workspace_repository import get_workspace_state
    from app.persistence.run_history_repository import get_run_history
    from app.workbook.service import WorkbookService
    from app.workbook.revenue_multistream import SNAPSHOT_KEY
    from app.services.export_service import resolve_canonical_last_run_from_workspace
    from app.services.v2_export_service import build_canonical_last_run_institutional_workbook_export
    from tests.test_model_decision_workspace_v2 import tokens
    from tests.test_wf07_revenue_workspace import contract_tokens

    record = env.create("solar")
    owner = "decision-user"
    project = {"project": record.project_code}
    def ws():
        return get_workspace_state(owner, record.project_id)
    def page():
        r = env.client.get("/v2/workbook", params=project)
        assert r.status_code == 200
        return r.text
    def save(payload):
        return env.client.post("/v2/workbook/revenue/contracts",
            headers={"HX-Request": "true"}, data={
                **project, **contract_tokens(page()), "contracts": json.dumps(payload)})
    pi = WorkbookService.build_draft_input_set_from_workspace(ws()).to_projectinputs()
    _, _, _, payload = book("solar", pi=pi)
    assert save(payload).status_code == 200
    raw_original = ws().draft_snapshot[SNAPSHOT_KEY]
    assert json.loads(raw_original) == payload
    run = env.client.post("/v2/workbook/run", headers={"HX-Request": "true"},
                          data={**project, **tokens(page())})
    assert 'id="v2-sheet-returns"' in run.text, run.text[:2000]
    assert not ws().dirty
    previous_run = copy.deepcopy(ws().last_runtime_summary)
    previous_history = get_run_history(owner, record.project_id)
    assert previous_run["revenue_multistream_audit"]
    edited = copy.deepcopy(payload)
    edited["streams"][0]["price_curve"][0]["price_eur_mwh"] = 100
    assert save(edited).status_code == 200
    assert ws().dirty
    assert get_run_history(owner, record.project_id) == previous_history
    assert ws().last_runtime_summary == previous_run
    authority = resolve_canonical_last_run_from_workspace(record, owner, ws())
    assert authority.working_changed_since_run
    assert json.loads(authority.project_inputs.revenue.multistream_config_json) == payload
    def no_engine(*args, **kwargs):
        pytest.fail("Export must never execute financial model")
    with monkeypatch.context() as patch:
        patch.setattr("app.runtime.model_execution.run_model_process", no_engine)
        patch.setattr("app.services.production_financial_authority.run_clean_production", no_engine)
        exported = build_canonical_last_run_institutional_workbook_export(
            "generic_solar_reference", safe_project=record.project_code,
            project_record=record, user_id=owner)
        assert exported.status_code == 200, exported.error_content
        workbook = load_workbook(BytesIO(exported.bytes_data), read_only=True, data_only=True)
        try:
            assert "Revenue_Streams" in workbook.sheetnames
            rows = list(workbook["Revenue_Streams"].iter_rows(min_row=8, values_only=True))
            assert any(r[1] == "alpha" and r[5] == 55 for r in rows)
            assert any(r[1] == "alpha" and r[5] == 47 for r in rows)
            assert not any(r[1] == "alpha" and r[5] == 100 for r in rows)
        finally:
            workbook.close()
