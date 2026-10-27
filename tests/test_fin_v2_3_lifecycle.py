"""Real clone, Save/CAS, HTTP Run, history and frozen export on disposable DBs."""
import copy
import io
import json
import sqlite3
import uuid

from bs4 import BeautifulSoup
from openpyxl import load_workbook
import pytest

from tests.test_fin_val_b2_c1_recovery_finite_opex import env, USER, _prime, _ws, _snapshot
from tests.fin_v2_3_support import config_json
from app.workbook.financial_core_config import FIELD_ID, SNAPSHOT_KEY
from app.workbook.registry import WORKBOOK
from app.workbook.input_set import ProjectInputSet
from app.workbook.update_service import WorkbookUpdateService, StaleContentError


def save(env,record,raw,mode,hash_override=None):
    projection=env.client.get('/v2/model-react/fields/project_setup',params={'project':record.project_code}).json()
    ws=_ws(record)
    h=hash_override or projection['save']['content_hash']
    if mode=='batch':
        WorkbookUpdateService.apply_batch_draft_update(ws=ws,updates=[(FIELD_ID,raw)],content_hash=h,
            workbook_version=WORKBOOK.version,expected_scenario_id=ws.active_scenario_id,
            project_record=record,actor_user_id=USER)
        return None
    return env.client.post('/v2/workbook/update',headers={'HX-Request':'true'},data={
        'project':record.project_code,'field_id':FIELD_ID,'value':raw,'content_hash':h,
        'workbook_version':WORKBOOK.version,'csrf_token':projection['save']['csrf_token']})


def history(env,record):
    with sqlite3.connect(env.db_path) as c:
        c.row_factory=sqlite3.Row
        return [dict(x) for x in c.execute('SELECT * FROM model_run_history WHERE project_id=? ORDER BY history_id',(record.project_id,))]


@pytest.mark.parametrize('kind',['solar','wind','data_center','ev_charging'])
@pytest.mark.parametrize('mode',['single','batch'])
def test_config_save_reload_run_history_export_is_frozen(env,kind,mode):
    record=env.create(kind,capacity=20);_prime(env,record)
    raw=config_json()
    before_history=history(env,record)
    r=save(env,record,raw,mode)
    assert r is None or r.status_code==200, r.text if r is not None else ''
    assert history(env,record)==before_history  # Save never creates a Run.
    snap=_snapshot(record)
    assert json.loads(snap[SNAPSHOT_KEY])==json.loads(raw)
    pi=ProjectInputSet.from_snapshot(snap).to_projectinputs()
    assert pi.valuation.project.annual_discount_rate==.08
    assert pi.valuation.sponsor.annual_discount_rate==.12
    page=env.client.get('/v2/workbook',params={'project':record.project_code})
    form=BeautifulSoup(page.text,'html.parser').select_one('#v2-canonical-run-form')
    data={i['name']:i.get('value','') for i in form.select('input[name]')}
    data['csrf_token']=data.pop('run_csrf_token');data['run_intent_id']=str(uuid.uuid4())
    response=env.client.post('/v2/workbook/run',headers={'Accept':'application/json','HX-Request':'true'},data=data)
    assert response.status_code==200 and response.json()['outcome']=='COMMITTED',response.text[:600]
    rows=history(env,record);assert len(rows)==1
    ws=_ws(record)
    core=ws.last_runtime_summary['financing_evidence']['financial_core_v1']
    assert core['project_npv_keur'] is not None and core['sponsor_npv_keur'] is not None
    # FIN-V2-2/#390 is now a real persisted, Run-bound authority. Its exact
    # contractual and payoff indices must feed FIN-V2-3's discount engine;
    # an OK string without matching values does NOT satisfy this test.
    debt_evidence=ws.last_runtime_summary['financing_evidence']['senior_debt_evidence']
    assert debt_evidence['schema_version']=='SENIOR_DEBT_EVIDENCE_V1'
    assert debt_evidence['authority']=='CANONICAL_SENIOR_DEBT_RESULT_READ_ONLY_PROJECTION'
    assert debt_evidence['status']=='RECORDED'
    ends=debt_evidence['llcr_inputs']['loan_life_end_candidates']
    llcr=core['loan_life_coverage']
    assert llcr['status']=='OK'
    assert llcr['contractual']['horizon_end_period_index']==ends['contractual_maturity_period_index']
    assert llcr['contractual']['horizon_end_date']==ends['contractual_maturity_date']
    assert llcr['actual_repayment']['horizon_end_period_index']==ends['actual_payoff_period_index']
    assert llcr['actual_repayment']['horizon_end_date']==ends['actual_payoff_date']
    assert llcr['contractual']['opening_senior_balance_keur']>0
    assert llcr['actual_repayment']['opening_senior_balance_keur']>0
    assert llcr['rolling'] and llcr['minimum']
    assert llcr['minimum']['ratio']==min(point['ratio'] for point in llcr['rolling'])
    assert core['llcr_contractual_maturity']==llcr['contractual']['ratio']
    assert core['llcr_actual_repayment']==llcr['actual_repayment']['ratio']
    assert core['min_rolling_llcr']==llcr['minimum']['ratio']
    assert core['valuation_policy']['policy_version']=='FIN_V2_3_POLICY_V1'
    assert ws.last_sponsor_schedule['contract_version']=='SPONSOR_PERIOD_EVIDENCE_V1'
    from app.persistence.run_history_repository import get_run_history_entry
    from app.export.run_evidence import evidence_from_history_entry
    from app.export.institutional_reporting_v3 import build_run_evidence_workbook_v3
    def evidence():
        entry=get_run_history_entry(USER,record.project_id,rows[0]['history_id'])
        return evidence_from_history_entry(entry,record,last_run_snapshot_id=ws.last_runtime_snapshot_id,
            last_run_is_stale=False,freshness_source='test committed identity')
    ev=evidence();digest=ev.evidence_digest
    from app.services.financial_core_evidence import from_runtime_summary
    mismatched=copy.deepcopy(ws.last_runtime_summary)
    mismatched['financing_evidence']['run_binding']['snapshot_id']='another-run'
    assert from_runtime_summary(mismatched,snapshot_id=ev.snapshot_id,composite_hash=ev.composite_hash)['status']=='RUN_BINDING_MISMATCH'
    book=load_workbook(io.BytesIO(build_run_evidence_workbook_v3(ev)),data_only=True)
    labels={str(r[0]):r[1] for r in book['Executive_Overview'].iter_rows(values_only=True) if len(r)>1}
    # Excel stores numeric cells with 15-digit decimal precision. JSON remains exact.
    assert labels['Total Sponsor NPV']==pytest.approx(core['sponsor_npv_keur'],rel=1e-12,abs=1e-9)
    assert labels['Generic tax policy']=='GENERIC_MVP_POLICY'
    old=copy.deepcopy(rows)
    c=json.loads(raw);c['sponsor']['discount_rate_pct']=15.0
    save(env,record,json.dumps(c),mode)
    assert history(env,record)==old
    assert evidence().evidence_digest==digest
    assert evidence().financial_core==ev.financial_core
    assert env.as_user('other-financial-owner').get('/v2/workbook',params={'project':record.project_code}).status_code in (403,404)


@pytest.mark.parametrize('mode',['single','batch'])
def test_stale_policy_save_and_cross_user_mutation_are_rejected(env,mode):
    record=env.create();_prime(env,record)
    projection=env.client.get('/v2/model-react/fields/project_setup',params={'project':record.project_code}).json()
    old=projection['save']['content_hash']
    save(env,record,config_json(),mode)
    before=_snapshot(record)
    c=json.loads(config_json());c['coverage']['discount_rate_pct']=9
    if mode=='batch':
        with pytest.raises((ValueError,StaleContentError)):save(env,record,json.dumps(c),mode,old)
    else:
        r=save(env,record,json.dumps(c),mode,old)
        # Canonical HTMX Save renders stale refusals as a 200 error fragment.
        assert r.status_code in (200,409) and 'edit was not saved' in r.text.lower()
    assert _snapshot(record)==before
    ws=_ws(record)
    from app.workbook.update_service import BatchApplyError
    with pytest.raises(BatchApplyError,match='OWNER_MISMATCH'):
        WorkbookUpdateService.apply_batch_draft_update(ws=ws,updates=[(FIELD_ID,config_json())],
            content_hash=old,workbook_version=WORKBOOK.version,expected_scenario_id=ws.active_scenario_id,
            project_record=record,actor_user_id='not-the-owner')
    assert _snapshot(record)==before
