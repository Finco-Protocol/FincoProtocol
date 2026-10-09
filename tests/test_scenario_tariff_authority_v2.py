"""Correction A: real scenario input binding, immutable lineage and export parity."""
from __future__ import annotations

import copy

import pytest

from tests.test_model_decision_workspace_v2 import env, tokens


def test_explicit_override_binds_alias_without_reversing_global_precedence(env):
    from app.persistence.scenarios_repository import resolve_scenario_snapshot
    from app.input_adapter import _snapshot_to_dict
    from app.persistence.workspace_repository import get_workspace_state
    record = env.create()
    base = dict(get_workspace_state("decision-user", record.project_id).draft_snapshot)
    base.update(rev_ppa_base_tariff=50, tariff_eur_mwh=40)
    original = copy.deepcopy(base)
    assert _snapshot_to_dict(base)["tariff_eur_mwh"] == 50
    for value in (85, 45, 0):
        resolved = resolve_scenario_snapshot(base, {"tariff_eur_mwh": value})
        assert _snapshot_to_dict(resolved)["tariff_eur_mwh"] == value
        assert resolved["rev_ppa_base_tariff"] == value
    assert resolve_scenario_snapshot(base, {}) == base == original
    assert resolve_scenario_snapshot({"tariff_eur_mwh": 50}, {"tariff_eur_mwh": 85}) == {"tariff_eur_mwh": 85}


@pytest.mark.parametrize("value", [True, None, "invalid", float("nan"), float("inf"), -1])
def test_invalid_tariff_binding_fails_closed(value):
    from app.workbook.scenario_revenue_authority import bind_scenario_tariff
    with pytest.raises(ValueError, match="finite non-negative"):
        bind_scenario_tariff({"project_type": "Wind", "rev_ppa_base_tariff": 50}, {"tariff_eur_mwh": value})


@pytest.mark.parametrize("kind", ["data_center", "ev_charging"])
def test_unsupported_vertical_tariff_cannot_be_injected_through_post(env, kind):
    from app.persistence.scenarios_repository import list_scenarios
    record = env.create(kind)
    env.client.get("/v2/workbook", params={"project": record.project_code})
    env.client.post("/v2/workbook/scenarios/create", data={"project": record.project_code, "scenario_name": "Alternative"})
    sc = next(s for s in list_scenarios("decision-user", record.project_id) if not s.is_base_case)
    response = env.client.post("/v2/workbook/scenarios/update-overrides", data={
        "project": record.project_code, "scenario_id": sc.scenario_id, "tariff_eur_mwh": 85})
    assert response.status_code == 422
    assert "only supported for Solar and Wind" in response.text
    assert next(s for s in list_scenarios("decision-user", record.project_id) if s.scenario_id == sc.scenario_id).overrides == {}


@pytest.mark.parametrize("value", ["nan", "inf", "-1"])
def test_tariff_validation_precedes_scenario_write(env, value):
    from app.persistence.scenarios_repository import list_scenarios
    record = env.create()
    env.client.get("/v2/workbook", params={"project": record.project_code})
    env.client.post("/v2/workbook/scenarios/create", data={"project": record.project_code, "scenario_name": "Alternative"})
    sc = next(s for s in list_scenarios("decision-user", record.project_id) if not s.is_base_case)
    response = env.client.post("/v2/workbook/scenarios/update-overrides", data={
        "project": record.project_code, "scenario_id": sc.scenario_id, "tariff_eur_mwh": value})
    assert response.status_code == 422
    assert next(s for s in list_scenarios("decision-user", record.project_id) if s.scenario_id == sc.scenario_id).overrides == {}


@pytest.mark.parametrize("kind", ["solar", "wind"])
def test_real_base_upside_downside_run_restore_reset_and_export(env, monkeypatch, kind):
    from app.persistence.scenarios_repository import get_scenario, list_scenarios
    from app.persistence.workspace_repository import get_workspace_state
    from app.persistence.run_history_repository import get_run_history
    from app.services.export_service import resolve_canonical_last_run_from_workspace
    from app.services.v2_export_service import build_canonical_last_run_institutional_workbook_export

    owner = "decision-user"
    record = env.create(kind)
    project = {"project": record.project_code}
    def page():
        return env.client.get("/v2/workbook", params=project)
    def run():
        response = env.client.post("/v2/workbook/run", headers={"HX-Request": "true"}, data={**project, **tokens(page().text)})
        assert response.status_code == 200 and 'id="v2-sheet-returns"' in response.text
        return get_workspace_state(owner, record.project_id)
    baseline = run()
    base = next(s for s in list_scenarios(owner, record.project_id) if s.is_base_case)
    original_draft = dict(baseline.draft_snapshot)
    evidence = {"Base": dict(baseline.last_runtime_summary)}
    alternatives = []
    for name, tariff in (("Upside", 85), ("Downside", 45)):
        created = env.client.post("/v2/workbook/scenarios/create", data={**project, "scenario_name": name})
        assert created.status_code == 200
        sc = next(s for s in list_scenarios(owner, record.project_id) if s.scenario_name == name)
        updated = env.client.post("/v2/workbook/scenarios/update-overrides", headers={"HX-Request": "true"},
                                  data={**project, "scenario_id": sc.scenario_id, "tariff_eur_mwh": tariff})
        assert updated.status_code == 200
        env.client.post("/v2/workbook/scenarios/select", data={**project, "scenario_id": sc.scenario_id})
        ws = run()
        assert dict(ws.draft_snapshot) == original_draft
        assert ws.last_runtime_summary["scenario_revenue_input"]["effective_tariff_eur_mwh"] == tariff
        bound = get_scenario(sc.scenario_id, owner).last_run_summary
        assert bound["scenario_overrides_at_run"]["tariff_eur_mwh"] == tariff
        assert bound["scenario_revenue_input"]["effective_tariff_eur_mwh"] == tariff
        assert bound["kpis"]["total_revenue_keur"] == ws.last_runtime_summary["total_revenue_keur"]
        assert not ws.dirty
        authority = resolve_canonical_last_run_from_workspace(record, owner, ws)
        assert authority.project_inputs.revenue.ppa_base_tariff == tariff
        assert authority.active_scenario_id == sc.scenario_id
        assert not authority.working_changed_since_run
        alternatives.append((sc, copy.deepcopy(bound)))
        evidence[name] = dict(ws.last_runtime_summary)

    assert evidence["Upside"]["total_revenue_keur"] > evidence["Base"]["total_revenue_keur"] > evidence["Downside"]["total_revenue_keur"]
    assert evidence["Upside"]["project_irr"] > evidence["Base"]["project_irr"] > evidence["Downside"]["project_irr"]
    before_history = get_run_history(owner, record.project_id)
    downside, _ = alternatives[-1]
    changed = env.client.post("/v2/workbook/scenarios/update-overrides", headers={"HX-Request": "true"},
                              data={**project, "scenario_id": downside.scenario_id, "tariff_eur_mwh": 46})
    assert changed.status_code == 200
    assert get_run_history(owner, record.project_id) == before_history
    ws = get_workspace_state(owner, record.project_id)
    authority = resolve_canonical_last_run_from_workspace(record, owner, ws)
    assert authority.working_changed_since_run
    assert authority.project_inputs.revenue.ppa_base_tariff == 45  # Last Run, not live 46
    tampered = copy.deepcopy(ws)
    tampered.last_runtime_summary["scenario_revenue_input"]["effective_tariff_eur_mwh"] = 46
    with pytest.raises(ValueError, match="effective tariff does not match"):
        resolve_canonical_last_run_from_workspace(record, owner, tampered)
    def forbidden(*args, **kwargs):
        raise AssertionError("Export must not run the engine")
    with monkeypatch.context() as patcher:
        patcher.setattr("app.runtime.model_execution.run_model_process", forbidden)
        patcher.setattr("app.api.project_runner.run_project", forbidden)
        patcher.setattr("app.services.production_waterfall_seam.execute_production_waterfall", forbidden)
        patcher.setattr("app.services.production_financial_authority.run_clean_production", forbidden)
        exported = build_canonical_last_run_institutional_workbook_export(
            "generic_" + kind + "_reference", safe_project=record.project_code,
            project_record=record, user_id=owner)
        assert exported.status_code == 200 and exported.bytes_data
        from io import BytesIO
        from openpyxl import load_workbook
        workbook = load_workbook(BytesIO(exported.bytes_data), read_only=True, data_only=True)
        try:
            assert any(row[0] == "PPA tariff EUR/MWh" and row[1] == 45
                       for sheet in workbook for row in sheet.iter_rows(values_only=True) if len(row) > 1)
        finally:
            workbook.close()
    assert get_run_history(owner, record.project_id) == before_history
    for sc, bound in alternatives:
        assert get_scenario(sc.scenario_id, owner).last_run_summary == bound
    env.client.post("/v2/workbook/scenarios/select", data={**project, "scenario_id": alternatives[0][0].scenario_id})
    restored = get_workspace_state(owner, record.project_id)
    assert restored.last_runtime_summary["scenario_revenue_input"]["effective_tariff_eur_mwh"] == 85
    assert resolve_canonical_last_run_from_workspace(record, owner, restored).project_inputs.revenue.ppa_base_tariff == 85
    env.client.post("/v2/workbook/scenarios/remove-override", data={**project,
        "scenario_id": alternatives[0][0].scenario_id, "field": "tariff_eur_mwh"})
    rerun = run()
    assert rerun.last_runtime_summary["scenario_revenue_input"]["effective_tariff_eur_mwh"] == evidence["Base"]["scenario_revenue_input"]["effective_tariff_eur_mwh"]
    env.client.post("/v2/workbook/scenarios/select", data={**project, "scenario_id": base.scenario_id})
    assert dict(get_workspace_state(owner, record.project_id).draft_snapshot) == original_draft
