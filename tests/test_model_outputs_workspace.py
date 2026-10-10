"""Statements & Debt output workspace: read-only presentation of the persisted Last Run."""
from __future__ import annotations

import html as htmllib
import re
from types import SimpleNamespace

import pytest

from app.workbook import multisenior_config as config
from tests.test_model_decision_workspace_v2 import env, tokens  # noqa: F401  (fixture)
from tests.test_model_financing_f2 import run, save
from tests.test_model_financing_f3_workspace import collection_for_inputs, state

OWNER = "decision-user"


def get_outputs(env, record):
    response = env.client.get("/v2/workbook/outputs", params={"project": record.project_code})
    assert response.status_code == 200, response.text[:500]
    return response.text


def activate_f3(env, record):
    from app.persistence.workspace_repository import get_workspace_state
    from app.workbook.input_set import ProjectInputSet
    ws = get_workspace_state(OWNER, record.project_id)
    pi = ProjectInputSet.from_snapshot(ws.draft_snapshot).to_projectinputs()
    response = save(env, record, state(active=True, collection=collection_for_inputs(pi)), field=config.FIELD_ID)
    assert "field-error-banner" not in response.text, response.text[:1500]


def tables(page: str) -> dict[str, dict]:
    """{table_id|caption: {row_label: [raw floats or None]}} parsed from the rendered HTML only."""
    out: dict[str, dict] = {}
    for match in re.finditer(r'<table class="out-table" data-out-table="([^"]*)">(.*?)</table>', page, re.S):
        key = match.group(1)
        body = match.group(2)
        suffix = 0
        while key in out:
            suffix += 1
            key = f"{match.group(1)}#{suffix}"
        rows: dict[str, list] = {"__dates__": re.findall(r'<th scope="col" class="out-num"[^>]*title="([^"]*)"', body)}
        for tr in re.findall(r"<tr[^>]*>(.*?)</tr>", body, re.S):
            label = re.search(r'<th scope="row" class="out-label">(.*?)(?:<small>|</th>)', tr, re.S)
            if not label:
                continue
            cells = re.findall(r'<td [^>]*data-raw="([^"]*)"', tr)
            rows[htmllib.unescape(re.sub(r"<[^>]+>", "", label.group(1))).strip()] = [float(c) if c != "" else None for c in cells]
        out[key] = rows
    return out


def persisted_rows(rr, section, key):
    from app.workbook.runtime_projection import thaw_runtime_payload
    fs = thaw_runtime_payload(rr.financial_statements)
    return [p.get(key) for p in fs[section]["periods"]]


@pytest.mark.parametrize("kind", ["solar", "wind"])
def test_f3_two_senior_output_reconciles_to_run_evidence(env, monkeypatch, kind):
    monkeypatch.setattr(config, "ACTIVATION_ENABLED", True)
    record = env.create(kind)
    activate_f3(env, record)
    ws, rr = run(env, record)
    page = get_outputs(env, record)
    assert 'data-testid="outputs-state" data-state="CURRENT"' in page
    evidence = rr.runtime_summary["financing_evidence"]
    schedules = evidence["facility_schedules"]
    assert len(schedules) == 2 and page.count('data-testid="out-facility"') == 2
    parsed = tables(page)
    # Statements: every rendered cell equals the persisted value (None stays unavailable, never zero).
    for table_id, section, key, label in (
        ("pnl", "pnl", "revenues_keur", "Revenues"), ("pnl", "pnl", "net_income_keur", "Net Income"),
        ("balance_sheet", "balance_sheet", "senior_balance_keur", "Senior Balance"),
        ("balance_sheet", "balance_sheet", "total_assets_keur", "Total Assets"),
        ("cash_flow", "pf_cash_waterfall", "fcf_banks_keur", "FCF to Banks"),
        ("cash_flow", "pf_cash_waterfall", "senior_total_ds_keur", "Senior Total DS"),
    ):
        expected = persisted_rows(rr, section, key)
        got = parsed[table_id][label]
        assert len(got) == len(expected)
        for g, e in zip(got, expected):
            assert (g is None) == (e is None)
            if e is not None:
                assert g == pytest.approx(e, abs=1e-9)
    # Facilities are separate and reconcile to the persisted facility evidence.
    for position, schedule in enumerate(schedules, start=1):
        op = parsed["debt-%d" % position]
        for label, key in (("Interest", "interest_keur"), ("Principal", "principal_keur"),
                           ("Opening balance", "opening_keur"), ("Closing balance", "closing_keur"),
                           ("Debt service", "debt_service_keur")):
            assert op[label] == pytest.approx(schedule[key], abs=1e-9)
        assert schedule["instrument_id"] in page
    assert page.count('data-testid="out-maturity"') == 2
    assert "UNPAID OBLIGATION" not in page and page.count("fully repaid") == 2
    # Σ facility service equals the PF waterfall Senior service (persisted-vs-persisted cross-check).
    assert 'data-testid="outputs-reconciliation" data-state="RECONCILED"' in page
    # CFADS section is the PF waterfall's fcf_banks_keur.
    assert parsed["cfads"]["CFADS to Senior"] == pytest.approx(persisted_rows(rr, "pf_cash_waterfall", "fcf_banks_keur"), abs=1e-9)
    assert "Run Integrity PASS" in page


@pytest.mark.parametrize("kind", ["data_center", "ev_charging", "solar"])
def test_legacy_single_senior_output_is_truthful_about_unexposed_terms(env, kind):
    record = env.create(kind)
    ws, rr = run(env, record)
    page = get_outputs(env, record)
    assert 'data-testid="outputs-state" data-state="CURRENT"' in page
    assert page.count('data-testid="out-facility"') == 1 and "Senior debt (aggregate)" in page
    assert "Fees and construction draws are not persisted" in page
    assert "(not exposed)" in page            # opening balance is not persisted for the aggregate schedule
    from app.workbook.runtime_projection import thaw_runtime_payload
    periods = thaw_runtime_payload(rr.debt_schedule)["periods"]
    senior = tables(page)["debt-1"]
    assert senior["Closing balance"] == pytest.approx([p["senior_balance_keur"] for p in periods], abs=1e-9)
    assert all(v is None for v in senior["Opening balance"])


# ───────────────────────── freshness, integrity, authority ─────────────────────────

def test_not_run_shows_no_tables_and_no_fabricated_zero(env):
    record = env.create("solar")
    page = get_outputs(env, record)
    assert 'data-testid="outputs-state" data-state="NOT_RUN"' in page
    assert "<table" not in page and 'data-raw="0' not in page
    assert "No Run has been committed" in page


def test_stale_view_identifies_the_previous_run_and_keeps_its_values(env):
    from app.persistence.workspace_repository import get_workspace_state
    record = env.create("solar")
    ws, rr = run(env, record)
    before = tables(get_outputs(env, record))
    prior_snapshot = ws.last_runtime_snapshot_id
    page = env.client.get("/v2/workbook", params={"project": record.project_code})
    changed = env.client.post("/v2/workbook/update", headers={"HX-Request": "true"}, data=dict(
        project=record.project_code, field_id="debt.senior.target_dscr", value="1.35", sheet_id="debt", **tokens(page.text)))
    assert changed.status_code == 200 and "field-error" not in changed.text, changed.text[:800]
    assert get_workspace_state(OWNER, record.project_id).dirty
    stale = get_outputs(env, record)
    assert 'data-testid="outputs-state" data-state="STALE"' in stale
    assert "STALE — showing the PRIOR Last Run" in stale and prior_snapshot[:22] in stale
    assert tables(stale) == before            # exactly the prior Run's values, nothing recomputed


def test_current_run_can_coexist_with_integrity_fail(env, monkeypatch):
    import app.api.v1_1.institutional as institutional
    record = env.create("solar")
    run(env, record)
    real = institutional.get_run_integrity_checks

    def failing(*args, **kwargs):
        state, report = real(*args, **kwargs)
        report = dict(report)
        report["overall"] = "FAIL"
        report["checks"] = [dict(report["checks"][0], status="FAIL", title="Synthetic integrity break")] + list(report["checks"][1:])
        return state, report

    monkeypatch.setattr(institutional, "get_run_integrity_checks", failing)
    page = get_outputs(env, record)
    assert 'data-testid="outputs-state" data-state="CURRENT"' in page
    assert 'data-testid="outputs-integrity" data-state="FAIL"' in page
    assert 'data-testid="outputs-integrity-fail"' in page and "Synthetic integrity break" in page
    assert "<table" in page                       # evidence still shown unchanged


def test_unavailable_values_are_not_rendered_as_zero(env):
    record = env.create("solar")
    ws, rr = run(env, record)
    page = get_outputs(env, record)
    parsed = tables(page)
    persisted = persisted_rows(rr, "balance_sheet", "total_assets_keur")
    assert any(v is None for v in persisted), "reference run is expected to leave some BS totals unavailable"
    rendered = parsed["balance_sheet"]["Total Assets"]
    assert [v is None for v in rendered] == [v is None for v in persisted]
    assert 'class="out-num out-na"' in page and "not available</span>" in page
    # A row that is unavailable in every period is disclosed in words, not only as dashes.
    assert 'data-testid="out-unavailable-balance_sheet"' in page and "Total Assets" in page.split('data-testid="out-unavailable-balance_sheet"')[1][:400]


def test_get_is_read_only_runs_no_engine_and_is_owner_scoped(env, monkeypatch):
    from app.persistence.workspace_repository import get_workspace_state
    from financial_engine import orchestrator
    record = env.create("solar")
    run(env, record)
    before = get_workspace_state(OWNER, record.project_id)

    def forbidden(*args, **kwargs):
        pytest.fail("Opening the output workspace executed a model")
    monkeypatch.setattr(orchestrator, "run_operating_model", forbidden)
    monkeypatch.setattr(orchestrator, "run_senior_debt_model", forbidden)
    for _ in range(2):
        get_outputs(env, record)
    after = get_workspace_state(OWNER, record.project_id)
    assert after == before
    foreign = env.client.get("/v2/workbook/outputs", params={"project": "does-not-exist"})
    assert foreign.status_code == 404
    env.client.cookies.clear()
    assert env.client.get("/v2/workbook/outputs", params={"project": record.project_code}).status_code == 401


def test_navigation_and_lazy_panel_are_additive(env):
    record = env.create("solar")
    page = env.client.get("/v2/workbook", params={"project": record.project_code}).text
    assert 'id="tab-outputs"' in page and 'id="panel-outputs"' in page and "Statements &amp; Debt" in page
    assert 'hx-get="/v2/workbook/outputs?project=' in page
    # The lazy panel carries no statement data until the tab is opened.
    panel = page[page.index('id="panel-outputs"'):page.index('id="panel-trust"')]
    assert "<table" not in panel


# ───────────────────────── accessibility (static contract) ─────────────────────────

def test_tables_are_keyboard_and_screen_reader_accessible(env):
    record = env.create("solar")
    run(env, record)
    page = get_outputs(env, record)
    n_tables = page.count('<table class="out-table"')
    assert n_tables >= 7
    assert page.count("<caption") == n_tables
    assert page.count('role="region"') == n_tables and page.count('tabindex="0"') >= n_tables
    assert page.count('aria-labelledby="out-cap-') >= n_tables
    assert 'scope="col"' in page and 'scope="row"' in page
    assert all(' scope="row"' in row for row in re.findall(r"<tr[^>]*>(.*?)</tr>", page, re.S)[1:3])
    assert 'aria-label="Decimal precision"' in page and "<legend" in page
    assert page.count('<h3') >= 6 and 'aria-label="Output sections"' in page
    # Year 0 column headers are announced, not just styled.
    assert "Year 0 / opening" in page
    # Unavailable cells are announced.
    assert "not available</span>" in page


# ───────────────────────── Returns / distributions ─────────────────────────

def test_returns_and_distributions_equal_persisted_sponsor_and_distribution_evidence(env):
    from app.workbook.runtime_projection import thaw_runtime_payload
    record = env.create("solar")
    ws, rr = run(env, record)
    page = get_outputs(env, record)
    parsed = tables(page)
    dist = thaw_runtime_payload(rr.distribution_schedule)["periods"]
    assert parsed["distributions"]["Distribution"] == pytest.approx([p["distribution_keur"] for p in dist], abs=1e-9)
    sponsor = thaw_runtime_payload(rr.sponsor_schedule)["periods"]
    assert parsed["sponsor"]["Legal equity distribution"] == pytest.approx(
        [p["legal_equity_distribution_keur"] for p in sponsor], abs=1e-9)
    summary = rr.runtime_summary
    assert f"{summary['project_irr'] * 100:,.2f}%" in page
    # The Sponsor net cash flow is not persisted by the canonical sponsor schedule: stays "NOT AVAILABLE".
    assert "Sponsor Net Cash Flow" in page and "NOT AVAILABLE" in page


# ───────────────────────── XLSX parity ─────────────────────────

XLSX_MAP = {
    "P&L": {"Revenue": ("pnl", "revenues_keur"), "OPEX": ("pnl", "operating_expenses_keur"),
            "Depreciation": ("pnl", "depreciation_keur"), "EBIT": ("pnl", "ebit_keur"),
            "Senior Interest": ("pnl", "senior_interest_expense_keur")},
    "Cash Flow": {"Revenue Cash": ("pf_cash_waterfall", "revenue_cash_keur"), "OPEX Cash": ("pf_cash_waterfall", "opex_cash_keur"),
                  "EBITDA Cash": ("pf_cash_waterfall", "ebitda_cash_keur"), "Cash Tax": ("pf_cash_waterfall", "cash_tax_keur"),
                  "FCF Banks": ("pf_cash_waterfall", "fcf_banks_keur"),
                  "Senior Debt Service": ("pf_cash_waterfall", "senior_total_ds_keur")},
    "Balance Sheet": {"Net Fixed Assets": ("balance_sheet", "net_fixed_assets_keur"),
                      "Senior Debt Balance": ("balance_sheet", "senior_balance_keur")},
}
SECTION_TABLE = {"P&L": "pnl", "Cash Flow": "cash_flow", "Balance Sheet": "balance_sheet"}
UI_LABEL = {"pnl": {"revenues_keur": "Revenues", "operating_expenses_keur": "Operating Expenses",
                    "depreciation_keur": "Depreciation", "ebit_keur": "EBIT",
                    "senior_interest_expense_keur": "Senior Interest"},
            "pf_cash_waterfall": {"revenue_cash_keur": "Revenue Cash", "opex_cash_keur": "OPEX Cash",
                                  "ebitda_cash_keur": "EBITDA Cash", "cash_tax_keur": "Cash Tax",
                                  "fcf_banks_keur": "FCF to Banks", "senior_total_ds_keur": "Senior Total DS"},
            "balance_sheet": {"net_fixed_assets_keur": "Net Fixed Assets", "senior_balance_keur": "Senior Balance"}}


@pytest.mark.parametrize("kind,f3", [("solar", True), ("wind", True), ("data_center", False), ("ev_charging", False)])
def test_xlsx_export_parity_with_the_output_workspace(env, monkeypatch, kind, f3):
    import io
    import openpyxl
    monkeypatch.setattr(config, "ACTIVATION_ENABLED", True)
    record = env.create(kind)
    if f3:
        activate_f3(env, record)
    ws, rr = run(env, record)
    page = get_outputs(env, record)
    parsed = tables(page)
    exported = env.client.post("/v2/workbook/export", params={"project": record.project_code}, data={"project": record.project_code})
    assert exported.status_code == 200
    book = openpyxl.load_workbook(io.BytesIO(exported.content), data_only=True)
    compared = 0
    for sheet, rows in XLSX_MAP.items():
        sh = book[sheet]
        header = next(r for r in sh.iter_rows(values_only=True) if r and r[0] == "Metric")
        dates = [str(d) for d in header[1:]]
        by_label = {r[0]: r[1:] for r in sh.iter_rows(values_only=True) if r and isinstance(r[0], str)}
        ui = parsed[SECTION_TABLE[sheet]]
        ui_dates = ui["__dates__"]
        for xlsx_label, (section, key) in rows.items():
            ui_row = ui[UI_LABEL[section][key]]
            ui_by_date = {}
            for d, v in zip(ui_dates, ui_row):
                ui_by_date.setdefault(d, []).append(v)
            for d, xv in zip(dates, by_label[xlsx_label]):
                values = ui_by_date.get(d)
                if values is None or xv is None:
                    continue
                if all(v is not None for v in values):
                    assert any(v == pytest.approx(float(xv), abs=0.0051) for v in values), (sheet, xlsx_label, d, values, xv)
                    compared += 1
    assert compared >= 100, compared


def test_f3_facility_service_equals_the_xlsx_senior_debt_service_row(env, monkeypatch):
    import io
    import openpyxl
    monkeypatch.setattr(config, "ACTIVATION_ENABLED", True)
    record = env.create("solar")
    activate_f3(env, record)
    ws, rr = run(env, record)
    page = get_outputs(env, record)
    parsed = tables(page)
    exported = env.client.post("/v2/workbook/export", params={"project": record.project_code}, data={"project": record.project_code})
    book = openpyxl.load_workbook(io.BytesIO(exported.content), data_only=True)
    row = next(r for r in book["Senior Debt"].iter_rows(values_only=True) if r and r[0] == "Senior Debt Service")
    xlsx_nonzero = [float(v) for v in row[1:] if isinstance(v, (int, float)) and v]
    service_a, service_b = parsed["debt-1"]["Debt service"], parsed["debt-2"]["Debt service"]
    # Facilities schedule different tenors: pair by the persisted period index through the evidence.
    schedules = rr.runtime_summary["financing_evidence"]["facility_schedules"]
    total = {}
    for sched in schedules:
        for index, service in zip(sched["period_indices"], sched["debt_service_keur"]):
            total[index] = total.get(index, 0.0) + service
    assert len(service_a) == len(schedules[0]["period_indices"]) and len(service_b) == len(schedules[1]["period_indices"])
    assert sorted(round(v, 2) for v in total.values()) == pytest.approx(sorted(round(v, 2) for v in xlsx_nonzero), abs=0.0101)


# ───────────────────────── projection contract (no workspace needed) ─────────────────────────

def test_projection_fails_closed_and_never_substitutes_zero():
    from app.v2.output_workspace_projection import build_output_workspace, build_output_workspace_safe
    freshness = SimpleNamespace(state=SimpleNamespace(value="CURRENT"))
    assert build_output_workspace(runtime_result=None, workspace=None, freshness=freshness)["state"] == "NOT_RUN"
    rr = SimpleNamespace(
        financial_statements={"pnl": {"periods": [{"period": 0, "date": "2030-01-01", "revenues_keur": None},
                                                    {"period": 1, "date": "2030-07-01", "revenues_keur": 5.0}]}},
        runtime_summary={}, debt_schedule={}, distribution_schedule={}, sponsor_schedule={},
        ran_at="2030-01-01T00:00:00", snapshot_id="S")
    ws = SimpleNamespace(last_runtime_identity={}, last_runtime_composite_hash="h", dirty=False)
    out = build_output_workspace(runtime_result=rr, workspace=ws, freshness=freshness)
    pnl = out["statements"][0]
    revenue = next(r for r in pnl["rows"] if r["key"] == "revenues_keur")
    assert [c["v"] for c in revenue["cells"]] == [None, 5.0]
    assert revenue["cells"][0]["text"] == "—"
    # Absent statements are "not available", not empty-zero tables.
    assert out["statements"][1]["available"] is False and out["statements"][2]["available"] is False
    assert out["reconciliation"]["overall"] == "UNVERIFIED"
    broken = build_output_workspace_safe(runtime_result=object(), workspace=ws, freshness=freshness)
    assert broken["state"] == "UNAVAILABLE" and not broken["statements"]


def test_stale_state_requires_the_prior_run_identity_in_the_banner():
    from app.v2.output_workspace_projection import build_output_workspace
    freshness = SimpleNamespace(state=SimpleNamespace(value="STALE"))
    rr = SimpleNamespace(financial_statements={}, runtime_summary={}, debt_schedule={}, distribution_schedule={},
                         sponsor_schedule={}, ran_at="2031-05-06T07:08:09", snapshot_id="20310506T070809.000000+0000")
    ws = SimpleNamespace(last_runtime_identity={"scenario_name": "Upside"}, last_runtime_composite_hash="h", dirty=True)
    out = build_output_workspace(runtime_result=rr, workspace=ws, freshness=freshness)
    assert out["state"] == "STALE"
    assert "PRIOR Last Run" in out["banner"] and "2031-05-06 07:08:09" in out["banner"] and "Upside" in out["banner"]
    assert "20310506T070809.000000" in out["banner"]


def test_refresh_is_a_plain_listener_not_a_csp_blocked_trigger_filter():
    """An hx-trigger filter expression is evaluated with eval, which the page CSP forbids; a failed filter fires on
    every request and ping-pongs with other sheets. The output workspace must not use one."""
    from pathlib import Path
    template = Path("app/templates/v2/partials/sheet_outputs.html").read_text(encoding="utf-8")
    script = Path("static/js/model_outputs_workspace.js").read_text(encoding="utf-8")
    assert "htmx:afterRequest[" not in template and 'hx-trigger="click from:#tab-outputs"' in template
    assert "addEventListener('htmx:afterRequest'" in script and "'post'" in script
