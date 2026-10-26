"""F3.8 uses real committed Runs, not fabricated facility or statement fixtures."""
from tests.model_mutation_csrf import post_valid_mutation

import io
from dataclasses import replace

import openpyxl
import pytest

from app.v2.financing_trace_projection import build_financing_trace
from app.v2.financing_sources_uses_projection import FinancingEvidenceInvalid
from app.workbook.runtime_projection import thaw_runtime_payload
from tests.test_model_decision_workspace_v2 import env  # noqa: F401
from tests.test_financing_f37_instrument_su import _two_senior_run, _junior_run
from tests.test_model_financing_f2 import run, save


def view(rr, ws, **kw):
    return build_financing_trace(runtime_result=rr, freshness="CURRENT", snapshot_id=rr.snapshot_id,
        composite_hash=ws.last_runtime_composite_hash, scenario_id=ws.last_runtime_scenario_id,
        integrity_evidence=ws.last_integrity_evidence, **kw)


def page(env, record, **kw):
    return env.client.get('/v2/workbook/financing/trace', params={'project': record.project_code, **kw})


@pytest.mark.parametrize('kind,junior', [('solar', False), ('wind', False), ('solar', True), ('wind', True)])
def test_canonical_instrument_statement_cash_tax_and_returns(env, monkeypatch, kind, junior):
    record, ws, rr = (_junior_run(env, monkeypatch, kind) if junior else _two_senior_run(env, monkeypatch, kind))
    trace = view(rr, ws)
    assert trace['status'] == 'RECONCILED', [c for c in trace['checks'] if c['status'] != 'RECONCILED']
    assert trace['balance_audit']['status'] == 'PASS'
    assert len(trace['facilities']) == 2
    ev = thaw_runtime_payload(rr.runtime_summary['financing_evidence'])
    raw = [ev['senior_facility'], ev['junior_facility']] if junior else ev['facility_schedules']
    for actual, source in zip(trace['facilities'], raw):
        closing = source['closing_keur'][-1]
        assert actual['repayment'] == ('REPAID' if closing == 0 else 'OUTSTANDING' if closing > 0 else 'NEGATIVE_BALANCE')
        assert actual['maturity'] == source['maturity_period_index']
        assert [r['values'][1].value for r in actual['rows']] == source['interest_keur']
        assert [r['values'][4].value for r in actual['rows']] == source['closing_keur']
        assert actual['drawn'].value == sum(source['construction_draws_keur'])
    fs = thaw_runtime_payload(rr.financial_statements)
    assert [r['values'][3].value for r in trace['cash']] == [r['fcf_banks_keur'] for r in fs['pf_cash_waterfall']['periods']]
    assert all(r['values'][5].available == junior for r in trace['cash'])
    ts = thaw_runtime_payload(rr.tax_schedule)
    assert [r['values'][0].value for r in trace['tax']] == [r['tax_depreciation_audit_keur'] for r in ts['periods']]
    metrics = {m.key: m.value for m in trace['returns']}
    assert metrics['project_irr'] == rr.runtime_summary['project_irr']
    assert metrics['equity_irr'] == rr.runtime_summary['equity_irr']
    assert metrics['total_sponsor_irr'] == rr.sponsor_schedule['summary']['total_sponsor_xirr']
    response = page(env, record)
    assert response.status_code == 200, response.text
    assert 'data-status="RECONCILED"' in response.text
    assert 'Facility-level deductible interest and tax allocation: UNAVAILABLE' in response.text
    assert 'Investor-specific returns remain F4 scope' in response.text
    assert 'scope="row"' in response.text and '<caption>' in response.text and 'tabindex="0"' in response.text


@pytest.mark.parametrize('kind', ['solar', 'wind', 'data_center', 'ev_charging'])
def test_inactive_single_senior_has_no_invented_split(env, kind):
    record = env.create(kind)
    ws, rr = run(env, record)
    trace = view(rr, ws)
    assert trace['facilities'] == [] and trace['status'] == 'UNAVAILABLE'
    assert trace['cash'] and trace['returns']
    assert page(env, record).status_code == 200


@pytest.fixture
def two(env, monkeypatch):
    return _two_senior_run(env, monkeypatch, 'solar')


@pytest.mark.parametrize('change', ['length', 'duplicate', 'gap', 'maturity', 'nan', 'dates', 'binding'])
def test_malformed_axis_and_binding_fail_closed(two, change):
    record, ws, rr = two
    rs = thaw_runtime_payload(rr.runtime_summary)
    fs = thaw_runtime_payload(rr.financial_statements)
    f = rs['financing_evidence']['facility_schedules'][0]
    if change == 'length':
        f['interest_keur'].pop()
    elif change == 'duplicate':
        f['period_indices'][1] = f['period_indices'][0]
    elif change == 'gap':
        f['period_indices'].pop(1)
        for k in ('opening_keur', 'interest_keur', 'principal_keur', 'debt_service_keur', 'closing_keur'):
            f[k].pop(1)
    elif change == 'maturity':
        f['maturity_period_index'] += 1
    elif change == 'nan':
        f['principal_keur'][0] = float('nan')
    elif change == 'dates':
        fs['pnl']['periods'][1]['date'] = '2099-01-01'
    else:
        rs['financing_evidence']['run_binding']['scenario_id'] = 'wrong-scenario'
    with pytest.raises(FinancingEvidenceInvalid):
        view(replace(rr, runtime_summary=rs, financial_statements=fs), ws)


@pytest.mark.parametrize('change', ['interest', 'closing', 'service', 'cash', 'statement'])
def test_tampered_values_are_not_repaired(two, change):
    _, ws, rr = two
    rs = thaw_runtime_payload(rr.runtime_summary)
    fs = thaw_runtime_payload(rr.financial_statements)
    audit = thaw_runtime_payload(ws.last_integrity_evidence)
    if change == 'statement':
        fs['pnl']['periods'][1]['senior_interest_expense_keur'] += 1
    elif change == 'cash':
        index = audit['senior_debt']['periods'][0]['period_index']
        next(r for r in audit['cash'] if r['period_index'] == index)['cash_after_senior_before_reserves'] += 1
    else:
        rs['financing_evidence']['facility_schedules'][0][change + '_keur' if change != 'service' else 'debt_service_keur'][0] += 1
    altered = replace(rr, runtime_summary=rs, financial_statements=fs)
    from app.run_integrity import evidence_digest
    audit['digest'] = evidence_digest(audit)
    w = replace(ws, last_integrity_evidence=audit)
    assert view(altered, w)['status'] == 'DIFFERS'
    assert altered.runtime_summary == rs


def test_not_run_context_and_existing_surface_loaders(env):
    record = env.create('solar')
    response = page(env, record)
    assert response.status_code == 200 and 'data-state="NOT_RUN"' in response.text
    assert 'data-testid="f38-facility"' not in response.text
    workbook = env.client.get('/v2/workbook', params={'project': record.project_code}).text
    assert workbook.count('data-testid="financing-trace-loader"') >= 2


def test_current_stale_history_and_last_run_immutability(two, env):
    from app.persistence.run_history_repository import get_run_history
    record, ws, rr = two
    h = get_run_history('decision-user', record.project_id, limit=1)[0]
    original = page(env, record, history_id=h.history_id).text
    assert 'data-state="HISTORICAL"' in original
    current = page(env, record).text
    assert 'data-state="CURRENT"' in current
    import json
    from app.workbook import multisenior_config as config
    from tests.test_model_financing_f3_workspace import collection_for_inputs, state
    from app.workbook.input_set import ProjectInputSet
    new = json.loads(state(active=True, collection=collection_for_inputs(ProjectInputSet.from_snapshot(ws.draft_snapshot).to_projectinputs())))
    new['scopes']['base']['proposal']['instruments'][0]['interest']['fixed_rate'] = .07
    new['scopes']['base']['activation']['proposal_digest'] = 'BIND_ON_SAVE'
    response = save(env, record, json.dumps(new), field=config.FIELD_ID)
    assert response.status_code == 200 and 'field-error' not in response.text
    stale = page(env, record).text
    assert 'data-state="STALE"' in stale and rr.snapshot_id in stale
    run(env, record)
    assert page(env, record, history_id=h.history_id).text.split('</nav>')[1] == original.split('</nav>')[1]


def test_owner_project_and_history_isolation_and_read_only(two, env, monkeypatch):
    from app.auth import COOKIE_NAME, create_session_token
    from app.persistence.run_history_repository import get_run_history
    from app.persistence.workspace_repository import get_workspace_state
    from app.services import production_financial_authority
    record, ws, _ = two
    h = get_run_history('decision-user', record.project_id, limit=1)[0]
    other = env.create('wind')
    assert page(env, other, history_id=h.history_id).status_code == 404
    monkeypatch.setattr(production_financial_authority, 'run_clean_production', lambda *a, **k: pytest.fail('GET ran engine'))
    assert page(env, record).status_code == 200
    assert get_workspace_state('decision-user', record.project_id) == ws
    env.client.cookies.set(COOKIE_NAME, create_session_token(user_id='f38-unrelated', username='admin'))
    assert page(env, record).status_code == 404
    env.client.cookies.clear()
    assert page(env, record).status_code == 401


def test_concurrent_context_change_cannot_serve_trace(two, env, monkeypatch):
    from app.v2.post_run_context import PostRunRequestContext, PostRunContextChanged
    monkeypatch.setattr(PostRunRequestContext, 'validate_current', lambda *a: (_ for _ in ()).throw(PostRunContextChanged('changed')))
    response = page(env, two[0])
    assert response.status_code == 200
    assert 'data-testid="financing-trace-unavailable"' in response.text
    assert 'UNAVAILABLE' in response.text and 'f38-facility' not in response.text and 'data-raw=' not in response.text


@pytest.mark.parametrize('junior', [False, True])
def test_xlsx_uses_same_persisted_statement_cells_and_no_engine(env, monkeypatch, junior):
    from app.services import production_financial_authority
    record, ws, rr = _junior_run(env, monkeypatch, 'solar') if junior else _two_senior_run(env, monkeypatch, 'solar')
    trace = view(rr, ws)
    monkeypatch.setattr(production_financial_authority, 'run_clean_production', lambda *a, **k: pytest.fail('export ran engine'))
    response = post_valid_mutation(env.client, '/v2/workbook/export', params={'project': record.project_code}, data={'project': record.project_code})
    assert response.status_code == 200
    book = openpyxl.load_workbook(io.BytesIO(response.content), data_only=True)
    rows = {r[0]: r[1:] for r in book['Cash Flow'].iter_rows(values_only=True) if r and isinstance(r[0], str)}
    headers = rows['Metric']
    assert [str(d) for d in headers[:len(trace['cash'])]] == [r['date'] for r in trace['cash']]
    # Existing XLSX statement binding uses exactly the same persisted 2dp values.
    for label, index in [('Revenue Cash', 0), ('EBITDA Cash', 1), ('Cash Tax', 2), ('FCF Banks', 3), ('Senior Debt Service', 4)]:
        exported = list(rows[label][:len(trace['cash'])])
        canonical = [r['values'][index].value for r in trace['cash']]
        assert exported == canonical
    fs = thaw_runtime_payload(rr.financial_statements)
    for sheet, section, mappings in (
        ('P&L', 'pnl', [('Senior Interest', 'senior_interest_expense_keur')] +
         ([('Junior Interest', 'junior_interest_expense_keur')] if junior else [])),
        ('Balance Sheet', 'balance_sheet', [('Senior Debt Balance', 'senior_balance_keur')] +
         ([('Junior Debt Balance', 'junior_balance_keur')] if junior else [])),
        ('Cash Flow', 'pf_cash_waterfall', [('Junior Debt Service', 'junior_debt_service_keur')] if junior else []),
    ):
        values = {r[0]: r[1:] for r in book[sheet].iter_rows(values_only=True) if r and isinstance(r[0], str)}
        for label, field in mappings:
            canonical = [r[field] for r in fs[section]['periods']]
            assert list(values[label][:len(canonical)]) == canonical


def test_projection_imports_no_engine_or_writer():
    import ast
    from pathlib import Path
    import app.v2.financing_trace_projection as module
    tree = ast.parse(Path(module.__file__).read_text(encoding='utf-8'))
    imports = [n.module for n in ast.walk(tree) if isinstance(n, ast.ImportFrom) and n.module]
    assert not [m for m in imports if m.startswith(('financial_engine', 'finco_core', 'domain', 'app.persistence'))]


def test_integrity_digest_is_checked_before_use(two):
    _, ws, rr = two
    evidence = thaw_runtime_payload(ws.last_integrity_evidence)
    evidence['cash'][0]['cash_after_senior_before_reserves'] += 1
    with pytest.raises(FinancingEvidenceInvalid, match='digest'):
        view(rr, replace(ws, last_integrity_evidence=evidence))


@pytest.mark.parametrize('section', ['financial_statements', 'tax_schedule', 'pnl', 'integrity'])
def test_wrong_evidence_containers_are_typed_unavailable_not_attribute_errors(two, section):
    _, ws, rr = two
    if section == 'integrity':
        ws = replace(ws, last_integrity_evidence=['invalid'])
    elif section == 'pnl':
        fs = thaw_runtime_payload(rr.financial_statements)
        fs['pnl'] = ['invalid']
        rr = replace(rr, financial_statements=fs)
    else:
        rr = replace(rr, **{section: ['invalid']})
    with pytest.raises(FinancingEvidenceInvalid, match='object'):
        view(rr, ws)


def test_strict_template_context_contract_for_all_run_states(two):
    from jinja2 import StrictUndefined
    from app.v2.router import _templates
    record, ws, rr = two
    template = _templates.env.overlay(undefined=StrictUndefined).get_template('partials/_financing_trace.html')
    for state in ('CURRENT', 'STALE', 'HISTORICAL', 'NOT_RUN'):
        trace = view(rr, ws, historical=state == 'HISTORICAL') if state != 'NOT_RUN' else build_financing_trace(
            runtime_result=None, freshness=state, snapshot_id='', composite_hash=None, scenario_id=None)
        if state == 'STALE':
            trace['state'] = state
        html = template.render(trace=trace, project_code=record.project_code,
            run_scenario='Base', run_at=rr.ran_at if state != 'NOT_RUN' else '', history=[], history_id=None)
        assert f'data-state="{state}"' in html
