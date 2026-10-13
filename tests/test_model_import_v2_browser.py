"""Import V2 in a REAL authenticated Chromium: intelligent mapping -> review -> approve -> atomic Apply -> manual Run.

Synthetic generic workbooks only; European numeric formats, several sheets, structured/unsupported rows, bulk
selection limited to validated rows, a user-modified mapping, and Apply -> STALE -> manual Run -> CURRENT.
"""
from __future__ import annotations

import io
import os
from pathlib import Path

import pytest

pytest.importorskip("playwright")
pytest.importorskip("uvicorn")
from openpyxl import Workbook  # noqa: E402

from tests.test_model_ai_import_browser_c import (  # noqa: E402,F401
    OWNER, _new_seeded_browser_project, _page, _run_in_real_browser_context, app_server, browser,
)

SHOTS = Path(os.environ.get("IMPORT_V2_SCREENSHOT_DIR", "artifacts/model-import-v2"))
SHOTS.mkdir(parents=True, exist_ok=True)          # the CI evidence contract requires these captures

# per vertical: (revenue-sheet row, expected canonical field, expected stored value)
REVENUE = {
    "solar": (("PPA tariff", "0,062", "EUR/kWh"), "revenue.ppa.base_tariff", 62.0),
    "wind": (("PPA tariff", "0,058", "EUR/kWh"), "revenue.ppa.base_tariff", 58.0),
    "data_center": (("Occupancy Y2", "70 %", ""), "revenue.data_center.occupancy_y2", 70.0),
    "ev_charging": (("Charging price", "0,38", "EUR/kWh"), "revenue.ev_charging.charging_price", 0.38),
}


def _workbook(kind, path):
    book = Workbook()
    capex = book.active
    capex.title = "CAPEX"
    capex.append(["Assumption", "Value", "Unit"])
    capex.append(["Grid interconnection cost", "1.200.000,50", "EUR"])        # ranked, European format, EUR -> kEUR
    capex.append(["Audit and legal", 45.5, "kEUR"])                          # ranked
    capex.append(["Transformer", 120, "kEUR"])                               # CAPEX line item: read-only
    opex = book.create_sheet("OPEX")
    opex.append(["Assumption", "Value", "Unit"])
    opex.append(["Insurance (annual)", "42,5", "kEUR/yr"])                    # ranked into OPEX, not CAPEX
    opex.append(["Guard services", "18000", "EUR/yr"])                        # unmapped: the user maps it by hand
    revenue = book.create_sheet("Revenue")
    revenue.append(["Assumption", "Value", "Unit"])
    revenue.append(list(REVENUE[kind][0]))
    revenue.append(["Merchant price curve", 55, "EUR/MWh"])                  # Revenue V2 contract: read-only
    financing = book.create_sheet("Financing")
    financing.append(["Assumption", "Value", "Unit"])
    financing.append(["Cost of debt", "5,25 %", ""])                          # ranked, percent suffix
    financing.append(["Junior debt amount", 500, "kEUR"])                    # F3.5 instrument: read-only
    technical = book.create_sheet("Technical")
    technical.append(["Assumption", "Value", "Unit"])
    technical.append(["Nameplate capacity", 20, "MW"])                       # recognised but not batch-applyable
    technical.append(["Weather station", "n/a", ""])                         # unsupported
    stream = io.BytesIO()
    book.save(stream)
    path.write_bytes(stream.getvalue())


def _snapshot_value(project, field_id):
    from app.persistence.workspace_repository import get_workspace_state
    from app.workbook.input_set import ProjectInputSet
    from app.workbook.registry import WORKBOOK
    ws = get_workspace_state(OWNER, project.project_id)
    return ProjectInputSet.from_snapshot(ws.draft_snapshot, workbook=WORKBOOK).get(field_id)


def _row(page, label):
    return page.locator("tr.review-row", has_text=label).first


@pytest.mark.parametrize("kind", ["solar", "wind", "data_center", "ev_charging"])
def test_intelligent_import_end_to_end(browser, app_server, tmp_path, kind):
    from app.persistence.run_history_repository import get_run_history
    from app.persistence.workspace_repository import get_workspace_state

    project = _new_seeded_browser_project(kind, f3_active=False)
    ctx, page = _page(browser, app_server["url"], width=1440)
    try:
        initial = _run_in_real_browser_context(ctx, page, app_server["url"], project)
        history_before = tuple(get_run_history(OWNER, project.project_id))
        before_values = {f: _snapshot_value(project, f) for f in (
            "capex.C.grid_connection", "capex.D.audit_legal", "opex.lines.insurance", "opex.lines.security",
            REVENUE[kind][1], "debt.senior.interest_rate_pct")}
        path = tmp_path / f"synthetic-{kind}.xlsx"
        _workbook(kind, path)

        page.goto(app_server["url"] + "/v2/workbook/import?project=" + project.project_code)
        page.locator("#upload").set_input_files(str(path))
        page.get_by_role("button", name="Extract and review").click()
        assert page.get_by_text("Review detected assumptions").count() == 1
        # Upload and review never change canonical state or start a Run.
        assert get_workspace_state(OWNER, project.project_id).draft_snapshot == initial.draft_snapshot

        # Suggestions are explained and categorised, never pre-approved.
        assert page.locator(".approve-box:checked").count() == 0
        grid = _row(page, "Grid interconnection cost")
        assert grid.get_attribute("data-category") == "CAPEX" and grid.get_attribute("data-status") == "needs_review"
        assert "ranked_candidate" in grid.inner_text() and "label terms matched" in grid.inner_text()
        assert _row(page, "Insurance (annual)").get_attribute("data-category") == "OPEX"
        for label, text in (("Transformer", "UNSUPPORTED_FOR_APPLY"), ("Merchant price curve", "Revenue contract editor"),
                            ("Junior debt amount", "Capital Stack"), ("Nameplate capacity", "not applyable")):
            row = _row(page, label)
            assert row.get_attribute("data-status") in ("unsupported_for_apply", "unsupported"), label
            assert text in row.inner_text(), (label, row.inner_text())
            assert row.locator(".approve-box").is_disabled() or row.get_attribute("data-status") == "unsupported"
        assert _row(page, "Weather station").get_attribute("data-status") == "unsupported"
        page.screenshot(path=str(SHOTS / f"{kind}-01-review-desktop.png"), full_page=True)

        # Category filter, attention queue, sort.
        page.locator('.cat-filter[data-category="OPEX"]').click()
        visible = page.locator("tr.review-row:visible")
        assert visible.count() >= 1 and all(c == "OPEX" for c in visible.evaluate_all("r => r.map(x => x.dataset.category)"))
        assert page.locator("tr.review-row", has_text="Guard services").get_attribute("data-category") == "UNMAPPED"
        page.locator('.cat-filter[data-category="all"]').click()
        page.locator('.filter[data-filter="attention"]').click()
        queued = page.locator("tr.review-row:visible").evaluate_all("r => r.map(x => x.dataset.attention)")
        assert queued and set(queued) == {"true"}
        page.locator('.filter[data-filter="all"]').click()
        page.locator("#sort-rows").select_option("confidence")
        order = page.locator("tr.review-row").evaluate_all("r => r.map(x => x.dataset.confidence)")
        rank = {"exact": 0, "exact_label": 1, "high": 2, "medium": 3, "low": 4, "none": 5}
        assert [rank[c] for c in order] == sorted(rank[c] for c in order)

        # Bulk selection takes ONLY independently validated rows; read-only / unmapped stay unchecked.
        page.get_by_test_id("select-validated").click()
        checked = page.locator(".approve-box:checked").count()
        assert checked >= 4
        for label in ("Transformer", "Merchant price curve", "Junior debt amount", "Nameplate capacity",
                      "Weather station", "Guard services"):
            assert not _row(page, label).locator(".approve-box").is_checked(), label

        # A user-modified mapping: Guard services -> Security (OPEX); Audit and legal is rejected by unchecking it.
        guard = _row(page, "Guard services")
        guard.locator('select[name^="field_"]').select_option("opex.lines.security")
        guard.locator(".approve-box").check()
        _row(page, "Audit and legal").locator(".approve-box").uncheck()
        approved_count = page.locator(".approve-box:checked").count()
        page.get_by_role("button", name="Review approved change set").click()
        assert page.get_by_text("Final confirmation", exact=False).count() == 1
        assert "manual_reviewed" in page.locator("table").first.inner_text()
        assert "->" in page.locator("table").first.inner_text()        # normalisation decision is shown
        # Review is not a write.
        assert get_workspace_state(OWNER, project.project_id).draft_snapshot == initial.draft_snapshot
        page.screenshot(path=str(SHOTS / f"{kind}-02-confirm.png"), full_page=True)

        page.get_by_role("button", name=f"Apply {approved_count} approved fields atomically").click()
        assert page.get_by_text("Import Apply result").count() == 1
        assert page.get_by_test_id("provenance").count() == 1
        assert "USER_CONFIRMED_REVIEW" in page.get_by_test_id("provenance").inner_text()
        page.screenshot(path=str(SHOTS / f"{kind}-03-applied-stale.png"), full_page=True)

        # Exact canonical values: only approved rows changed; normalisation (EUR->kEUR, kWh->MWh, %) applied.
        assert float(_snapshot_value(project, "capex.C.grid_connection")) == pytest.approx(1200.0005)
        assert float(_snapshot_value(project, "opex.lines.insurance")) == pytest.approx(42.5)
        assert float(_snapshot_value(project, "opex.lines.security")) == pytest.approx(18.0)
        assert float(_snapshot_value(project, REVENUE[kind][1])) == pytest.approx(REVENUE[kind][2])
        assert float(_snapshot_value(project, "debt.senior.interest_rate_pct")) == pytest.approx(5.25)
        assert _snapshot_value(project, "capex.D.audit_legal") == before_values["capex.D.audit_legal"]     # rejected
        stale = get_workspace_state(OWNER, project.project_id)
        assert stale.dirty and stale.last_runtime_snapshot_id == initial.last_runtime_snapshot_id
        assert len(get_run_history(OWNER, project.project_id)) == len(history_before)

        # Manual Run -> CURRENT; earlier Run History entries are immutable.
        current = _run_in_real_browser_context(ctx, page, app_server["url"], project)
        assert not current.dirty and current.last_runtime_snapshot_id != initial.last_runtime_snapshot_id
        history_after = tuple(get_run_history(OWNER, project.project_id))
        assert len(history_after) == len(history_before) + 1
        assert [h.history_id for h in history_after[-len(history_before):]] == [h.history_id for h in history_before] or \
            {h.history_id for h in history_before} <= {h.history_id for h in history_after}
        by_id = {h.history_id: h for h in history_after}
        for old in history_before:
            assert by_id[old.history_id].runtime_summary == old.runtime_summary
    finally:
        ctx.close()


def test_mobile_layout_has_no_page_overflow_and_keeps_bulk_controls(browser, app_server, tmp_path):
    project = _new_seeded_browser_project("solar", f3_active=False)
    ctx, page = _page(browser, app_server["url"], width=390)
    try:
        path = tmp_path / "synthetic-mobile.xlsx"
        _workbook("solar", path)
        page.goto(app_server["url"] + "/v2/workbook/import?project=" + project.project_code)
        page.locator("#upload").set_input_files(str(path))
        page.get_by_role("button", name="Extract and review").click()
        assert page.get_by_test_id("select-validated").is_visible()
        assert page.evaluate("document.documentElement.scrollWidth - document.documentElement.clientWidth") <= 1
        page.screenshot(path=str(SHOTS / "solar-04-review-mobile.png"), full_page=True)
        assert not page.browser_errors
    finally:
        ctx.close()


def test_user_unit_correction_is_shown_as_user_reviewed_from_review_to_result(browser, app_server, tmp_path):
    from app.persistence.run_history_repository import get_run_history
    from app.persistence.workspace_repository import get_workspace_state

    project = _new_seeded_browser_project("solar", f3_active=False)
    ctx, page = _page(browser, app_server["url"], width=1440)
    try:
        initial = _run_in_real_browser_context(ctx, page, app_server["url"], project)
        history = tuple(h.history_id for h in get_run_history(OWNER, project.project_id))
        path = tmp_path / "synthetic-unit-correction.xlsx"
        _workbook("solar", path)
        page.goto(app_server["url"] + "/v2/workbook/import?project=" + project.project_code)
        page.locator("#upload").set_input_files(str(path))
        page.get_by_role("button", name="Extract and review").click()

        grid = _row(page, "Grid interconnection cost")                    # extracted unit: EUR (value 1.200.000,50)
        audit = _row(page, "Audit and legal")                             # extracted unit: kEUR, left untouched
        grid.locator(".approve-box").check()
        audit.locator(".approve-box").check()
        grid.locator('select[name^="unit_"]').select_option("kEUR")        # the user overrides the extracted unit
        page.get_by_role("button", name="Review approved change set").click()
        assert page.get_by_text("Final confirmation", exact=False).count() == 1

        rows = page.locator("table").first.locator("tbody tr")
        corrected = rows.filter(has_text="Grid interconnection cost")
        untouched = rows.filter(has_text="Audit and legal")
        assert corrected.get_by_test_id("unit-corrected").count() == 1
        assert "extracted unit EUR" in corrected.inner_text() and "Unit used: kEUR" in corrected.inner_text()
        assert corrected.get_by_test_id("mapping-origin").inner_text().strip() == "user reviewed"
        assert "from EUR to kEUR" in corrected.get_by_test_id("correction-note").inner_text()
        assert untouched.get_by_test_id("unit-corrected").count() == 0
        assert untouched.get_by_test_id("mapping-origin").inner_text().strip() == "automatic"
        # Review is not a write.
        assert get_workspace_state(OWNER, project.project_id).draft_snapshot == initial.draft_snapshot
        page.screenshot(path=str(SHOTS / "solar-05-unit-correction-confirm.png"), full_page=True)

        page.get_by_role("button", name="Apply 2 approved fields atomically").click()
        assert page.get_by_text("Import Apply result").count() == 1
        result = page.locator("table").first.locator("tbody tr").filter(has_text="Grid interconnection cost")
        assert result.get_by_test_id("unit-corrected").count() == 1
        assert result.get_by_test_id("mapping-origin").inner_text().strip() == "user reviewed"
        page.screenshot(path=str(SHOTS / "solar-06-unit-correction-result.png"), full_page=True)

        # The canonical value is exactly what the user approved (kEUR), not the automatic EUR conversion.
        assert float(_snapshot_value(project, "capex.C.grid_connection")) == pytest.approx(1200000.50)
        stale = get_workspace_state(OWNER, project.project_id)
        assert stale.dirty and stale.last_runtime_snapshot_id == initial.last_runtime_snapshot_id
        assert tuple(h.history_id for h in get_run_history(OWNER, project.project_id)) == history
    finally:
        ctx.close()
