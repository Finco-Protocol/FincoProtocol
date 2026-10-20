"""Joint acceptance of three persisted-evidence readers in one application."""
import io
import sqlite3
from unittest import mock
from urllib.parse import parse_qs, urlsplit

import openpyxl
from bs4 import BeautifulSoup

from tests.test_model_decision_support_v1 import seeded_db, two_runs  # noqa: F401


def _stored_rows():
    from app.persistence import db
    with sqlite3.connect(db.DB_PATH) as connection:
        return {table: connection.execute(f"SELECT * FROM {table} ORDER BY {key}").fetchall()
                for table, key in (("projects", "project_id"),
                                   ("workspace_states", "workspace_id"),
                                   ("model_run_history", "history_id"))}


def test_all_three_routes_registered_once_in_same_application(two_runs):
    client, _, _, _ = two_runs
    for path in ("/v2/portfolio", "/v2/workbook/run-history/export",
                 "/v2/workbook/run-workspace-v3"):
        routes = [route for route in client.app.routes
                  if getattr(route, "path", None) == path and "GET" in route.methods]
        assert len(routes) == 1, path


def test_joint_history_export_portfolio_and_workspace_are_read_only(two_runs):
    from app.auth import COOKIE_NAME, create_session_token
    from app.persistence.run_history_repository import get_run_history_entry

    client, cookies, record, ids = two_runs
    before = _stored_rows()
    selected = get_run_history_entry("ds-user", record.project_id, ids[-1])
    with mock.patch("app.api.project_runner.run_project", side_effect=AssertionError("GET ran engine")), \
         mock.patch("app.services.production_financial_authority.run_clean_production",
                    side_effect=AssertionError("GET ran engine")), \
         mock.patch("app.persistence.workspace_repository.v2_atomic_run_commit",
                    side_effect=AssertionError("GET committed Run")):
        portfolio = client.get("/v2/portfolio", cookies=cookies)
        assert portfolio.status_code == 200
        doc = BeautifulSoup(portfolio.text, "html.parser")
        row = doc.select_one(f'[data-project="{record.project_code}"]')
        assert row is not None and doc.select_one('[data-testid="portfolio-capital"]')
        history_url = next(link["href"] for link in row.select("a[href]")
                           if "/run-history/detail?" in link["href"])
        credit_url = next(link["href"] for link in row.select("a[href]")
                          if "/credit-decision?" in link["href"])
        assert parse_qs(urlsplit(history_url).query)["history_id"] == [ids[0]]
        assert parse_qs(urlsplit(credit_url).query)["run"] == [ids[0]]
        assert client.get(history_url, cookies=cookies).status_code == 200
        assert client.get(credit_url, cookies=cookies).status_code == 200

        query = {"project": record.project_code, "history_id": ids[-1]}
        export = client.get("/v2/workbook/run-history/export", params=query, cookies=cookies)
        repeat = client.get("/v2/workbook/run-history/export", params=query, cookies=cookies)
        assert export.status_code == repeat.status_code == 200
        assert export.content == repeat.content
        book = openpyxl.load_workbook(io.BytesIO(export.content))
        assert "Executive_Overview" in book.sheetnames
        overview = {cells[0]: cells[1] for cells in book["Executive_Overview"].iter_rows(values_only=True)
                    if cells[0] is not None and len(cells) > 1}
        assert overview["Runtime snapshot ID"] == selected.runtime_snapshot_id
        assert str(overview["Run state"]).startswith("HISTORICAL")

        workspace = client.get("/v2/workbook/run-workspace-v3",
                               params={"project": record.project_code, "run": ids,
                                       "reference": ids[-1]}, cookies=cookies)
        assert workspace.status_code == 200
        chips = BeautifulSoup(workspace.text, "html.parser").select(".v3-rw-identity .v3-rw-chip")
        assert [chip.get_text(strip=True) for chip in chips] == ["HISTORICAL", "HISTORICAL"]
        for path, params in (("/v2/workbook/run-history/export", query),
                             ("/v2/workbook/run-workspace-v3", {"project": record.project_code})):
            foreign = {COOKIE_NAME: create_session_token(user_id="rc-other-owner", username="admin")}
            assert client.get(path, params=params, cookies=foreign).status_code == 404
    assert _stored_rows() == before
