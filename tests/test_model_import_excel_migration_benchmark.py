"""Synthetic Excel-migration benchmark (original data, generic labels).

Two parts:
  1. a plan-level corpus: classification / scenario outcomes against declared expectations;
  2. an applied benchmark per vertical: assess -> Apply -> manual Run -> reconcile, with exact canonical values.

It is a regression guard on a small synthetic set, NOT a claim about arbitrary institutional workbooks.
"""
from __future__ import annotations

from collections import Counter
from decimal import Decimal

import pytest
from openpyxl import Workbook

from app.persistence.run_history_repository import get_run_history
from app.model_import.migration_plan import DIRECT, FORMULA, OUTPUT, STRUCTURED, UNSUPPORTED, build_plan
from app.model_import.workbook_inspect import inspect_workbook
from app.workbook.input_set import ProjectInputSet
from app.workbook.registry import WORKBOOK
from tests import migration_fixtures as F
from tests.migration_routes_support import (
    FIELDS, assess, confirm, hidden, ids_for, recon_detail, recon_rows, reconcile, review_slice, run_model, slice_ticket, apply, ws_of,
)
from tests.test_model_decision_workspace_v2 import env  # noqa: F401

ARGS = dict(project_type="solar", template_source="generic_solar_reference", project_code="P-1",
            target_scenario_id=None, target_scenario_name=None)

REVENUE = {
    "solar": (("PPA tariff", "0,062", "EUR/kWh"), "revenue.ppa.base_tariff", 62.0),
    "wind": (("PPA tariff", "0,058", "EUR/kWh"), "revenue.ppa.base_tariff", 58.0),
    "data_center": (("Occupancy Y2", "70 %", ""), "revenue.data_center.occupancy_y2", 70.0),
    "ev_charging": (("Charging price", "0,38", "EUR/kWh"), "revenue.ev_charging.charging_price", 0.38),
}


def european_workbook(kind: str) -> bytes:
    book = Workbook()
    capex = book.active
    capex.title = "CAPEX"
    capex.append(["Assumption", "Value", "Unit"])
    capex.append(["Grid interconnection cost", "1.200.000,50", "EUR"])
    capex.append(["Audit and legal", 45.5, "kEUR"])
    capex.append(["Transformer", 120, "kEUR"])
    opex = book.create_sheet("OPEX")
    opex.append(["Assumption", "Value", "Unit"])
    opex.append(["Insurance (annual)", "42,5", "kEUR/yr"])
    opex.append(["Guard services", "18000", "EUR/yr"])
    revenue = book.create_sheet("Revenue")
    revenue.append(["Assumption", "Value", "Unit"])
    revenue.append(list(REVENUE[kind][0]))
    revenue.append(["Merchant price curve", 55, "EUR/MWh"])
    financing = book.create_sheet("Financing")
    financing.append(["Assumption", "Value", "Unit"])
    financing.append(["Cost of debt", "5,25 %", ""])
    financing.append(["Junior debt amount", 500, "kEUR"])
    technical = book.create_sheet("Technical")
    technical.append(["Assumption", "Value", "Unit"])
    technical.append(["Nameplate capacity", 20, "MW"])
    technical.append(["Weather station", "n/a", ""])
    outputs = book.create_sheet("Results")
    outputs.append(["Metric", "Value", "Unit"])
    outputs.append(["Total CAPEX", "=CAPEX!B2/1000+CAPEX!B3", "kEUR"])
    return F.add_cached_values(F.save(book), {("Results", "B2"): 1245.5005})


def corpus():
    """(name, bytes, filename, expectations).  expectations: overall scenario state, eligible slice keys, labels -> class."""
    base = F.save(F.scenario_columns_workbook())
    hidden_ext = Workbook()
    ws = hidden_ext.active
    ws.title = "Inputs"
    ws.append(["Assumption", "Value", "Unit"])
    ws.append(["Insurance (annual)", 40, "kEUR/yr"])
    calc = hidden_ext.create_sheet("Calc")
    calc.sheet_state = "hidden"
    calc.append(["Escalation factor", 1.02])
    formulas = Workbook()
    f = formulas.active
    f.title = "Inputs"
    f.append(["Assumption", "Value", "Unit"])
    f.append(["Insurance (annual)", "=2*20", "kEUR/yr"])
    f.append(["Cost of debt", "=B2/1000", "%"])
    f.append(["PPA tariff", "=SUM(B2:B3)", "EUR/MWh"])
    probability = Workbook()
    del probability["Sheet"]
    for title, tariff in (("Inputs_P50", 60), ("Inputs_P90", 55)):
        s = probability.create_sheet(title)
        s.append(["Assumption", "Value", "Unit"])
        s.append(["PPA tariff", tariff, "EUR/MWh"])
    periodic = Workbook()
    p = periodic.active
    p.title = "Inputs"
    p.append(["Assumption", "Value", "Unit"])
    p.append(["Insurance (annual)", 40, "kEUR/yr"])
    series = periodic.create_sheet("Series")
    series.append(["Year", 2025, 2026, 2027])
    series.append(["Generation", 100, 101, 102])
    xlsm = F.to_xlsm(base, F.build_vba_container(["ThisWorkbook", "Module1"]))
    return [
        ("scenario_columns", base, "a.xlsx", dict(overall="STATIC_OBSERVABLE", slices=["base_case", "upside", "downside"],
                                                  classes={"Cost of debt": DIRECT, "Audit and legal": DIRECT, "Transformer": STRUCTURED,
                                                           "Merchant price curve": STRUCTURED, "Target leverage": FORMULA,
                                                           "Weather station": UNSUPPORTED, "Total CAPEX": OUTPUT, "Equity IRR": OUTPUT})),
        ("scenario_columns_xlsm", xlsm, "a.xlsm", dict(overall="STATIC_OBSERVABLE", slices=["base_case", "upside", "downside"],
                                                       classes={"Cost of debt": DIRECT, "Transformer": STRUCTURED, "Total CAPEX": OUTPUT})),
        ("selector_with_case_columns", F.save(F.selector_workbook()), "b.xlsx", dict(overall="SCENARIO_LOGIC_UNRESOLVED",
                                                                                     slices=["case_1", "case_2", "case_3"], classes={})),
        ("selector_only", F.save(F.selector_only_workbook()), "c.xlsx", dict(overall="SCENARIO_LOGIC_UNRESOLVED", slices=["COMMON"],
                                                                             classes={"Cost of debt": DIRECT, "Scenario selector": UNSUPPORTED})),
        ("macro_selector", F.to_xlsm(F.save(F.selector_only_workbook()), F.build_vba_container(["Module1"])), "c.xlsm",
         dict(overall="SCENARIO_LOGIC_UNRESOLVED", slices=["COMMON"], classes={"Cost of debt": DIRECT})),
        ("hidden_external_named", F.add_external_link(F.add_defined_names(F.save(hidden_ext), {"Esc": "Calc!$B$1"})), "d.xlsx",
         dict(overall="NONE_DETECTED", slices=["SINGLE_CASE"], classes={"Insurance (annual)": DIRECT})),
        ("formula_only_inputs", F.save(formulas), "e.xlsx", dict(overall="NONE_DETECTED", slices=[], classes={
            "Insurance (annual)": FORMULA, "Cost of debt": FORMULA, "PPA tariff": FORMULA})),
        ("probability_sheets", F.save(probability), "f.xlsx", dict(overall="STATIC_OBSERVABLE", slices=["p50", "p90"], classes={})),
        ("period_tables", F.save(periodic), "g.xlsx", dict(overall="NONE_DETECTED", slices=["SINGLE_CASE"], classes={"Insurance (annual)": DIRECT})),
    ] + [(f"european_{k}", european_workbook(k), "h.xlsx", dict(overall="NONE_DETECTED", slices=["SINGLE_CASE"], classes={
        "Grid interconnection cost": DIRECT, "Transformer": STRUCTURED, "Merchant price curve": STRUCTURED, "Junior debt amount": STRUCTURED,
        "Guard services": UNSUPPORTED, "Weather station": UNSUPPORTED, "Total CAPEX": OUTPUT, "Cost of debt": DIRECT})) for k in REVENUE]


def test_plan_corpus_matches_declared_expectations_and_no_unsafe_promotion():
    totals, violations, checked, mismatches = Counter(), [], 0, []
    for name, data, filename, expect in corpus():
        insp = inspect_workbook(data, filename)
        plan = build_plan(insp, **ARGS)
        assert plan["scenarios"]["overall"] == expect["overall"], (name, plan["scenarios"]["overall"])
        eligible = [s["key"] for s in plan["slices"] if s["apply_eligible"]]
        assert eligible == expect["slices"], (name, eligible)
        for label, cls in expect["classes"].items():
            hits = [i for i in plan["items"] if i["label"] == label]
            checked += 1
            if not hits or any(i["class"] != cls for i in hits):
                mismatches.append((name, label, cls, [i["class"] for i in hits]))
        for i in plan["items"]:
            totals[i["class"]] += 1
            if i["class"] == FORMULA and i["proposed_field"]:
                violations.append((name, "formula promoted", i["label"]))
            if i["class"] != DIRECT and i["supported_writer"] and i["class"] != STRUCTURED:
                violations.append((name, "writer on non-direct", i["label"]))
        for sl in plan["slices"]:
            if sl["scope"] == "scenario" and sl["resolution"] == "SCENARIO_LOGIC_UNRESOLVED" and sl["apply_eligible"]:
                violations.append((name, "unresolved scenario applyable", sl["key"]))
            for pr in sl["_proposals"]:
                if pr["sheet"] == "Results" or pr["sheet"] == "Summary":
                    violations.append((name, "output sheet row in slice", pr["label"]))
                if pr["status"] in ("ready", "needs_review") and (not pr["field_id"] or pr["original_value"].startswith("=")):
                    violations.append((name, "non-scalar row approvable", pr["label"]))
        assert not insp.report["executed"]["vba"]
    print(f"\nBENCHMARK plan-level: workbooks={len(corpus())} classified_items={sum(totals.values())} classes={dict(totals)} "
          f"declared_label_checks={checked} mismatches={len(mismatches)} safety_violations={len(violations)}")
    assert not mismatches, mismatches
    assert not violations, violations
    assert checked >= 40


@pytest.fixture(autouse=True)
def _routes(env):
    from app.v2.import_router import router
    if not any(getattr(r, "path", "") == "/v2/workbook/import/assess" for r in env.client.app.routes):
        env.client.app.include_router(router, prefix="/v2")


EXPECTED = {"capex.C.grid_connection": 1200.0005, "capex.D.audit_legal": 45.5, "opex.lines.insurance": 42.5,
            "debt.senior.interest_rate_pct": 5.25}


def apply_benchmark(env, kind, approve_audit):
    project = env.create(kind)
    assessed = assess(env, project, european_workbook(kind))
    assert assessed.status_code == 200
    reviewed = review_slice(env, project, slice_ticket(assessed, "SINGLE_CASE"), hidden(assessed, "manifest_token"))
    revenue_label, revenue_field = REVENUE[kind][0][0], REVENUE[kind][1]
    labels = {"Grid interconnection cost": "capex.C.grid_connection", "Insurance (annual)": "opex.lines.insurance",
              "Cost of debt": "debt.senior.interest_rate_pct", revenue_label: revenue_field}
    if approve_audit:
        labels["Audit and legal"] = "capex.D.audit_legal"
    sel = {}
    for label, field in labels.items():
        rid = ids_for(reviewed, label)[0]
        sel[f"keep_{rid}"], sel[f"field_{rid}"] = "on", field
    confirmed = confirm(env, project, hidden(reviewed, "preview_ticket"), **sel)
    assert confirmed.status_code == 200
    applied = apply(env, project, hidden(confirmed, "approved_ticket"))
    assert applied.status_code == 200 and "CHANGED — NOT RUN" in applied.text
    pis = ProjectInputSet.from_snapshot(ws_of(project).draft_snapshot, workbook=WORKBOOK)
    expected = {**{k: v for k, v in EXPECTED.items() if approve_audit or k != "capex.D.audit_legal"}, revenue_field: REVENUE[kind][2]}
    for field, value in expected.items():
        assert float(pis.get(field)) == pytest.approx(value), field
    run_model(env, project)
    return project, applied


@pytest.mark.parametrize("kind", ["solar", "wind", "data_center", "ev_charging"])
def test_applied_benchmark_exact_values_manual_run_and_real_difference(env, kind):
    project, applied = apply_benchmark(env, kind, approve_audit=True)
    compared = reconcile(env, project, hidden(applied, "applied_ticket"))
    detail = recon_detail(compared)["Results!B2"]
    entry = get_run_history("decision-user", project.project_id)[0]
    finco = Decimal(str(entry.runtime_summary["total_capex_keur"]))
    # every precedent was mapped AND applied and the user attests scope: the REAL residual is reported (no parity claimed)
    assert detail["status"] in ("PASS", "FAIL")
    assert Decimal(detail["cells"][3]) == finco and Decimal(detail["cells"][4].split()[0]) == finco - Decimal("1245.5005")
    print(f"\nBENCHMARK applied[{kind}]: canonical_values_exact=5 manual_run=1 source_total_capex=1245.5005 "
          f"finco_total_capex={finco} residual={finco - Decimal('1245.5005')} status={detail['status']} (scope attested, not a parity claim)")


def test_applied_benchmark_leaving_a_precedent_unapplied_is_not_comparable(env):
    project, applied = apply_benchmark(env, "solar", approve_audit=False)
    status, reasons = recon_rows(reconcile(env, project, hidden(applied, "applied_ticket")))["Results!B2"]
    assert status == "NOT_COMPARABLE" and "SOURCE_OUTPUT_DEPENDS_ON_UNMAPPED_OR_UNAPPLIED_INPUTS" in reasons
