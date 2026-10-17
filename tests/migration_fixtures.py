"""Original, synthetic workbook builders for the Excel migration tests (no real company, project or macro code).

``add_cached_values`` patches stored formula results into the sheet XML (openpyxl cannot compute them), so the
"cached value vs formula" paths are exercised against realistic containers.  ``build_vba_container`` writes a
minimal OLE container that has only stream NAMES - there is no VBA source anywhere in this repository.
"""
from __future__ import annotations

import io
import re
import struct
import zipfile

from openpyxl import Workbook

ENDOFCHAIN = 0xFFFFFFFE
FREESECT = 0xFFFFFFFF
FATSECT = 0xFFFFFFFD


def save(book: Workbook) -> bytes:
    stream = io.BytesIO()
    book.save(stream)
    return stream.getvalue()


def _rewrite(data: bytes, edit) -> bytes:
    src = zipfile.ZipFile(io.BytesIO(data))
    out = io.BytesIO()
    with zipfile.ZipFile(out, "w", zipfile.ZIP_DEFLATED) as dst:
        for info in src.infolist():
            dst.writestr(info.filename, edit(info.filename, src.read(info.filename)))
        extra = edit(None, None)
        for name, blob in (extra or {}).items():
            dst.writestr(name, blob)
    return out.getvalue()


def sheet_part(data: bytes, title: str) -> str:
    with zipfile.ZipFile(io.BytesIO(data)) as z:
        wb = z.read("xl/workbook.xml").decode()
        names = re.findall(r'<sheet [^>]*name="([^"]+)"', wb)
        return f"xl/worksheets/sheet{names.index(title) + 1}.xml"


def add_cached_values(data: bytes, cached: dict[tuple[str, str], float]) -> bytes:
    """cached = {(sheet title, 'B7'): 0.123} -> stores <v> next to the formula, like a saved Excel file."""
    targets = {sheet_part(data, t): {} for t, _ in cached}
    for (t, ref), value in cached.items():
        targets[sheet_part(data, t)][ref] = value

    def edit(name, blob):
        if name is None:
            return None
        if name in targets:
            text = blob.decode()
            for ref, value in targets[name].items():
                text, n = re.subn(r'(<c r="%s"[^>]*>\s*<f>[^<]*</f>\s*)<v\s*/>|(<c r="%s"[^>]*>\s*<f>[^<]*</f>\s*)<v></v>' % (ref, ref),
                                  lambda m: (m.group(1) or m.group(2)) + f"<v>{value}</v>", text)
                assert n == 1, (name, ref, text[:400])
            return text.encode()
        return blob
    return _rewrite(data, edit)


def build_vba_container(modules: list[str], *, project_stream: bool = True) -> bytes:
    """Minimal OLE compound file: Root / VBA storage / {dir, _VBA_PROJECT, *modules} and a PROJECT stream."""
    names = [("Root Entry", 5), ("VBA", 1), ("dir", 2), ("_VBA_PROJECT", 2)] + [(m, 2) for m in modules]
    if project_stream:
        names.append(("PROJECT", 2))
    count = len(names)
    entries = []
    for i, (name, kind) in enumerate(names):
        raw = name.encode("utf-16-le")
        body = bytearray(128)
        body[:len(raw)] = raw
        struct.pack_into("<H", body, 64, len(raw) + 2)
        body[66] = kind
        body[67] = 1
        left = right = child = FREESECT
        if i == 0:
            child = 1
        elif i == 1:                                                  # VBA storage: first stream, sibling = PROJECT
            child = 2
            right = count - 1 if project_stream else FREESECT
        elif kind == 2 and i < (count - 1 if project_stream else count) and i + 1 < (count - 1 if project_stream else count):
            right = i + 1
        struct.pack_into("<III", body, 68, left, right, child)
        struct.pack_into("<I", body, 116, ENDOFCHAIN)
        struct.pack_into("<Q", body, 120, 0 if kind != 2 else 100 + 10 * i)
        entries.append(bytes(body))
    while len(entries) % 4:
        entries.append(bytes(128))
    dir_sectors = len(entries) // 4
    fat = [FREESECT] * 128
    fat[0] = FATSECT
    for s in range(1, dir_sectors + 1):
        fat[s] = s + 1 if s < dir_sectors else ENDOFCHAIN
    header = bytearray(512)
    header[:8] = b"\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1"
    struct.pack_into("<HH", header, 24, 0x3E, 3)
    struct.pack_into("<H", header, 28, 0xFFFE)
    struct.pack_into("<HH", header, 30, 9, 6)
    struct.pack_into("<III", header, 44, 1, 1, 0)
    struct.pack_into("<II", header, 56, 4096, ENDOFCHAIN)
    struct.pack_into("<II", header, 64, 0, ENDOFCHAIN)
    struct.pack_into("<I", header, 72, 0)
    difat = [0] + [FREESECT] * 108
    struct.pack_into("<109I", header, 76, *difat)
    return bytes(header) + struct.pack("<128I", *fat) + b"".join(entries)


def to_xlsm(data: bytes, vba: bytes | None, *, signed: bool = False) -> bytes:
    """Re-label an .xlsx container as macro-enabled and attach ``vba`` as xl/vbaProject.bin (never executed)."""
    def edit(name, blob):
        if name is None:
            extra = {}
            if vba is not None:
                extra["xl/vbaProject.bin"] = vba
            if signed:
                extra["xl/vbaProjectSignature.bin"] = b"\x00" * 16
            return extra
        if name == "[Content_Types].xml":
            text = blob.decode()
            text = text.replace("application/vnd.openxmlformats-officedocument.spreadsheetml.sheet.main+xml",
                                "application/vnd.ms-excel.sheet.macroEnabled.main+xml")
            if vba is not None:
                text = text.replace("</Types>", '<Default Extension="bin" ContentType="application/vnd.ms-office.vbaProject"/></Types>')
            return text.encode()
        return blob
    return _rewrite(data, edit)


def add_external_link(data: bytes, target: str = "file:///C:/private/other.xlsx") -> bytes:
    def edit(name, blob):
        if name is None:
            return {
                "xl/externalLinks/externalLink1.xml":
                    '<?xml version="1.0"?><externalLink xmlns="http://schemas.openxmlformats.org/spreadsheetml/2006/main"/>',
                "xl/externalLinks/_rels/externalLink1.xml.rels":
                    '<?xml version="1.0"?><Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships">'
                    f'<Relationship Id="rId1" Type="x" Target="{target}" TargetMode="External"/></Relationships>',
            }
        return blob
    return _rewrite(data, edit)


def add_defined_names(data: bytes, names: dict[str, str]) -> bytes:
    def edit(name, blob):
        if name == "xl/workbook.xml":
            text = blob.decode()
            block = "<definedNames>" + "".join(f'<definedName name="{k}">{v}</definedName>' for k, v in names.items()) + "</definedNames>"
            if "<definedNames>" in text:
                text = text.replace("</definedNames>", block[len("<definedNames>"):])
            else:
                text = re.sub(r"(</sheets>)", r"\1" + block, text, count=1)
            return text.encode()
        return blob if name is not None else None
    return _rewrite(data, edit)


def set_calc_iterative(data: bytes) -> bytes:
    def edit(name, blob):
        if name == "xl/workbook.xml":
            text = blob.decode()
            text = re.sub(r"<calcPr[^>]*/>", '<calcPr calcId="1" iterate="1"/>', text)
            if "<calcPr" not in text:
                text = text.replace("</workbook>", '<calcPr calcId="1" iterate="1"/></workbook>')
            return text.encode()
        return blob if name is not None else None
    return _rewrite(data, edit)


# ── workbook shapes ─────────────────────────────────────────────────────────────────────────────────────────
def scenario_columns_workbook() -> Workbook:
    """Static scenario columns (Base/Upside/Downside), common inputs, structured rows, a formula input and outputs."""
    book = Workbook()
    inputs = book.active
    inputs.title = "Inputs"
    inputs.append(["Assumption", "Base Case", "Upside", "Downside", "Unit"])
    inputs.append(["PPA tariff", 60, 66, 54, "EUR/MWh"])
    inputs.append(["Insurance (annual)", 40, 38, 44, "kEUR/yr"])
    inputs.append(["Grid interconnection cost", 1200000, 1100000, 1300000, "EUR"])
    inputs.append(["Nameplate capacity", 20, 20, 20, "MW"])
    common = book.create_sheet("Common")
    common.append(["Assumption", "Value", "Unit"])
    common.append(["Cost of debt", "5,25 %", ""])
    common.append(["Audit and legal", 45.5, "kEUR"])
    common.append(["Transformer", 120, "kEUR"])
    common.append(["Merchant price curve", 55, "EUR/MWh"])
    common.append(["Weather station", "n/a", ""])
    common.append(["Target leverage", "=B2*2", ""])
    summary = book.create_sheet("Summary")
    summary.append(["Metric", "Value", "Unit"])
    summary.append(["Total CAPEX", "=Common!B3+Common!B4", "kEUR"])
    summary.append(["Equity IRR", "=Common!B2*3", ""])
    summary.append(["Project IRR", 0.1, ""])
    summary["B3"].number_format = "0.0%"
    return book


def selector_workbook() -> Workbook:
    """A switch cell read by CHOOSE: only the case selected at save time is observable."""
    book = Workbook()
    ws = book.active
    ws.title = "Model"
    ws.append(["Scenario selector", 2])
    ws.append(["Assumption", "Value", "Case 1", "Case 2", "Case 3"])
    ws.append(["PPA tariff", "=CHOOSE($B$1,C3,D3,E3)", 60, 66, 54])
    ws.append(["Insurance (annual)", "=CHOOSE($B$1,C4,D4,E4)", 40, 38, 44])
    ws.append(["Cost of debt", 5.25, None, None, None])
    return book


def selector_only_workbook() -> Workbook:
    """Switch + CHOOSE over unlabeled value ranges: scenario identity cannot be resolved statically."""
    book = Workbook()
    ws = book.active
    ws.title = "Model"
    ws.append(["Scenario selector", 2])
    ws.append(["Assumption", "Value", "Unit"])
    ws.append(["PPA tariff", "=CHOOSE($B$1,Data!B1,Data!C1,Data!D1)", "EUR/MWh"])
    ws.append(["Cost of debt", 5.25, "%"])
    data = book.create_sheet("Data")
    data.append(["PPA tariff by case", 60, 66, 54])
    return book
