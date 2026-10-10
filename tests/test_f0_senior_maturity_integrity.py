"""Production maturity protection; no alternative financial calculation."""
from __future__ import annotations

import asyncio
import copy
import re
from contextlib import closing
from dataclasses import replace
from types import SimpleNamespace as NS

import pytest

from app import project_factories as factories
from app.services import production_financial_authority as authority
from finco_core.inputs._models import DebtSizingMode, GearingCapRepaymentMethod, SponsorFundingMode


def project(kind, *, gearing=.95, tenor=8):
    pi = getattr(factories, f"create_generic_{kind}_reference")()
    return replace(pi, financing=replace(
        pi.financing, debt_sizing_mode=DebtSizingMode.GEARING_CAP,
        gearing_cap_repayment_method=GearingCapRepaymentMethod.DSCR_SCULPTED,
        gearing_ratio=gearing, senior_tenor_years=tenor,
    ))


@pytest.fixture(scope="module")
def valid_run():
    return authority.run_clean_production(factories.create_generic_solar_reference())


@pytest.fixture(autouse=True)
def executor_cleanup():
    from app.runtime.model_execution import reset_model_executor_for_tests
    yield
    assert reset_model_executor_for_tests(None) == []


@pytest.mark.parametrize("kind,balance,maturity", [
    ("solar", 20912.234045226287, 18), ("wind", 23644.893677596927, 19),
])
def test_original_p0_rejected_before_publication_even_on_cache_hit(monkeypatch, kind, balance, maturity):
    import financial_engine.financial_statements as statements
    pi = project(kind)
    before = copy.deepcopy(pi)
    original = authority._memoised_policy_run
    calculated = []
    authority._POLICY_RUN_CACHE.clear()

    def capture(inputs, policy, compute):
        def execute():
            result = compute()
            calculated.append(result)
            return result
        return original(inputs, policy, execute)

    def forbidden(*args, **kwargs):
        pytest.fail("An unpaid maturity liability reached financial statement publication")

    monkeypatch.setattr(authority, "_memoised_policy_run", capture)
    monkeypatch.setattr(statements, "assemble_decision_complete_financial_statements", forbidden)
    for _ in range(2):
        with pytest.raises(authority.CleanProductionRunUnavailable) as exc:
            authority.run_clean_production(pi)
        assert exc.value.reason_code == "SENIOR_MATURITY_UNSETTLED_LIABILITY"
        assert f"contractual_maturity_period={maturity}" in exc.value.detail
        reported = float(re.search(r"outstanding_principal_keur=([^;]+)", exc.value.detail).group(1))
        assert reported == pytest.approx(balance, abs=1e-9)
        assert "no complete canonical settlement/refinancing/accounting authority" in exc.value.detail
    assert len(calculated) == 1  # One policy entry, including its existing fixed points.
    g2c = calculated[0][0]
    senior = g2c.financing_result.project_model_result.senior_debt
    assert senior.senior_debt_closing_keur[-1] == pytest.approx(balance, abs=1e-9)
    assert senior.debt_size_keur - sum(senior.senior_principal_keur) == pytest.approx(balance, abs=1e-8)
    assert pi == before


def test_cache_cannot_publish_balloon_by_zeroing_terminal_copies(monkeypatch):
    from financial_engine.financing.generic_product_policy import run_with_generic_financing_policy
    from financial_engine.project_returns.contracts import DebtTerminalStatus
    from financial_engine.shareholder_waterfall import run_project_shareholder_waterfall_model
    import financial_engine.financial_statements as statements
    pi = project("solar")
    g2c, effective, policy_evidence = run_with_generic_financing_policy(
        pi, run_project_shareholder_waterfall_model,
    )
    model = g2c.financing_result.project_model_result
    sd = model.senior_debt
    assert sd.senior_debt_closing_keur[-1] > 20000
    fake_sd = replace(sd, senior_debt_closing_keur=sd.senior_debt_closing_keur[:-1] + (0.,))
    fake_term = replace(g2c.return_summary.terminal.senior,
        balance_at_contractual_maturity_keur=0., terminal_model_horizon_balance_keur=0.,
        status=DebtTerminalStatus.REPAID)
    fake_g2c = replace(g2c,
        financing_result=replace(g2c.financing_result,
            project_model_result=replace(model, senior_debt=fake_sd)),
        return_summary=replace(g2c.return_summary,
            terminal=replace(g2c.return_summary.terminal, senior=fake_term)))
    monkeypatch.setattr(authority, "_memoised_policy_run", lambda *args: (fake_g2c, effective, policy_evidence))
    monkeypatch.setattr(statements, "assemble_decision_complete_financial_statements",
                        lambda *args: pytest.fail("Corrupt cached repayment evidence reached publication"))
    with pytest.raises(authority.CleanProductionRunUnavailable) as exc:
        authority.run_clean_production(pi)
    assert exc.value.reason_code == "SENIOR_MATURITY_EVIDENCE_INVALID"
    assert "roll-forward" in exc.value.detail


def evidence(run, *, residual=None, **changes):
    g2c = run.g2c_result
    model = g2c.financing_result.project_model_result
    senior = model.senior_debt
    terminal = g2c.return_summary.terminal.senior
    if residual is not None:
        from financial_engine.project_returns.contracts import DebtTerminalStatus
        principal = list(senior.senior_principal_keur)
        opening = list(senior.senior_debt_opening_keur)
        closing = list(senior.senior_debt_closing_keur)
        service = list(senior.senior_debt_service_keur)
        last_payment = max(i for i, p in enumerate(principal) if p > 0)
        principal[last_payment] -= residual
        service[last_payment] -= residual
        for i in range(last_payment, len(closing)):
            closing[i] += residual
            if i > last_payment:
                opening[i] += residual
        senior = replace(senior, senior_debt_opening_keur=tuple(opening),
            senior_principal_keur=tuple(principal), senior_debt_closing_keur=tuple(closing),
            senior_debt_service_keur=tuple(service))
        terminal = replace(terminal, balance_at_contractual_maturity_keur=residual,
            terminal_model_horizon_balance_keur=residual,
            status=DebtTerminalStatus.OUTSTANDING_AT_MATURITY)
    senior = changes.get("senior", senior)
    terminal = changes.get("terminal", terminal)
    model = replace(model, senior_debt=senior, axis_contract=changes.get("axis", model.axis_contract))
    return NS(financing_result=NS(project_model_result=model, final_senior_commitment_keur=senior.debt_size_keur if senior else None),
        return_summary=NS(terminal=NS(senior=terminal)))


@pytest.mark.parametrize("residual", [0.0, 1e-8, 2.3221237597681466e-5, 1e-4])
def test_existing_solver_precision_does_not_modify_balances_or_status(valid_run, residual):
    g2c = evidence(valid_run, residual=residual)
    before = copy.deepcopy(g2c)
    authority._require_settled_senior_maturity(g2c, valid_run.project_inputs)
    assert g2c == before
    assert g2c.return_summary.terminal.senior.balance_at_contractual_maturity_keur == residual


def test_above_existing_absolute_solver_precision_is_rejected(valid_run):
    with pytest.raises(authority.CleanProductionRunUnavailable, match="SENIOR_MATURITY_UNSETTLED_LIABILITY"):
        authority._require_settled_senior_maturity(evidence(valid_run, residual=1.00001e-4), valid_run.project_inputs)


def test_real_small_solver_residual_is_preserved_in_post_maturity_balance_sheet():
    from app.run_integrity import build_run_integrity_evidence, run_integrity_checks
    pi = factories.create_generic_solar_reference()
    pi = replace(pi, financing=replace(pi.financing, gearing_ratio=.9, target_dscr=1.3))
    run = authority.run_clean_production(pi)
    closing = run.g2c_result.financing_result.project_model_result.senior_debt.senior_debt_closing_keur[-1]
    assert closing == pytest.approx(2.3221237597681466e-5, abs=1e-12)
    report = run_integrity_checks(build_run_integrity_evidence(run)).to_dict()
    assert report["overall"] == "PASS"
    assert all(c["reason_code"] != "BALANCE_SHEET_IMBALANCE" for c in report["checks"])
    senior = run.g2c_result.financing_result.project_model_result.senior_debt
    post_maturity = [b for b in run.financial_statements_result.balance_sheet_periods
                     if b.period_index > senior.period_indices[-1]]
    assert post_maturity
    assert all(b.senior_debt_balance_keur == closing for b in post_maturity)
    assert run.g2c_result.return_summary.terminal.senior.status.value == "OUTSTANDING_AT_MATURITY"


def test_zero_funding_not_applicable_requires_actual_zero_schedule(valid_run):
    from financial_engine.project_returns.contracts import DebtTerminalStatus
    model = valid_run.g2c_result.financing_result.project_model_result
    zero = tuple(0. for _ in model.senior_debt.period_indices)
    sd = replace(model.senior_debt, debt_size_keur=0., senior_debt_opening_keur=zero,
                 senior_principal_keur=zero, senior_debt_closing_keur=zero,
                 senior_interest_keur=zero, senior_debt_service_keur=zero)
    term = replace(valid_run.g2c_result.return_summary.terminal.senior,
        contractual_maturity_period_index=None, contractual_maturity_date=None,
        balance_at_contractual_maturity_keur=0., terminal_model_horizon_balance_keur=0.,
        status=DebtTerminalStatus.NOT_APPLICABLE)
    authority._require_settled_senior_maturity(evidence(valid_run, senior=sd, terminal=term), valid_run.project_inputs)
    with pytest.raises(authority.CleanProductionRunUnavailable, match="SENIOR_MATURITY_EVIDENCE_INVALID"):
        authority._require_settled_senior_maturity(evidence(valid_run, terminal=term), valid_run.project_inputs)


@pytest.mark.parametrize("gearing", [1e-12, 2e-13])
def test_real_microscopic_funding_preserves_canonical_not_applicable_precision(gearing):
    from financial_engine.project_returns.contracts import DebtTerminalStatus
    pi = factories.create_generic_solar_reference()
    pi = replace(pi, financing=replace(pi.financing, debt_sizing_mode=DebtSizingMode.GEARING_CAP,
        gearing_ratio=gearing, sponsor_funding_mode=SponsorFundingMode.EQUITY_ONLY))
    run = authority.run_clean_production(pi)
    sd = run.g2c_result.financing_result.project_model_result.senior_debt
    assert 0 < sd.debt_size_keur < 1e-7
    assert run.g2c_result.return_summary.terminal.senior.status is DebtTerminalStatus.NOT_APPLICABLE
    assert sd.senior_debt_closing_keur[-1] < 1e-7
    assert run.financial_statements_result is not None


@pytest.mark.parametrize("fault", ["missing_schedule", "missing_maturity", "wrong_maturity",
    "missing_date", "wrong_date", "negative", "nan", "infinite", "none_balance", "bool_balance",
    "short_vector", "wrong_axis", "duplicate_period", "concealed_balance", "unknown_status"])
def test_incomplete_or_malformed_evidence_never_defaults_to_zero(valid_run, fault):
    from datetime import date
    model = valid_run.g2c_result.financing_result.project_model_result
    sd = model.senior_debt
    term = valid_run.g2c_result.return_summary.terminal.senior
    kwargs = {}
    if fault == "missing_schedule":
        kwargs["senior"] = None
    elif fault in ("missing_maturity", "wrong_maturity"):
        kwargs["terminal"] = replace(term, contractual_maturity_period_index=None if fault == "missing_maturity" else 999)
    elif fault in ("missing_date", "wrong_date"):
        kwargs["terminal"] = replace(term, contractual_maturity_date=None if fault == "missing_date" else date(1999, 1, 1))
    elif fault in ("negative", "nan", "infinite", "none_balance", "bool_balance"):
        kwargs["residual"] = {"negative": -1., "nan": float("nan"), "infinite": float("inf"),
                              "none_balance": None, "bool_balance": True}[fault]
        if fault == "none_balance":
            kwargs = {"senior": replace(sd, senior_debt_closing_keur=sd.senior_debt_closing_keur[:-1] + (None,))}
    elif fault == "short_vector":
        kwargs["senior"] = replace(sd, senior_debt_closing_keur=sd.senior_debt_closing_keur[:-1])
    elif fault == "wrong_axis":
        kwargs["axis"] = replace(model.axis_contract, senior_axis=model.axis_contract.senior_axis[:-1])
    elif fault == "duplicate_period":
        kwargs["senior"] = replace(sd, period_indices=sd.period_indices[:-1] + (sd.period_indices[0],))
    elif fault == "concealed_balance":
        kwargs["senior"] = replace(sd, senior_debt_closing_keur=sd.senior_debt_closing_keur[:-1] + (100.,))
    else:
        kwargs["terminal"] = replace(term, status="UNKNOWN")
    with pytest.raises(authority.CleanProductionRunUnavailable) as exc:
        authority._require_settled_senior_maturity(evidence(valid_run, **kwargs), valid_run.project_inputs)
    assert exc.value.reason_code == "SENIOR_MATURITY_EVIDENCE_INVALID"


def test_unsupported_repayment_contract_fails_closed(valid_run):
    # Deliberately corrupt a copied typed input: normal constructors already reject it.
    pi = copy.deepcopy(valid_run.project_inputs)
    object.__setattr__(pi.financing, "gearing_cap_repayment_method", "UNSUPPORTED")
    with pytest.raises(authority.CleanProductionRunUnavailable, match="SENIOR_MATURITY_EVIDENCE_INVALID"):
        authority._require_settled_senior_maturity(valid_run.g2c_result, pi)


@pytest.mark.parametrize("field", ["senior_debt_opening_keur", "senior_principal_keur", "senior_debt_service_keur"])
def test_changed_repayment_amounts_fail_existing_rollforward_contract(valid_run, field):
    sd = valid_run.g2c_result.financing_result.project_model_result.senior_debt
    values = list(getattr(sd, field))
    values[0] += 100.
    corrupt = evidence(valid_run, senior=replace(sd, **{field: tuple(values)}))
    with pytest.raises(authority.CleanProductionRunUnavailable, match="SENIOR_MATURITY_EVIDENCE_INVALID"):
        authority._require_settled_senior_maturity(corrupt, valid_run.project_inputs)


def test_funded_commitment_cannot_disagree_with_opening_schedule(valid_run):
    corrupt = evidence(valid_run)
    corrupt.financing_result.final_senior_commitment_keur += 100.
    with pytest.raises(authority.CleanProductionRunUnavailable, match="SENIOR_MATURITY_EVIDENCE_INVALID"):
        authority._require_settled_senior_maturity(corrupt, valid_run.project_inputs)


@pytest.mark.parametrize("kind", ["solar", "wind", "data_center", "ev_charging"])
def test_fully_amortizing_reference_remains_publishable(kind):
    run = authority.run_clean_production(getattr(factories, f"create_generic_{kind}_reference")())
    assert run.financial_statements_result is not None
    assert run.g2c_result.financing_result.project_model_result.senior_debt.senior_debt_closing_keur[-1] <= 1e-4


@pytest.mark.parametrize("consumer", ["api_runner", "waterfall", "recalculating_export"])
def test_shared_production_consumers_cannot_bypass_gate(consumer):
    pi = project("solar")
    with pytest.raises(authority.CleanProductionRunUnavailable, match="SENIOR_MATURITY_UNSETTLED_LIABILITY"):
        if consumer == "api_runner":
            from app.api.project_runner import run_project
            run_project("Generic Solar Reference", "Base", project_inputs_override=pi)
        elif consumer == "waterfall":
            from app.services.production_waterfall_seam import execute_production_waterfall
            execute_production_waterfall(pi)
        else:
            from app.services.production_waterfall_seam import execute_production_demo
            execute_production_demo("Generic Solar Reference", project_inputs_override=pi)


def test_sensitivity_and_goal_seek_have_no_success_metric_for_balloon():
    from app.runtime.model_execution import ModelExecutor, ModelExecutionConfig, reset_model_executor_for_tests
    from app.services.goal_seek import GOAL_SEEK_METRICS, make_canonical_evaluator
    from app.services.sensitivity_execution import run_sensitivity_points
    pi = project("solar")
    point = run_sensitivity_points("Generic Solar Reference", [pi])[0]
    assert "SENIOR_MATURITY_UNSETTLED_LIABILITY" in point["error"]
    assert "kpis" not in point
    reset_model_executor_for_tests(ModelExecutor(ModelExecutionConfig(mode="thread")))
    evaluate = make_canonical_evaluator("Generic Solar Reference", pi, GOAL_SEEK_METRICS["project_irr"])
    assert asyncio.run(evaluate([pi.revenue.ppa_base_tariff])) == [None]


def test_authenticated_failed_run_preserves_last_run_history_and_export(tmp_path, monkeypatch, caplog):
    from fastapi.testclient import TestClient
    from app.auth import COOKIE_NAME, create_session_token
    from app.persistence import db
    from app.persistence.run_history_repository import get_run_history
    from app.persistence.workspace_repository import get_workspace_state
    from app.runtime.model_execution import ModelExecutor, ModelExecutionConfig, reset_model_executor_for_tests
    from app.services.reference_seed_service import create_reference_seeded_project
    from app.workbook.service import WorkbookService
    from app.workbook.workbook_identity import assemble_consistent_for_get
    from main_web import app

    monkeypatch.setattr(db, "DB_PATH", str(tmp_path / "maturity.db"))
    db.init_db()
    original_factory = factories.create_generic_solar_reference
    # A supported typed gearing-cap/sculpted baseline. Only fixture input policy
    # differs; all adapters, persistence, worker and financial calculations are real.
    def seed():
        base = original_factory()
        return replace(base, financing=replace(base.financing,
            debt_sizing_mode=DebtSizingMode.GEARING_CAP, gearing_ratio=.25,
            sponsor_funding_mode=SponsorFundingMode.EQUITY_ONLY,
            gearing_cap_repayment_method=GearingCapRepaymentMethod.DSCR_SCULPTED))
    monkeypatch.setattr(factories, "create_generic_solar_reference", seed)
    record = create_reference_seeded_project(user_id="maturity-owner", requested_name="Maturity safety",
        template_source="generic_solar_reference", capacity_mw=64)
    reset_model_executor_for_tests(ModelExecutor(ModelExecutionConfig(mode="thread")))

    def tokens():
        ws = get_workspace_state("maturity-owner", record.project_id)
        pis = WorkbookService.build_draft_input_set_from_workspace(ws)
        identity = assemble_consistent_for_get(user_id="maturity-owner", project_id=record.project_id,
                                              workbook_version=pis.workbook_version)
        return ws, {"project": record.project_code, "workbook_version": pis.workbook_version,
                    "content_hash": identity.composite_hash}

    with closing(TestClient(app)) as client:
        client.cookies.set(COOKIE_NAME, create_session_token(user_id="maturity-owner", username="admin"))
        assert client.get("/v2/workbook", params={"project": record.project_code}).status_code == 200
        _, data = tokens()
        response = client.post("/v2/workbook/run", headers={"HX-Request": "true"}, data=data)
        assert response.status_code == 200
        valid, _ = tokens()
        assert WorkbookService.get_runtime_result(valid) is not None, response.text[:1000]
        history = get_run_history("maturity-owner", record.project_id)
        assert len(history) == 1
        original_export = client.post("/v2/workbook/export", data={"project": record.project_code})
        assert original_export.status_code == 200
        for field, value in (("debt.senior.gearing_pct", "95"), ("debt.senior.tenor_years", "8")):
            _, data = tokens()
            response = client.post("/v2/workbook/update", headers={"HX-Request": "true"},
                data=data | {"sheet_id": "debt", "field_id": field, "value": value})
            assert '"workbook-field-saved"' in response.headers.get("HX-Trigger", ""), response.text[:600]
        stale, data = tokens()
        assert stale.dirty
        assert "STALE" in client.get("/v2/workbook", params={"project": record.project_code}).text
        effective = WorkbookService.build_draft_input_set_from_workspace(stale).to_projectinputs()
        assert effective.financing.gearing_ratio == .95
        assert effective.financing.senior_tenor_years == 8
        response = client.post("/v2/workbook/run", headers={"HX-Request": "true"}, data=data)
        assert response.status_code == 200
        assert 'data-testid="banner-error"' in response.text
        # The route's existing safe banner does not print raw engine internals.
        # Verify the real typed exception caught there, not a fabricated response.
        failures = [r.exc_info[1] for r in caplog.records if r.exc_info]
        assert any(getattr(e, "reason_code", None) == "SENIOR_MATURITY_UNSETTLED_LIABILITY"
                   for e in failures)
        after, _ = tokens()
        assert after == stale
        assert after.last_runtime_snapshot_id == valid.last_runtime_snapshot_id
        assert after.last_runtime_summary == valid.last_runtime_summary
        assert after.last_runtime_snapshot == valid.last_runtime_snapshot
        assert get_run_history("maturity-owner", record.project_id) == history
        assert WorkbookService.get_runtime_result(after) == WorkbookService.get_runtime_result(valid)
        def no_calculation(*args, **kwargs):
            pytest.fail("Committed Last Run export attempted financial execution")
        monkeypatch.setattr(authority, "run_clean_production", no_calculation)
        exported = client.post("/v2/workbook/export", data={"project": record.project_code})
        assert exported.status_code == 200 and exported.content[:2] == b"PK"
        from io import BytesIO
        from openpyxl import load_workbook
        old_book = load_workbook(BytesIO(original_export.content), data_only=True)
        new_book = load_workbook(BytesIO(exported.content), data_only=True)
        # Export metadata may truthfully report changed Working state. Financial
        # statement cells still belong exclusively to the immutable old Last Run.
        statement_sheets = ("P&L", "Balance Sheet", "Cash Flow")
        for sheet in statement_sheets:
            def financial_cells(book):
                return tuple(row for row in book[sheet].values if row[0] != "Export generated at")
            assert financial_cells(old_book) == financial_cells(new_book)
        assert get_run_history("maturity-owner", record.project_id) == history
        assert get_workspace_state("maturity-owner", record.project_id).last_runtime_summary == valid.last_runtime_summary


def test_stateless_model_api_rejects_new_balloon_calculation(monkeypatch):
    from fastapi.testclient import TestClient
    from app.api.v1 import model_run
    from app.runtime.model_execution import ModelExecutor, ModelExecutionConfig, reset_model_executor_for_tests
    from main_web import app
    pi = project("solar")
    monkeypatch.setattr(model_run, "get_pi", lambda key: pi)
    reset_model_executor_for_tests(ModelExecutor(ModelExecutionConfig(mode="thread")))
    with closing(TestClient(app)) as client:
        response = client.post("/api/v1/model/references/generic_solar_reference/run",
                               json={"capacity_mw": pi.technical.capacity_mw})
    assert response.status_code == 503
    reported = float(re.search(r"outstanding_principal_keur=([^;]+)", response.json()["detail"]).group(1))
    assert reported == pytest.approx(20912.234045226287, abs=1e-9)
    assert "data" not in response.json()
