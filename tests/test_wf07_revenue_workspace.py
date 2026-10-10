"""Real authorized Save / Scenario / Run / Last Run export acceptance."""
import copy
import json
from io import BytesIO

import pytest

from tests.test_model_decision_workspace_v2 import env, tokens
from tests.test_wf07_revenue_multistream import contracts
from app.workbook.revenue_multistream import SNAPSHOT_KEY


@pytest.mark.parametrize("kind", ["solar", "wind"])
def test_contract_save_reload_scenario_run_and_immutable_export(env, monkeypatch, kind):
    from app.persistence.workspace_repository import get_workspace_state
    from app.persistence.scenarios_repository import list_scenarios, get_scenario, duplicate_scenario
    from app.persistence.run_history_repository import get_run_history
    from app.workbook.service import WorkbookService
    from app.services.export_service import resolve_canonical_last_run_from_workspace
    from app.services.v2_export_service import build_canonical_last_run_institutional_workbook_export
    owner = "decision-user"
    record = env.create(kind)
    project = {"project": record.project_code}
    def page():
        response = env.client.get("/v2/workbook", params=project)
        assert response.status_code == 200, response.text
        return response
    def ws():
        return get_workspace_state(owner, record.project_id)
    pi = WorkbookService.build_draft_input_set_from_workspace(ws()).to_projectinputs()
    payload = contracts(pi, kind, price=95, merchant_price=90, strike=100)
    def save(raw, **extra):
        return env.client.post("/v2/workbook/revenue/contracts", headers={"HX-Request": "true"},
                               data={**project, **tokens(page().text), "contracts": json.dumps(raw) if isinstance(raw, dict) else raw, **extra})
    def run():
        response = env.client.post("/v2/workbook/run", headers={"HX-Request": "true"}, data={**project, **tokens(page().text)})
        assert 'id="v2-sheet-returns"' in response.text, response.text[:3000]
        if (ws().last_runtime_summary or {}).get("revenue_multistream_audit"):
            assert 'data-testid="revenue-contracts-audit"' in response.text
        assert not ws().dirty
        return copy.deepcopy(ws().last_runtime_summary)
    assert save(payload).status_code == 200
    assert ws().dirty
    assert json.loads(ws().draft_snapshot[SNAPSHOT_KEY]) == payload
    assert "synthetic-executed-contract" in page().text
    baseline = run()
    assert baseline["revenue_multistream_audit"]
    assert "multistream_config_json" in baseline["scenario_revenue_input"]
    assert 'data-testid="revenue-contracts-audit"' in page().text
    assert baseline["total_revenue_keur"] > 0
    created = env.client.post("/v2/workbook/scenarios/create", data={**project, "scenario_name": "Contract alternative"})
    assert created.status_code == 200
    sc = next(s for s in list_scenarios(owner, record.project_id) if not s.is_base_case)
    selected = env.client.post("/v2/workbook/scenarios/select", data={**project, "scenario_id": sc.scenario_id})
    assert selected.status_code == 200
    original_draft = copy.deepcopy(ws().draft_snapshot)
    different = copy.deepcopy(payload)
    different["streams"][0]["price_eur_mwh"] = 105
    assert save(different).status_code == 200
    assert ws().draft_snapshot == original_draft
    assert json.loads(get_scenario(sc.scenario_id, owner).overrides[SNAPSHOT_KEY]) == different
    alternative = run()
    assert alternative["total_revenue_keur"] > baseline["total_revenue_keur"]
    assert alternative["total_ebitda_keur"] > baseline["total_ebitda_keur"]
    assert alternative["project_irr"] != baseline["project_irr"]
    copied = duplicate_scenario(owner, sc.scenario_id, "Contract copy")
    assert copied.overrides == get_scenario(sc.scenario_id, owner).overrides
    history_before = get_run_history(owner, record.project_id)
    third = copy.deepcopy(different)
    third["streams"][0]["price_eur_mwh"] = 115
    assert save(third).status_code == 200
    assert ws().dirty
    assert get_run_history(owner, record.project_id) == history_before
    authority = resolve_canonical_last_run_from_workspace(record, owner, ws())
    assert authority.working_changed_since_run
    assert json.loads(authority.project_inputs.revenue.multistream_config_json) == different
    tampered = copy.deepcopy(ws())
    tampered.last_runtime_summary["scenario_revenue_input"]["multistream_config_json"] = json.dumps(third)
    with pytest.raises(ValueError, match="contracts do not match"):
        resolve_canonical_last_run_from_workspace(record, owner, tampered)
    def forbidden(*args, **kwargs):
        raise AssertionError("Export cannot execute the engine")
    with monkeypatch.context() as patcher:
        patcher.setattr("app.runtime.model_execution.run_model_process", forbidden)
        patcher.setattr("app.services.production_financial_authority.run_clean_production", forbidden)
        exported = build_canonical_last_run_institutional_workbook_export("generic_" + kind + "_reference",
            safe_project=record.project_code, project_record=record, user_id=owner)
        assert exported.status_code == 200, exported.error_content
        from openpyxl import load_workbook
        book = load_workbook(BytesIO(exported.bytes_data), read_only=True, data_only=True)
        try:
            assert "Revenue_Streams" in book.sheetnames
            rows = list(book["Revenue_Streams"].iter_rows(min_row=8, values_only=True))
            assert any(row[1] == "ppa-1" and row[5] == 105 for row in rows)
            assert not any(row[1] == "ppa-1" and row[5] == 115 for row in rows)
        finally:
            book.close()
    assert save("", action="inherit").status_code == 200
    assert SNAPSHOT_KEY not in get_scenario(sc.scenario_id, owner).overrides
    inherited = run()
    assert inherited["total_revenue_keur"] == baseline["total_revenue_keur"]


@pytest.mark.parametrize("kind", ["solar", "wind", "data_center", "ev_charging"])
def test_route_context_and_sector_contract(env, kind):
    record = env.create(kind)
    response = env.client.get("/v2/workbook", params={"project": record.project_code})
    assert response.status_code == 200
    assert ('data-testid="revenue-contracts-form"' in response.text) == (kind in ("solar", "wind"))
    if kind in ("data_center", "ev_charging"):
        from app.project_factories import create_generic_solar_reference
        raw = contracts(create_generic_solar_reference())
        result = env.client.post("/v2/workbook/revenue/contracts", data={"project": record.project_code,
            **tokens(response.text), "contracts": json.dumps(raw)})
        assert result.status_code == 422


def test_cas_owner_and_invalid_write_are_atomic(env):
    from app.persistence.workspace_repository import get_workspace_state
    from app.workbook.service import WorkbookService
    record = env.create()
    project = {"project": record.project_code}
    html = env.client.get("/v2/workbook", params=project).text
    ws = get_workspace_state("decision-user", record.project_id)
    raw = json.dumps(contracts(WorkbookService.build_draft_input_set_from_workspace(ws).to_projectinputs()))
    invalid = env.client.post("/v2/workbook/revenue/contracts", data={**project, **tokens(html), "contracts": "{}"})
    assert invalid.status_code == 422
    assert get_workspace_state("decision-user", record.project_id).draft_snapshot == ws.draft_snapshot
    valid = env.client.post("/v2/workbook/revenue/contracts", data={**project, **tokens(html), "contracts": raw})
    assert valid.status_code == 200
    stale = env.client.post("/v2/workbook/revenue/contracts", data={**project, **tokens(html), "contracts": ""})
    assert stale.status_code == 409
    from app.auth import COOKIE_NAME, create_session_token
    env.client.cookies.set(COOKIE_NAME, create_session_token(user_id="other-contract-owner", username="other"))
    denied = env.client.post("/v2/workbook/revenue/contracts", data={**project, **tokens(html), "contracts": raw})
    assert denied.status_code == 404


def test_scenario_pricing_alias_conflict_never_silently_ignores_a_price():
    from app.workbook.scenario_revenue_authority import bind_scenario_tariff
    from app.project_factories import create_generic_solar_reference
    raw = json.dumps(contracts(create_generic_solar_reference()))
    with pytest.raises(ValueError, match="replace the explicit contracts"):
        bind_scenario_tariff({"project_type": "solar", SNAPSHOT_KEY: raw}, {"tariff_eur_mwh": 90})


def test_selected_scenario_contracts_reject_legacy_price_edits_and_old_selection_token(env):
    from app.persistence.workspace_repository import get_workspace_state
    from app.workbook.service import WorkbookService
    record = env.create()
    project = {"project": record.project_code}
    base_page = env.client.get("/v2/workbook", params=project).text
    created = env.client.post("/v2/workbook/scenarios/create", data={**project, "scenario_name": "Explicit contracts"})
    assert created.status_code == 200
    ws = get_workspace_state("decision-user", record.project_id)
    pi = WorkbookService.build_draft_input_set_from_workspace(ws).to_projectinputs()
    raw = json.dumps(contracts(pi))
    old_selection = env.client.post("/v2/workbook/revenue/contracts",
        data={**project, **tokens(base_page), "contracts": raw})
    assert old_selection.status_code == 409
    page = env.client.get("/v2/workbook", params=project).text
    saved = env.client.post("/v2/workbook/revenue/contracts", data={**project, **tokens(page), "contracts": raw})
    assert saved.status_code == 200
    before = get_workspace_state("decision-user", record.project_id)
    assert not before.draft_snapshot.get(SNAPSHOT_KEY)
    page = env.client.get("/v2/workbook", params=project).text
    rejected = env.client.post("/v2/workbook/update", headers={"HX-Request": "true"},
        data={**project, **tokens(page), "sheet_id": "revenue", "field_id": "revenue.ppa.base_tariff", "value": "123"})
    assert rejected.status_code == 200
    assert "workbook-field-error" in rejected.headers.get("HX-Trigger", "")
    assert "REVENUE_V2_PRICING_AUTHORITY" in rejected.text
    assert get_workspace_state("decision-user", record.project_id) == before


def test_transaction_rechecks_project_authority(env):
    from app.persistence.db import get_connection
    from app.workbook.revenue_contracts_service import save_contracts, RevenueContractConflict
    record = env.create()
    page = env.client.get("/v2/workbook", params={"project": record.project_code})
    conn = get_connection()
    try:
        conn.execute("UPDATE projects SET archived=1 WHERE project_id=?", (record.project_id,))
        conn.commit()
    finally:
        conn.close()
    with pytest.raises(RevenueContractConflict, match="unavailable or read-only"):
        save_contracts(owner="decision-user", project_id=record.project_id,
                       expected_hash=tokens(page.text)["content_hash"], raw="")
