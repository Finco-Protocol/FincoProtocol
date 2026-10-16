"""Excel migration V3 — non-executing inspection, VBA metadata, scenario detection and classification (module level)."""
from __future__ import annotations

import io
import struct
import zipfile

import pytest
from openpyxl import Workbook

from app.model_import.intake import IntakeError, extract_proposals
from app.model_import.migration_plan import (
    DIRECT, FORMULA, OUTPUT, STRUCTURED, UNSUPPORTED, build_plan,
)
from app.model_import.workbook_inspect import (
    InspectionError, analyse_formula, expand_ref, inspect_workbook, parse_vba_container,
)
from tests import migration_fixtures as F

PLAN_ARGS = dict(project_type="solar", template_source="generic_solar_reference", project_code="P-1",
                 target_scenario_id=None, target_scenario_name=None)


def plan_of(data: bytes, name: str = "m.xlsx"):
    insp = inspect_workbook(data, name)
    return insp, build_plan(insp, **PLAN_ARGS)


def by_label(plan, label, scenario=None):
    hits = [i for i in plan["items"] if i["label"] == label and (scenario is None or i["scenario_key"] == scenario)]
    assert hits, (label, [i["label"] for i in plan["items"]])
    return hits[0]


# ── Phase A: structure ────────────────────────────────────────────────────────────────────────────────────────
def test_normal_xlsx_structure_and_cached_vs_formula():
    data = F.add_cached_values(F.save(F.scenario_columns_workbook()), {("Summary", "B2"): 165.5, ("Summary", "B3"): 0.1575})
    insp, plan = plan_of(data)
    r = insp.report
    assert r["format"] == "xlsx" and not r["macro_bearing"] and r["vba"]["status"] == "ABSENT"
    assert [s["name"] for s in r["sheets"]] == ["Inputs", "Common", "Summary"]
    assert r["formulas"]["total"] == 3 and r["formulas"]["with_cached_value"] == 2
    cached = insp.formulas[("Summary", "B2")]
    assert cached["text"].startswith("=Common!B3") and cached["cached"] == 165.5 and cached["has_cached"]
    assert insp.formulas[("Common", "B7")]["has_cached"] is False                   # formula without a stored result
    assert "FORMULAS_WITHOUT_CACHED_VALUES" in r["warnings"]
    assert r["executed"] == {"vba": False, "formulas": False, "external_links": False, "dde": False}
    # a cached result is evidence, never an authoritative input
    out = by_label(plan, "Total CAPEX")
    assert out["class"] == OUTPUT and out["formula"]["cached_is_authoritative"] is False and out["proposed_field"] is None


def test_hidden_sheets_named_ranges_and_calc_flags():
    book = Workbook()
    ws = book.active
    ws.title = "Inputs"
    ws.append(["Assumption", "Value", "Unit"])
    ws.append(["Cost of debt", 5.25, "%"])
    hidden = book.create_sheet("Calc")
    hidden["A1"], hidden["B1"] = "secret factor", 3
    hidden.sheet_state = "veryHidden"
    data = F.set_calc_iterative(F.add_defined_names(F.save(book), {
        "DebtRate": "Inputs!$B$2", "Const": "0.5", "Ext": "[1]Sheet1!$A$1", "Calc_Rate": "Inputs!$B$2*2"}))
    insp = inspect_workbook(data, "m.xlsx")
    sheets = {s["name"]: s for s in insp.report["sheets"]}
    assert sheets["Calc"]["hidden"] and sheets["Calc"]["state"] == "veryHidden" and not sheets["Inputs"]["hidden"]
    names = {n["name"]: n for n in insp.report["defined_names"]}
    assert names["DebtRate"]["kind"] == "cell" and names["DebtRate"]["target"] == {"sheet": "Inputs", "start": "B2", "end": None}
    assert names["Const"]["kind"] == "constant" and names["Ext"]["kind"] == "external" and names["Calc_Rate"]["kind"] == "formula"
    assert {"HIDDEN_SHEETS_PRESENT", "ITERATIVE_CALCULATION_ENABLED"} <= set(insp.report["warnings"])


def test_formula_references_functions_and_active_flags():
    info = analyse_formula('=IF(Sel!$B$1=2,SUM(Data!B2:D2),OFFSET(A1,1,0))+INDIRECT("Z9")+WEBSERVICE("http://x")', "Home", {})
    assert {"IF", "SUM", "OFFSET", "INDIRECT", "WEBSERVICE"} <= set(info["functions"])
    assert "ACTIVE_OR_NETWORK_FUNCTION" in info["flags"] and "DYNAMIC_OR_VOLATILE_FUNCTION" in info["flags"]
    refs = {(r["sheet"], r["start"], r["end"]) for r in info["refs"]}
    assert ("Sel", "B1", None) in refs and ("Data", "B2", "D2") in refs
    assert "Z9" not in str(info["refs"])                      # string literals are never parsed as references
    assert "EXTERNAL_WORKBOOK_REFERENCE" in analyse_formula("=[1]Sheet1!A1+1", "Home", {})["flags"]
    assert expand_ref({"sheet": "S", "start": "A1", "end": "C2"}) == [("S", c) for c in ("A1", "B1", "C1", "A2", "B2", "C2")]
    assert expand_ref({"sheet": "S", "start": "A1", "end": "ZZ9999"}) is None             # unexpanded, never unbounded


def test_external_links_are_described_but_never_followed():
    data = F.add_external_link(F.save(F.scenario_columns_workbook()), "file:///C:/private/other.xlsx")
    insp, plan = plan_of(data)
    assert insp.report["external_links"] == [{"kind": "file_path", "target_mode": "External", "target_hint": "other.xlsx"}]
    assert "EXTERNAL_LINKS_PRESENT" in insp.report["warnings"]
    assert any("EXTERNAL_LINKS" in b for b in plan["macro_assessment"]["prevents_accurate_migration"])
    # the strict apply/preview path still refuses the same container
    with pytest.raises(IntakeError) as exc:
        extract_proposals(data, "m.xlsx", project_type="solar", template_source="generic_solar_reference")
    assert exc.value.code == "IMPORT_XLSX_ACTIVE_OR_EXTERNAL_CONTENT"


# ── macro-enabled inspection ──────────────────────────────────────────────────────────────────────────────────
def test_xlsm_is_inspected_without_execution_and_vba_names_are_metadata_only():
    vba = F.build_vba_container(["ThisWorkbook", "Module1", "ScenarioSwitch"])
    data = F.to_xlsm(F.save(F.scenario_columns_workbook()), vba, signed=True)
    insp, plan = plan_of(data, "m.xlsm")
    r = insp.report
    assert r["format"] == "xlsm" and r["macro_bearing"] and r["container"]["macro_enabled_content_type"]
    assert r["vba"]["status"] == "PARSED" and [m["name"] for m in r["vba"]["modules"]] == ["Module1", "ScenarioSwitch", "ThisWorkbook"]
    assert r["vba"]["source_interpreted"] is False and r["vba"]["executed"] is False
    assert "VBA_PROJECT_SIGNED" in r["warnings"]
    assert "code" not in str(r["vba"]).lower().replace("source_interpreted", "")        # no source in the report
    # the strict path refuses .xlsm; its scalar values reach Apply only through the same rules (slices)
    with pytest.raises(IntakeError) as exc:
        extract_proposals(data, "m.xlsm", project_type="solar", template_source="generic_solar_reference")
    assert exc.value.code == "IMPORT_FORMAT_UNSUPPORTED"
    assert plan["slices"] and any(s["apply_eligible"] for s in plan["slices"])
    assert by_label(plan, "PPA tariff", "upside")["proposed_field"] == "revenue.ppa.base_tariff"
    macro = plan["macro_assessment"]
    assert macro["macro_bearing"] and any("MACRO_BEHAVIOUR" in x for x in macro["requires_separate_approved_feature"])
    assert any("MACRO_PRESENT" in b for b in macro["prevents_accurate_migration"])


@pytest.mark.parametrize("blob,reason", [
    (b"", "VBA_CONTAINER_SIGNATURE_INVALID"),
    (b"not an ole file" * 100, "VBA_CONTAINER_SIGNATURE_INVALID"),
    (b"\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1" + b"\x00" * 600, "VBA_CONTAINER_SECTOR_SIZE_INVALID"),
])
def test_malformed_vba_container_is_reported_not_raised(blob, reason):
    parsed = parse_vba_container(blob)
    assert parsed["status"] == "MALFORMED_CONTAINER" and parsed["reason"] == reason and parsed["modules"] == []


def test_truncated_and_looping_vba_containers_fail_closed():
    good = F.build_vba_container(["Module1"])
    assert parse_vba_container(good)["status"] == "PARSED"
    assert parse_vba_container(good[:700])["status"] == "MALFORMED_CONTAINER"             # truncated
    looped = bytearray(good)
    struct.pack_into("<I", looped, 512 + 4, 1)                  # FAT entry of the directory sector points at itself
    assert parse_vba_container(bytes(looped))["status"] == "MALFORMED_CONTAINER"
    bad_name = bytearray(good)
    struct.pack_into("<H", bad_name, 512 + 512 + 128 + 64, 4000)                          # impossible name length
    assert parse_vba_container(bytes(bad_name))["status"] == "MALFORMED_CONTAINER"


def test_workbook_with_malformed_vba_still_inspects_and_warns():
    data = F.to_xlsm(F.save(F.scenario_columns_workbook()), b"garbage" * 100)
    insp = inspect_workbook(data, "m.xlsm")
    assert insp.report["vba"]["status"] == "MALFORMED_CONTAINER" and "VBA_CONTAINER_MALFORMED" in insp.report["warnings"]


def test_extension_content_mismatch_is_flagged():
    data = F.to_xlsm(F.save(F.scenario_columns_workbook()), F.build_vba_container(["Module1"]))
    assert "EXTENSION_MACRO_MISMATCH" in inspect_workbook(data, "renamed.xlsx").report["warnings"]
    assert "XLSM_WITHOUT_VBA_PROJECT" in inspect_workbook(F.to_xlsm(F.save(F.scenario_columns_workbook()), None), "m.xlsm").report["warnings"]


# ── container safety (retained guards) ────────────────────────────────────────────────────────────────────────
def _zip_with(extra: dict[str, bytes], base: bytes | None = None) -> bytes:
    base = base or F.save(F.scenario_columns_workbook())
    out = io.BytesIO()
    with zipfile.ZipFile(io.BytesIO(base)) as src, zipfile.ZipFile(out, "w", zipfile.ZIP_DEFLATED) as dst:
        for info in src.infolist():
            dst.writestr(info.filename, src.read(info.filename))
        for name, blob in extra.items():
            dst.writestr(name, blob)
    return out.getvalue()


def test_zip_bomb_and_unsafe_archives_are_rejected():
    bomb = _zip_with({"xl/worksheets/bomb.xml": b"0" * (6 * 1024 * 1024)})
    with pytest.raises(IntakeError) as exc:
        inspect_workbook(bomb, "m.xlsx")
    assert exc.value.code == "IMPORT_XLSX_UNSAFE_COMPRESSION"
    with pytest.raises(IntakeError) as exc:
        inspect_workbook(_zip_with({"../evil.xml": b"x"}), "m.xlsx")
    assert exc.value.code == "IMPORT_XLSX_UNSAFE_ARCHIVE"
    many = _zip_with({f"xl/extra/{i}.txt": b"x" for i in range(130)})
    with pytest.raises(IntakeError) as exc:
        inspect_workbook(many, "m.xlsx")
    assert exc.value.code == "IMPORT_ZIP_MEMBER_LIMIT_EXCEEDED"
    with pytest.raises(IntakeError) as exc:
        inspect_workbook(b"PK\x03\x04 definitely not a zip", "m.xlsx")
    assert exc.value.code == "IMPORT_XLSX_MALFORMED"
    with pytest.raises(IntakeError) as exc:
        inspect_workbook(b"x" * 10, "m.csv")
    assert exc.value.code == "IMPORT_FORMAT_UNSUPPORTED"
    with pytest.raises(IntakeError) as exc:
        inspect_workbook(b"x" * (2 * 1024 * 1024 + 1), "m.xlsx")
    assert exc.value.code == "IMPORT_FILE_SIZE_INVALID"


def test_xml_entity_expansion_and_dtd_are_rejected():
    base = F.save(F.scenario_columns_workbook())
    evil = ('<?xml version="1.0"?><!DOCTYPE lol [<!ENTITY a "aaaaaaaaaa"><!ENTITY b "&a;&a;&a;&a;&a;&a;&a;&a;">]>'
            '<workbook xmlns="http://schemas.openxmlformats.org/spreadsheetml/2006/main"><sheets><sheet name="&b;" sheetId="1"/></sheets></workbook>')
    out = io.BytesIO()
    with zipfile.ZipFile(io.BytesIO(base)) as src, zipfile.ZipFile(out, "w", zipfile.ZIP_DEFLATED) as dst:
        for info in src.infolist():
            dst.writestr(info.filename, evil if info.filename == "xl/workbook.xml" else src.read(info.filename))
    with pytest.raises(InspectionError) as exc:
        inspect_workbook(out.getvalue(), "m.xlsx")
    assert exc.value.code == "IMPORT_XLSX_XML_UNSAFE_OR_MALFORMED"


def test_cell_and_sheet_limits_hold():
    book = Workbook()
    for i in range(13):
        book.create_sheet(f"S{i}")
    with pytest.raises(IntakeError) as exc:
        inspect_workbook(F.save(book), "m.xlsx")
    assert exc.value.code == "IMPORT_SHEET_LIMIT_EXCEEDED"
    wide = Workbook()
    wide.active.append(list(range(30)))
    with pytest.raises(IntakeError) as exc:
        inspect_workbook(F.save(wide), "m.xlsx")
    assert exc.value.code == "IMPORT_TABLE_LIMIT_EXCEEDED"


def test_untrusted_labels_formulas_and_names_are_bounded_and_clean():
    book = Workbook()
    ws = book.active
    ws.title = "Inputs"
    ws.append(["Assumption", "Value", "Unit"])
    ws.append(["<script>alert(1)</script>\u202eInsurance (annual)" + "x" * 500, 40, "kEUR/yr"])
    ws.append(["Target", "=" + "+1" * 400, ""])
    insp, plan = plan_of(F.save(book))
    for item in plan["items"]:
        assert "\u202e" not in item["label"] and len(item["label"]) <= 160
        if item["formula"]:
            assert len(item["formula"]["text"]) <= 240


# ── Phase B: scenarios ────────────────────────────────────────────────────────────────────────────────────────
def test_static_scenario_columns_are_observable_but_label_only():
    insp, plan = plan_of(F.save(F.scenario_columns_workbook()))
    scen = plan["scenarios"]
    assert scen["overall"] == "STATIC_OBSERVABLE"
    keys = [s["key"] for st in scen["structures"] for s in st["scenarios"]]
    assert keys == ["base_case", "upside", "downside"]
    assert all(s["semantics"] == "LABEL_ONLY_UNVERIFIED" for st in scen["structures"] for s in st["scenarios"])
    up = by_label(plan, "PPA tariff", "upside")
    assert up["source"]["cell"] == "C2" and up["source_value"] == "66" and up["source_unit"] == "EUR/MWh"
    assert up["target"] == {"project": "P-1", "scenario_id": None, "scenario_name": "(default Working Copy)"}
    assert up["supported_writer"].startswith("Canonical batch writer")
    assert by_label(plan, "Grid interconnection cost", "downside")["normalized_value"] == "1300"           # EUR -> kEUR
    # one slice per scenario, each including the case-independent inputs, never mixing scenarios
    assert [s["key"] for s in plan["slices"]] == ["base_case", "upside", "downside"]
    for sl in plan["slices"]:
        assert sl["apply_eligible"] and all(not p["sheet"] == "Summary" for p in sl["_proposals"])
        values = {p["label"]: p["original_value"] for p in sl["_proposals"] if p["sheet"] == "Inputs"}
        assert values["PPA tariff"] == {"base_case": "60", "upside": "66", "downside": "54"}[sl["key"]]
        assert len([p for p in sl["_proposals"] if p["label"] == "PPA tariff"]) == 1


def test_selector_with_literal_case_columns_records_the_switch_and_its_dependents():
    insp, plan = plan_of(F.save(F.selector_workbook()))
    sel = plan["scenarios"]["selectors"]
    assert len(sel) == 1 and sel[0]["cell"] == "B1" and sel[0]["dependent_formula_count"] == 2
    assert sel[0]["currently_selected_only"] and sel[0]["resolution"] == "SCENARIO_LOGIC_UNRESOLVED"
    assert plan["scenarios"]["overall"] == "SCENARIO_LOGIC_UNRESOLVED"
    # the literal case columns remain statically observable; the live (formula) column is a dependency
    assert by_label(plan, "PPA tariff", "case_2")["source_value"] == "66"
    assert not any(i["class"] == DIRECT and i["source"]["cell"] in ("B3", "B4") for i in plan["items"])


def test_selector_only_workbook_is_scenario_logic_unresolved_and_blocks_scenario_apply():
    insp, plan = plan_of(F.save(F.selector_only_workbook()))
    assert plan["scenarios"]["overall"] == "SCENARIO_LOGIC_UNRESOLVED"
    stuck = [i for i in plan["items"] if i["scenario_key"] == "SCENARIO_LOGIC_UNRESOLVED"]
    assert {i["source"]["cell"] for i in stuck} >= {"B1", "C1", "D1"}
    assert by_label(plan, "Scenario selector")["reason"] == "SCENARIO_SELECTOR_CELL"
    assert by_label(plan, "PPA tariff", "COMMON")["class"] == FORMULA
    blocked = [s for s in plan["slices"] if s["resolution"] == "SCENARIO_LOGIC_UNRESOLVED" and s["kind"] != "COMMON"]
    assert blocked and all(not s["apply_eligible"] and "SCENARIO_LOGIC_UNRESOLVED" in s["blockers"] for s in blocked)
    # only the case-independent input is offered, and the macro assessment lists the dependency
    common = next(s for s in plan["slices"] if s["kind"] == "COMMON")
    assert common["apply_eligible"] and [p["label"] for p in common["_proposals"] if p["field_id"]] == ["Cost of debt"]
    assert plan["macro_assessment"]["scenarios_requiring_macro_or_recalculation"] == ["Currently selected case"]


def test_macro_bearing_selector_adds_the_macro_switch_reason():
    data = F.to_xlsm(F.save(F.selector_only_workbook()), F.build_vba_container(["ThisWorkbook", "Switch"]))
    insp, plan = plan_of(data, "m.xlsm")
    assert "MACRO_BEARING_WORKBOOK_SWITCH_MAY_BE_MACRO_DRIVEN" in plan["scenarios"]["selectors"][0]["reasons"]


def test_scenario_name_alone_is_not_semantics_and_never_degrades_to_common():
    book = Workbook()
    ws = book.active
    ws.title = "Inputs"
    ws.append(["Assumption", "Value", "Unit"])
    ws.append(["Cost of debt", "5,25 %", ""])
    up = book.create_sheet("Upside")
    up.append(["Assumption", "Value", "Unit"])
    up.append(["Insurance (annual)", 99, "kEUR/yr"])
    insp, plan = plan_of(F.save(book))
    assert plan["scenarios"]["structures"][0]["evidence"] == "NAME_ONLY"
    item = by_label(plan, "Insurance (annual)")
    assert item["scenario_key"] == "upside" and item["item_status"] == "BLOCKED_SCENARIO_LOGIC_UNRESOLVED" and item["supported_writer"] is None
    assert not any(p["label"] == "Insurance (annual)" for s in plan["slices"] for p in s["_proposals"])


def test_partially_observable_scenario_keeps_literals_and_reports_formula_cells():
    book = Workbook()
    ws = book.active
    ws.title = "Inputs"
    ws.append(["Assumption", "Base Case", "Upside", "Unit"])
    ws.append(["PPA tariff", 60, "=B2*1.1", "EUR/MWh"])
    ws.append(["Insurance (annual)", 40, 38, "kEUR/yr"])
    insp, plan = plan_of(F.save(book))
    states = {s["key"]: s["resolution"] for st in plan["scenarios"]["structures"] for s in st["scenarios"]}
    assert states == {"base_case": "STATIC_OBSERVABLE", "upside": "PARTIALLY_OBSERVABLE"}
    assert plan["scenarios"]["overall"] == "PARTIALLY_OBSERVABLE"
    assert by_label(plan, "PPA tariff", "upside")["class"] == FORMULA
    assert by_label(plan, "Insurance (annual)", "upside")["class"] == DIRECT


def test_sheet_per_scenario_structure():
    book = Workbook()
    for title, tariff in (("Inputs_Base", 60), ("Inputs_Downside", 54)):
        ws = book.create_sheet(title)
        ws.append(["Assumption", "Value", "Unit"])
        ws.append(["PPA tariff", tariff, "EUR/MWh"])
    del book["Sheet"]
    insp, plan = plan_of(F.save(book))
    assert plan["scenarios"]["structures"][0]["evidence"] == "SHEET_PER_SCENARIO"
    assert {s["key"]: next(p["original_value"] for p in s["_proposals"]) for s in plan["slices"]} == {"base": "60", "downside": "54"}


# ── Phase C: classification ───────────────────────────────────────────────────────────────────────────────────
def test_every_class_is_produced_and_nothing_is_promoted():
    insp, plan = plan_of(F.add_cached_values(F.save(F.scenario_columns_workbook()), {("Summary", "B2"): 165.5, ("Summary", "B3"): 0.1575}))
    classes = {c: [i["label"] for i in plan["items"] if i["class"] == c] for c in plan["class_counts"]}
    assert "Cost of debt" in classes[DIRECT] and "Audit and legal" in classes[DIRECT]
    assert {"Transformer", "Merchant price curve", "Nameplate capacity"} <= set(classes[STRUCTURED])
    assert classes[FORMULA] == ["Target leverage"]
    assert {"Total CAPEX", "Equity IRR", "Project IRR"} <= set(classes[OUTPUT])
    assert classes[UNSUPPORTED] == ["Weather station"]
    formula_item = by_label(plan, "Target leverage")
    assert formula_item["proposed_field"] is None and formula_item["normalized_value"] is None and formula_item["supported_writer"] is None
    for item in plan["items"]:
        if item["class"] != DIRECT:
            assert item["proposed_field"] is None and item["item_status"] == "NOT_APPLYABLE" and item["supported_writer"] in (None,) or item["class"] == STRUCTURED
    # structured data is named with its owner editor and never enters a slice as applyable
    assert "Revenue contract editor" in by_label(plan, "Merchant price curve")["supported_writer"]
    for sl in plan["slices"]:
        assert not any(p["field_id"] is None and p["status"] in ("ready", "needs_review") for p in sl["_proposals"])


def test_mapping_completeness_is_separate_from_reconciliation():
    insp, plan = plan_of(F.save(F.scenario_columns_workbook()))
    m = plan["mapping"]
    assert m["input_candidates"] == m["direct_total"] + sum(plan["class_counts"][c] for c in (STRUCTURED, FORMULA, UNSUPPORTED))
    assert 0 < m["input_mapping_completeness"] < 1 and "parity" in m["definition"]
    assert "reconcil" not in " ".join(m.keys())


def test_ambiguous_units_and_numbers_are_blocked_with_a_required_decision():
    book = Workbook()
    ws = book.active
    ws.title = "Inputs"
    ws.append(["Assumption", "Value", "Unit"])
    ws.append(["Insurance (annual)", 40, ""])                    # currency unit missing
    ws.append(["Cost of debt", 5.25, ""])                         # percent scale unknown
    ws.append(["Audit and legal", "1,234", "kEUR"])               # thousands or decimal?
    insp, plan = plan_of(F.save(book))
    for label, reason in (("Insurance (annual)", "IMPORT_CURRENCY_UNIT_REQUIRED"), ("Cost of debt", "IMPORT_PERCENT_SCALE_AMBIGUOUS"),
                          ("Audit and legal", "IMPORT_NUMBER_FORMAT_AMBIGUOUS")):
        item = by_label(plan, label)
        assert item["class"] == DIRECT and item["item_status"] == "BLOCKED" and item["reason"] == reason
        assert item["required_unit_decision"] and item["normalized_value"] is None


def test_duplicate_mapping_targets_are_flagged_for_review():
    book = Workbook()
    ws = book.active
    ws.title = "Inputs"
    ws.append(["Assumption", "Value", "Unit"])
    ws.append(["Insurance (annual)", 40, "kEUR/yr"])
    ws.append(["Insurance (annual)", 41, "kEUR/yr"])
    insp, plan = plan_of(F.save(book))
    items = [i for i in plan["items"] if i["label"] == "Insurance (annual)"]
    assert len(items) == 2 and all(i["reason"] == "IMPORT_DUPLICATE_TARGET_REQUIRES_REVIEW" for i in items)
    assert all(i["item_status"] == "REVIEW_REQUIRED" for i in items)


def test_output_sheets_never_feed_scalar_apply_even_for_exact_field_ids():
    book = Workbook()
    ws = book.active
    ws.title = "Inputs"
    ws.append(["Assumption", "Value", "Unit"])
    ws.append(["Cost of debt", 5.25, "%"])
    out = book.create_sheet("Results")
    out.append(["Assumption", "Value", "Unit"])
    out.append(["debt.senior.interest_rate_pct", 4.5, "%"])                  # an echo of an input on an output sheet
    insp, plan = plan_of(F.save(book))
    echo = by_label(plan, "debt.senior.interest_rate_pct")
    assert echo["class"] == OUTPUT and echo["supported_writer"] is None
    for sl in plan["slices"]:
        assert all(p["sheet"] != "Results" for p in sl["_proposals"])


def test_wide_period_tables_are_reported_as_uninterpreted():
    book = Workbook()
    ws = book.active
    ws.title = "Inputs"
    ws.append(["Assumption", "Value", "Unit"])
    ws.append(["Insurance (annual)", 40, "kEUR/yr"])
    series = book.create_sheet("Series")
    series.append(["Year", 2025, 2026, 2027, 2028])
    series.append(["Generation", 100, 101, 102, 103])
    insp, plan = plan_of(F.save(book))
    assert plan["coverage"]["numeric_cells_not_interpreted"].get("Series", 0) >= 4
    assert any("UNINTERPRETED_REGIONS" in b for b in plan["macro_assessment"]["prevents_accurate_migration"])
